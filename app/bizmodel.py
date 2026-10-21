"""Business-model pages (E35, ADR-014): /business-models/ and /business-models/<key>.

A page explains how a kind of business makes money; its tables and charts are computed here from the
template's sample brief (Base scenario, ``sample_stage``, Full tier, monthly, 5 years = the committed sample
``samples/<key>-full-5y.xlsx`` the page offers) by the engine's plain-Python reference (``reference_run``,
``annualise``, ``engine.goals``) -- the same reference the challenger uses -- so a page never shows a figure the
workbook would not. Every table cell also carries the address(es) of the workbook cell it must equal (taken
from the build plan, never hard-coded rows); ``tests/test_bizmodel.py`` rebuilds and recalculates the sample
and compares every cell.

Copy lives in ``content/business-models/<template key>.md``: the SEO front matter (``slug``, ``title``,
``description``, ``keyword``, ``template``, ``sample``) and a body that may contain

    {{table:<name>}}          a table on its own line: drivers, startup, pnl, costs, unit, breakeven
    {{chart:<name>}}          a chart on its own line: revenue_costs, cost_structure
    {{num:<table>.<row>[.<column>][|abs]}}   one figure of a table, formatted as in the table

An unknown placeholder is a build error (the page fails to load), so copy can never drift from the data.
"""

from __future__ import annotations

import html
import math
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from functools import cache, lru_cache
from pathlib import Path

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import HTMLResponse
from markupsafe import Markup
from openpyxl.utils import get_column_letter as L

from app import calculators, seo
from app.web import templates
from engine.brief import Brief
from engine.builder import Plan, manifest_for, plan, q
from engine.core import LAND_LINE, annualise, reference_run
from engine.dashboard import SHEET as DASH_SHEET
from engine.goals import apply_goal
from engine.goals import reference as goal_reference
from engine.samples import SampleSpec
from engine.spec import Timing
from engine.templates import TEMPLATES

ROOT = Path(__file__).resolve().parents[1]
CONTENT = ROOT / "content" / "business-models"

SECTION = "business-models"
SECTION_TITLE = "Business models"
SECTION_INTRO = (
    "How different kinds of business make money: revenue drivers, costs, start-up investment, unit economics "
    "and break-even, with tables from a working financial model."
)
CAPTION = "Illustrative Base case of the sample model — your model uses your numbers"
SLUG_RE = re.compile(r"^/business-models/[a-z][a-z0-9_]*$")
BASE_SCENARIO = 2
HORIZON_YEARS = 5

TABLE_NAMES = ("drivers", "startup", "pnl", "costs", "unit", "breakeven")
CHART_NAMES = ("revenue_costs", "cost_structure")

DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
PLACEHOLDER_RE = re.compile(r"\{\{\s*([a-z]+)\s*:\s*([^{}]*?)\s*\}\}")
NUM_RE = re.compile(r"^([a-z0-9_]+)\.([a-z0-9_]+)(?:\.([a-z0-9_]+))?(\|abs)?$")


class PlaceholderError(ValueError):
    """Unknown or misplaced placeholder in a business-model page."""


def sample_spec(key: str) -> SampleSpec:
    """The sample every page of template ``key`` is computed from and offers for download."""
    return SampleSpec(key=key, tier="full", granularity="monthly", horizon_years=HORIZON_YEARS)


# --------------------------------------------------------------------------- data model
@dataclass(frozen=True)
class Ref:
    """Workbook cells a table cell must equal: ``sum(coef * cell for num) / sum(coef * cell for den)``."""

    num: tuple[tuple[float, str], ...]
    den: tuple[tuple[float, str], ...] = ()


def at(addr: str, coef: float = 1.0) -> Ref:
    return Ref(((coef, addr),))


@dataclass(frozen=True)
class Cell:
    value: float | str
    ref: Ref | None = None


@dataclass(frozen=True)
class Row:
    key: str
    label: str
    cells: tuple[Cell, ...]
    fmt: str = "usd"
    bold: bool = False
    total: bool = False
    unit: str = ""


@dataclass(frozen=True)
class Table:
    name: str
    title: str
    columns: tuple[tuple[str, str], ...]  # (key, header)
    rows: tuple[Row, ...]
    unit_column: bool = False
    row_header: str = ""

    def row(self, key: str) -> Row:
        for r in self.rows:
            if r.key == key:
                return r
        raise KeyError(key)

    def cell(self, row: str, col: str | None = None) -> tuple[Row, Cell]:
        keys = [k for k, _ in self.columns]
        col = col or keys[0]
        if col not in keys:
            raise KeyError(col)
        r = self.row(row)
        return r, r.cells[keys.index(col)]


@dataclass(frozen=True)
class BizModel:
    key: str
    title: str
    spec: SampleSpec
    word: str  # "month" / "year"
    tables: dict[str, Table]
    charts: dict[str, dict]  # chart data (see _chart_data)
    years: int


# --------------------------------------------------------------------------- number formats
def _q(v: float, places: int) -> str:
    d = Decimal(repr(abs(float(v)))).quantize(Decimal(1).scaleb(-places), ROUND_HALF_UP)
    return f"{d:,.{places}f}"


def format_value(v: float | str, fmt: str, *, prose: bool = False) -> str:
    """A table value the way the workbook shows it: costs in brackets, zero as a dash."""
    if isinstance(v, str):
        return v
    zero = abs(v) < 0.5 if fmt in ("usd", "int", "days") else abs(v) < 0.05 if fmt in ("usd2", "dec1") else v == 0
    if fmt == "pct":
        zero = abs(v) < 0.0005
    if zero:
        return "-"
    neg = v < 0
    if fmt == "usd":
        s = "$" + _q(v, 0)
    elif fmt == "usd2":
        s = "$" + _q(v, 2)
    elif fmt == "pct":
        s = _q(v * 100, 1) + "%"
    elif fmt in ("int", "days"):
        s = _q(v, 0)
    elif fmt == "dec1":
        s = _q(v, 1)
    elif fmt == "dec3":
        s = _q(v, 3)
    elif fmt == "x":
        s = _q(v, 1) + "x"
    elif fmt == "period":
        s = ("month " if prose else "Month ") + _q(v, 0)
    else:
        raise ValueError(f"unknown format {fmt!r}")
    return f"({s})" if neg else s


def tolerance(fmt: str) -> float:
    """How far a table cell may sit from its workbook cell: $1 for dollars, 1e-6 for ratios and counts."""
    return {"usd": 1.0, "usd2": 0.01}.get(fmt, 1e-6)


# --------------------------------------------------------------------------- KPI formulas (plain Python)
_TOKEN = re.compile(r"\s*(\{[^{}]+\}|\d+\.?\d*|[A-Za-z_][A-Za-z_0-9]*|<=|>=|<>|[-+*/^(),=<>])")


def _tokens(text: str) -> list[str]:
    out, pos = [], 0
    text = text.strip()
    while pos < len(text):
        m = _TOKEN.match(text, pos)
        if not m:
            raise ValueError(f"cannot read formula {text!r} at {pos}")
        out.append(m.group(1))
        pos = m.end()
    return out


class _Parser:
    """Recursive-descent parser for the small formula language of KPI rows (no ``eval``)."""

    def __init__(self, tokens: list[str]):
        self.t, self.i = tokens, 0

    def peek(self) -> str:
        return self.t[self.i] if self.i < len(self.t) else ""

    def take(self) -> str:
        tok = self.peek()
        self.i += 1
        return tok

    def parse(self):
        node = self.compare()
        if self.i != len(self.t):
            raise ValueError(f"unexpected {self.peek()!r}")
        return node

    def compare(self):
        left = self.add()
        if self.peek() in ("=", "<>", "<", ">", "<=", ">="):
            op = self.take()
            return ("cmp", op, left, self.add())
        return left

    def add(self):
        node = self.mul()
        while self.peek() in ("+", "-"):
            op = self.take()
            node = ("bin", op, node, self.mul())
        return node

    def mul(self):
        node = self.unary()
        while self.peek() in ("*", "/"):
            op = self.take()
            node = ("bin", op, node, self.unary())
        return node

    def unary(self):
        if self.peek() == "-":
            self.take()
            return ("neg", self.unary())
        node = self.atom()
        if self.peek() == "^":
            self.take()
            node = ("bin", "^", node, self.unary())
        return node

    def atom(self):
        tok = self.take()
        if tok == "(":
            node = self.compare()
            if self.take() != ")":
                raise ValueError("missing )")
            return node
        if tok.startswith("{"):
            return ("ref", tok[1:-1])
        if re.fullmatch(r"\d+\.?\d*", tok):
            return ("num", float(tok))
        if re.fullmatch(r"[A-Za-z_]\w*", tok) and self.peek() == "(":
            self.take()
            args = []
            if self.peek() != ")":
                args.append(self.compare())
                while self.peek() == ",":
                    self.take()
                    args.append(self.compare())
            if self.take() != ")":
                raise ValueError("missing )")
            return ("fn", tok.upper(), args)
        raise ValueError(f"unexpected {tok!r}")


def _evaluate(node, ref: Callable[[str], float]) -> float:
    kind = node[0]
    if kind == "num":
        return node[1]
    if kind == "ref":
        return ref(node[1])
    if kind == "neg":
        return -_evaluate(node[1], ref)
    if kind == "bin":
        a, b = _evaluate(node[2], ref), _evaluate(node[3], ref)
        op = node[1]
        return a + b if op == "+" else a - b if op == "-" else a * b if op == "*" else a / b if op == "/" else a**b
    if kind == "cmp":
        a, b = _evaluate(node[2], ref), _evaluate(node[3], ref)
        return float({"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[node[1]])
    name, args = node[1], node[2]
    if name == "IF":  # lazy: the unchosen branch may divide by zero
        return _evaluate(args[1] if _evaluate(args[0], ref) else args[2], ref)
    vals = [_evaluate(a, ref) for a in args]
    if name == "MAX":
        return max(vals)
    if name == "MIN":
        return min(vals)
    if name == "AND":
        return float(all(vals))
    if name == "OR":
        return float(any(vals))
    if name == "ABS":
        return abs(vals[0])
    raise ValueError(f"unsupported function {name}")


# --------------------------------------------------------------------------- computation
class _Recorder:
    """Wraps a template's ``reference_top`` and keeps what ``reference_run`` does not return (opex by line)."""

    def __init__(self, fn):
        self.fn, self.tops = fn, []

    def __call__(self, d, p, y, t, hc):
        top = self.fn(d, p, y, t, hc)
        self.tops.append(top)
        return top


@dataclass
class _Ctx:
    key: str
    brief: Brief
    pl: Plan
    t: Timing
    d: dict
    series: dict[str, list[float]]
    rec: _Recorder
    goal: dict
    goal_years: dict
    man: object

    @property
    def H(self) -> int:
        return self.t.years

    def chunk(self, vals: list[float], y: int) -> list[float]:
        return vals[(y - 1) * self.t.ppy : y * self.t.ppy]

    def annual(self, vals: list[float], y: int, agg: str = "sum") -> float:
        c = self.chunk(vals, y)
        return sum(c) if agg == "sum" else c[0] if agg == "first" else c[-1]

    def opex(self, key: str) -> list[float]:
        return [top.opex[key] for top in self.rec.tops]

    def pl_cell(self, key: str, y: int) -> str:
        return self.pl.stmt_cell("pl", key, y)

    def grid(self, key: str, p: int) -> str:
        return self.pl.period_cell(key, p, "")


def _year_cols(H: int) -> tuple[tuple[str, str], ...]:
    return tuple((f"y{y}", f"Year {y}") for y in range(1, H + 1))


def _t_drivers(c: _Ctx) -> Table:
    specs = {s.key: s for s in c.pl.specs}
    rows = []
    for key in c.pl.template.key_numbers:
        s, dv = specs[key], c.brief.driver(key)
        cells = tuple(
            Cell(getattr(dv, col), at(c.pl.input_col(key, letter)))
            for col, letter in (("low", "B"), ("base", "C"), ("high", "D"))
        )
        rows.append(Row(key, s.label, cells, s.fmt, unit=s.unit))
    cols = (("low", "Low"), ("base", "Base"), ("high", "High"))
    return Table("drivers", "Key numbers: Low, Base and High", cols, tuple(rows), unit_column=True, row_header="Driver")


def _capex_items(c: _Ctx, low_p: int) -> list[Row]:
    """Capital spending by item up to the lowest cash point; empty when the template has a single line."""
    lines = [ln for ln in c.pl.template.capex_lines(c.t) if ln.key != LAND_LINE]
    parts = [ln for ln in lines if ln.key != "capex"]
    if not parts:
        return []
    tops = c.rec.tops
    total = [top.capex for top in tops]
    resolved = {ln.key: [top.kpis[ln.key] for top in tops] for ln in parts if ln.key in tops[0].kpis}
    rows = []

    def cells_to_low(key: str, coef: float = 1.0) -> tuple[tuple[float, str], ...]:
        return tuple((coef, c.grid(key, p)) for p in range(1, low_p + 1))

    for ln in parts:
        if ln.key in resolved:
            rows.append(
                Row(
                    ln.key,
                    ln.label,
                    (Cell(sum(resolved[ln.key][:low_p]), Ref(cells_to_low(ln.key))),),
                )
            )
    rest = [ln for ln in parts if ln.key not in resolved]
    if rest:
        other = [t - sum(r[i] for r in resolved.values()) for i, t in enumerate(total)]
        num = cells_to_low("capex")
        for key in resolved:
            num += cells_to_low(key, -1.0)
        label = rest[0].label if len(rest) == 1 else "Other capital spending"
        rows.append(Row("capex_other", label, (Cell(sum(other[:low_p]), Ref(num)),)))
    return rows


def _t_startup(c: _Ctx) -> Table:
    lay = c.pl.modules["breakeven"]
    row_label = {r.row: r.label for r in lay.rows}

    def goal(name: str, value: float, **kw) -> Row:
        addr = lay.local[name]
        label = row_label[int(re.sub(r"\D", "", addr))]
        return Row(name, label, (Cell(value, at(f"{q(lay.sheet)}!{addr}")),), "usd", **kw)

    g = c.goal
    rows = _capex_items(c, g["low_p"])
    rows += [
        goal("capex_to_low", g["capex_to_low"], bold=bool(rows), total=bool(rows)),
        goal("burn_to_low", g["burn_to_low"]),
        goal("repay_to_low", g["repay_to_low"]),
        goal("uses", g["uses"], bold=True, total=True),
        goal("buffer", g["buffer"]),
        goal("required", g["required"], bold=True, total=True),
        goal("equity", c.d["equity"]),
        goal("loan", c.d["loan_amount"]),
        goal("sources", g["sources"], bold=True, total=True),
        goal("gap", g["gap"], bold=True, total=True),
    ]
    return Table("startup", "Start-up costs and funding need", (("value", "Base case"),), tuple(rows))


def _t_pnl(c: _Ctx) -> Table:
    H, tpl = c.H, c.pl.template
    s = c.series
    ann = lambda key: {y: c.annual(s[key], y) for y in range(1, H + 1)}  # noqa: E731
    rev, cogs, gp, opex_t, ebitda, ni = (ann(k) for k in ("rev", "cogs", "gp", "opex_total", "ebitda", "ni"))
    below = {y: sum(c.annual(s[k], y) for k in ("dep", "interest", "tax")) for y in range(1, H + 1)}

    def row(key, label, vals, fmt="usd", **kw) -> Row:
        return Row(key, label, tuple(Cell(vals[y], at(c.pl_cell(key, y))) for y in range(1, H + 1)), fmt, **kw)

    ratio = lambda num: {y: (num[y] / rev[y] if rev[y] else 0.0) for y in range(1, H + 1)}  # noqa: E731
    rows = [
        row("rev", "Revenue", rev, bold=True),
        row("cogs", c.pl.lines["cogs"].label, cogs),
        row("gp", "Gross profit", gp, bold=True, total=True),
        row("gm", "Gross margin", ratio(gp), "pct"),
    ]
    for key in tpl.opex_keys(c.t):
        vals = {y: c.annual(c.opex(key), y) for y in range(1, H + 1)}
        rows.append(row(key, c.pl.lines[key].label, vals))
    rows += [
        row("opex_total", "Total operating expenses", opex_t, bold=True, total=True),
        row("ebitda", "EBITDA", ebitda, bold=True, total=True),
        row("ebitda_m", "EBITDA margin", ratio(ebitda), "pct"),
        Row(
            "below_ebitda",
            "Depreciation, interest and tax",
            tuple(
                Cell(
                    below[y],
                    Ref(tuple((1.0, c.pl_cell(k, y)) for k in ("dep", "interest", "tax"))),
                )
                for y in range(1, H + 1)
            ),
        ),
        row("ni", "Net income", ni, bold=True, total=True),
    ]
    return Table("pnl", "Profit and loss by year", _year_cols(H), tuple(rows), row_header="US dollars")


def _t_costs(c: _Ctx) -> Table:
    H, s = c.H, c.series
    first, last = 1, H
    rev = {y: c.annual(s["rev"], y) for y in (first, last)}

    def share(vals: dict[int, float], terms: list[str], key: str, label: str, **kw) -> Row:
        cells = []
        for y in (first, last):
            v = -vals[y] / rev[y] if rev[y] else 0.0
            cells.append(Cell(v, Ref(tuple((-1.0, c.pl_cell(t, y)) for t in terms), ((1.0, c.pl_cell("rev", y)),))))
        return Row(key, label, tuple(cells), "pct", **kw)

    cogs = {y: c.annual(s["cogs"], y) for y in (first, last)}
    rows = [share(cogs, ["cogs"], "cogs", c.pl.lines["cogs"].label)]
    opex_keys = c.pl.template.opex_keys(c.t)
    for key in opex_keys:
        vals = {y: c.annual(c.opex(key), y) for y in (first, last)}
        rows.append(share(vals, [key], key, c.pl.lines[key].label))
    total = {y: cogs[y] + c.annual(s["opex_total"], y) for y in (first, last)}
    rows.append(
        share(total, ["cogs", "opex_total"], "total", "Total costs of sales and operations", bold=True, total=True)
    )
    margin = []
    for y in (first, last):
        v = c.annual(s["ebitda"], y) / rev[y] if rev[y] else 0.0
        margin.append(Cell(v, Ref(((1.0, c.pl_cell("ebitda", y)),), ((1.0, c.pl_cell("rev", y)),))))
    rows.append(Row("ebitda_m", "EBITDA margin", tuple(margin), "pct", bold=True))
    cols = ((f"y{first}", f"Year {first}"), (f"y{last}", f"Year {last}"))
    return Table("costs", "Cost structure, % of revenue", cols, tuple(rows), row_header="% of revenue")


def _resolve_kpis(c: _Ctx) -> Callable[[str], dict[int, float] | None]:
    """Annual value of a KPI row of the template (or None when the reference cannot compute it)."""
    tpl, H = c.pl.template, c.H
    rows = {r.key: r for r in tpl.kpi_rows(c.t)}
    cache: dict[str, dict[int, float] | None] = {}

    def series_for(src: str) -> list[float] | None:
        if src in c.series:
            return c.series[src]
        if c.rec.tops and src in c.rec.tops[0].opex:
            return c.opex(src)
        return None

    def resolve(key: str) -> dict[int, float] | None:
        if key in cache:
            return cache[key]
        cache[key] = None  # guards against a cycle
        row = rows.get(key)
        out: dict[int, float] | None = None
        if row is not None and row.kind == "line":
            vals, line = series_for(row.src), c.pl.lines.get(row.src)
            if vals is not None and line is not None and line.agg in ("sum", "last", "first"):
                sign = -1.0 if row.neg else 1.0
                out = {y: sign * c.annual(vals, y, line.agg) for y in range(1, H + 1)}
        elif row is not None and row.kind == "formula":
            try:
                out = _formula(row.src)
            except (ValueError, ZeroDivisionError, KeyError):
                out = None
        cache[key] = out
        return out

    def _formula(text: str) -> dict[int, float]:
        tree = _Parser(_tokens(text.lstrip("="))).parse()
        refs = {m for m in re.findall(r"\{([^{}]+)\}", text) if not m.startswith("in:") and m != "year"}
        if any(resolve(r) is None for r in refs):
            raise KeyError("unresolved")
        out = {}
        for y in range(1, H + 1):

            def ref(name: str, y=y) -> float:
                if name.startswith("in:"):
                    return float(c.d[name[3:]])
                if name == "year":
                    return float(y)
                return resolve(name)[y]  # type: ignore[index]

            out[y] = _evaluate(tree, ref)
        return out

    return resolve


_UNIT_FMT = {"usdk": "usd"}


def _t_unit(c: _Ctx) -> Table:
    H, pl = c.H, c.pl
    resolve = _resolve_kpis(c)
    kpi_block = pl.blocks.get("kpi")
    rows: list[Row] = []
    seen_kpi: set[str] = set()
    for m, drow in pl.dash.metrics:
        if m.src.startswith("kpi:"):
            vals = resolve(m.src[4:])
            seen_kpi.add(m.src[4:])
        else:
            line = c.series.get(m.src)
            vals = None if line is None else {y: c.annual(line, y, m.agg) for y in range(1, H + 1)}
        if vals is None:
            continue
        cells = tuple(Cell(vals[y], at(f"{DASH_SHEET}!{L(1 + y)}{drow}")) for y in range(1, H + 1))
        rows.append(Row(m.key, m.label, cells, _UNIT_FMT.get(m.fmt, m.fmt)))
    if kpi_block is not None:
        for r in pl.template.kpi_rows(c.t):
            if r.kind == "section" or r.key in seen_kpi or any(x.key == r.key for x in rows):
                continue
            vals = resolve(r.key)
            if vals is None:
                continue
            cells = tuple(Cell(vals[y], at(pl.stmt_cell("kpi", r.key, y))) for y in range(1, H + 1))
            rows.append(Row(r.key, r.label, cells, _UNIT_FMT.get(r.fmt, r.fmt), bold=r.bold))
    title = pl.template.kpi_title + " by year"
    return Table("unit", title, _year_cols(H), tuple(rows), row_header="Metric")


def _t_breakeven(c: _Ctx) -> Table:
    g, word, H = c.goal, c.t.period_word, c.H
    cells = c.man.goal_cells
    be_years = c.man.goal_years
    cap = word.capitalize()

    def one(key, label, fmt, **kw) -> Row:
        return Row(key, label, (Cell(g[key], at(cells[key])),), fmt, **kw)

    def period(key, label) -> Row:
        v = g[key]
        return Row(key, label, (Cell(v, at(cells[key])),), "period" if isinstance(v, int | float) else "text")

    rows = [
        period("first_ebitda", f"First {word} with positive EBITDA"),
        period("first_ni", f"First {word} with positive net income"),
        period("low_p", f"{cap} of the lowest cash balance"),
        one("min_cash", "Lowest cash balance in the plan", "usd"),
        one("required", "Funding required (uses + cash buffer)", "usd", bold=True),
        period("payback", "Cash invested paid back"),
    ]
    yrs = c.goal_years
    for y in (1, H):
        if "be_rev" in be_years and y in be_years["be_rev"]:
            v = yrs["be_rev"][y]
            fmt = "text" if isinstance(v, str) else "usd"
            rows.append(
                Row(f"be_rev_y{y}", f"Break-even revenue, year {y}", (Cell(v, at(be_years["be_rev"][y])),), fmt)
            )
    return Table("breakeven", "Break-even and cash", (("value", "Base case"),), tuple(rows))


TABLE_BUILDERS = {
    "drivers": _t_drivers,
    "startup": _t_startup,
    "pnl": _t_pnl,
    "costs": _t_costs,
    "unit": _t_unit,
    "breakeven": _t_breakeven,
}


def _chart_data(tables: dict[str, Table], c: _Ctx) -> dict[str, dict]:
    pnl, costs = tables["pnl"], tables["costs"]
    H = c.H
    keys = [k for k, _ in pnl.columns]
    rev = [pnl.row("rev").cells[i].value for i in range(H)]
    cost = [-(pnl.row("cogs").cells[i].value + pnl.row("opex_total").cells[i].value) for i in range(H)]
    cats = [r for r in costs.rows if r.key not in ("total", "ebitda_m")]
    return {
        "revenue_costs": {"years": keys, "revenue": rev, "costs": cost},
        "cost_structure": {
            "labels": [r.label for r in cats],
            "first": [r.cells[0].value for r in cats],
            "last": [r.cells[1].value for r in cats],
            "first_label": costs.columns[0][1],
            "last_label": costs.columns[1][1],
        },
    }


def sample_brief(key: str) -> Brief:
    """The brief of template ``key``'s sample: sample stage defaults, Full, monthly, 5 years, Base selected."""
    tpl = TEMPLATES[key]
    spec = sample_spec(key)
    brief = tpl.default_brief(tpl.sample_stage).with_options(
        tier=spec.tier, granularity=spec.granularity, horizon_years=spec.horizon_years, default_scenario=BASE_SCENARIO
    )
    return apply_goal(brief, tpl=tpl)


@cache
def compute(key: str) -> BizModel:
    """Tables and chart data of template ``key``'s sample (Base case, Full, monthly, 5 years)."""
    tpl = TEMPLATES[key]
    spec = sample_spec(key)
    brief = sample_brief(key)
    t = Timing(brief.granularity, brief.horizon_years)
    d = brief.values(BASE_SCENARIO)
    rec = _Recorder(tpl.reference_top)
    series = reference_run(d, brief.roles, t, rec, tpl.loan_style)
    pl = plan(brief, tpl)
    goal = goal_reference(brief, d, series, t)
    ctx = _Ctx(key, brief, pl, t, d, series, rec, goal["cells"], goal["years"], manifest_for(pl, verdict=""))
    annualise(series, t, tpl.kpi_flows)  # the reference's own roll-up must not fail on this brief
    tables = {name: TABLE_BUILDERS[name](ctx) for name in TABLE_NAMES}
    return BizModel(key, tpl.title, spec, t.period_word, tables, _chart_data(tables, ctx), t.years)


# --------------------------------------------------------------------------- HTML: tables
def _cell_html(r: Row, cell: Cell) -> str:
    return html.escape(format_value(cell.value, r.fmt))


def table_html(tbl: Table) -> str:
    head = "".join(f'<th scope="col" class="num">{html.escape(h)}</th>' for _, h in tbl.columns)
    if tbl.unit_column:
        head += '<th scope="col">Unit</th>'
    corner = html.escape(tbl.row_header) if tbl.row_header else '<span class="sr-only">Line</span>'
    body = []
    for r in tbl.rows:
        cls = " ".join(x for x in ("bm-total" if r.total else "", "bm-bold" if r.bold else "") if x)
        tds = "".join(
            f'<td class="num{" bm-text" if isinstance(c.value, str) else ""}">{_cell_html(r, c)}</td>' for c in r.cells
        )
        if tbl.unit_column:
            tds += f"<td>{html.escape(r.unit)}</td>"
        attr = f' class="{cls}"' if cls else ""
        body.append(f'<tr{attr}><th scope="row">{html.escape(r.label)}</th>{tds}</tr>')
    return (
        f'<figure class="bm-table" id="table-{tbl.name}">'
        # sighted readers get the caption above the scroll box (it wraps on a phone); the table's own caption
        # carries the same words for screen readers
        f'<figcaption aria-hidden="true"><strong>{html.escape(tbl.title)}</strong> {html.escape(CAPTION)}</figcaption>'
        f'<div class="table-wrap" role="region" aria-label="{html.escape(tbl.title)}" tabindex="0">'
        f'<table><caption class="sr-only">{html.escape(tbl.title)}. {html.escape(CAPTION)}</caption>'
        f'<thead><tr><th scope="col">{corner}</th>{head}</tr></thead>'
        f"<tbody>{''.join(body)}</tbody></table></div></figure>"
    )


# --------------------------------------------------------------------------- HTML: charts (inline SVG)
def _nice_max(v: float) -> tuple[float, float]:
    """(axis maximum, step) with 4 intervals and round numbers."""
    if v <= 0:
        return 1.0, 0.25
    raw = v / 4
    mag = 10 ** math.floor(math.log10(raw))
    step = next(m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw)
    return step * 4, step


def _compact(v: float, money: bool = True) -> str:
    a = abs(v)
    s = f"{a / 1e6:g}M" if a >= 1e6 else f"{a / 1e3:g}k" if a >= 1e3 else f"{a:g}"
    return ("$" if money else "") + s if v else ("$0" if money else "0")


def _short(text: str, n: int = 26) -> str:
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def _bar_path(x: float, y: float, w: float, h: float, r: float = 3.0, horizontal: bool = False) -> str:
    """Bar with a rounded data end and a flat base on the axis (vertical: top rounded; horizontal: right rounded)."""
    if h <= 0 or w <= 0:
        return ""
    if horizontal:
        r = min(r, h / 2, w)
        return f"M{x:.1f},{y:.1f}H{x + w - r:.1f}a{r},{r} 0 0 1 {r},{r}V{y + h - r:.1f}a{r},{r} 0 0 1 -{r},{r}H{x:.1f}Z"
    r = min(r, w / 2, h)
    return f"M{x:.1f},{y + h:.1f}V{y + r:.1f}a{r},{r} 0 0 1 {r},-{r}H{x + w - r:.1f}a{r},{r} 0 0 1 {r},{r}V{y + h:.1f}Z"


def chart_revenue_costs(data: dict, title: str) -> str:
    years, rev, cost = data["years"], data["revenue"], data["costs"]
    n = len(years)
    W, H = 520, 300
    left, right, top, bottom = 58, 12, 40, 34
    pw, ph = W - left - right, H - top - bottom
    vmax, step = _nice_max(max([*rev, *cost, 0]))
    sy = lambda v: top + ph - v / vmax * ph  # noqa: E731
    group = pw / n
    bw = min(34.0, group * 0.36)
    out = [
        f'<svg class="bm-svg" viewBox="0 0 {W} {H}" role="img" aria-labelledby="bm-rc-t bm-rc-d" focusable="false">',
        f'<title id="bm-rc-t">{html.escape(title)}</title>',
        '<desc id="bm-rc-d">'
        + html.escape(
            "Grouped bars by year. "
            + " ".join(
                f"Year {i + 1}: revenue {format_value(rev[i], 'usd')}, costs {format_value(cost[i], 'usd')}."
                for i in range(n)
            )
        )
        + "</desc>",
        '<g class="bm-legend">'
        '<rect class="bm-s1" x="58" y="10" width="12" height="12" rx="2"/><text x="76" y="20">Revenue</text>'
        '<rect class="bm-s2" x="146" y="10" width="12" height="12" rx="2"/>'
        '<text x="164" y="20">Cost of sales + operating expenses</text></g>',
    ]
    for i in range(5):
        v = step * i
        y = sy(v)
        cls = "bm-axis" if i == 0 else "bm-grid"
        out.append(f'<line class="{cls}" x1="{left}" x2="{W - right}" y1="{y:.1f}" y2="{y:.1f}"/>')
        out.append(f'<text class="bm-tick" x="{left - 8}" y="{y + 4:.1f}" text-anchor="end">{_compact(v)}</text>')
    for i in range(n):
        gx = left + group * i + group / 2
        for j, (vals, cls, name) in enumerate(((rev, "bm-s1", "Revenue"), (cost, "bm-s2", "Costs"))):
            x = gx - bw - 1 + j * (bw + 2)
            h = vals[i] / vmax * ph
            d = _bar_path(x, sy(vals[i]), bw, h)
            if d:
                out.append(
                    f'<path class="{cls}" d="{d}"><title>Year {i + 1} {name.lower()}: '
                    f"{html.escape(format_value(vals[i], 'usd'))}</title></path>"
                )
        out.append(f'<text class="bm-tick" x="{gx:.1f}" y="{H - 12}" text-anchor="middle">Year {i + 1}</text>')
    out.append("</svg>")
    return "".join(out)


def chart_cost_structure(data: dict, title: str) -> str:
    labels, first, last = data["labels"], data["first"], data["last"]
    fl, ll = data["first_label"], data["last_label"]
    n = len(labels)
    W = 520
    left, right, top = 150, 52, 40
    bar, gap, pad = 14, 2, 14
    row_h = 2 * bar + gap + pad
    H = top + n * row_h + 8
    vmax, _ = _nice_max(max([*first, *last, 0.01]))
    vmax = max(vmax, 0.01)
    pw = W - left - right
    sx = lambda v: v / vmax * pw  # noqa: E731
    out = [
        f'<svg class="bm-svg" viewBox="0 0 {W} {H:.0f}" role="img" aria-labelledby="bm-cs-t bm-cs-d" '
        'focusable="false">',
        f'<title id="bm-cs-t">{html.escape(title)}</title>',
        '<desc id="bm-cs-d">'
        + html.escape(
            "Horizontal bars, share of revenue. "
            + " ".join(
                f"{labels[i]}: {fl} {format_value(first[i], 'pct')}, {ll} {format_value(last[i], 'pct')}."
                for i in range(n)
            )
        )
        + "</desc>",
        '<g class="bm-legend">'
        f'<rect class="bm-y1" x="{left}" y="10" width="12" height="12" rx="2"/>'
        f'<text x="{left + 18}" y="20">{html.escape(fl)}</text>'
        f'<rect class="bm-y5" x="{left + 90}" y="10" width="12" height="12" rx="2"/>'
        f'<text x="{left + 108}" y="20">{html.escape(ll)}</text></g>',
        f'<line class="bm-axis" x1="{left}" x2="{left}" y1="{top - 4}" y2="{H - 4:.0f}"/>',
    ]
    for i in range(n):
        y0 = top + i * row_h
        out.append(
            f'<text class="bm-cat" x="{left - 8}" y="{y0 + bar + 4:.1f}" text-anchor="end">'
            f"<title>{html.escape(labels[i])}</title>{html.escape(_short(labels[i]))}</text>"
        )
        for j, (vals, cls, name) in enumerate(((first, "bm-y1", fl), (last, "bm-y5", ll))):
            y = y0 + j * (bar + gap)
            w = sx(max(vals[i], 0.0))
            d = _bar_path(left, y, w, bar, horizontal=True)
            share = html.escape(format_value(vals[i], "pct"))
            if d:
                tip = f"{html.escape(labels[i])}, {html.escape(name)}: {share}"
                out.append(f'<path class="{cls}" d="{d}"><title>{tip}</title></path>')
            out.append(f'<text class="bm-val" x="{left + w + 6:.1f}" y="{y + bar - 3:.1f}">{share}</text>')
    out.append("</svg>")
    return "".join(out)


CHARTS = {
    "revenue_costs": (
        "Revenue against costs by year",
        "Bars compare revenue with cost of sales plus operating expenses; the gap is EBITDA.",
        chart_revenue_costs,
    ),
    "cost_structure": (
        "Where each dollar of revenue goes",
        "Cost lines as a share of revenue in the first and last year of the plan.",
        chart_cost_structure,
    ),
}


def chart_html(name: str, model: BizModel) -> str:
    title, note, fn = CHARTS[name]
    return (
        f'<figure class="bm-chart" id="chart-{name}"><div class="bm-chart__scroll" role="region" '
        f'aria-label="{html.escape(title)}" tabindex="0">{fn(model.charts[name], title)}</div>'
        f"<figcaption><strong>{html.escape(title)}.</strong> {html.escape(note)} "
        f"{html.escape(CAPTION)}.</figcaption></figure>"
    )


# --------------------------------------------------------------------------- copy: placeholders
def num_text(model: BizModel, spec: str) -> str:
    m = NUM_RE.match(spec)
    if not m:
        raise PlaceholderError(f"bad num placeholder {spec!r}; use <table>.<row>[.<column>][|abs]")
    table, row, col, absflag = m.groups()
    if table not in model.tables:
        raise PlaceholderError(f"num placeholder {spec!r}: unknown table {table!r}")
    try:
        r, cell = model.tables[table].cell(row, col)
    except KeyError as exc:
        raise PlaceholderError(f"num placeholder {spec!r}: unknown row or column {exc}") from None
    v = cell.value
    if absflag and not isinstance(v, str):
        v = abs(v)
    return format_value(v, r.fmt, prose=True)


def placeholders(md: str) -> list[tuple[str, str]]:
    return [(m.group(1), m.group(2)) for m in PLACEHOLDER_RE.finditer(md)]


def check_placeholders(md: str, model: BizModel) -> None:
    """Raise PlaceholderError for any unknown placeholder (also used by the copy test)."""
    for kind, arg in placeholders(md):
        if kind == "table" and arg not in TABLE_NAMES:
            raise PlaceholderError(f"unknown table {arg!r}; known: {TABLE_NAMES}")
        elif kind == "chart" and arg not in CHART_NAMES:
            raise PlaceholderError(f"unknown chart {arg!r}; known: {CHART_NAMES}")
        elif kind == "num":
            num_text(model, arg)
        elif kind not in ("table", "chart"):
            raise PlaceholderError(f"unknown placeholder kind {kind!r} in {{{{{kind}:{arg}}}}}")


def expand(md: str, model: BizModel) -> tuple[str, dict[str, str]]:
    """(markdown with num placeholders replaced and block placeholders turned into tokens, token -> HTML)."""
    check_placeholders(md, model)
    blocks: dict[str, str] = {}

    def sub(m: re.Match) -> str:
        kind, arg = m.group(1), m.group(2)
        if kind == "num":
            return num_text(model, arg)
        token = f"BMBLOCK{len(blocks)}X"
        blocks[token] = table_html(model.tables[arg]) if kind == "table" else chart_html(arg, model)
        return token

    return PLACEHOLDER_RE.sub(sub, md), blocks


# --------------------------------------------------------------------------- pages
@dataclass(frozen=True)
class BizPage:
    slug: str
    key: str
    title: str
    description: str
    keyword: str
    h1: str
    body: str  # rendered HTML with tables and charts in place
    words: int
    sample: SampleSpec
    published: str = ""  # ISO date from the front matter
    modified: str = ""
    meta: dict = field(default_factory=dict)

    @property
    def name(self) -> str:
        return self.key

    @property
    def template(self) -> str:
        return self.key


#: a dollar amount or a percentage typed into the copy (``$4,500``, ``33%``, ``33 percent``)
FIGURE_RE = re.compile(r"\$\s?\d|\d\s?%|\d\s?(?:percent|per cent)\b", re.I)
EXTERNAL_LINK_RE = re.compile(r"\]\(https?://")


def unsourced_figures(body_md: str) -> list[str]:
    """Sentences of the copy that state a dollar or percent figure without an external source link.

    ``{{num:...}}`` and the table and chart placeholders are data, not typed figures, so they are ignored.
    """
    plain = PLACEHOLDER_RE.sub("", body_md)
    sentences = re.split(r"(?<=[.!?])\s+|\n+", plain)
    return [s.strip() for s in sentences if FIGURE_RE.search(s) and not EXTERNAL_LINK_RE.search(s)]


def strip_blocks(md: str) -> str:
    return PLACEHOLDER_RE.sub(lambda m: "" if m.group(1) in ("table", "chart") else "0", md)


def load_page(path: Path) -> BizPage:
    meta, body_md = seo.parse_front_matter(path.read_text(encoding="utf-8"))
    for k in ("slug", "title", "description", "keyword", "template"):
        if not meta.get(k):
            raise ValueError(f"{path.name}: front matter needs {k!r}")
    key = path.stem
    if key not in TEMPLATES:
        raise ValueError(f"{path.name}: the file name must be a registered template key; known: {sorted(TEMPLATES)}")
    if meta["slug"] != f"/business-models/{key}" or not SLUG_RE.match(meta["slug"]):
        raise ValueError(f"{path.name}: slug must be /business-models/{key}")
    if seo.page_template(meta) != key:
        raise ValueError(f"{path.name}: template must be {key!r}")
    if seo.page_sample(meta) != sample_spec(key):
        raise ValueError(f"{path.name}: sample must be '{key}, full, monthly, {HORIZON_YEARS} years'")
    published = meta.get("published", "")
    modified = meta.get("modified", published)
    for label, value in (("published", published), ("modified", modified)):
        if not DATE_RE.match(value):
            raise ValueError(f"{path.name}: front matter needs {label!r} as YYYY-MM-DD")
    if modified < published:
        raise ValueError(f"{path.name}: modified is earlier than published")
    model = compute(key)
    md, blocks = expand(body_md, model)
    rendered, h1 = seo.render_markdown(md)
    if not h1:
        raise ValueError(f"{path.name}: the body must start with a # heading")
    for token, block in blocks.items():
        para = f"<p>{token}</p>"
        if para not in rendered:
            raise PlaceholderError(f"{path.name}: table and chart placeholders must stand alone on their own line")
        rendered = rendered.replace(para, block)
    if "{{" in rendered or "BMBLOCK" in rendered:
        raise PlaceholderError(f"{path.name}: unresolved placeholder left in the page")
    words = len(re.findall(r"\w+", strip_blocks(body_md)))
    return BizPage(
        slug=meta["slug"],
        key=key,
        title=meta["title"],
        description=meta["description"],
        keyword=meta["keyword"],
        h1=h1,
        body=rendered,
        words=words,
        sample=sample_spec(key),
        published=published,
        modified=modified,
        meta=meta,
    )


@lru_cache(maxsize=1)
def pages() -> dict[str, BizPage]:
    """Every page in content/business-models by slug. A broken page fails loudly at startup / in tests."""
    out: dict[str, BizPage] = {}
    for path in sorted(CONTENT.glob("*.md")):
        page = load_page(path)
        if page.slug in out:
            raise ValueError(f"duplicate slug {page.slug}")
        out[page.slug] = page
    return out


def page_for(template: str) -> BizPage | None:
    return pages().get(f"/business-models/{template}")


def sitemap_paths() -> list[str]:
    return [f"/{SECTION}", *pages()]


def calculators_for(template: str) -> list[dict]:
    """Calculator pages of the template (``content/calculators``, loaded by ``app.calculators``); [] if none."""
    return [{"slug": c.slug, "title": c.h1} for c in calculators.calculators().values() if c.template == template]


def seo_pages_for(template: str) -> list[seo.SeoPage]:
    return [p for p in seo.pages().values() if p.template == template and p.section == "models"]


# --------------------------------------------------------------------------- routes
router = APIRouter()


def structured_data(page: BizPage, base: str) -> list[dict]:
    url = base + page.slug
    crumbs = {
        "@context": "https://schema.org",
        "@type": "BreadcrumbList",
        "itemListElement": [
            {"@type": "ListItem", "position": 1, "name": "Home", "item": base + "/"},
            {"@type": "ListItem", "position": 2, "name": SECTION_TITLE, "item": f"{base}/{SECTION}"},
            {"@type": "ListItem", "position": 3, "name": page.h1, "item": url},
        ],
    }
    publisher = {"@type": "Organization", "name": "StartupModel", "url": base + "/"}
    article = {
        "@context": "https://schema.org",
        "@type": "Article",
        "headline": page.h1,
        "description": page.description,
        "mainEntityOfPage": url,
        "datePublished": page.published,
        "dateModified": page.modified,
        "author": publisher,
        "publisher": publisher,
    }
    return [crumbs, article]


def _price(cents: int) -> str:
    return f"${cents / 100:,.0f}"


def _page_ctx(request: Request) -> dict:
    cfg = request.app.state.services.cfg
    return {
        "base_url": cfg.public_base_url,
        "support_email": cfg.support_email,
        "price_light": _price(cfg.prices["light"]),
        "price_full": _price(cfg.prices["full"]),
    }


@router.get("/business-models", response_class=HTMLResponse)
def index(request: Request) -> HTMLResponse:
    cards = []
    for page in pages().values():
        model = compute(page.key)
        cards.append(
            {
                "page": page,
                "template_title": TEMPLATES[page.key].title,
                "facts": [
                    ("Funding required", num_text(model, "startup.required")),
                    (f"First {model.word} with positive EBITDA", num_text(model, "breakeven.first_ebitda")),
                ],
            }
        )
    return templates.TemplateResponse(
        request,
        "bizmodel_index.html",
        {"title": SECTION_TITLE, "intro": SECTION_INTRO, "cards": cards, "caption": CAPTION, **_page_ctx(request)},
    )


@router.get("/business-models/{name}", response_class=HTMLResponse)
def page(name: str, request: Request) -> HTMLResponse:
    bp = pages().get(f"/business-models/{name}")
    if bp is None:
        raise HTTPException(404, "not found")
    ctx = _page_ctx(request)
    sample = bp.sample if (ROOT / "samples" / bp.sample.filename).is_file() else None
    tpl = TEMPLATES[bp.key]
    return templates.TemplateResponse(
        request,
        "bizmodel_page.html",
        {
            "page": bp,
            "body": Markup(bp.body),
            "section_title": SECTION_TITLE,
            "template_title": tpl.title,
            "start_href": f"/start?type={tpl.business_types[0]}",
            "sample": sample,
            "calculators": calculators_for(bp.key),
            "model_pages": seo_pages_for(bp.key),
            "ld_json": [seo._json_script(x) for x in structured_data(bp, ctx["base_url"])],
            **ctx,
        },
    )
