"""Independent reviewer for a generated model.

1. Recalculates the workbook once per scenario (1/2/3) in LibreOffice.
2. Reads the results and recomputes the same P&L in plain Python from the
   brief, so a wrong cell reference that still "evaluates" gets caught.
3. Checks Low < Base < High revenue and that the in-file checks pass.

Usage: python challenger.py brief.json model.xlsx
"""
import json
import math
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

RECALC = "/mnt/skills/public/xlsx/scripts/recalc.py"


def python_model(brief, scen):
    k = {"1": "low", "2": "base", "3": "high"}[str(scen)]
    d = {x["key"]: x[k] for x in brief["drivers"]}
    years = brief["horizon_years"]
    out = {y: {"rev": 0.0, "ebitda": 0.0, "ni": 0.0} for y in range(1, years + 1)}
    cash = d["start_cash"] - d["capex"]
    year_end_cash = {}
    for m in range(1, years * 12 + 1):
        y = math.ceil(m / 12)
        ramp = min(1, d["open_pct"] + (1 - d["open_pct"]) * (m - 1) / max(1, d["ramp_months"]))
        cust = d["cust_full"] * ramp * (1 + d["traffic_g"]) ** (y - 1)
        ticket = d["ticket"] * (1 + d["price_g"]) ** (y - 1)
        rev = cust * ticket * d["days"]
        gp = rev * (1 - d["cogs_pct"])
        payroll = d["fte"] * d["salary"] / 12 * (1 + d["payroll_tax"]) * (1 + d["wage_g"]) ** (y - 1)
        opex = d["rent"] + payroll + rev * d["mkt_pct"] + d["other_opex"]
        ebitda = gp - opex
        dep = d["capex"] / (d["dep_years"] * 12) if m <= d["dep_years"] * 12 else 0
        ebit = ebitda - dep
        ni = ebit - max(0, ebit) * d["tax_rate"]
        cash += ni + dep
        out[y]["rev"] += rev
        out[y]["ebitda"] += ebitda
        out[y]["ni"] += ni
        if m % 12 == 0:
            year_end_cash[y] = cash
    for y in out:
        out[y]["cash"] = year_end_cash[y]
    return out


def run_scenario(src, scen, workdir):
    path = Path(workdir) / f"s{scen}.xlsx"
    shutil.copy(src, path)
    wb = load_workbook(path)
    wb["Dashboard"]["B4"] = scen
    wb.save(path)
    res = json.loads(subprocess.run([sys.executable, RECALC, str(path), "60"],
                                    capture_output=True, text=True).stdout)
    vals = load_workbook(path, data_only=True)
    return res, vals


def main(brief_path, model_path):
    brief = json.load(open(brief_path, encoding="utf-8"))
    years = brief["horizon_years"]
    report = {"scenarios": {}, "problems": []}
    revenue_by_scen = {}
    with tempfile.TemporaryDirectory() as tmp:
        for scen in (1, 2, 3):
            res, wb = run_scenario(model_path, scen, tmp)
            pl, ck, dash = wb["P&L"], wb["Checks"], wb["Dashboard"]
            if res.get("status") != "success":
                report["problems"].append(f"scenario {scen}: recalc {res}")
            expected = python_model(brief, scen)
            rows = {"rev": 5, "ebitda": 14, "ni": 19, "cash": 21}
            scen_out = {"label": dash["C4"].value, "verdict": ck["B13"].value,
                        "first_positive_ebitda_month": dash["B16"].value,
                        "lowest_cash": round(dash["B17"].value or 0), "years": {}}
            for y in range(1, years + 1):
                col = chr(ord("A") + y)
                got = {k: pl[f"{col}{r}"].value for k, r in rows.items()}
                for k in rows:
                    if abs(got[k] - expected[y][k]) > 1:
                        report["problems"].append(
                            f"scenario {scen} year {y} {k}: workbook {got[k]:.0f} vs python {expected[y][k]:.0f}")
                scen_out["years"][y] = {k: round(v) for k, v in got.items()}
            if scen_out["verdict"] == "FAIL":
                report["problems"].append(f"scenario {scen}: in-file checks FAIL")
            revenue_by_scen[scen] = [scen_out["years"][y]["rev"] for y in range(1, years + 1)]
            report["scenarios"][scen] = scen_out
    for y in range(years):
        if not (revenue_by_scen[1][y] < revenue_by_scen[2][y] < revenue_by_scen[3][y]):
            report["problems"].append(f"year {y+1}: revenue not Low < Base < High")
    report["verdict"] = "PASS" if not report["problems"] else "FAIL"
    print(json.dumps(report, indent=2, ensure_ascii=False))
    return report


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
