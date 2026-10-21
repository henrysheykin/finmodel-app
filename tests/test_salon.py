"""Hair salon / barbershop template: economics in the plain-Python twin, then the workbook through LibreOffice."""

import itertools

import pytest
from openpyxl import load_workbook

from engine.challenger import challenge, python_model
from engine.core import reference_run
from engine.recalc import recalc
from engine.spec import Timing
from engine.templates import get_template

SALON = get_template("salon")
STAGES = ["planning", "early", "growing", "established"]
COMBOS = list(itertools.product(["light", "full"], ["monthly", "annual"], [1, 3, 5]))


@pytest.fixture
def salon_brief():
    return SALON.default_brief("early")


def run(brief, scen=2, g="monthly", years=3):
    return reference_run(brief.values(scen), brief.roles, Timing(g, years), SALON.reference_top)


def with_values(brief, **vals):
    d = brief.model_dump()
    for drv in d["drivers"]:
        if drv["key"] in vals:
            v = vals[drv["key"]]
            drv.update(low=v, base=v, high=v)
    return brief.model_validate(d)


# ---------------------------------------------------------------------- plain-Python twin
@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_are_valid_ordered_and_funded(stage):
    b = SALON.default_brief(stage)
    b = b.with_options(horizon_years=5)
    low, base, high = (python_model(b, s) for s in (1, 2, 3))
    for y in range(1, 6):
        assert low["rev"][y] < base["rev"][y] < high["rev"][y]
        assert 0 < base["gp"][y] / base["rev"][y] < 1
    assert min(base["cash_close"].values()) > 0  # the Base case is funded; Low is the stress case


@pytest.mark.parametrize("stage", STAGES)
def test_monthly_and_annual_agree_exactly(stage):
    """Nothing in the salon model is seasonal or compounds monthly, so annual = sum of months."""
    b = SALON.default_brief(stage)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    for k in ("rev", "gp", "ebitda", "services", "rev_chair", "hours_booked", "chairs_total", "capex"):
        for y in (1, 2, 3):
            assert m[k][y] == pytest.approx(a[k][y], abs=0.01), (k, y)


def test_services_are_chairs_times_utilisation_hours_and_rate(salon_brief):
    d = salon_brief.values(2)
    s = run(salon_brief)
    employed = d["chairs"] * (1 - d["rented_share"])
    hours = employed * d["days_open"] * d["hrs_day"]
    assert s["hours_avail"][0] == pytest.approx(hours)
    assert s["hours_booked"][0] == pytest.approx(hours * d["util"])
    assert s["services"][0] == pytest.approx(hours * d["util"] * d["svc_hr"])
    assert s["rev_svc"][0] == pytest.approx(s["services"][0] * d["price"])


def test_retail_is_attach_rate_times_product_ticket(salon_brief):
    d = salon_brief.values(2)
    s = run(salon_brief)
    assert s["rev_retail"][0] == pytest.approx(s["services"][0] * d["attach"] * d["product_ticket"])
    none = run(with_values(salon_brief, attach=0))
    assert max(none["rev_retail"]) == 0
    assert none["rev"][0] == pytest.approx(s["rev"][0] - s["rev_retail"][0])


def test_utilisation_ramps_each_year_up_to_the_ceiling(salon_brief):
    d = salon_brief.values(2)
    s = run(salon_brief, years=3)
    util = [s["hours_booked"][i] / s["hours_avail"][i] for i in (0, 12, 24)]
    assert util[0] == pytest.approx(d["util"])
    assert util[1] == pytest.approx(d["util"] + d["util_gain"])
    capped = run(with_values(salon_brief, util_max=d["util"] + d["util_gain"] / 2), years=3)
    assert capped["hours_booked"][24] / capped["hours_avail"][24] == pytest.approx(d["util"] + d["util_gain"] / 2)


def test_chair_rent_scales_with_rented_chairs_and_occupancy(salon_brief):
    d = salon_brief.values(2)
    s = run(salon_brief)
    rented = d["chairs"] * d["rented_share"]
    assert s["chairs_rented"][0] == pytest.approx(rented)
    assert s["rev_chair"][0] == pytest.approx(rented * d["chair_rent_wk"] * 52 / 12 * d["rent_occ"])
    assert s["rev_chair"][0] > 0
    none = run(with_values(salon_brief, rented_share=0))
    assert max(none["rev_chair"]) == 0
    assert max(none["chairs_rented"]) == 0
    # every chair now earns service revenue instead of rent
    assert none["services"][0] > s["services"][0]


def test_stylist_pay_models_trade_commission_against_rent(salon_brief):
    """Renting a chair out moves it from commission-paid service revenue to rent income."""
    d = salon_brief.values(2)
    all_emp = run(with_values(salon_brief, rented_share=0, attach=0, backbar=0, card_pct=0))
    expected = -all_emp["rev_svc"][0] * d["comm_pct"] * (1 + d["payroll_tax"])
    assert all_emp["cogs"][0] == pytest.approx(expected)


def test_cost_of_sales_lines(salon_brief):
    d = salon_brief.values(2)
    s = run(salon_brief)
    cogs = (
        -s["rev_svc"][0] * d["comm_pct"] * (1 + d["payroll_tax"])
        - s["rev_retail"][0] * d["prod_cost_pct"]
        - s["services"][0] * d["backbar"]
        - (s["rev_svc"][0] + s["rev_retail"][0]) * d["card_pct"]
    )
    assert s["cogs"][0] == pytest.approx(cogs)


def test_added_chairs_raise_capacity_and_buy_stations_from_year_two():
    b = SALON.default_brief("growing")
    d = b.values(2)
    s = run(b, years=3)
    assert s["chairs_total"][0] == d["chairs"]
    assert s["chairs_total"][12] == d["chairs"] + d["chairs_add"]
    assert s["chairs_total"][24] == d["chairs"] + 2 * d["chairs_add"]
    assert s["capex"][12] == pytest.approx(d["chairs_add"] * d["station_cost"])
    assert s["capex"][13] == 0
    assert s["capex"][0] == pytest.approx(d["capex_fitout"] + d["wash_new"] * d["wash_cost"])


def test_opening_capex_buys_only_the_unequipped_chairs(salon_brief):
    d = salon_brief.values(2)
    new = run(with_values(salon_brief, chairs_equipped=0, wash_new=2))
    assert new["capex"][0] == pytest.approx(d["capex_fitout"] + d["chairs"] * d["station_cost"] + 2 * d["wash_cost"])


@pytest.mark.parametrize("g", ["monthly", "annual"])
@pytest.mark.parametrize("scen", [1, 2, 3])
def test_balance_sheet_balances(salon_brief, g, scen):
    s = run(salon_brief, scen, g)
    for ta, tle in zip(s["total_assets"], s["total_le"], strict=True):
        assert ta == pytest.approx(tle, abs=1e-6)


# ---------------------------------------------------------------------- workbook (LibreOffice + challenger)
@pytest.mark.libreoffice
@pytest.mark.parametrize("tier,g,years", COMBOS, ids=[f"{t}-{g}-{y}y" for t, g, y in COMBOS])
def test_matrix(build, tier, g, years):
    brief = SALON.default_brief("growing").with_options(tier=tier, granularity=g, horizon_years=years)
    path = build(brief)
    res = recalc(path)
    assert res.ok, res.errors
    vals = load_workbook(path, data_only=True)
    verdict = [r[1].value for r in vals["Checks"].iter_rows(min_col=1, max_col=2) if r[0].value == "Overall verdict"]
    assert verdict and verdict[0] in ("PASS", "PASS WITH WARNINGS")
    report = challenge(brief, path)
    assert report["problems"] == []
    assert report["verdict"] == "PASS"


@pytest.mark.libreoffice
@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_pass_the_challenger(build, stage):
    brief = SALON.default_brief(stage)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
    assert report["verdict"] == "PASS"
    assert any(k.startswith("kpi:") for k in report["scenarios"][2]["years"][1])


@pytest.mark.libreoffice
def test_monthly_workbook_equals_annual_workbook(build):
    brief = SALON.default_brief("early")
    monthly = challenge(brief, build(brief, "m.xlsx"))
    annual_brief = brief.with_options(granularity="annual")
    annual = challenge(annual_brief, build(annual_brief, "a.xlsx"))
    for s in (1, 2, 3):
        for y in (1, 2, 3):
            m, a = monthly["scenarios"][s]["years"][y], annual["scenarios"][s]["years"][y]
            for k in ("rev", "gp", "ebitda", "kpi:services", "kpi:rev_chair"):
                assert m[k] == pytest.approx(a[k], abs=1), (s, y, k)


@pytest.mark.libreoffice
def test_revenue_rises_from_low_to_high_in_the_workbook(build, salon_brief):
    report = challenge(salon_brief, build(salon_brief))
    for y in (1, 2, 3):
        revs = [report["scenarios"][s]["years"][y]["rev"] for s in (1, 2, 3)]
        assert revs[0] < revs[1] < revs[2]
