"""Tier modules (E07, ADR-013): Break-even (both tiers) and Valuation (Full)."""

import pytest
from openpyxl import load_workbook

from app.briefing import Answers, PrefillBriefSource
from engine.brief import BriefError
from engine.builder import plan
from engine.challenger import challenge
from engine.core import reference_run
from engine.goals import (
    ALL_GOAL_KEYS,
    NOT_WITHIN,
    NOTHING_TO_PAY_BACK,
    apply_goal,
    goal_specs,
    modules_for,
    reference,
)
from engine.recalc import recalc
from engine.spec import Timing
from engine.templates import TEMPLATES, get_template

GOALS = ["startup_breakeven", "dcf_valuation", "investor_pitch"]
BE_KEYS = {"cash_buffer"}


# --------------------------------------------------------------------------- drivers & validation
def test_modules_per_tier():
    # ADR-013: the tier decides the modules; the goal only sets the Cover's purpose line
    assert modules_for("light") == ("breakeven",)
    assert modules_for("full") == ("breakeven", "dcf")
    assert modules_for("investor_pitch") == ()


@pytest.mark.parametrize("tier", ["light", "full"])
@pytest.mark.parametrize("goal", GOALS + ["revenue_profit"])
def test_apply_goal_adds_exactly_the_tier_drivers(cafe_brief, goal, tier):
    b = apply_goal(cafe_brief.with_options(tier=tier), goal)
    keys = {d.key for d in b.drivers}
    assert b.goal == goal
    assert keys & ALL_GOAL_KEYS == {s.key for s in goal_specs(tier)}
    get_template(b.template).validate(b)
    other = "full" if tier == "light" else "light"
    switched = apply_goal(b.with_options(tier=other))
    assert {d.key for d in switched.drivers} & ALL_GOAL_KEYS == {s.key for s in goal_specs(other)}


def test_apply_goal_keeps_existing_values_and_buffer_is_three_months(cafe_brief):
    b = apply_goal(cafe_brief.with_options(tier="full"), "investor_pitch")
    assert b.driver("cash_buffer").base > 0 and b.driver("cash_buffer").base % 1000 == 0
    d = b.model_dump()
    for drv in d["drivers"]:
        if drv["key"] == "wacc":
            drv.update(low=0.4, base=0.3, high=0.2)
    again = apply_goal(type(b).model_validate(d), "investor_pitch")
    assert again.driver("wacc").base == 0.3


def test_goal_drivers_are_completed_and_bounded(cafe_brief):
    tpl = get_template("cafe")
    bare = cafe_brief.model_dump()
    bare["drivers"] = [d for d in bare["drivers"] if d["key"] not in ALL_GOAL_KEYS]
    tpl.validate(type(cafe_brief).model_validate(bare))  # missing module drivers are added from defaults
    b = apply_goal(cafe_brief.with_options(tier="full"), "dcf_valuation").model_dump()
    for drv in b["drivers"]:
        if drv["key"] == "tg":
            drv.update(low=0.01, base=0.02, high=0.06)
        if drv["key"] == "wacc":
            drv.update(low=0.2, base=0.1, high=0.06)
    with pytest.raises(BriefError, match="WACC must exceed terminal growth"):
        tpl.validate(type(cafe_brief).model_validate(b))
    for drv in b["drivers"]:
        if drv["key"] == "wacc":
            drv.update(low=0.08, base=0.1, high=0.12)  # a higher WACC is the pessimistic case
    with pytest.raises(BriefError, match="wacc: needs Low >= Base >= High"):
        tpl.validate(type(cafe_brief).model_validate(b))


@pytest.mark.parametrize("tier", ["light", "full"])
@pytest.mark.parametrize("goal", GOALS + ["revenue_profit"])
def test_wizard_prefill_builds_valid_briefs_for_every_goal(goal, tier):
    for bt in ("saas", "restaurant_cafe"):
        b = PrefillBriefSource().build_brief(Answers(business_type=bt, goal=goal, stage="early", tier=tier))
        assert b.goal == goal
        assert {d.key for d in b.drivers} & ALL_GOAL_KEYS == {s.key for s in goal_specs(tier)}


@pytest.mark.parametrize("goal", GOALS + ["revenue_profit"])
@pytest.mark.parametrize(
    "tier,sheets",
    [
        ("light", ["Cover", "Dashboard", "Break-even", "Inputs", "P&L", "Checks"]),
        (
            "full",
            [
                "Cover",
                "Dashboard",
                "Summary",
                "Break-even",
                "Valuation",
                "Inputs",
                "P&L",
                "Cash Flow",
                "Balance Sheet",
                "Monthly",
                "Headcount",
                "Checks",
            ],
        ),
    ],
)
def test_sheet_order(cafe_brief, goal, tier, sheets):
    b = apply_goal(cafe_brief.with_options(tier=tier), goal)
    assert plan(b, get_template("cafe")).sheets == sheets


# --------------------------------------------------------------------------- python twin
def _series(fcf_like: list[float], **extra):
    n = len(fcf_like)
    base = dict.fromkeys(
        ("rev", "gp", "opex_total", "ebitda", "ni", "capex", "cfi", "interest", "loan_repay", "cash_close"), None
    )
    s = {k: [0.0] * n for k in base}
    s["cfo"] = list(fcf_like)
    s.update(extra)
    return s


def test_dcf_twin_matches_closed_form(cafe_brief):
    b = apply_goal(cafe_brief.with_options(tier="full", horizon_years=3, granularity="annual"), "dcf_valuation")
    d = {"wacc": 0.10, "tg": 0.02, "exit_mult": 5.0, "tax_rate": 0.25, "cash_buffer": 0, "equity": 0, "loan_amount": 0}
    s = _series([100.0, 100.0, 100.0], ebitda=[200.0, 200.0, 200.0])
    v = reference(b, d, s, Timing("annual", 3))["cells"]
    annuity = 100 / 1.1 + 100 / 1.1**2 + 100 / 1.1**3
    assert v["sum_pv"] == pytest.approx(annuity)
    assert v["tv_g"] == pytest.approx(100 * 1.02 / 0.08)
    assert v["ev_g"] == pytest.approx(annuity + 100 * 1.02 / 0.08 / 1.1**3)
    assert v["ev_m"] == pytest.approx(annuity + 1000 / 1.1**3)
    assert v["implied_g"] == pytest.approx((1000 * 0.1 - 100) / 1100)


def test_interest_tax_shield_in_free_cash_flow(cafe_brief):
    b = apply_goal(cafe_brief.with_options(tier="full", horizon_years=1, granularity="annual"), "dcf_valuation")
    d = {"wacc": 0.10, "tg": 0.02, "exit_mult": 5.0, "tax_rate": 0.25, "cash_buffer": 0, "equity": 0, "loan_amount": 0}
    s = _series([100.0], interest=[-40.0], cfi=[-10.0])
    assert reference(b, d, s, Timing("annual", 1))["years"]["fcff"][1] == pytest.approx(100 + 30 - 10)


def test_breakeven_twin_on_the_cafe(cafe_brief):
    b = apply_goal(cafe_brief, "startup_breakeven")
    d = b.values(2)
    t = Timing("monthly", 3)
    s = reference_run(d, b.roles, t, get_template("cafe").reference_top)
    v = reference(b, d, s, t)["cells"]
    assert v["min_cash"] == pytest.approx(v["sources"] - v["uses"])
    assert v["gap"] == pytest.approx(v["sources"] - v["uses"] - d["cash_buffer"])
    assert 1 <= v["first_ebitda"] <= v["first_ni"] <= 36
    assert v["payback"] == NOT_WITHIN or isinstance(v["payback"], int)


def test_payback_texts(cafe_brief):
    b = apply_goal(cafe_brief.with_options(tier="light", horizon_years=1, granularity="annual"), "startup_breakeven")
    d = {"cash_buffer": 0, "equity": 0, "loan_amount": 0}
    one = Timing("annual", 1)
    s = _series([5.0], cash_close=[5.0], rev=[10.0], gp=[5.0])
    assert reference(b, d, s, one)["cells"]["payback"] == NOTHING_TO_PAY_BACK
    s = _series([-5.0], cash_close=[-5.0], rev=[10.0], gp=[5.0])
    assert reference(b, d, s, one)["cells"]["payback"] == NOT_WITHIN


# --------------------------------------------------------------------------- LibreOffice
def _verdict(path):
    vals = load_workbook(path, data_only=True)
    return next(r[1].value for r in vals["Checks"].iter_rows(min_col=1, max_col=2) if r[0].value == "Overall verdict")


def _passes(build, brief):
    path = build(brief)
    res = recalc(path)
    assert res.ok, res.errors
    assert _verdict(path) in ("PASS", "PASS WITH WARNINGS")
    report = challenge(brief, path)
    assert report["problems"] == []
    return report


@pytest.mark.libreoffice
@pytest.mark.parametrize("template", sorted(TEMPLATES))
def test_full_modules_in_every_template(build, template):
    tpl = get_template(template)
    brief = apply_goal(tpl.default_brief(tpl.sample_stage).with_options(tier="full"), "investor_pitch")
    report = _passes(build, brief)
    goal = report["scenarios"][2]["goal"]
    assert {"gap", "ev_g", "ev_m", "ev_avg", "sens_centre"} <= set(goal)


CAFE_COMBOS = [
    ("investor_pitch", "light", "monthly", 1),
    ("revenue_profit", "light", "annual", 5),
    ("investor_pitch", "full", "annual", 1),
    ("revenue_profit", "full", "monthly", 5),
    ("startup_breakeven", "light", "monthly", 3),
    ("dcf_valuation", "full", "annual", 3),
]


@pytest.mark.libreoffice
@pytest.mark.parametrize("goal,tier,g,years", CAFE_COMBOS, ids=["-".join(map(str, c)) for c in CAFE_COMBOS])
def test_cafe_goal_matrix(build, cafe_brief, goal, tier, g, years):
    brief = apply_goal(cafe_brief.with_options(tier=tier, granularity=g, horizon_years=years), goal)
    _passes(build, brief)


@pytest.mark.libreoffice
def test_tampered_free_cash_flow_is_caught(build, cafe_brief):
    brief = apply_goal(cafe_brief.with_options(tier="full"), "dcf_valuation")
    path = build(brief)
    pl = plan(brief, get_template("cafe"))
    sheet, row = pl.line_at["fcff"]
    wb = load_workbook(path)
    for col in range(3, 3 + pl.t.periods):
        c = wb[sheet].cell(row=row, column=col)
        c.value = c.value.replace("*(1-", "*(1.5-")
    wb.save(path)
    report = challenge(brief, path)
    assert report["verdict"] == "FAIL"
    assert any("goal fcff year" in p for p in report["problems"])
    assert any("goal ev_g" in p for p in report["problems"])


@pytest.mark.libreoffice
def test_brief_mismatch_on_goal_driver_is_caught(build, cafe_brief):
    brief = apply_goal(cafe_brief.with_options(tier="full"), "investor_pitch")
    path = build(brief)
    d = brief.model_dump()
    for drv in d["drivers"]:
        if drv["key"] == "cash_buffer":
            drv["base"] += 5000
    report = challenge(d, path)
    assert report["verdict"] == "FAIL"
    assert any("scenario 2 goal gap" in p for p in report["problems"])


def test_build_cli_goal_option(tmp_path, capsys):
    from engine.build import main

    out = tmp_path / "m.xlsx"
    assert main(["default:saas:early", str(out), "--goal", "dcf_valuation", "--tier", "light"]) == 0
    assert load_workbook(out).sheetnames == ["Cover", "Dashboard", "Break-even", "Inputs", "P&L", "Checks"]
