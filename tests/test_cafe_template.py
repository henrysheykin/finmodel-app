import re

from openpyxl import load_workbook

FORBIDDEN = re.compile(r"\b(XLOOKUP|XMATCH|FILTER|SORT|UNIQUE|SEQUENCE|LET|LAMBDA)\s*\(", re.I)


def test_sheets_and_scenario_selector(cafe_xlsx):
    wb = load_workbook(cafe_xlsx)
    assert wb.sheetnames == ["Dashboard", "Inputs", "Monthly", "P&L", "Checks"]
    b4 = wb["Dashboard"]["B4"]
    assert b4.value == 2
    assert b4.fill.fgColor.rgb.endswith("FFFF00")


def test_excel_rules(cafe_xlsx, cafe_brief):
    wb = load_workbook(cafe_xlsx)
    yellow = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                assert c.font.name == "Arial", f"{ws.title}!{c.coordinate}"
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert not FORBIDDEN.search(c.value), f"{ws.title}!{c.coordinate}: {c.value}"
                if c.fill.fill_type == "solid" and c.fill.fgColor.rgb.endswith("FFFF00"):
                    yellow.append(f"{ws.title}!{c.coordinate}")
    assert yellow == ["Dashboard!B4"]
    inp = wb["Inputs"]
    for r in range(5, 5 + len(cafe_brief["drivers"])):
        for col in "BCD":
            assert inp[f"{col}{r}"].font.color.rgb.endswith("0000FF")


def test_same_formula_across_months(cafe_xlsx):
    """Every month after the first uses the same formula shape (R1C1-style comparison)."""
    wb = load_workbook(cafe_xlsx)
    ws = wb["Monthly"]

    def shape(v, col):
        from openpyxl.utils import get_column_letter

        letter = get_column_letter(col)
        prev = get_column_letter(col - 1)
        v = re.sub(rf"(?<![A-Z$]){prev}(?=\d)", "<P>", v)
        return re.sub(rf"(?<![A-Z$]){letter}(?=\d)", "<C>", v)

    for row in (10, 11, 15, 19, 23, 27):
        shapes = {shape(ws.cell(row=row, column=c).value, c) for c in range(4, ws.max_column + 1)}
        assert len(shapes) == 1, (row, shapes)
