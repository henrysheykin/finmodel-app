import pytest
from openpyxl import Workbook, load_workbook

from engine.recalc import RecalcError, recalc

pytestmark = pytest.mark.libreoffice


def test_cafe_recalculates_with_zero_errors(cafe_xlsx):
    res = recalc(cafe_xlsx)
    assert res.ok, res.errors
    assert res.total_formulas > 1000
    assert res.total_errors == 0
    vals = load_workbook(cafe_xlsx, data_only=True)
    assert vals["P&L"]["B5"].value == pytest.approx(319800)
    assert vals["Dashboard"]["B8"].value == pytest.approx(319800)


def test_formulas_survive_recalc(cafe_xlsx):
    recalc(cafe_xlsx)
    wb = load_workbook(cafe_xlsx)
    assert str(wb["Monthly"]["D10"].value).startswith("=")
    assert wb["Dashboard"]["B4"].value == 2


def test_errors_are_reported_with_locations(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 0
    ws["B1"] = "=1/A1"
    ws["B2"] = "=NOSUCHFUNCTION(1)"
    ws["B3"] = "=A1+1"
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    res = recalc(p)
    assert res.status == "errors_found"
    assert res.total_formulas == 3
    assert res.total_errors == 2
    assert "S!B1" in res.errors["#DIV/0!"]
    assert load_workbook(p, data_only=True)["S"]["B3"].value == 1


def test_missing_file_raises(tmp_path):
    with pytest.raises(RecalcError):
        recalc(tmp_path / "nope.xlsx")
