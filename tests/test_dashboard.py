"""Dashboard v2 and the chart kit (E27, docs/WORKBOOK_STANDARD.md §2-§3)."""

import re
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from engine import charts as ck
from engine.builder import PURE_LINK, manifest_for, plan
from engine.challenger import challenge
from engine.dashboard import CORE, LEFT, TILES
from engine.goals import dashboard_reference
from engine.recalc import recalc
from engine.spec import Timing
from engine.templates import TEMPLATES, get_template

FMT_K = '$#,##0,;($#,##0,);"-"'


def _brief(key="saas", **kw):
    tpl = TEMPLATES[key]
    return tpl.default_brief(tpl.sample_stage).with_options(**kw)


def _plan(key="saas", **kw):
    return plan(_brief(key, **kw), get_template(key))


def _labels(ws, col=1, first=1):
    return {
        r: ws.cell(row=r, column=col).value for r in range(first, ws.max_row + 1) if ws.cell(row=r, column=col).value
    }


def _cell(ref: str) -> str:
    return ref.split("!")[1]


def _widths(ws) -> dict[str, float]:
    """Column widths by letter (openpyxl stores runs of equal columns as one entry)."""
    out = {}
    for dim in ws.column_dimensions.values():
        for i in range(dim.min or 1, (dim.max or dim.min or 1) + 1):
            out[get_column_letter(i)] = dim.width
    return out


@pytest.fixture
def saas_dash(build):
    return load_workbook(build(_brief("saas", tier="full", horizon_years=5)))["Dashboard"]


# --------------------------------------------------------------------------- layout (standard §2)
def test_grid_selector_verdict_and_tiles(saas_dash):
    ws = saas_dash
    widths = _widths(ws)
    assert widths["A"] == pytest.approx(28.4) and widths["G"] == 2
    for col in "BCDEFHIJKLMNOPQRSTUVW":
        assert widths[col] == 9, col
    assert ws["A4"].value == "Active scenario" and ws["B4"].value == 2
    assert ws["B4"].fill.fgColor.rgb[-6:] == "FFE08A"
    assert ws["C4"].value == '=CHOOSE(B4,"Low","Base","High")' and ws["D4"].value.startswith("1 = Low")
    assert any("B4" in dv.sqref for dv in ws.data_validations.dataValidation)
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert ws["S4"].value == "Model checks" and "T4:W4" in merged
    assert ws["T4"].value.startswith("=Checks!B")
    for a, b in TILES:
        assert {f"{a}6:{b}6", f"{a}7:{b}7", f"{a}8:{b}8"} <= merged
    assert [ws[f"{a}6"].value for a, _ in TILES] == [
        "REVENUE · YEAR 5",
        "EBITDA MARGIN · YEAR 5",
        "NET INCOME · YEAR 5",
        "LOWEST CASH BALANCE",
        "FUNDING REQUIRED",
        "ENTERPRISE VALUE (DCF)",
    ]
    assert ws["A7"].number_format == '$#,##0.0,,"M";($#,##0.0,,"M");"-"'
    assert ws["A7"].font.size == 18 and ws["A6"].font.size == 7


def test_core_financials_are_the_same_13_lines(saas_dash):
    lab = _labels(saas_dash)
    assert lab[10].startswith("Core financials ($'000)")
    assert [lab[r] for r in range(12, 25)] == [
        "Revenue",
        "Gross profit",
        "Gross margin",
        "Operating expenses",
        "EBITDA",
        "EBITDA margin",
        "Net income",
        "Net margin",
        "Operating cash flow",
        "Investing cash flow",
        "Financing cash flow",
        "Net change in cash",
        "Year-end cash",
    ]
    assert len(CORE) == 13
    assert saas_dash["B12"].value.startswith("=SUMIF(Monthly!")  # rolled up from the grid, not linked to a statement
    assert saas_dash["B12"].number_format == FMT_K
    assert saas_dash["F24"].value.startswith("=INDEX(Monthly!")  # balances at year end
    assert saas_dash["B11"].value == 1 and saas_dash["B11"].number_format == '"Year "0'
    assert saas_dash["B14"].font.italic and saas_dash["A14"].font.italic  # margins: italic grey
    assert saas_dash["B18"].border.bottom.style == "double"  # net income: key total


def test_sections_in_standard_order_and_template_metrics(saas_dash):
    lab = _labels(saas_dash)
    bands = [
        r for r, v in lab.items() if v.startswith(("Core financials", "Funding & valuation")) or v == "SaaS metrics"
    ]
    assert [lab[r][:8] for r in bands] == ["Core fin", "SaaS met", "Funding "]
    assert [lab[r] for r in range(28, 32)] == [
        "Paying customers (year end)",
        "ARR (year-end MRR × 12)",
        "LTV ÷ CAC",
        "Team (FTE, year end)",
    ]
    fv = [lab[r] for r in range(bands[2] + 1, bands[2] + 11)]
    assert fv[:4] == [
        "Equity raised",
        "Loan drawn",
        "Funding required (incl. cash buffer)",
        "Funding surplus / (shortfall)",
    ]
    assert fv[4].startswith("=\"Lowest cash balance (month \"&'Break-even'!")
    assert fv[5:] == [
        "First month with positive EBITDA",
        "Cash invested paid back",
        "Enterprise value — perpetual growth",
        "Enterprise value — exit multiple",
        "Enterprise value — average",
    ]


def test_four_charts_in_the_standard_grid(saas_dash):
    anchors = [(c.anchor._from.col, c.anchor._from.row) for c in saas_dash._charts]
    assert anchors == [(7, 9), (15, 9), (7, 26), (15, 26)]  # H10 P10 H27 P27


def test_template_kpi_table_moved_to_the_kpis_sheet(build):
    """E31: the template's KPI table lives on the KPIs sheet; the Dashboard keeps its industry block (standard §2)."""
    brief = _brief("saas", tier="full")
    wb = load_workbook(build(brief))
    pl = plan(brief, get_template("saas"))
    kpi = pl.blocks["kpi"]
    assert kpi.sheet == "KPIs"
    lab = _labels(wb["Dashboard"])
    for row in get_template("saas").kpi_rows(pl.t):
        assert row.label in _labels(wb["KPIs"]).values()
    assert "Runway at year-end burn (months; - = not burning)" not in lab.values()  # KPI-table-only row
    assert max(lab) == pl.dash.end  # the Dashboard ends with Funding & valuation
    for m, row in pl.dash.metrics:  # the industry block links to the KPIs sheet
        if m.src.startswith("kpi:"):
            assert wb["Dashboard"].cell(row=row, column=2).value == "=" + pl.stmt_cell("kpi", m.src[4:], 1)


def test_light_dashboard_has_no_cf_bs_and_no_ev(build):
    wb = load_workbook(build(_brief("saas", tier="light", horizon_years=3)))
    ws = wb["Dashboard"]
    assert "Cash Flow" not in wb.sheetnames and "Balance Sheet" not in wb.sheetnames
    assert ws["T6"].value == "FIRST MONTH WITH POSITIVE EBITDA"
    assert ws["T7"].value.startswith("=F") and ws["T7"].number_format == '"Month "0'
    lab = _labels(ws)
    assert not any("Enterprise value" in str(v) for v in lab.values())
    assert "Cash invested paid back" in lab.values()
    assert "'P&L'!" in ws["B12"].value and "'P&L'!" in ws["D24"].value  # rolled up from the grid on the P&L sheet
    assert "SaaS metrics" in [c.value for c in wb["P&L"]["A"]]  # the KPI table of Light stays on the P&L sheet


def test_annual_models_say_year_not_month(build):
    brief = _brief("cafe", tier="light", granularity="annual", horizon_years=3)
    ws = load_workbook(build(brief))["Dashboard"]
    pl = plan(brief, get_template("cafe"))
    assert ws["T6"].value == "FIRST YEAR WITH POSITIVE EBITDA"
    assert ws["L8"].number_format == '"in year "0'
    low = ws[_cell(pl.dash.cells["dash_fv_low_label"])].value
    assert '(year "&' in low
    assert ws[_cell(pl.dash.cells["dash_fv_first"])].number_format == '"Year "0'
    assert "First year with positive EBITDA" in _labels(ws).values()


def test_cagr_is_n_a_for_one_year_and_a_formula_otherwise(build):
    one = load_workbook(build(_brief("cafe", tier="full", horizon_years=1)))["Dashboard"]
    assert one["A8"].value == "CAGR n/a" and one["A6"].value == "REVENUE · YEAR 1"
    three = load_workbook(build(_brief("cafe", tier="full", horizon_years=3), "t.xlsx"))["Dashboard"]
    assert three["A8"].value.startswith("=IF(AND(B12>0,D12>=0),(D12/B12)^(1/2)-1,")
    assert three["A6"].value == "REVENUE · YEAR 3"


@pytest.mark.parametrize("key", sorted(TEMPLATES))
@pytest.mark.parametrize("tier", ["light", "full"])
def test_every_template_has_industry_metrics_and_chart(build, key, tier):
    brief = _brief(key, tier=tier)
    tpl = get_template(key)
    pl = plan(brief, tpl)
    metrics = tpl.dashboard_metrics(pl.t)
    assert 2 <= len(pl.dash.metrics) <= 5  # three to five lines with the team line (two for the cafe)
    assert all(m.fmt in ("usdk", "usd2", "int", "dec1", "pct", "x") for m in metrics)
    ws = load_workbook(build(brief))["Dashboard"]
    assert len(ws._charts) == 4
    lab = _labels(ws)
    assert tpl.kpi_title in lab.values()
    for m, row in pl.dash.metrics:
        assert lab[row] == m.label
    spec = tpl.dashboard_chart(pl.t)
    assert spec is not None and set(spec.bars + spec.lines) <= {m.key for m in metrics}
    assert ws._charts[3].title.tx.rich.p[0].r[0].t == spec.title


def test_default_chart_when_a_template_names_none(build, monkeypatch):
    monkeypatch.setattr(type(get_template("cafe")), "dashboard_chart", lambda self, t: None)
    ws = load_workbook(build(_brief("cafe", tier="full")))["Dashboard"]
    assert ws._charts[3].title.tx.rich.p[0].r[0].t == "Gross profit & gross margin"


# --------------------------------------------------------------------------- styles (standard §4)
def test_report_sheet_links_are_ink_not_green(build):
    wb = load_workbook(build(_brief("saas", tier="full")))
    for row in wb["Dashboard"].iter_rows(min_row=6, max_row=45, max_col=23):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                assert c.font.color.rgb[-6:] != "008000", c.coordinate
    verdict = next(c for c in wb["Cover"]["D"] if isinstance(c.value, str) and c.value.startswith("=Checks!"))
    assert verdict.font.color.rgb[-6:] == "1B2A2F"


def test_module_scenario_label_and_percentage_lines_follow_the_tokens(build):
    brief = _brief("cafe", tier="full", horizon_years=3)
    wb = load_workbook(build(brief))
    for name in ("Break-even", "Valuation"):
        c = wb[name]["B3"]
        assert c.font.color.rgb[-6:] == "0F6E56" and c.font.bold, name
    pl = plan(brief, get_template("cafe"))
    key = next(
        k
        for k, ln in pl.lines.items()
        if ln.fmt == "pct" and k in pl.line_at and not PURE_LINK.match(pl.render(ln, 2, pl.line_at[k][0]))
    )
    sheet, row = pl.line_at[key]
    for col in (1, 4):  # grid sheet: label and a period value of a percentage line (links stay green)
        c = wb[sheet].cell(row=row, column=col)
        assert c.font.italic and c.font.color.rgb[-6:] == "6B7280", (key, col)
    val = wb["Valuation"]
    row = next(r for r in range(1, 60) if val.cell(row=r, column=1).value == "Terminal value share of enterprise value")
    assert val.cell(row=row, column=1).font.italic and val.cell(row=row, column=2).font.italic
    assert val.cell(row=row, column=2).font.color.rgb[-6:] == "6B7280"


# --------------------------------------------------------------------------- the Python twin
def test_twin_texts_and_cagr_rules():
    brief = _brief("cafe", tier="light", horizon_years=3)
    t = Timing("monthly", 3)
    keys = ("rev", "gp", "opex_total", "ebitda", "ni", "cfo", "cfi", "cff", "cash_close", "capex")
    series = {k: [100.0] * t.periods for k in keys}
    goal = {
        "gap": -1_250_000.0,
        "min_cash": -5.0,
        "low_p": 4,
        "required": 9.0,
        "first_ebitda": 2,
        "first_ni": "Not within horizon",
        "payback": "No investment to pay back",
    }
    d = {"equity": 1.0, "loan_amount": 2.0}
    out = dashboard_reference(brief, d, series, t, goal)["cells"]
    assert out["dash_t5_sub"] == 1_250_000.0  # numeric (shown as "Shortfall $1.3M" by the cell format): no rounding
    assert out["dash_t1_sub"] == 0.0  # flat revenue: CAGR 0
    assert out["dash_t6_first_sub"] == "Net income not positive within horizon"
    assert out["dash_fv_low_label"] == "Lowest cash balance (month 4)"
    assert "dash_t6_ev" not in out and "dash_fv_ev_avg" not in out
    one = dashboard_reference(brief.with_options(horizon_years=1), d, series, Timing("monthly", 1), goal)["cells"]
    assert one["dash_t1_sub"] == "CAGR n/a"
    goal["gap"] = 0.0
    assert dashboard_reference(brief, d, series, t, goal)["cells"]["dash_t5_sub"] == "Covered by funding raised"
    goal["first_ni"] = 7
    assert (
        dashboard_reference(brief, d, series, t, goal)["cells"]["dash_t6_first_sub"]
        == "Net income positive from month 7"
    )


def test_manifest_lists_every_tile_and_funding_cell():
    for tier, ev in (("light", False), ("full", True)):
        brief = _brief("saas", tier=tier)
        pl = plan(brief, get_template("saas"))
        man = manifest_for(pl, "")
        names = set(man.dash_cells)
        assert {f"dash_t{i}" for i in range(1, 6)} <= names and {f"dash_t{i}_sub" for i in range(1, 6)} <= names
        assert ("dash_fv_ev_avg" in names) is ev and ("dash_t6_ev" in names) is ev
        assert ("dash_t6_first" in names) is (not ev)
        assert {"dash_fv_equity", "dash_fv_loan", "dash_fv_required", "dash_fv_gap", "dash_fv_low"} <= names
        assert {"dash_fv_first", "dash_fv_payback", "dash_fv_low_label"} <= names
        # 13 Core financials lines + one per industry line (E31: the challenger recomputes those too)
        assert len(man.dash_years) == 13 + len(pl.dash.metrics) and all(
            len(v) == brief.horizon_years for v in man.dash_years.values()
        )


def test_dashboard_error_range_joins_the_no_error_check(build):
    ws = load_workbook(build(_brief("saas", tier="full")))["Checks"]
    text = " ".join(str(c.value) for c in ws["B"] if c.value)
    assert "Dashboard!$A$6:$W$" in text and "Dashboard!$A$4" not in text  # row 4 holds the verdict: no circle


def test_left_column_constants():
    assert LEFT == 6 and TILES[0] == ("A", "B") and TILES[-1] == ("T", "W")


# --------------------------------------------------------------------------- chart kit
def _kit_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["", "Year 1", "Year 2", "Year 3"])
    ws.append(["Revenue", 100, 200, 300])
    ws.append(["Cost", -50, -60, -70])
    ws.append(["Margin", 0.1, 0.2, 0.3])
    cats = "'Data'!$B$1:$D$1"
    rev, cost, mar = (ck.Ser(f"'Data'!$B${r}:$D${r}", f"'Data'!$A${r}") for r in (2, 3, 4))
    ck.place(ws, ck.bar("Bars", cats, [rev, cost], labels=True), "F2")
    ck.place(ws, ck.line("Lines", cats, [rev, cost]), "F20")
    ck.place(ws, ck.combo("Combo", cats, [rev, cost], [mar]), "F38")
    ck.place(ws, ck.doughnut("Doughnut", cats, rev, points=3), "F56")
    ck.place(ws, ck.bar("Horizontal", cats, [rev], horizontal=True), "F74")
    wb.save(path)


def _norm(xml: str) -> str:
    """Chart XML without namespace prefixes or spacing quirks (openpyxl and LibreOffice write it differently)."""
    return re.sub(r"\s+/>", "/>", re.sub(r"<(/?)\w+:", r"<\1", xml))


def _chart_xml(path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        names = sorted(
            (n for n in z.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", n)),
            key=lambda n: int(re.findall(r"\d+", n)[0]),
        )
        return [_norm(z.read(n).decode()) for n in names]


def test_chart_kit_style_flags(tmp_path):
    path = tmp_path / "kit.xlsx"
    _kit_workbook(path)
    xml = _chart_xml(path)
    assert len(xml) == 5
    bars, lines, combo, donut, hbar = xml
    for x in xml:
        assert 'sz="1000"' in x and 'b="1"' in x and "1B2A2F" in x and "Arial" in x  # title: 10 pt bold ink
    for x in (bars, lines, combo, hbar):
        assert 'sz="800"' in x and "6B7280" in x  # axes and legend 8 pt grey
        assert '<tickLblPos val="low"/>' in x  # labels stay at the bottom when values go negative
    assert '<legendPos val="b"/>' in bars and '<legendPos val="b"/>' in donut
    assert "<legend>" not in hbar  # one series: no legend
    flags = re.search(r"<dLbls>.*?</dLbls>", bars, re.S).group(0)  # the value on, every other flag off
    for flag, want in (
        ("showVal", 1),
        ("showSerName", 0),
        ("showCatName", 0),
        ("showLegendKey", 0),
        ("showPercent", 0),
    ):
        assert f'<{flag} val="{want}"/>' in flags, flag
    assert "<dLbls>" not in lines
    assert '<invertIfNegative val="0"/>' in bars  # Excel would paint negative bars white
    assert '<axId val="200"/>' in combo and '<crosses val="max"/>' in combo  # secondary axis
    assert "<holeSize" in donut and '<showPercent val="1"/>' in donut
    assert '<barDir val="bar"/>' in hbar
    assert bars.index("0F6E56") < bars.index("7FB5A3")  # standard order: pine, sage
    assert re.search(r"<spPr><ln><noFill/>", bars.split("</chart>")[-1])  # no chart border


def test_slot_anchors_follow_the_two_column_grid():
    assert [ck.slot_anchor(i, 10) for i in range(4)] == ["H10", "P10", "H27", "P27"]
    assert ck.slot_anchor(4, 10) == "H44" and ck.LARGE_CM == (20.0, 10.5)
    assert (ck.WIDTH_CM, ck.HEIGHT_CM) == (14.2, 8.3)


@pytest.mark.libreoffice
def test_chart_kit_survives_the_libreoffice_recalculation(tmp_path):
    path = tmp_path / "kit.xlsx"
    _kit_workbook(path)
    assert recalc(path).ok
    xml = _chart_xml(path)
    assert len(xml) == 5
    assert all("Data!$B$2:$D$2" in x for x in xml)
    donut = next(x for x in xml if "<doughnutChart>" in x)
    assert '<showPercent val="1"/>' in donut
    labelled = next(x for x in xml if "<dLbls>" in x and "<barChart>" in x)
    assert '<showSerName val="0"/>' in labelled and '<showCatName val="0"/>' in labelled


# --------------------------------------------------------------------------- LibreOffice: Dashboard end to end
def _dashboard_chart_refs(path) -> dict[str, list[str]]:
    with zipfile.ZipFile(path) as z:
        sheet_no = load_workbook(path).sheetnames.index("Dashboard") + 1
        rels = z.read(f"xl/worksheets/_rels/sheet{sheet_no}.xml.rels").decode()
        drawing = re.search(r'Target="[^"]*/(drawing\d+\.xml)"', rels).group(1)
        drel = z.read(f"xl/drawings/_rels/{drawing}.rels").decode()
        charts = re.findall(r'Target="[^"]*/(chart\d+\.xml)"', drel)
        return {c: re.findall(r"<c:f>([^<]+)</c:f>", z.read(f"xl/charts/{c}").decode()) for c in charts}


@pytest.mark.libreoffice
@pytest.mark.parametrize("tier,gran,years", [("full", "monthly", 5), ("light", "annual", 3)])
def test_recalculated_file_keeps_four_live_charts(build, tier, gran, years):
    path = build(_brief("saas", tier=tier, granularity=gran, horizon_years=years))
    assert recalc(path).ok
    refs = _dashboard_chart_refs(path)
    assert len(refs) == 4
    for chart, fs in refs.items():
        assert fs and all(f.startswith("Dashboard!") for f in fs), (chart, fs)  # every series reads Dashboard cells
    assert len(load_workbook(path)["Dashboard"]._charts) == 4
    vals = load_workbook(path, data_only=True)["Dashboard"]  # the cells the charts read hold values
    assert isinstance(vals["B12"].value, int | float) and vals["B11"].value == 1
    assert vals["T4"].value in ("PASS", "PASS WITH WARNINGS")


@pytest.mark.libreoffice
def test_dashboard_values_after_recalc(build):
    brief = _brief("saas", tier="full", horizon_years=3)
    path = build(brief)
    assert recalc(path).ok
    cells = _plan("saas", tier="full", horizon_years=3).dash.cells
    wb = load_workbook(path, data_only=True)
    ws = wb["Dashboard"]
    assert ws["A7"].value == pytest.approx(ws["D12"].value)  # tile 1 = revenue of the last year
    assert ws["A8"].value == pytest.approx((ws["D12"].value / ws["B12"].value) ** 0.5 - 1)
    assert ws["P7"].value == pytest.approx(ws[_cell(cells["dash_fv_required"])].value)
    assert ws["T7"].value == pytest.approx(ws[_cell(cells["dash_fv_ev_avg"])].value)
    assert ws["C4"].value == "Base" and ws["T4"].value == wb["Cover"]["D7"].value
    for row in (12, 13, 15, 17, 19, 21, 22, 23, 24):
        assert all(isinstance(ws.cell(row=row, column=c).value, int | float) for c in (2, 3, 4)), row


@pytest.mark.libreoffice
@pytest.mark.parametrize("key", ["cafe", "real_estate"])
def test_challenger_covers_the_dashboard(build, key):
    brief = _brief(key, tier="full", horizon_years=3)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
    assert {"dash_t1", "dash_t5_sub", "dash_fv_ev_avg"} <= set(report["scenarios"][2]["dash"])


TAMPER = [
    ("full", "dash_t1", "=D12*1.05", "dash_t1"),  # a wrong tile
    ("full", "dash_t1_sub", "=0.5", "dash_t1_sub"),  # a wrong CAGR
    ("full", "dash_t5_sub", "Tampered text", "dash_t5_sub"),  # a wrong funding text
    ("full", "dash_fv_gap", "=12345", "dash_fv_gap"),  # a wrong funding gap
    ("full", "dash_fv_ev_avg", "=12345", "dash_fv_ev_avg"),  # a wrong EV average
    ("full", "dash_ebitda_m year 3", "=12345", "dash_ebitda_m year 3"),  # a wrong Core financials line
    ("light", "dash_t6_first", "=12345", "dash_t6_first"),  # Light: first positive-EBITDA period
]


@pytest.mark.libreoffice
@pytest.mark.parametrize("tier,name,formula,needle", TAMPER, ids=[t[1] for t in TAMPER])
def test_challenger_catches_a_tampered_dashboard(build, tier, name, formula, needle):
    brief = _brief("cafe", tier=tier, horizon_years=3)
    path = build(brief)
    man = manifest_for(plan(brief, get_template("cafe")), "")
    if " year " in name:
        stem, y = name.split(" year ")
        ref = man.dash_years[stem][int(y)]
    else:
        ref = man.dash_cells[name]
    wb = load_workbook(path)
    wb["Dashboard"][_cell(ref)] = formula
    wb.save(path)
    report = challenge(brief, path)
    assert report["verdict"] == "FAIL"
    assert any(needle in p for p in report["problems"]), report["problems"][:4]
