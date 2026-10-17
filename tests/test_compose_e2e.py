"""End to end against the running Docker stack (docker-compose.yml: production image + Postgres 16).

Skipped unless COMPOSE_BASE_URL is set; CI's `compose` job starts the stack and sets it:

    docker compose up --build --wait
    COMPOSE_BASE_URL=http://localhost:8000 pytest -m compose tests/test_compose_e2e.py

The path is the customer's: landing → wizard → checkout → (fake) payment → worker builds with
LibreOffice → challenger → ready; the file is then fetched through the admin panel and its Checks
sheet must say PASS.
"""

from __future__ import annotations

import io
import os
import time

import httpx
import pytest
from openpyxl import load_workbook

BASE = os.environ.get("COMPOSE_BASE_URL", "").rstrip("/")
ADMIN = (os.environ.get("ADMIN_USER", "admin"), os.environ.get("ADMIN_PASSWORD", "local-admin-only"))
BUILD_TIMEOUT = float(os.environ.get("COMPOSE_BUILD_TIMEOUT", "240"))

pytestmark = [
    pytest.mark.compose,
    pytest.mark.skipif(not BASE, reason="COMPOSE_BASE_URL not set (needs `docker compose up`)"),
]

FORM = {
    "business_type": "professional_services",
    "business_name": "Compose Check",
    "goal": "revenue_profit",
    "stage": "early",
    "horizon_years": "1",
    "granularity": "monthly",
    "tier": "light",
    "email": "founder@example.com",
    "sample_data": "1",
}


@pytest.fixture
def http():
    with httpx.Client(base_url=BASE, timeout=30, follow_redirects=False) as c:
        yield c


def _location(r: httpx.Response) -> str:
    assert r.status_code == 303, (r.status_code, r.text[:500])
    return r.headers["location"].replace(BASE, "")


def test_stack_is_up(http):
    assert http.get("/healthz").json() == {"status": "ok"}
    landing = http.get("/")
    assert landing.status_code == 200 and "StartupModel" in landing.text
    assert http.get("/start").status_code == 200


def test_paid_order_is_built_checked_and_downloadable(http):
    checkout = _location(http.post("/start", data=FORM))
    assert checkout.startswith("/checkout/")
    oid = checkout.rsplit("/", 1)[1]
    assert http.get(checkout).status_code == 200

    dev_pay = _location(http.post(f"{checkout}/pay"))
    assert dev_pay.startswith("/dev/pay/")
    success = _location(http.post(dev_pay, data={"action": "pay"}))
    assert success == f"/checkout/{oid}/success"

    deadline, status = time.monotonic() + BUILD_TIMEOUT, ""
    while time.monotonic() < deadline:
        status = http.get(f"/api/orders/{oid}").json()["status"]
        if status in ("ready", "failed"):
            break
        time.sleep(2)
    assert status == "ready", f"order {oid} ended as {status!r}"

    r = http.post("/admin/login", data={"user": ADMIN[0], "password": ADMIN[1]})
    assert r.status_code == 303, r.text[:300]
    f = http.get(f"/admin/orders/{oid}/download")
    assert f.status_code == 200
    assert f.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(f.content), data_only=True)
    verdict = [r[1].value for r in wb["Checks"].iter_rows(min_col=1, max_col=2) if r[0].value == "Overall verdict"]
    assert verdict and verdict[0] in ("PASS", "PASS WITH WARNINGS")
