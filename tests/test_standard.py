"""Workbook standard v1 (ADR-013, docs/WORKBOOK_STANDARD.md): sheet set by tier, selector, Cover, look."""

import re
from datetime import date

import pytest
from openpyxl import load_workbook

from engine.brief import Brief
from engine.build import build_model
from engine.builder import plan, sheet_order
from engine.challenger import challenge
from engine.goals import ALL_GOAL_KEYS, DCF_SPECS, apply_goal, purpose
from engine.templates import TEMPLATES, get_template

NAVY = ("1F3864", "2F5597")
LIGHT = ["Cover", "Dashboard", "Break-even", "Inputs", "P&L", "Checks"]
FULL = [
    "Cover",
    "Dashboard",
    "Summary",
    "Break-even",
    "Valuation",
    "KPIs",
    "Inputs",
    "P&L",
    "Cash Flow",
    "Balance Sheet",
    "Monthly",
    "Headcount",
    "Checks",
]


def _brief(key="saas", **kw):
    tpl = TEMPLATES[key]
    return tpl.default_brief(tpl.sample_stage).with_options(**kw)


@pytest.mark.parametrize("tier,expected", [("light", LIGHT), ("full", FULL)])
@pytest.mark.parametrize("goal", ["revenue_profit", "startup_breakeven", "dcf_valuation", "investor_pitch"])
def test_sheet_set_follows_the_tier_not_the_goal(tier, expected, goal):
    b = _brief(tier=tier).model_copy(update={"goal": goal})
    pl = plan(b, get_template(b.template))
    assert pl.sheets == expected
    assert pl.selector == ("Dashboard", "B4")


def test_sheet_order_annual_and_industry_sheets():
    assert sheet_order("full", "Annual")[10] == "Annual"
    assert sheet_order("light", returns=True) == [
        "Cover",
        "Dashboard",
        "Break-even",
        "Returns",
        "Inputs",
        "P&L",
        "Checks",
    ]
    re_full = plan(_brief("real_estate", tier="full"), TEMPLATES["real_estate"]).sheets
    assert re_full.index("Returns") == re_full.index("KPIs") + 1  # industry sheets come after KPIs


def test_module_drivers_follow_the_tier():
    light = apply_goal(_brief(tier="light"), "investor_pitch")
    full = apply_goal(_brief(tier="full"), "revenue_profit")
    assert {d.key for d in light.drivers} & ALL_GOAL_KEYS == {"cash_buffer"}
    assert {d.key for d in full.drivers} & ALL_GOAL_KEYS == ALL_GOAL_KEYS
    assert apply_goal(full) == full  # idempotent


def test_purpose_line():
    assert purpose("investor_pitch") == "Investor-ready plan"
    assert purpose("anything else") == "Financial plan"


@pytest.mark.parametrize("key,tier", [("saas", "light"), ("real_estate", "full"), ("real_estate", "light")])
def test_no_navy_anywhere(tmp_path, key, tier):
    from engine.builder import build

    p = tmp_path / "m.xlsx"
    build(apply_goal(_brief(key, tier=tier)), TEMPLATES[key], p)
    wb = load_workbook(p)
    for ws in wb.worksheets:
        tab = ws.sheet_properties.tabColor
        assert not (tab and tab.rgb[-6:] in NAVY), ws.title
        for row in ws.iter_rows():
            for c in row:
                b = c.border
                colours = [c.fill.fgColor, c.font.color if c.font else None]
                colours += [getattr(s, "color", None) for s in (b.left, b.right, b.top, b.bottom)]
                for col in colours:
                    rgb = getattr(col, "rgb", None)
                    assert not (isinstance(rgb, str) and rgb[-6:] in NAVY), (ws.title, c.coordinate)


@pytest.fixture(scope="module")
def saas_full(tmp_path_factory):
    p = tmp_path_factory.mktemp("std") / "saas.xlsx"
    from engine.builder import build

    build(apply_goal(_brief(tier="full")), TEMPLATES["saas"], p, built_on=date(2026, 10, 20))
    return load_workbook(p)


def test_cover_links_every_sheet_and_shows_status(saas_full):
    ws = saas_full["Cover"]
    links = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value]
    targets = {re.search(r"#\'?([^\'!]+)\'?!A1", v).group(1) for v in links}
    assert targets == set(saas_full.sheetnames) - {"Cover"}
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "20 October 2026" in values and "Revenue and profit plan" in values
    assert any(isinstance(v, str) and v.startswith("=Checks!B") for v in values)
    assert saas_full.sheetnames[0] == "Cover" and saas_full.active.title == "Cover"


def test_every_sheet_has_a_contents_link_and_tab_colour(saas_full):
    for ws in saas_full.worksheets:
        assert ws.sheet_properties.tabColor is not None, ws.title
        assert not ws.sheet_view.showGridLines, ws.title
        if ws.title != "Cover":
            row1 = [c.value for c in ws[1] if isinstance(c.value, str)]
            assert any("◂ Contents" in v for v in row1), ws.title


def test_style_tokens_and_no_navy(saas_full):
    inp = saas_full["Inputs"]
    cells = [
        inp.cell(r, c)
        for r in range(6, inp.max_row + 1)
        for c in (2, 3, 4)
        if isinstance(inp.cell(r, c).value, (int, float))
    ]
    assert cells
    for c in cells:
        assert c.fill.fgColor.rgb[-6:] == "FFF2CC" and c.font.color.rgb[-6:] == "0000FF", c.coordinate
    sel = saas_full["Dashboard"]["B4"]
    assert sel.value == 2 and sel.fill.fgColor.rgb[-6:] == "FFE08A"
    for ws in saas_full.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.font is not None and c.value is not None:
                    assert c.font.name == "Arial", (ws.title, c.coordinate)
                for col in (c.fill.fgColor, c.font.color if c.font else None):
                    rgb = getattr(col, "rgb", None)
                    assert not (isinstance(rgb, str) and rgb[-6:] in NAVY), (ws.title, c.coordinate)
    assert saas_full["Inputs"].sheet_properties.tabColor.rgb[-6:] == "E0B43C"
    assert saas_full["Checks"].sheet_properties.tabColor.rgb[-6:] == "B8322A"


def _strip(brief: Brief, keys) -> Brief:
    data = brief.model_dump()
    data["drivers"] = [d for d in data["drivers"] if d["key"] not in keys]
    return Brief.model_validate(data)


@pytest.mark.libreoffice
@pytest.mark.parametrize(
    "tier,goal,drop",
    [
        ("light", "revenue_profit", {"cash_buffer"}),  # stored before ADR-013: Light without the break-even module
        ("full", "startup_breakeven", {s.key for s in DCF_SPECS}),  # Full without the DCF drivers
    ],
)
def test_orders_stored_before_the_standard_still_build(tmp_path, tier, goal, drop):
    old = _strip(_brief("cafe", tier=tier, horizon_years=3, granularity="annual"), ALL_GOAL_KEYS)
    old = old.model_copy(update={"goal": goal})
    assert not {d.key for d in old.drivers} & drop
    path = tmp_path / "old.xlsx"
    build_model(old, path)
    rep = challenge(old.model_dump(), path)
    assert rep["verdict"].startswith("PASS"), rep["problems"][:3]


@pytest.mark.libreoffice
def test_stored_full_brief_keeps_its_own_buffer(tmp_path):
    # an old break-even-goal Full order: has its own cash buffer, lacks the DCF drivers
    b = apply_goal(_brief("cafe", tier="full", horizon_years=3, granularity="annual"))
    data = b.model_dump()
    data["drivers"] = [d for d in data["drivers"] if d["key"] not in {s.key for s in DCF_SPECS}]
    for d in data["drivers"]:
        if d["key"] == "cash_buffer":
            d.update(low=123000, base=123000, high=123000)
    old = Brief.model_validate(data)
    path = tmp_path / "old.xlsx"
    build_model(old, path)
    pl = plan(old, TEMPLATES["cafe"])
    r = pl.input_row["cash_buffer"]
    assert load_workbook(path)["Inputs"][f"C{r}"].value == 123000
    assert challenge(old.model_dump(), path)["verdict"].startswith("PASS")
