"""Car wash template: economics in the plain-Python twin, plus the LibreOffice + challenger matrix."""

import itertools

import pytest
from openpyxl import load_workbook

from engine.challenger import challenge, python_model
from engine.core import reference_run
from engine.recalc import recalc
from engine.spec import Timing
from engine.templates import get_template
from engine.templates.car_wash import _ramp

CW = get_template("car_wash")
STAGES = ["planning", "early", "growing", "established"]
#: no cap, no winter dip: the volume path is then exactly additive over months
ROOMY = {"cph": 1e4, "winter_dip": 0}


@pytest.fixture
def cw_brief():
    return CW.default_brief("early")


def run(brief, scen=2, g="monthly", years=3):
    return reference_run(brief.values(scen), brief.roles, Timing(g, years), CW.reference_top)


def with_values(brief, **vals):
    d = brief.model_dump()
    for drv in d["drivers"]:
        if drv["key"] in vals:
            v = vals[drv["key"]]
            drv.update(low=v, base=v, high=v)
    return brief.model_validate(d)


def test_registered_with_wizard_hooks():
    assert CW.order == 110 and CW.business_types == ("car_wash",)
    assert "equity" in CW.key_numbers and 5 <= len(CW.key_numbers) <= 8
    assert CW.dashboard_metrics(Timing("monthly", 3)) and CW.dashboard_chart(Timing("monthly", 3))


@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_are_valid_ordered_and_funded(stage):
    b = CW.default_brief(stage)
    low, base, high = (python_model(b, s) for s in (1, 2, 3))
    for y in (1, 2, 3):
        assert low["rev"][y] < base["rev"][y] < high["rev"][y]
        assert 0 < base["gp"][y] / base["rev"][y] < 1
    monthly = reference_run(b.values(2), b.roles, Timing("monthly", 3), CW.reference_top)
    assert min(monthly["cash_close"]) > 0  # the Base case is funded in every month; Low is the stress case


def test_ramp_up_is_linear_and_annual_is_the_mean_of_its_months(cw_brief):
    d = {**cw_brief.values(2), "ramp_months": 18, "ramp_start": 0.4}
    m, a = Timing("monthly", 3), Timing("annual", 3)
    factors = [_ramp(d, p, m) for p in range(1, 37)]
    assert factors[0] == pytest.approx(0.4)
    assert factors[9] == pytest.approx(0.4 + 0.6 * 9 / 18)
    assert factors[18:] == pytest.approx([1.0] * 18)
    for y in (1, 2, 3):
        assert _ramp(d, y, a) == pytest.approx(sum(factors[(y - 1) * 12 : y * 12]) / 12)
    assert _ramp({**d, "ramp_months": 0}, 1, m) == 1.0


def test_retail_volume_follows_cars_per_day_days_and_ramp(cw_brief):
    d = cw_brief.values(2)
    b = with_values(cw_brief, **ROOMY)
    s = run(b)
    # roomy tunnel: every retail car demanded is washed; month 1 = start share of mature volume
    assert s["ret_cars"][0] == pytest.approx(d["cars_day"] * d["days_open"] * d["ramp_start"])
    assert s["lost_washes"][0] == pytest.approx(0)


def test_member_roll_forward_start_plus_adds_minus_churn(cw_brief):
    d = cw_brief.values(2)
    b = with_values(cw_brief, **ROOMY, ramp_months=0)
    s = run(b)
    retail = d["cars_day"] * d["days_open"]
    new = retail * d["conv"] + d["mem_new"]
    m1 = d["members_start"] * (1 - d["churn_m"]) + new
    assert s["new_total"][0] == pytest.approx(new)
    assert s["members_close"][0] == pytest.approx(m1)
    assert s["members_close"][1] == pytest.approx(m1 * (1 - d["churn_m"]) + new)
    assert s["churned"][0] == pytest.approx(-d["members_start"] * d["churn_m"])


def test_revenue_is_retail_price_plus_addons_plus_plan_fees(cw_brief):
    d = cw_brief.values(2)
    b = with_values(cw_brief, **ROOMY, ramp_months=0)
    s = run(b)
    avg = (d["members_start"] + s["members_close"][0]) / 2
    expected = s["ret_cars"][0] * (d["price"] + d["addon"]) + avg * d["mem_fee"]
    assert s["rev"][0] == pytest.approx(expected)
    assert s["rev_mem"][0] == pytest.approx(avg * d["mem_fee"])


def test_member_washes_cost_money_but_earn_no_extra_revenue(cw_brief):
    base = run(with_values(cw_brief, **ROOMY))
    heavy = run(with_values(cw_brief, **ROOMY, wash_pm=7))
    for r0, r1 in zip(base["rev"], heavy["rev"], strict=True):
        assert r1 == pytest.approx(r0)  # the plan fee is the same however often members wash
    assert heavy["mem_washes"][5] > 1.9 * base["mem_washes"][5]
    assert heavy["cogs"][5] < base["cogs"][5]  # more chemicals, water and electricity


def test_cost_of_sales_is_variable_cost_per_wash_plus_card_fees(cw_brief):
    d = cw_brief.values(2)
    s = run(with_values(cw_brief, **ROOMY))
    per_wash = d["chem_car"] + d["water_gal"] / 1000 * d["water_rate"] + d["elec_car"]
    assert s["cogs"][0] == pytest.approx(-s["washes"][0] * per_wash - s["rev"][0] * d["card_pct"])


def test_tunnel_capacity_caps_total_washes_and_scales_both_groups(cw_brief):
    d = cw_brief.values(2)
    tight = run(with_values(cw_brief, cph=40))
    cap = 40 * d["hours_day"] * d["cap_util"] * d["days_open"]
    assert max(tight["washes"]) == pytest.approx(cap)
    assert tight["lost_washes"][-1] > 0
    roomy = run(with_values(cw_brief, cph=1e4))
    assert tight["rev"][-1] < roomy["rev"][-1]
    # a member pays the fee whatever the queue; only the retail side loses revenue
    assert tight["rev_mem"][-1] == pytest.approx(roomy["rev_mem"][-1], rel=0.2)
    share = tight["mem_washes"][-1] / tight["washes"][-1]
    share_roomy = roomy["mem_washes"][-1] / roomy["washes"][-1]
    assert share == pytest.approx(share_roomy, rel=0.05)


def test_retail_cars_grow_each_year(cw_brief):
    d = cw_brief.values(2)
    s = run(with_values(cw_brief, **ROOMY, ramp_months=0, cars_g=0.1), g="annual")
    assert s["ret_cars"][0] == pytest.approx(d["cars_day"] * d["days_open"] * 12)
    assert s["ret_cars"][1] == pytest.approx(s["ret_cars"][0] * 1.1)
    assert s["ret_cars"][2] == pytest.approx(s["ret_cars"][0] * 1.21)


def test_land_is_not_depreciated(cw_brief):
    b = with_values(cw_brief, capex_land=500_000, capex_refresh=0)
    no_land = with_values(b, capex_land=0)
    s, n = run(b, g="annual"), run(no_land, g="annual")
    assert s["capex"][0] - n["capex"][0] == pytest.approx(500_000)
    assert s["dep"] == pytest.approx(n["dep"])  # land adds no depreciation
    full = run(with_values(b, dep_years=1), g="annual")
    assert full["ppe_net"][-1] == pytest.approx(500_000)  # still carried at cost once the rest is written off


def test_winter_dip_moves_volume_but_keeps_the_year(cw_brief):
    b = with_values(cw_brief, cph=1e4, winter_dip=0.3, ramp_months=0)
    flat = run(with_values(cw_brief, **ROOMY, ramp_months=0))
    s = run(b)
    assert s["ret_cars"][0] < s["ret_cars"][5]  # January below June
    assert sum(s["ret_cars"][:12]) == pytest.approx(sum(flat["ret_cars"][:12]))


@pytest.mark.parametrize("stage", STAGES)
def test_monthly_and_annual_agree_where_volume_is_additive(stage):
    b = with_values(CW.default_brief(stage), **ROOMY)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    for k in ("ret_cars", "cap_washes", "payroll_total", "capex"):
        for y in (1, 2, 3):
            assert m[k][y] == pytest.approx(a[k][y], abs=0.01), (k, y)
    for y in (1, 2, 3):  # members: annual compounds churn and joins mid-year (same convention as fitness)
        assert a["members_close"][y] == pytest.approx(m["members_close"][y], rel=0.06)
        assert a["rev"][y] == pytest.approx(m["rev"][y], rel=0.08)


def test_annual_is_approximate_with_a_winter_dip_during_the_ramp(cw_brief):
    """Exact without seasonality (above); with a dip the ramp and the season interact, so only close."""
    b = with_values(cw_brief, cph=1e4, winter_dip=0.3, ramp_months=18)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    for y in (1, 2, 3):
        assert a["ret_cars"][y] == pytest.approx(m["ret_cars"][y], rel=0.05)
    assert a["ret_cars"][1] != pytest.approx(m["ret_cars"][1], rel=1e-6)


def test_capex_items_in_month_one_and_refresh_from_year_two(cw_brief):
    d = cw_brief.values(2)
    s = run(cw_brief)
    assert s["capex"][0] == pytest.approx(d["capex_equip"] + d["capex_site"] + d["capex_pay"])
    assert s["capex"][1] == 0 and s["capex"][11] == 0
    assert s["capex"][12] == pytest.approx(d["capex_refresh"])
    assert s["capex"][24] == pytest.approx(d["capex_refresh"])
    annual = run(cw_brief, g="annual")
    assert annual["capex"][1] == pytest.approx(d["capex_refresh"])


@pytest.mark.parametrize("g", ["monthly", "annual"])
@pytest.mark.parametrize("scen", [1, 2, 3])
def test_balance_sheet_balances(cw_brief, g, scen):
    s = run(cw_brief, scen, g)
    for ta, tle in zip(s["total_assets"], s["total_le"], strict=True):
        assert ta == pytest.approx(tle, abs=1e-6)


def test_degenerate_inputs_are_guarded(cw_brief):
    s = run(with_values(cw_brief, cars_day=0, members_start=0, mem_new=0, ramp_months=0, wash_pm=0))
    assert all(v == 0 for v in s["washes"]) and all(v == 0 for v in s["rev"])
    assert all(v >= 0 for v in s["members_close"])


# ---------------------------------------------------------------------- LibreOffice + challenger
lo = pytest.mark.libreoffice
COMBOS = list(itertools.product(["light", "full"], ["monthly", "annual"], [1, 3, 5]))


@lo
@pytest.mark.parametrize("tier,g,years", COMBOS, ids=[f"{t}-{g}-{y}y" for t, g, y in COMBOS])
def test_matrix_passes_the_challenger(build, cw_brief, tier, g, years):
    brief = cw_brief.with_options(tier=tier, granularity=g, horizon_years=years)
    path = build(brief)
    res = recalc(path)
    assert res.ok, res.errors
    vals = load_workbook(path, data_only=True)
    verdict = [r[1].value for r in vals["Checks"].iter_rows(min_col=1, max_col=2) if r[0].value == "Overall verdict"]
    assert verdict and verdict[0] in ("PASS", "PASS WITH WARNINGS")
    report = challenge(brief, path)
    assert report["problems"] == []
    assert report["verdict"] == "PASS"
    assert any(k.startswith("kpi:") for k in report["scenarios"][2]["years"][1])


@lo
@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_pass_the_challenger(build, stage):
    brief = CW.default_brief(stage)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
    assert report["verdict"].startswith("PASS")


@lo
def test_a_binding_tunnel_and_a_winter_dip_pass_the_challenger(build, cw_brief):
    brief = with_values(cw_brief, cph=60, winter_dip=0.3, capex_land=600_000, cars_g=0.05)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
