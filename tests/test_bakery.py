"""Bakery template: economics in the plain-Python twin, plus the LibreOffice + challenger matrix."""

import itertools

import pytest
from openpyxl import load_workbook

from engine.challenger import challenge, python_model
from engine.core import reference_run
from engine.recalc import recalc
from engine.spec import Timing
from engine.templates import get_template

BAKERY = get_template("bakery")
STAGES = ["planning", "early", "growing", "established"]


@pytest.fixture
def bk_brief():
    return BAKERY.default_brief("growing")


def run(brief, scen=2, g="monthly", years=3):
    return reference_run(brief.values(scen), brief.roles, Timing(g, years), BAKERY.reference_top)


def with_values(brief, **vals):
    d = brief.model_dump()
    for drv in d["drivers"]:
        if drv["key"] in vals:
            v = vals[drv["key"]]
            drv.update(low=v, base=v, high=v)
    return brief.model_validate(d)


def flat(brief, **vals):
    """No ramp, no holiday peak, no growth: one steady-state month repeated."""
    base = dict(
        open_pct=1,
        ramp_months=1,
        peak_uplift=0,
        traffic_g=0,
        price_g=0,
        ws_g=0,
        capex_repl=0,
        wage_g=0,
    )
    return with_values(brief, **{**base, **vals})


# --------------------------------------------------------------------------- stage defaults, scenarios
@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_are_valid_ordered_and_funded(stage):
    b = BAKERY.default_brief(stage)
    low, base, high = (python_model(b, s) for s in (1, 2, 3))
    for y in (1, 2, 3):
        assert low["rev"][y] < base["rev"][y] < high["rev"][y]
        assert 0.5 < base["gp"][y] / base["rev"][y] < 0.85
    assert min(base["cash_close"].values()) > 0  # the Base case is funded; Low is the stress case


@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_have_the_expected_roles(stage):
    b = BAKERY.default_brief(stage)
    assert [r.key for r in b.roles] == ["baker", "counter", "manager"]
    assert b.business_type == "bakery" and b.template == "bakery"


# --------------------------------------------------------------------------- revenue logic
def test_retail_revenue_is_customers_times_ticket_times_opening_days(bk_brief):
    d = bk_brief.values(2)
    s = run(flat(bk_brief, ws_accounts=0))
    days = d["open_days"] * 52 / 12
    assert s["rev_retail"][0] == pytest.approx(d["cust_full"] * d["ticket"] * days)
    assert s["rev"][0] == pytest.approx(s["rev_retail"][0])
    assert s["tx"][0] == pytest.approx(d["cust_full"] * days)


def test_opening_days_per_week_scale_retail_but_not_wholesale(bk_brief):
    six, five = run(flat(bk_brief, open_days=6)), run(flat(bk_brief, open_days=5))
    assert five["rev_retail"][0] == pytest.approx(six["rev_retail"][0] * 5 / 6)
    assert five["rev_ws"][0] == pytest.approx(six["rev_ws"][0])


def test_wholesale_revenue_is_accounts_times_weekly_order(bk_brief):
    d = bk_brief.values(2)
    s = run(flat(bk_brief))
    assert s["rev_ws"][0] == pytest.approx(d["ws_accounts"] * d["ws_order"] * 52 / 12)
    assert s["accounts"][0] == pytest.approx(d["ws_accounts"])
    none = run(flat(bk_brief, ws_accounts=0))
    assert max(none["rev_ws"]) == 0 and max(none["deliv"]) == 0
    assert s["rev"][0] == pytest.approx(s["rev_retail"][0] + s["rev_ws"][0])


def test_ramp_starts_at_opening_traffic_and_reaches_full(bk_brief):
    b = with_values(bk_brief, open_pct=0.4, ramp_months=6, peak_uplift=0, ws_accounts=0)
    s = run(b)
    full = s["tx"][11]
    assert s["tx"][0] == pytest.approx(full * 0.4)
    assert s["tx"][6] == pytest.approx(full)
    assert s["tx"][1] > s["tx"][0]


def test_holiday_peak_moves_sales_into_nov_dec_but_keeps_the_year(bk_brief):
    b = flat(bk_brief, ws_accounts=0)
    peak, even = run(with_values(b, peak_uplift=0.3)), run(b)
    assert peak["rev"][10] > even["rev"][10] > peak["rev"][0]  # November busy, January quiet
    assert sum(peak["rev"][:12]) == pytest.approx(sum(even["rev"][:12]))


def test_prices_and_traffic_grow_once_a_year(bk_brief):
    d = bk_brief.values(2)
    b = flat(bk_brief, ws_accounts=0, price_g=0.05, traffic_g=0.1)
    s = run(b)
    assert s["rev_retail"][12] / s["rev_retail"][0] == pytest.approx(1.05 * 1.1)
    assert s["rev_retail"][11] == pytest.approx(s["rev_retail"][0])
    assert d["ticket"] > 0


# --------------------------------------------------------------------------- cost of sales
def test_cost_of_sales_is_ingredients_waste_and_packaging(bk_brief):
    d = bk_brief.values(2)
    s = run(flat(bk_brief))
    rr, rw = s["rev_retail"][0], s["rev_ws"][0]
    ingredients = rr * d["ingr_pct"] + rw * d["ws_ingr_pct"]
    waste = rr * d["ingr_pct"] * d["waste_pct"] / (1 - d["waste_pct"])
    packaging = (rr + rw) * d["pack_pct"]
    assert -s["cogs"][0] == pytest.approx(ingredients + waste + packaging)
    assert s["waste"][0] == pytest.approx(waste)


def test_waste_is_unsold_share_of_what_was_baked(bk_brief):
    """Waste share w of units baked: baked = sold / (1-w), so unsold ingredient cost = sold cost × w/(1-w)."""
    b = flat(bk_brief, ws_accounts=0, pack_pct=0)
    none, tenth = run(with_values(b, waste_pct=0)), run(with_values(b, waste_pct=0.1))
    assert tenth["cogs"][0] / none["cogs"][0] == pytest.approx(1 + 0.1 / 0.9)
    assert tenth["rev"][0] == pytest.approx(none["rev"][0])  # waste is a cost, not lost revenue


def test_waste_hits_retail_only(bk_brief):
    b = flat(bk_brief, ingr_pct=0)  # no retail ingredient cost -> no waste cost either
    assert max(run(b)["waste"]) == 0


def test_more_waste_lowers_gross_profit(bk_brief):
    low_w, high_w = run(with_values(bk_brief, waste_pct=0.05)), run(with_values(bk_brief, waste_pct=0.2))
    assert sum(high_w["gp"]) < sum(low_w["gp"])


# --------------------------------------------------------------------------- operating costs and capex
def test_delivery_cost_is_per_drop(bk_brief):
    d = bk_brief.values(2)
    s = run(flat(bk_brief))
    assert s["deliv"][0] == pytest.approx(d["ws_accounts"] * d["ws_deliv_wk"] * 52 / 12 * d["cost_drop"])
    double = run(flat(bk_brief, ws_deliv_wk=2 * d["ws_deliv_wk"]))
    assert double["deliv"][0] == pytest.approx(2 * s["deliv"][0])


def test_bakers_and_counter_staff_are_payroll(bk_brief):
    d = bk_brief.values(2)
    s = run(flat(bk_brief))
    monthly = sum(d[f"{r}_fte"] * d[f"{r}_salary"] / 12 for r in ("baker", "counter", "manager"))
    assert -s["payroll_total"][0] == pytest.approx(monthly * (1 + d["payroll_tax"]))
    late = run(flat(bk_brief, baker_start=7))  # bakers hired in month 7 add no cost before
    assert late["payroll_total"][5] > s["payroll_total"][5]


def test_capex_is_fit_out_then_yearly_replacement(bk_brief):
    d = bk_brief.values(2)
    ann = run(bk_brief, g="annual")
    assert ann["capex"][0] == pytest.approx(d["capex"])
    assert ann["capex"][1] == pytest.approx(d["capex_repl"])
    mon = run(bk_brief)
    assert mon["capex"][0] == pytest.approx(d["capex"])
    assert sum(mon["capex"][12:24]) == pytest.approx(d["capex_repl"])
    assert sum(mon["capex"][:12]) == pytest.approx(d["capex"])


# --------------------------------------------------------------------------- consistency
@pytest.mark.parametrize("stage", STAGES)
def test_monthly_equals_annual_without_season(stage):
    """The annual ramp factor is the exact average of the months, so only the holiday peak separates them."""
    b = with_values(BAKERY.default_brief(stage), peak_uplift=0)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    for k in ("rev", "gp", "ebitda", "payroll_total", "rev_retail", "rev_ws", "tx", "waste", "deliv", "capex"):
        for y in (1, 2, 3):
            assert m[k][y] == pytest.approx(a[k][y], rel=1e-9, abs=0.01), (k, y)


@pytest.mark.parametrize("stage", STAGES)
def test_monthly_equals_annual_with_season(stage):
    """The peak redistributes within a year (average factor 1), but the ramp makes it not exactly neutral."""
    b = BAKERY.default_brief(stage)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    for y in (2, 3):
        assert m["rev"][y] == pytest.approx(a["rev"][y], rel=0.002)


@pytest.mark.parametrize("g", ["monthly", "annual"])
@pytest.mark.parametrize("scen", [1, 2, 3])
def test_balance_sheet_balances(bk_brief, g, scen):
    s = run(bk_brief, scen, g)
    for ta, tle in zip(s["total_assets"], s["total_le"], strict=True):
        assert ta == pytest.approx(tle, abs=1e-6)


def test_waste_at_the_limit_is_guarded(bk_brief):
    s = run(with_values(bk_brief, waste_pct=0.5))
    assert all(v >= 0 for v in s["waste"])


# --------------------------------------------------------------------------- LibreOffice + challenger
COMBOS = list(itertools.product(["light", "full"], ["monthly", "annual"], [1, 3, 5]))


@pytest.mark.libreoffice
@pytest.mark.parametrize("tier,g,years", COMBOS, ids=[f"{t}-{g}-{y}y" for t, g, y in COMBOS])
def test_matrix_passes_the_challenger(build, tier, g, years):
    brief = BAKERY.default_brief("early").with_options(tier=tier, granularity=g, horizon_years=years)
    path = build(brief)
    res = recalc(path)
    assert res.ok, res.errors
    vals = load_workbook(path, data_only=True)
    verdict = [r[1].value for r in vals["Checks"].iter_rows(min_col=1, max_col=2) if r[0].value == "Overall verdict"]
    assert verdict and verdict[0] in ("PASS", "PASS WITH WARNINGS")
    report = challenge(brief, path)
    assert report["problems"] == []
    assert report["verdict"] == "PASS"


@pytest.mark.libreoffice
@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_pass_the_challenger_and_show_kpis(build, stage):
    brief = BAKERY.default_brief(stage)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
    kpis = {k for k in report["scenarios"][2]["years"][1] if k.startswith("kpi:")}
    assert {f"kpi:{k}" for k in BAKERY.kpi_checked} <= kpis
