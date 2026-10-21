"""Scenarios sheet (E29, docs/WORKBOOK_STANDARD.md §1): Low / Base / High chains, live, challenger-covered."""

import shutil

import pytest
from openpyxl import load_workbook

from engine.builder import plan
from engine.challenger import challenge
from engine.goals import apply_goal
from engine.recalc import recalc
from engine.scenarios import CALC, COLS, SHEET, reference
from engine.templates import TEMPLATES


def _brief(key, **kw):
    tpl = TEMPLATES[key]
    return apply_goal(tpl.default_brief(tpl.sample_stage).with_options(**kw))


def _check(path, name):
    ws = load_workbook(path, data_only=True)["Checks"]
    return next(r[1].value for r in ws.iter_rows(min_col=1, max_col=2) if r[0].value == name)


def test_full_only_and_order():
    full = plan(_brief("cafe", tier="full"), TEMPLATES["cafe"])
    assert full.sheets[full.sheets.index("Summary") + 1] == SHEET
    assert full.sheets[full.sheets.index("Headcount") + 1] == CALC
    light = plan(_brief("cafe", tier="light"), TEMPLATES["cafe"])
    assert light.scen is None and SHEET not in light.sheets and CALC not in light.sheets


def test_chains_read_their_own_inputs_column():
    pl = plan(_brief("saas", tier="full"), TEMPLATES["saas"])
    for s, g in pl.scen.grids.items():
        assert g.input_col == COLS[s]
        rev = pl.lines["rev"]
        f = pl.render(rev, 2, CALC, g)
        assert "Inputs!$E$" not in f  # never the active column
        assert set(g.line_at.values()) == {(CALC, r) for r in g.rows.values()}
    tops = [g.start_row for g in pl.scen.grids.values()]
    ends = [g.end_row for g in pl.scen.grids.values()]
    assert tops[1] > ends[0] and tops[2] > ends[1]  # stacked, no overlap


def test_python_twin_matches_the_scenario():
    b = _brief("cafe", tier="full")
    lo, hi = reference(b, 1)["cells"], reference(b, 3)["cells"]
    assert lo["rev_total"] < reference(b, 2)["cells"]["rev_total"] < hi["rev_total"]


@pytest.mark.libreoffice
@pytest.mark.parametrize("key", sorted(TEMPLATES))
@pytest.mark.parametrize("g,years", [("monthly", 5), ("annual", 3)])
def test_every_template_passes(build, key, g, years):
    b = _brief(key, tier="full", granularity=g, horizon_years=years)
    path = build(b)
    rep = challenge(b, path)
    assert rep["problems"] == [], rep["problems"][:3]
    assert recalc(path).ok
    assert _check(path, "Scenarios: Active = selected scenario") == "PASS"


@pytest.mark.libreoffice
@pytest.mark.parametrize("sel", [1, 2, 3])
def test_active_equals_the_selected_column(build, tmp_path, sel):
    b = _brief("real_estate", tier="full")
    path = build(b)
    pl = plan(b, TEMPLATES["real_estate"])
    edited = tmp_path / f"s{sel}.xlsx"
    shutil.copy(path, edited)
    wb = load_workbook(edited)
    wb["Dashboard"]["B4"] = sel
    wb.save(edited)
    assert recalc(edited).ok
    ws = load_workbook(edited, data_only=True)[SHEET]
    for name, r in pl.scen.outputs.items():
        assert ws[f"E{r}"].value == pytest.approx(ws[f"{COLS[sel]}{r}"].value, abs=1e-6), name
    assert _check(edited, "Scenarios: Active = selected scenario") == "PASS"


@pytest.mark.libreoffice
def test_an_input_edit_updates_all_three_columns(build, tmp_path):
    b = _brief("saas", tier="full")
    path = build(b)
    pl = plan(b, TEMPLATES["saas"])
    assert recalc(path).ok
    before = load_workbook(path, data_only=True)
    row = pl.input_row["arpa"]
    edited = tmp_path / "edited.xlsx"
    shutil.copy(path, edited)
    wb = load_workbook(edited)
    for col in "BCD":
        wb["Inputs"][f"{col}{row}"] = wb["Inputs"][f"{col}{row}"].value * 1.2
    wb.save(edited)
    assert recalc(edited).ok
    after = load_workbook(edited, data_only=True)
    r = pl.scen.outputs["rev_total"]
    for col in "BCDE":
        assert after[SHEET][f"{col}{r}"].value > before[SHEET][f"{col}{r}"].value * 1.05, col
    drivers = [
        d.model_copy(update={"low": d.low * 1.2, "base": d.base * 1.2, "high": d.high * 1.2}) if d.key == "arpa" else d
        for d in b.drivers
    ]
    rep = challenge(b.model_copy(update={"drivers": drivers}), edited)
    assert rep["problems"] == [], rep["problems"][:3]


@pytest.mark.libreoffice
def test_tampered_chain_is_caught(build):
    b = _brief("cafe", tier="full")
    path = build(b)
    pl = plan(b, TEMPLATES["cafe"])
    g = pl.scen.grids[1]  # Low: never the selected scenario by default, so only the challenger can see it
    wb = load_workbook(path)
    ref = f"E{g.rows['rev']}"
    wb[CALC][ref] = wb[CALC][ref].value + "*1.1"
    wb.save(path)
    rep = challenge(b, path)
    assert rep["verdict"] == "FAIL"
    assert any("scenarios rev" in p for p in rep["problems"])


@pytest.mark.libreoffice
def test_four_charts_survive_recalc(build):
    path = build(_brief("fitness", tier="full"))
    assert recalc(path).ok
    ws = load_workbook(path)[SHEET]
    assert len(ws._charts) == 4
    assert sorted((c.anchor._from.col, c.anchor._from.row) for c in ws._charts) == [(7, 4), (7, 21), (15, 4), (15, 21)]
