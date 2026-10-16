import copy

import pytest
from openpyxl import load_workbook

from engine.challenger import challenge, python_model

pytestmark = pytest.mark.libreoffice


def test_cafe_passes_all_three_scenarios(cafe_brief, cafe_xlsx):
    report = challenge(cafe_brief, cafe_xlsx)
    assert report["problems"] == []
    assert report["verdict"] == "PASS"
    assert set(report["scenarios"]) == {1, 2, 3}
    assert [report["scenarios"][s]["label"] for s in (1, 2, 3)] == ["Low", "Base", "High"]
    years = cafe_brief["horizon_years"]
    for s in (1, 2, 3):
        assert len(report["scenarios"][s]["years"]) == years  # 3 scenarios x 3 years = 9 matches


def test_wrong_formula_is_caught(cafe_brief, cafe_xlsx):
    wb = load_workbook(cafe_xlsx)
    ws = wb["Monthly"]
    ws["C10"] = ws["C10"].value + "*1.01"  # one month of revenue off by 1%
    wb.save(cafe_xlsx)
    report = challenge(cafe_brief, cafe_xlsx)
    assert report["verdict"] == "FAIL"
    assert any("year 1 rev" in p for p in report["problems"])


def test_brief_mismatch_is_caught(cafe_brief, cafe_xlsx):
    other = copy.deepcopy(cafe_brief)
    for d in other["drivers"]:
        if d["key"] == "ticket":
            d["base"] += 0.5
    report = challenge(other, cafe_xlsx)
    assert report["verdict"] == "FAIL"
    assert any(p.startswith("scenario 2") for p in report["problems"])


def test_formula_error_fails(cafe_brief, cafe_xlsx):
    wb = load_workbook(cafe_xlsx)
    wb["Monthly"]["C8"] = "=1/0"
    wb.save(cafe_xlsx)
    report = challenge(cafe_brief, cafe_xlsx)
    assert report["verdict"] == "FAIL"
    assert any("recalc" in p for p in report["problems"])


def test_python_model_orders_scenarios(cafe_brief):
    low, base, high = (python_model(cafe_brief, s) for s in (1, 2, 3))
    for y in range(1, cafe_brief["horizon_years"] + 1):
        assert low[y]["rev"] < base[y]["rev"] < high[y]["rev"]
