"""E16: content/seo pages rendered with a shared layout, sample workbooks, sitemap and SEO checks."""

import json
import re
from html.parser import HTMLParser
from pathlib import Path

import openpyxl
import pytest

from app import seo
from engine.samples import SampleSpec, parse_description, parse_name

ROOT = Path(__file__).resolve().parents[1]
PAGES = list(seo.pages().values())
ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")


class _Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.hrefs: list[str] = []
        self.ids: set[str] = set()
        self.headings: list[int] = []
        self.imgs_without_alt = 0

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "a" and a.get("href"):
            self.hrefs.append(a["href"])
        if a.get("id"):
            self.ids.add(a["id"])
        if re.fullmatch(r"h[1-6]", tag):
            self.headings.append(int(tag[1]))
        if tag == "img" and not a.get("alt"):
            self.imgs_without_alt += 1


def parse(html: str) -> _Links:
    p = _Links()
    p.feed(html)
    return p


def test_every_content_file_is_a_page():
    assert len(PAGES) == len(list((ROOT / "content" / "seo").glob("*.md"))) >= 15
    assert len({p.slug for p in PAGES}) == len(PAGES)


@pytest.mark.parametrize("page", PAGES, ids=lambda p: p.name)
def test_page_renders_with_seo_basics(client, page):
    r = client.get(page.slug)
    assert r.status_code == 200
    html = r.text
    doc = parse(html)
    # one h1 and no skipped heading levels
    assert doc.headings.count(1) == 1
    for prev, cur in zip(doc.headings, doc.headings[1:], strict=False):
        assert cur <= prev + 1, (page.slug, doc.headings)
    assert f"<title>{page.title.replace('&', '&amp;')}</title>" in html
    assert 30 <= len(page.title) <= 90 and page.title.endswith("| StartupModel")
    assert 70 <= len(page.description) <= 200
    assert f'<link rel="canonical" href="http://testserver{page.slug}">' in html
    assert 'name="description"' in html and 'property="og:title"' in html and "noindex" not in html
    assert doc.imgs_without_alt == 0
    # body text survived rendering, no raw markdown or placeholders left
    assert not re.search(r"\[[A-Z][A-Z_]{3,}\]", html)
    assert "{#" not in html and "](/" not in html and "**" not in html
    assert "FinModelsLab" not in html and "Ledgerline" not in html
    assert 800 <= page.words <= 1600


@pytest.mark.parametrize("page", PAGES, ids=lambda p: p.name)
def test_internal_links_and_anchors_resolve(client, page):
    doc = parse(client.get(page.slug).text)
    for href in doc.hrefs:
        if href.startswith("mailto:") or href.startswith("http"):
            continue
        if href.startswith("#"):
            assert href[1:] in doc.ids, (page.slug, href)
            continue
        path = href.split("#")[0].split("?")[0]
        r = client.get(path)
        assert r.status_code == 200, (page.slug, href)
        if "#" in href:
            assert href.split("#")[1] in parse(r.text).ids, (page.slug, href)


@pytest.mark.parametrize("page", PAGES, ids=lambda p: p.name)
def test_structured_data(client, services, page):
    html = client.get(page.slug).text
    blocks = [json.loads(b) for b in re.findall(r'<script type="application/ld\+json">(.*?)</script>', html, re.S)]
    types = {b["@type"] for b in blocks}
    assert "BreadcrumbList" in types
    crumbs = next(b for b in blocks if b["@type"] == "BreadcrumbList")["itemListElement"]
    assert crumbs[-1]["item"] == "http://testserver" + page.slug
    if page.section == "guides":
        article = next(b for b in blocks if b["@type"] == "Article")
        assert article["headline"] == page.h1
    else:
        product = next(b for b in blocks if b["@type"] == "Product")
        assert {o["price"] for o in product["offers"]} == {
            f"{services.cfg.prices[t] / 100:.2f}" for t in ("light", "full")
        }


@pytest.mark.parametrize("page", PAGES, ids=lambda p: p.name)
def test_cta_and_sample_download(client, page):
    html = client.get(page.slug).text
    assert 'class="cta' in html and seo.start_link(page) in html
    assert page.sample is not None, f"{page.slug}: every page offers an engine-built sample"
    href = f"/samples/{page.sample.filename}"
    assert f'href="{href}"' in html
    r = client.get(href)
    assert r.status_code == 200 and r.content[:2] == b"PK"


def test_cta_opens_the_wizard_on_the_page_template(client):
    page = seo.pages()["/models/saas-financial-model-with-scenarios"]
    assert seo.start_link(page) == "/start?type=saas"
    html = client.get(seo.start_link(page)).text
    assert re.search(r'name="business_type" value="saas"\s+checked', html)
    assert seo.start_link(seo.pages()["/models/custom-financial-model"]) == "/start"


def test_copy_prices_match_the_configured_prices(services):
    """Pages quote tier prices in prose; they must equal the prices the checkout charges."""
    light, full = (f"${services.cfg.prices[t] / 100:,.0f}" for t in ("light", "full"))
    found = 0
    for page in PAGES:
        text = re.sub(r"<[^>]+>", "", page.body)
        for tier, price in re.findall(r"\b(Light|Full)[\s:(]{1,3}(\$[\d,]+)", text):
            assert price == {"Light": light, "Full": full}[tier], (page.slug, tier, price)
            found += 1
        for price in re.findall(r"\b[Ff]rom (\$[\d,]+)", text):
            assert price == light, (page.slug, price)
            found += 1
    assert found >= 15


def test_related_links_prefer_the_same_template():
    page = seo.pages()["/guides/restaurant-pro-forma"]
    rel = seo.related(page)
    assert rel and all(p.slug != page.slug for p in rel)
    assert rel[0].template == "cafe"


def test_indexes_sitemap_and_llms_list_every_page(client):
    sm = client.get("/sitemap.xml").text
    llms = client.get("/llms.txt").text
    for section in seo.SECTIONS:
        idx = client.get(f"/{section}")
        assert idx.status_code == 200 and f"<loc>http://testserver/{section}</loc>" in sm
        for page in PAGES:
            if page.section == section:
                assert f'href="{page.slug}"' in idx.text
    for page in PAGES:
        assert f"<loc>http://testserver{page.slug}</loc>" in sm
        assert f"http://testserver{page.slug}" in llms
    footer = client.get("/").text
    assert 'href="/models"' in footer and 'href="/guides"' in footer


def test_unknown_page_is_404(client):
    assert client.get("/models/nope").status_code == 404
    assert client.get("/guides/custom-financial-model").status_code == 404  # wrong section


def test_markdown_rendering_is_safe_and_anchored():
    html, h1 = seo.render_markdown(
        "# Title\n\n## Sample {#sample}\n\n## Same\n\n## Same\n\n<script>x</script>\n\n"
        "[bad](javascript:alert(1))\n\n| a | b |\n|---|---|\n| 1 | 2 |\n"
    )
    assert h1 == "Title"
    assert '<h2 id="sample">Sample</h2>' in html
    assert '<h2 id="same">' in html and '<h2 id="same-2">' in html
    assert "<script>" not in html and "&lt;script&gt;" in html
    assert 'href="javascript:' not in html
    assert '<div class="table-wrap" role="region" aria-label="Table" tabindex="0"><table>' in html


def test_front_matter_errors_are_loud(tmp_path):
    bad = tmp_path / "x.md"
    bad.write_text("---\nslug: /blog/x\ntitle: t\ndescription: d\nkeyword: k\n---\n# H\n", encoding="utf-8")
    with pytest.raises(ValueError, match="slug"):
        seo.load_page(bad)
    bad.write_text("no front matter", encoding="utf-8")
    with pytest.raises(ValueError, match="front matter"):
        seo.load_page(bad)
    bad.write_text(
        "---\nslug: /guides/x\ntitle: t\ndescription: d\nkeyword: k\nsample: cafe, huge, monthly, 3 years\n---\n# H\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="sample"):
        seo.load_page(bad)


def test_sample_names_round_trip():
    for spec in (
        SampleSpec("cafe"),
        SampleSpec("cafe", "light"),
        SampleSpec("cafe", "full", "monthly", 5),
        SampleSpec("saas", "light", "annual", 1),
    ):
        assert parse_name(spec.name) == spec
    assert SampleSpec("cafe", "full", "monthly", 5).name == "cafe-full-5y"
    for bad in ("cafe", "cafe-full-3y", "nope-full", "cafe-full-2y", "cafe-investor_pitch", "../cafe-full"):
        assert parse_name(bad) is None, bad
    assert parse_description("cafe, light, monthly, 3 years") == SampleSpec("cafe", "light")
    assert parse_description("retail, full, annual, 1 year") == SampleSpec("retail", "full", "annual", 1)
    assert parse_description("cafe, full, weekly, 3 years") is None


def test_named_samples_are_served_and_review_files_are_not(client):
    r = client.get("/samples/cafe-full-5y.xlsx")
    assert r.status_code == 200 and "startupmodel-sample-cafe-full-5y.xlsx" in r.headers["content-disposition"]
    assert client.get("/samples/cafe.xlsx").status_code == 200  # landing alias of cafe-full
    assert client.get("/samples/cafe-investor_pitch.xlsx").status_code == 404
    assert client.get("/samples/saas-light.xlsx").status_code == 404  # valid name, file not built


@pytest.mark.parametrize("spec", sorted({p.sample for p in PAGES}, key=lambda s: s.name), ids=lambda s: s.name)
def test_committed_samples_match_their_spec_and_are_clean(spec):
    """The file offered on a page is the workbook its name promises, with recalculated values and no errors."""
    path = ROOT / "samples" / spec.filename
    wb = openpyxl.load_workbook(path, data_only=True)
    names = wb.sheetnames
    if spec.tier == "light":
        assert names == ["Inputs", "P&L", "Checks"]
    else:
        assert {"Dashboard", "Inputs", "P&L", "Cash Flow", "Balance Sheet", "Checks"} <= set(names)
    periods = spec.horizon_years * (12 if spec.granularity == "monthly" else 1)
    detail = wb["Monthly"] if spec.tier == "full" and spec.granularity == "monthly" else wb["P&L"]
    assert detail.max_column >= periods
    for ws in wb:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                assert not (isinstance(v, str) and v in ERRORS), (spec.name, ws.title, v)
    checks = [r[1] for r in wb["Checks"].iter_rows(min_row=1, values_only=True) if isinstance(r[1], str)]
    assert any(c.startswith("PASS") for c in checks)
    assert not any(c.startswith("FAIL") for c in checks), checks
