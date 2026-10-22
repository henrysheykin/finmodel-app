"""Auto repair template: economics in the plain-Python twin, plus the LibreOffice + challenger matrix."""

import itertools

import pytest
from openpyxl import load_workbook

from engine.brief import BriefError
from engine.challenger import challenge, python_model
from engine.core import reference_run
from engine.recalc import recalc
from engine.spec import Timing
from engine.templates import get_template
from engine.templates.auto_repair import _ramp

AR = get_template("auto_repair")
STAGES = ["planning", "early", "growing", "established"]
#: no ramp-up, plenty of bays and technicians: demand is always served
ROOMY = {"bays": 100, "tech_eff": 3, "ramp_months": 0}


@pytest.fixture
def ar_brief():
    return AR.default_brief("early")


def run(brief, scen=2, g="monthly", years=3):
    return reference_run(brief.values(scen), brief.roles, Timing(g, years), AR.reference_top)


def with_values(brief, **vals):
    d = brief.model_dump()
    for drv in d["drivers"]:
        if drv["key"] in vals:
            v = vals[drv["key"]]
            drv.update(low=v, base=v, high=v)
    return brief.model_validate(d)


def test_registered_with_wizard_hooks():
    assert AR.order == 120 and AR.business_types == ("auto_repair",)
    assert "equity" in AR.key_numbers and 5 <= len(AR.key_numbers) <= 8
    assert AR.dashboard_metrics(Timing("monthly", 3)) and AR.dashboard_chart(Timing("monthly", 3))


@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_are_valid_ordered_and_funded(stage):
    b = AR.default_brief(stage)
    low, base, high = (python_model(b, s) for s in (1, 2, 3))
    for y in (1, 2, 3):
        assert low["rev"][y] < base["rev"][y] < high["rev"][y]
        assert 0 < base["gp"][y] / base["rev"][y] < 1
    monthly = reference_run(b.values(2), b.roles, Timing("monthly", 3), AR.reference_top)
    assert min(monthly["cash_close"]) > 0  # the Base case is funded in every month; Low is the stress case


def test_needs_the_technician_role(ar_brief):
    no_tech = ar_brief.model_copy(update={"roles": [r for r in ar_brief.roles if r.key != "tech"]})
    with pytest.raises(BriefError, match="tech"):
        AR.validate(no_tech)


def test_ramp_up_is_linear_and_annual_is_the_mean_of_its_months(ar_brief):
    d = {**ar_brief.values(2), "ramp_months": 18, "ramp_start": 0.4}
    m, a = Timing("monthly", 3), Timing("annual", 3)
    factors = [_ramp(d, p, m) for p in range(1, 37)]
    assert factors[0] == pytest.approx(0.4)
    assert factors[9] == pytest.approx(0.4 + 0.6 * 9 / 18)
    assert factors[18:] == pytest.approx([1.0] * 18)
    for y in (1, 2, 3):
        assert _ramp(d, y, a) == pytest.approx(sum(factors[(y - 1) * 12 : y * 12]) / 12)
    assert _ramp({**d, "ramp_months": 0}, 1, m) == 1.0


def test_billed_hours_follow_orders_hours_per_order_and_ramp(ar_brief):
    d = ar_brief.values(2)
    s = run(with_values(ar_brief, **{**ROOMY, "ramp_months": 12}))
    assert s["billed"][0] == pytest.approx(d["ro_day"] * d["hrs_ro"] * d["days_open"] * d["ramp_start"])
    assert s["lost_hours"][0] == pytest.approx(0)
    assert s["ro"][0] == pytest.approx(d["ro_day"] * d["days_open"] * d["ramp_start"])


def test_revenue_is_labor_plus_parts_plus_supplies_fee(ar_brief):
    d = ar_brief.values(2)
    s = run(with_values(ar_brief, **ROOMY))
    labor = s["billed"][0] * d["labor_rate"]
    expected = labor * (1 + d["supplies_pct"]) + s["billed"][0] * d["parts_hr"]
    assert s["rev"][0] == pytest.approx(expected)
    assert s["rev_labor"][0] == pytest.approx(labor)


def test_parts_margin_sets_parts_cost(ar_brief):
    d = ar_brief.values(2)
    s = run(with_values(ar_brief, **ROOMY))
    assert s["parts_cost"][0] == pytest.approx(-s["rev_parts"][0] * (1 - d["parts_margin"]))
    better = run(with_values(ar_brief, **ROOMY, parts_margin=0.6))
    assert better["rev"][0] == pytest.approx(s["rev"][0])  # the margin moves cost, not revenue
    assert better["gp"][0] - s["gp"][0] == pytest.approx(s["rev_parts"][0] * (0.6 - d["parts_margin"]))


def test_cost_of_sales_is_tech_payroll_parts_supplies_and_comebacks(ar_brief):
    d = ar_brief.values(2)
    s = run(with_values(ar_brief, **ROOMY))
    tech = d["tech_fte"] * d["tech_salary"] / 12 * (1 + d["payroll_tax"])
    expected = (
        -tech
        - s["rev_parts"][0] * (1 - d["parts_margin"])
        - s["rev_labor"][0] * (d["supplies_cost_pct"] + d["warranty_pct"])
    )
    assert s["cogs"][0] == pytest.approx(expected)


def test_bay_capacity_caps_billed_hours(ar_brief):
    d = ar_brief.values(2)
    b = with_values(ar_brief, bays=2, tech_eff=3, ramp_months=0)
    s = run(b)
    cap = 2 * d["bay_hours"] * d["days_open"]
    assert max(s["billed"]) == pytest.approx(cap)
    assert s["cap_hours"][5] == pytest.approx(cap)  # bays are the lower limit
    assert s["lost_hours"][5] > 0
    roomy = run(with_values(ar_brief, **ROOMY))
    assert s["rev"][5] < roomy["rev"][5]


def test_technician_capacity_caps_billed_hours_when_below_the_bays(ar_brief):
    d = ar_brief.values(2)
    s = run(with_values(ar_brief, bays=100, ro_day=100, ramp_months=0))
    cap = d["tech_fte"] * d["tech_hours"] * d["tech_eff"] * d["days_open"]
    assert s["billed"][5] == pytest.approx(cap)
    more = run(with_values(ar_brief, bays=100, ro_day=100, ramp_months=0, tech_eff=d["tech_eff"] * 1.1))
    assert more["billed"][5] == pytest.approx(cap * 1.1)


def test_orders_grow_each_year_and_prices_rise(ar_brief):
    d = ar_brief.values(2)
    s = run(with_values(ar_brief, **ROOMY, ro_g=0.1), g="annual")
    assert s["ro"][0] == pytest.approx(d["ro_day"] * d["days_open"] * 12)
    assert s["ro"][1] == pytest.approx(s["ro"][0] * 1.1)
    assert s["ro"][2] == pytest.approx(s["ro"][0] * 1.21)
    assert s["rev_labor"][1] / s["billed"][1] == pytest.approx(d["labor_rate"] * (1 + d["price_g"]))


def test_land_is_not_depreciated(ar_brief):
    b = with_values(ar_brief, capex_land=500_000, capex_refresh=0)
    no_land = with_values(b, capex_land=0)
    s, n = run(b, g="annual"), run(no_land, g="annual")
    assert s["capex"][0] - n["capex"][0] == pytest.approx(500_000)
    assert s["dep"] == pytest.approx(n["dep"])
    full = run(with_values(b, dep_years=1), g="annual")
    assert full["ppe_net"][-1] == pytest.approx(500_000)


@pytest.mark.parametrize("stage", STAGES)
def test_monthly_and_annual_agree_where_volume_is_additive(stage):
    b = with_values(AR.default_brief(stage), **ROOMY)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    for k in ("billed", "ro", "rev", "cogs", "payroll_total", "capex"):
        for y in (1, 2, 3):
            assert m[k][y] == pytest.approx(a[k][y], abs=0.01), (k, y)


@pytest.mark.parametrize("ramp", [12, 18])
def test_monthly_and_annual_agree_through_the_ramp(ramp):
    b = with_values(AR.default_brief("planning"), bays=100, tech_eff=3, ramp_months=ramp)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    assert m["billed"][1] < m["billed"][2]
    for k in ("billed", "ro", "rev", "cogs"):
        for y in (1, 2, 3):
            assert m[k][y] == pytest.approx(a[k][y], abs=0.01), (k, y)


def test_capex_scales_with_bays_in_month_one_and_refresh_from_year_two(ar_brief):
    d = ar_brief.values(2)
    s = run(ar_brief)
    assert s["capex"][0] == pytest.approx(d["bays"] * d["capex_bay"] + d["capex_diag"] + d["capex_fit"])
    assert s["capex"][1] == 0 and s["capex"][11] == 0
    assert s["capex"][12] == pytest.approx(d["capex_refresh"])
    assert run(ar_brief, g="annual")["capex"][1] == pytest.approx(d["capex_refresh"])


@pytest.mark.parametrize("g", ["monthly", "annual"])
@pytest.mark.parametrize("scen", [1, 2, 3])
def test_balance_sheet_balances(ar_brief, g, scen):
    s = run(ar_brief, scen, g)
    for ta, tle in zip(s["total_assets"], s["total_le"], strict=True):
        assert ta == pytest.approx(tle, abs=1e-6)


def test_degenerate_inputs_are_guarded(ar_brief):
    s = run(with_values(ar_brief, ro_day=0, ramp_months=0))
    assert all(v == 0 for v in s["billed"]) and all(v == 0 for v in s["rev"])
    assert all(v == 0 for v in s["ro"])


# ---------------------------------------------------------------------- LibreOffice + challenger
lo = pytest.mark.libreoffice
COMBOS = list(itertools.product(["light", "full"], ["monthly", "annual"], [1, 3, 5]))


@lo
@pytest.mark.parametrize("tier,g,years", COMBOS, ids=[f"{t}-{g}-{y}y" for t, g, y in COMBOS])
def test_matrix_passes_the_challenger(build, ar_brief, tier, g, years):
    brief = ar_brief.with_options(tier=tier, granularity=g, horizon_years=years)
    path = build(brief)
    res = recalc(path)
    assert res.ok, res.errors
    vals = load_workbook(path, data_only=True)
    verdict = [r[1].value for r in vals["Checks"].iter_rows(min_col=1, max_col=2) if r[0].value == "Overall verdict"]
    assert verdict and verdict[0] in ("PASS", "PASS WITH WARNINGS")
    report = challenge(brief, path)
    assert report["problems"] == []
    assert report["verdict"] == "PASS"
    assert any(k.startswith("kpi:") for k in report["scenarios"][2]["years"][1])


@lo
@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_pass_the_challenger(build, stage):
    brief = AR.default_brief(stage)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
    assert report["verdict"].startswith("PASS")


@lo
def test_a_binding_bay_limit_and_land_pass_the_challenger(build, ar_brief):
    brief = with_values(ar_brief, bays=3, capex_land=600_000, ro_g=0.05)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
