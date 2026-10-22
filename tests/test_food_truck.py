"""Food truck template: economics in the plain-Python twin, plus the LibreOffice + challenger matrix."""

import itertools

import pytest
from openpyxl import load_workbook

from engine.challenger import challenge, python_model
from engine.core import reference_run
from engine.recalc import recalc
from engine.spec import Timing
from engine.templates import get_template
from engine.templates.food_truck import _ramp

FT = get_template("food_truck")
STAGES = ["planning", "early", "growing", "established"]
#: no winter dip: the volume path is then exactly additive over months
FLAT = {"winter_dip": 0}


@pytest.fixture
def ft_brief():
    return FT.default_brief("early")


def run(brief, scen=2, g="monthly", years=3):
    return reference_run(brief.values(scen), brief.roles, Timing(g, years), FT.reference_top)


def with_values(brief, **vals):
    d = brief.model_dump()
    for drv in d["drivers"]:
        if drv["key"] in vals:
            v = vals[drv["key"]]
            drv.update(low=v, base=v, high=v)
    return brief.model_validate(d)


def test_registered_with_wizard_hooks():
    assert FT.order == 150 and FT.business_types == ("food_truck",)
    assert "equity" in FT.key_numbers and 5 <= len(FT.key_numbers) <= 8
    assert FT.dashboard_metrics(Timing("monthly", 3)) and FT.dashboard_chart(Timing("monthly", 3))


@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_are_valid_ordered_and_funded(stage):
    b = FT.default_brief(stage)
    low, base, high = (python_model(b, s) for s in (1, 2, 3))
    for y in (1, 2, 3):
        assert low["rev"][y] < base["rev"][y] < high["rev"][y]
        assert 0 < base["gp"][y] / base["rev"][y] < 1
    monthly = reference_run(b.values(2), b.roles, Timing("monthly", 3), FT.reference_top)
    assert min(monthly["cash_close"]) > 0  # the Base case is funded in every month; Low is the stress case


def test_ramp_up_is_linear_and_annual_is_the_mean_of_its_months(ft_brief):
    d = {**ft_brief.values(2), "ramp_months": 18, "ramp_start": 0.4}
    m, a = Timing("monthly", 3), Timing("annual", 3)
    factors = [_ramp(d, p, m) for p in range(1, 37)]
    assert factors[0] == pytest.approx(0.4)
    assert factors[9] == pytest.approx(0.4 + 0.6 * 9 / 18)
    assert factors[18:] == pytest.approx([1.0] * 18)
    for y in (1, 2, 3):
        assert _ramp(d, y, a) == pytest.approx(sum(factors[(y - 1) * 12 : y * 12]) / 12)
    assert _ramp({**d, "ramp_months": 0}, 1, m) == 1.0


def test_street_revenue_is_days_times_meals_times_ticket(ft_brief):
    d = ft_brief.values(2)
    s = run(with_values(ft_brief, **FLAT, ramp_months=0))
    days = d["trucks"] * d["street_days"]
    assert s["street_days"][0] == pytest.approx(days)
    assert s["meals_street"][0] == pytest.approx(days * d["meals_day"])
    assert s["rev_street"][0] == pytest.approx(days * d["meals_day"] * d["ticket"])


def test_event_revenue_is_event_days_times_revenue_per_event(ft_brief):
    d = ft_brief.values(2)
    s = run(with_values(ft_brief, **FLAT, ramp_months=0))
    assert s["event_days"][0] == pytest.approx(d["trucks"] * d["events_pm"])
    assert s["rev_event"][0] == pytest.approx(d["trucks"] * d["events_pm"] * d["event_rev"])
    assert s["rev"][0] == pytest.approx(s["rev_street"][0] + s["rev_event"][0])
    no_events = run(with_values(ft_brief, **FLAT, ramp_months=0, events_pm=0))
    assert all(v == 0 for v in no_events["rev_event"]) and all(v == 0 for v in no_events["event_fees"])


def test_ramp_scales_meals_per_day_and_event_bookings_not_street_days(ft_brief):
    d = ft_brief.values(2)
    s = run(with_values(ft_brief, **FLAT, ramp_months=12, ramp_start=0.5))
    assert s["street_days"][0] == pytest.approx(d["trucks"] * d["street_days"])  # the truck is out every day
    assert s["meals_street"][0] == pytest.approx(s["street_days"][0] * d["meals_day"] * 0.5)
    assert s["event_days"][0] == pytest.approx(d["trucks"] * d["events_pm"] * 0.5)
    assert s["meals_street"][12] == pytest.approx(s["street_days"][12] * d["meals_day"] * (1 + d["meals_g"]))


def test_winter_dip_moves_street_days_but_keeps_the_year(ft_brief):
    flat = run(with_values(ft_brief, **FLAT, ramp_months=0))
    s = run(with_values(ft_brief, winter_dip=0.3, ramp_months=0))
    assert s["street_days"][0] < s["street_days"][5]  # January below June
    assert s["street_days"][11] < s["street_days"][6]  # December below July
    assert sum(s["street_days"][:12]) == pytest.approx(sum(flat["street_days"][:12]))
    assert sum(s["event_days"][:12]) == pytest.approx(sum(flat["event_days"][:12]))  # events carry no season


def test_cost_of_sales_is_food_packaging_card_and_event_fees(ft_brief):
    d = ft_brief.values(2)
    s = run(with_values(ft_brief, **FLAT, ramp_months=0))
    rev, meals = s["rev"][0], s["meals"][0]
    fees = s["rev_event"][0] * d["event_fee_pct"] + s["event_days"][0] * d["event_fee_day"]
    expected = rev * (d["food_pct"] + d["card_pct"]) + meals * d["pack"] + fees
    assert s["cogs"][0] == pytest.approx(-expected)
    assert s["meals"][0] == pytest.approx(s["meals_street"][0] + s["event_days"][0] * d["event_meals"])


def test_fuel_follows_service_days_and_fixed_costs_follow_the_fleet(ft_brief):
    d = ft_brief.values(2)
    one = run(with_values(ft_brief, **FLAT, ramp_months=0, trucks=1))
    three = run(with_values(ft_brief, **FLAT, ramp_months=0, trucks=3))
    assert one["opex_fuel"][0] == pytest.approx(-d["fuel"] * one["service_days"][0])
    for k in ("fuel", "commissary", "permits", "insurance", "maint"):
        assert three[f"opex_{k}"][0] == pytest.approx(3 * one[f"opex_{k}"][0]), k
    assert three["rev"][0] == pytest.approx(3 * one["rev"][0])


def test_trucks_added_each_year_scale_volume_and_capex(ft_brief):
    d = ft_brief.values(2)
    b = with_values(ft_brief, **FLAT, ramp_months=0, trucks_add=1, capex_refresh=0)
    s = run(b, g="annual")
    assert s["trucks_total"][:3] == pytest.approx([1, 2, 3])
    assert s["street_days"][1] == pytest.approx(2 * 12 * d["street_days"])
    assert s["capex"][0] == pytest.approx(d["capex_pos"])  # the early stage owns its truck
    assert s["capex"][1] == pytest.approx(d["capex_truck"] + d["capex_build"])
    assert s["capex"][2] == pytest.approx(d["capex_truck"] + d["capex_build"])
    # the yearly refresh covers the trucks already running, not the one bought that year
    r = run(with_values(b, capex_refresh=2000), g="annual")
    assert r["capex"][1] == pytest.approx(d["capex_truck"] + d["capex_build"] + 2000 * 1)
    assert r["capex"][2] == pytest.approx(d["capex_truck"] + d["capex_build"] + 2000 * 2)


def test_capex_month_one_then_refresh_per_truck_from_year_two():
    b = FT.default_brief("planning")
    d = b.values(2)
    s = run(b)
    assert s["capex"][0] == pytest.approx(d["capex_truck"] + d["capex_build"] + d["capex_pos"])
    assert s["capex"][1] == 0 and s["capex"][11] == 0
    assert s["capex"][12] == pytest.approx(d["capex_refresh"] * d["trucks"])
    assert s["capex"][24] == pytest.approx(d["capex_refresh"] * d["trucks"])
    assert run(b, g="annual")["capex"][1] == pytest.approx(d["capex_refresh"] * d["trucks"])


def test_street_meals_and_prices_grow_each_year(ft_brief):
    d = ft_brief.values(2)
    s = run(with_values(ft_brief, **FLAT, ramp_months=0, meals_g=0.1, price_g=0.05, events_pm=0), g="annual")
    assert s["meals_street"][1] == pytest.approx(s["meals_street"][0] * 1.1)
    assert s["rev_street"][1] == pytest.approx(s["rev_street"][0] * 1.1 * 1.05)
    assert s["rev_street"][0] == pytest.approx(12 * d["street_days"] * d["meals_day"] * d["ticket"])


@pytest.mark.parametrize("stage", STAGES)
def test_monthly_and_annual_agree_without_a_winter_dip(stage):
    b = with_values(FT.default_brief(stage), **FLAT)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    for k in ("rev", "cogs", "gp", "payroll_total", "capex", "meals", "service_days", "ebitda"):
        for y in (1, 2, 3):
            assert m[k][y] == pytest.approx(a[k][y], abs=0.01), (k, y)


@pytest.mark.parametrize("stage", ["planning", "early"])
def test_annual_is_close_with_a_winter_dip_during_the_ramp(stage):
    """Exact without a dip (above); the dip and the ramp interact in year 1, so annual is only close there."""
    b = FT.default_brief(stage)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    assert a["rev"][1] == pytest.approx(m["rev"][1], rel=0.05)
    assert a["rev"][1] != pytest.approx(m["rev"][1], rel=1e-6)
    for y in (2, 3):  # the ramp is over: a neutral season leaves the year unchanged
        assert a["rev"][y] == pytest.approx(m["rev"][y], abs=0.01), y


@pytest.mark.parametrize("g", ["monthly", "annual"])
@pytest.mark.parametrize("scen", [1, 2, 3])
def test_balance_sheet_balances(ft_brief, g, scen):
    s = run(ft_brief, scen, g)
    for ta, tle in zip(s["total_assets"], s["total_le"], strict=True):
        assert ta == pytest.approx(tle, abs=1e-6)


def test_degenerate_inputs_are_guarded(ft_brief):
    s = run(with_values(ft_brief, meals_day=0, street_days=0, events_pm=0, ramp_months=0))
    assert all(v == 0 for v in s["meals"]) and all(v == 0 for v in s["rev"])


# ---------------------------------------------------------------------- LibreOffice + challenger
lo = pytest.mark.libreoffice
COMBOS = list(itertools.product(["light", "full"], ["monthly", "annual"], [1, 3, 5]))


@lo
@pytest.mark.parametrize("tier,g,years", COMBOS, ids=[f"{t}-{g}-{y}y" for t, g, y in COMBOS])
def test_matrix_passes_the_challenger(build, ft_brief, tier, g, years):
    brief = ft_brief.with_options(tier=tier, granularity=g, horizon_years=years)
    path = build(brief)
    res = recalc(path)
    assert res.ok, res.errors
    vals = load_workbook(path, data_only=True)
    verdict = [r[1].value for r in vals["Checks"].iter_rows(min_col=1, max_col=2) if r[0].value == "Overall verdict"]
    assert verdict and verdict[0] in ("PASS", "PASS WITH WARNINGS")
    report = challenge(brief, path)
    assert report["problems"] == []
    assert report["verdict"] == "PASS"
    assert any(k.startswith("kpi:") for k in report["scenarios"][2]["years"][1])


@lo
@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_pass_the_challenger(build, stage):
    brief = FT.default_brief(stage)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
    assert report["verdict"].startswith("PASS")


@lo
def test_a_growing_fleet_with_events_and_a_deep_winter_dip_passes_the_challenger(build, ft_brief):
    brief = with_values(ft_brief, trucks_add=1, winter_dip=0.5, events_g=0.2, ramp_months=12, mkt_g=0.1)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
