"""End to end with the real engine: paid order → worker → LibreOffice → challenger → download.

A template whose formulas disagree with its Python twin must end in `failed`, alert the admin
and never produce a download.
"""

import io
from dataclasses import replace

import pytest
from openpyxl import load_workbook

from app.briefing import Answers, PrefillBriefSource
from app.db import session_scope
from app.models import Order
from app.orders import create_order, mark_paid
from engine import templates
from engine.core import Top
from engine.templates.services import ServicesTemplate
from tests.conftest import ANSWERS
from worker.pipeline import Pipeline

pytestmark = pytest.mark.libreoffice


def test_paid_order_is_built_checked_and_downloadable(client, services, notifier):
    r = client.post("/api/orders", json={"email": "founder@example.com", "answers": ANSWERS})
    oid = r.json()["order_id"]
    with session_scope(services.sessions) as s:
        mark_paid(s, services.cfg, oid, "cs_test_e2e", 14900)
    Pipeline(services.sessions, services.cfg, services.storage, notifier).run_once()
    assert client.get(f"/api/orders/{oid}").json()["status"] == "ready"
    ((_, _, url),) = notifier.ready
    resp = client.get(url.replace("http://testserver", ""))
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content), data_only=True)
    verdict = [
        row[1].value for row in wb["Checks"].iter_rows(min_col=1, max_col=2) if row[0].value == "Overall verdict"
    ]
    assert verdict[0] in ("PASS", "PASS WITH WARNINGS")
    assert wb.sheetnames == ["Cover", "Dashboard", "Break-even", "Inputs", "P&L", "Checks"]  # light tier (ADR-013)
    with session_scope(services.sessions) as s:
        rep = s.get(Order, oid).artifact("check_report")
        assert b'"verdict": "PASS"' in services.storage.get(rep.storage_key)


class BrokenServices(ServicesTemplate):
    """Python twin overstates revenue by 1%: the challenger must catch it."""

    key = "broken_services"

    def reference_top(self, d, p, y, t, hc) -> Top:
        top = super().reference_top(d, p, y, t, hc)
        return replace(top, rev=top.rev * 1.01)


def test_broken_template_ends_in_failed(services, notifier, monkeypatch):
    monkeypatch.setitem(templates.TEMPLATES, "broken_services", BrokenServices())
    brief = PrefillBriefSource().build_brief(Answers(**ANSWERS)).model_copy(update={"template": "broken_services"})
    with session_scope(services.sessions) as s:
        oid = create_order(s, services.cfg, "a@b.co", brief).id
    with session_scope(services.sessions) as s:
        mark_paid(s, services.cfg, oid, "cs_broken")
    Pipeline(services.sessions, services.cfg, services.storage, notifier).run_once()
    with session_scope(services.sessions) as s:
        o = s.get(Order, oid)
        assert o.status == "failed" and o.artifacts == []
        assert "challenger" in o.failure_reason and "workbook" in o.failure_reason
        assert [j.status for j in o.jobs] == ["failed"]
    assert notifier.ready == [] and len(notifier.alerts) == 1
