"""Sources & Uses sheet (E30, docs/WORKBOOK_STANDARD.md §1): funding structure and cash sources / uses."""

import shutil

import pytest
from openpyxl import load_workbook

from engine.builder import plan
from engine.challenger import challenge
from engine.core import reference_run
from engine.goal_sheets import su_default_periods
from engine.goals import apply_goal
from engine.recalc import recalc
from engine.spec import Timing
from engine.templates import TEMPLATES

SHEET = "Sources & Uses"


def _brief(key, **kw):
    tpl = TEMPLATES[key]
    return apply_goal(tpl.default_brief(tpl.sample_stage).with_options(**kw))


def test_full_only_and_after_valuation():
    full = plan(_brief("cafe", tier="full"), TEMPLATES["cafe"]).sheets
    assert full[full.index("Valuation") + 1] == SHEET
    assert SHEET not in plan(_brief("cafe", tier="light"), TEMPLATES["cafe"]).sheets


@pytest.mark.parametrize("g,years,n", [("monthly", 5, 12), ("monthly", 1, 12), ("annual", 3, 1), ("annual", 1, 1)])
def test_default_period(g, years, n):
    pl = plan(_brief("cafe", tier="full", granularity=g, horizon_years=years), TEMPLATES["cafe"])
    assert su_default_periods(pl) == n
    lay = pl.modules["sources_uses"]
    assert {"su_sources", "su_uses", "su_n", "su_cash_end", "su_gap"} <= set(lay.cells)


def _expected(brief, k):
    tpl = TEMPLATES[brief.template]
    t = Timing(brief.granularity, brief.horizon_years)
    d = brief.values(2)
    s = reference_run(d, brief.roles, t, tpl.reference_top, tpl.loan_style)
    k = min(max(1, k), t.periods)
    cfo = sum(s["cfo"][:k])
    src = d["equity"] + d["loan_amount"] + max(0.0, cfo)
    use = max(0.0, -cfo) - sum(s["cfi"][:k]) - sum(s["loan_repay"][:k]) + s["cash_close"][k - 1]
    return k, src, use, s["cash_close"][k - 1]


def _cells(path):
    ws = load_workbook(path, data_only=True)[SHEET]
    return {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}


@pytest.mark.libreoffice
@pytest.mark.parametrize("key", sorted(TEMPLATES))
@pytest.mark.parametrize("g,years", [("monthly", 5), ("annual", 3)])
def test_every_template_balances_and_passes_the_challenger(build, key, g, years):
    b = _brief(key, tier="full", granularity=g, horizon_years=years)
    path = build(b)
    rep = challenge(b, path)
    assert rep["problems"] == [], rep["problems"][:3]
    assert rep["scenarios"][2]["goal"]["su_uses"] == pytest.approx(rep["scenarios"][2]["goal"]["su_sources"], abs=0.5)
    assert recalc(path).ok
    vals = load_workbook(path, data_only=True)["Checks"]
    row = next(r for r in vals.iter_rows(min_col=1, max_col=2) if r[0].value == "Sources & uses balance")
    assert row[1].value == "PASS"


@pytest.mark.libreoffice
@pytest.mark.parametrize("key", sorted(TEMPLATES))
@pytest.mark.parametrize("g,typed", [("monthly", 1), ("monthly", 999), ("annual", 2), ("annual", 50)])
def test_period_input_is_capped_and_ties(build, tmp_path, key, g, typed):
    b = _brief(key, tier="full", granularity=g, horizon_years=3)
    path = build(b)
    pl = plan(b, TEMPLATES[key])
    ref = pl.modules["sources_uses"].local["su_n_in"]
    edited = tmp_path / "edited.xlsx"
    shutil.copy(path, edited)
    wb = load_workbook(edited)
    wb[SHEET][ref] = typed
    wb.save(edited)
    assert recalc(edited).ok
    got = _cells(edited)
    k, src, use, cash_end = _expected(b, typed)
    used = next(v for lab, v in got.items() if lab.endswith("s used"))
    assert used == k
    assert got["Total sources"] == pytest.approx(src, abs=0.5)
    assert got["Total uses"] == pytest.approx(use, abs=0.5)
    assert got["Cash at the end of the period"] == pytest.approx(cash_end, abs=0.5)
    assert src == pytest.approx(use, abs=0.5)
    checks = load_workbook(edited, data_only=True)["Checks"]
    row = next(r for r in checks.iter_rows(min_col=1, max_col=2) if r[0].value == "Sources & uses balance")
    assert row[1].value == "PASS"
    band = next(v for v in _labels(edited) if v.startswith("Sources and uses of cash in the first "))
    assert band.startswith(f"Sources and uses of cash in the first {k} {pl.t.period_word}")


def _labels(path):
    ws = load_workbook(path, data_only=True)[SHEET]
    return [ws.cell(r, 1).value for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, str)]


@pytest.mark.libreoffice
def test_tampered_sources_and_uses_is_caught(build):
    b = _brief("cafe", tier="full")
    path = build(b)
    pl = plan(b, TEMPLATES["cafe"])
    ref = pl.modules["sources_uses"].local["su_capex"]
    wb = load_workbook(path)
    wb[SHEET][ref] = wb[SHEET][ref].value + "*1.1"
    wb.save(path)
    rep = challenge(b, path)
    assert rep["verdict"] == "FAIL"
    assert any("su_capex" in p for p in rep["problems"])


@pytest.mark.libreoffice
def test_tampered_share_is_caught(build):
    b = _brief("cafe", tier="full")
    path = build(b)
    ref = plan(b, TEMPLATES["cafe"]).modules["sources_uses"].local["su_capex_c"]
    wb = load_workbook(path)
    wb[SHEET][ref] = wb[SHEET][ref].value + "*1.1"
    wb.save(path)
    rep = challenge(b, path)
    assert rep["verdict"] == "FAIL"
    assert any("su_capex_c" in p for p in rep["problems"])


def test_gearing_is_na_for_debt_without_equity():
    from engine.goals import NA, reference

    b = _brief("real_estate", tier="full")
    tpl = TEMPLATES["real_estate"]
    t = Timing(b.granularity, b.horizon_years)
    d = b.values(2) | {"equity": 0.0}
    s = reference_run(d, b.roles, t, tpl.reference_top, tpl.loan_style)
    assert d["loan_amount"] > 0
    assert reference(b, d, s, t)["cells"]["su_gearing"] == NA


def test_charts_and_input_cell(build):
    b = _brief("saas", tier="full")
    path = build(b)
    ws = load_workbook(path)[SHEET]
    assert len(ws._charts) == 2
    # standard slots side by side (docs/WORKBOOK_STANDARD.md §3); the size is pinned in the chart kit
    assert [c.anchor._from.col for c in ws._charts] == [5, 14]  # F and O
    assert {c.anchor._from.row for c in ws._charts} == {4}
    pl = plan(b, TEMPLATES["saas"])
    c = ws[pl.modules["sources_uses"].local["su_n_in"]]
    assert c.value == 12 and c.fill.fgColor.rgb.endswith("FFF2CC") and c.font.color.rgb.endswith("0000FF")
    assert any(dv.sqref and str(dv.sqref) == c.coordinate for dv in ws.data_validations.dataValidation)
