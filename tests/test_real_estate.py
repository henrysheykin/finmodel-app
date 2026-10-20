"""Real estate template (E05): rent roll, annuity mortgage, land, and the property Returns sheet.

Plain-Python checks run everywhere; the LibreOffice ones compare Excel's IRR / NPV / PMT with the twin.
The full tier × granularity × horizon matrix and the four stage defaults run in tests/test_matrix.py.
"""

import pytest
from openpyxl import load_workbook

from engine.builder import plan
from engine.challenger import challenge, python_goal_model, python_model
from engine.core import reference_run
from engine.goals import NA
from engine.returns import irr, sign_changes
from engine.spec import Timing
from engine.templates import get_template, planned
from engine.templates.base import Template

RE = get_template("real_estate")
STAGES = ["planning", "early", "growing", "established"]


@pytest.fixture
def re_brief():
    return RE.default_brief("early")


def run(brief, scen=2, g="monthly", years=3):
    return reference_run(brief.values(scen), brief.roles, Timing(g, years), RE.reference_top, RE.loan_style)


def with_values(brief, **vals):
    d = brief.model_dump()
    for drv in d["drivers"]:
        if drv["key"] in vals:
            v = vals[drv["key"]]
            drv.update(low=v, base=v, high=v)
    return brief.model_validate(d)


def returns(brief, scen=2):
    return python_goal_model(brief, scen)["cells"]


def test_registered_and_no_longer_coming_soon():
    assert "real_estate" not in planned()
    assert RE.business_types == ("real_estate",) and RE.returns and RE.loan_style == "annuity"


@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_are_ordered_funded_and_sensible(stage):
    b = RE.default_brief(stage)
    low, base, high = (python_model(b, s) for s in (1, 2, 3))
    for y in (1, 2, 3):
        assert low["rev"][y] < base["rev"][y] < high["rev"][y]
    for s in (1, 2, 3):
        assert min(python_model(b, s)["cash_close"].values()) > 0  # reserves + debt service covered
    r = returns(b)
    assert 0.05 < r["cap_in"] < 0.08  # illustrative going-in cap rate
    assert r["dscr_1"] > 1.2 and 0.6 <= r["ltv"] <= 0.7
    assert returns(b, 1)["irr_l"] < r["irr_l"] < returns(b, 3)["irr_l"]


def test_rent_roll(re_brief):
    d = re_brief.values(2)
    s = run(re_brief)
    gpr = d["units"] * d["rent"]
    assert s["gpr"][0] == pytest.approx(gpr)
    assert s["rev"][0] == pytest.approx(gpr * (1 - d["vacancy"]) + d["units"] * d["other_inc"] * (1 - d["vacancy"]))
    assert s["gpr"][12] == pytest.approx(gpr * (1 + d["rent_g"]))  # rent steps up once a year


def test_annuity_payment_is_level_and_matches_the_closed_form(re_brief):
    d = re_brief.values(2)
    s = run(re_brief, years=5)
    r, n, loan = d["loan_rate"] / 12, d["loan_term_years"] * 12, d["loan_amount"]
    pmt = loan * r / (1 - (1 + r) ** -n)
    for p in (0, 1, 30, 59):
        assert -s["debt_service"][p] == pytest.approx(pmt)
    assert s["interest"][0] == pytest.approx(-loan * r)  # interest on the balance after drawdown
    k = 60
    remaining = loan * (1 + r) ** k - pmt * ((1 + r) ** k - 1) / r
    assert s["loan_close"][-1] == pytest.approx(remaining)


def test_annuity_at_zero_rate_and_short_term_repays_in_full(re_brief):
    b = with_values(re_brief, loan_rate=0, loan_term_years=2)
    s = run(b, years=3)
    assert s["interest"][0] == 0
    assert -s["loan_repay"][0] == pytest.approx(b.values(2)["loan_amount"] / 24)
    assert s["loan_close"][23] == pytest.approx(0, abs=1e-6) and s["loan_close"][-1] == pytest.approx(0, abs=1e-6)
    assert all(v >= -1e-9 for v in s["loan_close"])


def test_land_is_not_depreciated(re_brief):
    b = with_values(re_brief, reserve=0)
    d = b.values(2)
    s = run(b, g="annual")
    building = d["price"] * (1 + d["closing_pct"]) - d["price"] * d["land_pct"]
    assert -s["dep"][0] == pytest.approx(building / d["dep_years"])
    no_land = run(with_values(b, land_pct=0), g="annual")
    assert -no_land["dep"][0] > -s["dep"][0]
    # depreciation stops at the building value: land stays on the balance sheet at cost
    full = run(with_values(b, dep_years=1), g="annual")
    assert full["ppe_net"][-1] == pytest.approx(d["price"] * d["land_pct"])


@pytest.mark.parametrize("g", ["monthly", "annual"])
@pytest.mark.parametrize("scen", [1, 2, 3])
def test_balance_sheet_balances(re_brief, g, scen):
    s = run(re_brief, scen, g)
    for ta, tle in zip(s["total_assets"], s["total_le"], strict=True):
        assert ta == pytest.approx(tle, abs=1e-6)


def test_annual_approximates_monthly(re_brief):
    m, a = python_model(re_brief, 2), python_model(re_brief.with_options(granularity="annual"), 2)
    for y in (1, 2, 3):
        assert a["rev"][y] == pytest.approx(m["rev"][y])  # no compounding inside the rent roll
        assert a["loan_close"][y] == pytest.approx(m["loan_close"][y], rel=0.01)


def test_returns_cash_flows(re_brief):
    d = re_brief.values(2)
    r = returns(re_brief)
    cost = d["price"] * (1 + d["closing_pct"])
    assert r["ucf_0"] == pytest.approx(-cost) and r["lcf_0"] == pytest.approx(-(cost - d["loan_amount"]))
    gross = r["noi_3"] * (1 + d["rent_g"]) / d["exit_cap"]
    assert r["net_sale"] == pytest.approx(gross * (1 - d["sell_pct"]))
    assert r["ucf_3"] == pytest.approx(r["noi_3"] + r["reserves_3"] + r["net_sale"])
    assert r["loan_in_0"] + sum(r[f"principal_{y}"] for y in (1, 2, 3)) + r["payoff_3"] == pytest.approx(0, abs=1e-6)
    flows = [r[f"lcf_{y}"] for y in range(4)]
    assert sum(f / (1 + r["irr_l"]) ** i for i, f in enumerate(flows)) == pytest.approx(0, abs=1e-3)
    npv = flows[0] + sum(f / (1 + d["disc_rate"]) ** i for i, f in enumerate(flows) if i)
    assert r["npv_l"] == pytest.approx(npv)
    assert r["dscr_0"] == NA and r["coc_0"] == NA


def test_dscr_after_the_loan_is_repaid_is_na(re_brief):
    r = returns(with_values(re_brief, loan_term_years=1).with_options(granularity="annual"))
    assert isinstance(r["dscr_1"], float) and r["dscr_2"] == NA and r["dscr_3"] == NA


def test_no_loan_means_levered_equals_unlevered(re_brief):
    r = returns(with_values(re_brief, loan_amount=0, loan_rate=0))
    assert r["irr_l"] == pytest.approx(r["irr_u"])
    assert r["dscr_1"] == NA


def test_irr_helper():
    assert irr([-100, 110]) == pytest.approx(0.10)
    assert irr([-1000, 50, 50, 1050]) == pytest.approx(0.05)
    assert irr([-100, 0, 121]) == pytest.approx(0.10)  # a zero year keeps the sign
    assert irr([-100, 50, 40]) < 0
    assert irr([100, 10]) == NA and irr([-100, -10]) == NA


def test_irr_is_not_shown_when_it_is_not_unique_or_a_near_total_loss():
    assert sign_changes([-100, 0, 50, -20, 80]) == 3
    assert irr([-100, 230, -132]) == NA  # roots at 10% and 20%: no single answer
    assert irr([-100, 24]) == NA and irr([-100, 25]) == pytest.approx(-0.75)  # recovers < / = 25%


def test_sale_below_the_loan_makes_levered_irr_na(re_brief):
    r = returns(with_values(re_brief, exit_cap=0.3, sell_pct=0.2))
    assert r["lcf_3"] < 0 < r["lcf_1"]  # equity gets cash, then pays in at the sale
    assert r["irr_l"] == NA and isinstance(r["irr_u"], float)


def test_returns_needs_its_drivers(re_brief):
    class Broken(type(RE)):
        key = ""  # not registered

        def driver_specs(self):
            return [s for s in super().driver_specs() if s.key != "exit_cap"]

    b = re_brief.model_validate(
        {**re_brief.model_dump(), "drivers": [d for d in re_brief.model_dump()["drivers"] if d["key"] != "exit_cap"]}
    )
    with pytest.raises(ValueError, match="exit_cap"):
        plan(b, Broken())


def test_other_templates_keep_straight_loans_and_no_returns_sheet(any_brief):
    tpl = get_template(any_brief.template)
    pl = plan(any_brief, tpl)
    if tpl.key == "real_estate":
        assert "Returns" in pl.sheets and "loan_pmt" in pl.lines and "land" in pl.lines
    else:
        assert type(tpl).loan_style is Template.loan_style and "Returns" not in pl.sheets
        assert "loan_pmt" not in pl.lines and "land" not in pl.lines


def test_unknown_loan_style_is_refused(re_brief):
    with pytest.raises(ValueError, match="loan style"):
        reference_run(re_brief.values(2), [], Timing("monthly", 1), RE.reference_top, "balloon")


# --------------------------------------------------------------------------- LibreOffice
@pytest.mark.libreoffice
@pytest.mark.parametrize("tier", ["light", "full"])
def test_workbook_returns_match_python(build, re_brief, tier):
    b = re_brief.with_options(tier=tier, horizon_years=5)
    path = build(b)
    report = challenge(b, path)
    assert report["problems"] == [] and report["verdict"] == "PASS"
    for scen in (1, 2, 3):
        got, exp = report["scenarios"][scen]["goal"], returns(b, scen)
        assert got["irr_u"] == pytest.approx(exp["irr_u"], abs=1e-4)  # Excel IRR vs bisection, 0.01 pp
    wb = load_workbook(path)
    assert "Returns" in wb.sheetnames and wb.sheetnames[-1] == "Checks"
    text = " ".join(str(c.value) for row in wb["Returns"].iter_rows() for c in row if c.value)
    assert "IRR(" in text and "NPV(" in text


@pytest.mark.libreoffice
def test_thin_coverage_warns_but_still_passes(build, re_brief):
    b = with_values(re_brief, loan_rate=0.11, equity=900000)
    assert returns(b)["dscr_1"] < 1
    path = build(b)
    report = challenge(b, path)
    assert report["problems"] == []
    assert report["scenarios"][2]["verdict"] == "PASS WITH WARNINGS"


@pytest.mark.libreoffice
def test_investor_pitch_goal_modules_work_with_returns(build, re_brief):
    from engine.goals import apply_goal

    b = apply_goal(re_brief.with_options(horizon_years=5), "investor_pitch")
    path = build(b)
    report = challenge(b, path)
    assert report["problems"] == []
    sheets = load_workbook(path).sheetnames
    assert sheets.index("Break-even") < sheets.index("Valuation") < sheets.index("Returns") < sheets.index("Inputs")
