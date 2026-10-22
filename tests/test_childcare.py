"""Childcare centre template: economics in the plain-Python twin, then the workbook through LibreOffice."""

import itertools
import math

import pytest
from openpyxl import load_workbook

from engine.challenger import challenge, python_model
from engine.core import reference_run
from engine.recalc import recalc
from engine.spec import Timing
from engine.templates import get_template

CC = get_template("childcare")
STAGES = ["planning", "early", "growing", "established"]
COMBOS = list(itertools.product(["light", "full"], ["monthly", "annual"], [1, 3, 5]))


@pytest.fixture
def cc_brief():
    return CC.default_brief("early")


def run(brief, scen=2, g="monthly", years=3):
    return reference_run(brief.values(scen), brief.roles, Timing(g, years), CC.reference_top)


def with_values(brief, **vals):
    d = brief.model_dump()
    for drv in d["drivers"]:
        if drv["key"] in vals:
            v = vals[drv["key"]]
            drv.update(low=v, base=v, high=v)
    return brief.model_validate(d)


# ---------------------------------------------------------------------- plain-Python twin
@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_are_valid_ordered_and_funded(stage):
    b = CC.default_brief(stage).with_options(horizon_years=5)
    low, base, high = (python_model(b, s) for s in (1, 2, 3))
    for y in range(1, 6):
        assert low["rev"][y] < base["rev"][y] < high["rev"][y]
        assert 0 < base["gp"][y] / base["rev"][y] < 1
    # Base is funded in every MONTH; Low is the stress case and may run out of cash
    cash = run(b, 2, years=5)["cash_close"]
    assert min(cash) > 0


@pytest.mark.parametrize("stage", STAGES)
def test_mature_base_is_profitable_with_plausible_labour_share(stage):
    b = CC.default_brief(stage).with_options(horizon_years=3)
    m = python_model(b, 2)
    assert m["ebitda"][3] / m["rev"][3] > 0.03
    assert 0.3 < -m["teach_pay"][3] / m["rev"][3] < 0.55


@pytest.mark.parametrize("stage", STAGES)
def test_monthly_and_annual_revenue_agree_exactly(stage):
    """Revenue is linear in the ramp (exact period average) and registrations count the end-of-period fill."""
    b = CC.default_brief(stage)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    for k in ("rev", "rev_tuition", "rev_sub", "rev_reg", "child_months", "new_kids", "cap_total", "capex"):
        for y in (1, 2, 3):
            assert m[k][y] == pytest.approx(a[k][y], abs=0.01), (k, y)


@pytest.mark.parametrize("stage", STAGES)
def test_monthly_and_annual_differ_only_by_teacher_rounding(stage):
    """Teachers are rounded up per period, so an annual model (one rounding per year) can differ from the monthly
    one only on teacher payroll, and by less than one teacher per band; every other line is identical. Once enrolment
    is full within a year the two are identical."""
    b = CC.default_brief(stage)
    d = b.values(2)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    for y in (1, 2, 3):
        for k in ("rev", "rev_tuition", "rev_sub", "rev_reg", "new_kids", "child_months", "capex", "opex_total"):
            assert m[k][y] == pytest.approx(a[k][y], abs=0.01), (k, y)
        dt = a["teach_pay"][y] - m["teach_pay"][y]
        assert a["cogs"][y] - m["cogs"][y] == pytest.approx(dt, abs=0.01)  # food and supplies are exact
        assert a["gp"][y] - m["gp"][y] == pytest.approx(dt, abs=0.01)
        assert a["ebitda"][y] - m["ebitda"][y] == pytest.approx(dt, abs=0.01)
        band_year = d["coverage"] * d["wage_teach"] * (1 + d["payroll_tax"]) * (1 + d["wage_g"]) ** (y - 1)
        assert abs(dt) <= 2 * band_year + 0.01  # less than one teacher per band
        assert abs(dt) < 0.06 * m["rev"][y]  # sanity: well inside 6% of revenue
        if d["ramp_months"] <= 12 * (y - 1):
            assert dt == pytest.approx(0, abs=0.01), y


def test_enrolment_ramps_in_a_straight_line_from_the_start_level_to_full():
    b = CC.default_brief("planning")
    d = b.values(2)
    s = run(b, years=3)
    n = int(d["ramp_months"])
    places = d["cap"]
    kids = [s["enr_end_kids"][i] for i in range(36)]
    assert kids[0] == pytest.approx(places * d["enr_start"])
    assert kids[n] == pytest.approx(places * d["enr_full"])  # month n + 1 is the first full month
    assert kids[n - 1] < kids[n]
    step = places * (d["enr_full"] - d["enr_start"]) / n
    assert kids[1] - kids[0] == pytest.approx(step)
    assert kids[-1] == pytest.approx(places * d["enr_full"])
    assert s["child_months"][0] == pytest.approx(kids[0])


def test_no_ramp_means_full_enrolment_from_month_one():
    b = with_values(CC.default_brief("planning"), ramp_months=0)
    d = b.values(2)
    s = run(b)
    assert s["child_months"][0] == pytest.approx(d["cap"] * d["enr_full"])
    assert s["enr_end_kids"][0] == pytest.approx(d["cap"] * d["enr_full"])


def test_tuition_by_age_band_discounts_and_subsidy(cc_brief):
    d = cc_brief.values(2)
    b = with_values(cc_brief, ramp_months=0)
    d = b.values(2)
    s = run(b)
    inf = d["cap"] * d["inf_share"] * d["enr_full"]
    pre = d["cap"] * (1 - d["inf_share"]) * d["enr_full"]
    gross = inf * d["tuit_inf"] + pre * d["tuit_pre"]
    assert s["rev_tuition"][0] == pytest.approx(gross * (1 - d["subsidy_share"]) * (1 - d["disc_pct"]))
    assert s["rev_sub"][0] == pytest.approx(gross * d["subsidy_share"] * d["sub_rate"])
    none = run(with_values(b, subsidy_share=0))
    assert max(none["rev_sub"]) == 0
    assert none["rev_tuition"][0] == pytest.approx(gross * (1 - d["disc_pct"]))


def test_tuition_rises_with_the_yearly_increase():
    b = with_values(CC.default_brief("established"), ramp_months=0, cap_add=0)
    d = b.values(2)
    s = run(b, years=3)
    ratio = (s["rev_tuition"][12] + s["rev_sub"][12]) / (s["rev_tuition"][0] + s["rev_sub"][0])
    assert ratio == pytest.approx(1 + d["tuition_g"])


def test_registration_fees_follow_new_children():
    b = with_values(CC.default_brief("established"), cap_add=0, turnover=0.24, kids_start=0)
    d = b.values(2)
    s = run(b, years=2)
    full = d["cap"] * d["enr_full"]
    assert s["new_kids"][0] == pytest.approx(full + full * 0.24 / 12)  # opening fill + replacement of leavers
    assert s["new_kids"][1] == pytest.approx(full * 0.24 / 12)
    assert s["rev_reg"][1] == pytest.approx(s["new_kids"][1] * d["reg_fee"])
    none = run(with_values(b, reg_fee=0))
    assert max(none["rev_reg"]) == 0


def test_children_already_enrolled_at_start_pay_no_registration_fee():
    b = with_values(CC.default_brief("established"), cap_add=0, turnover=0.24, ramp_months=0)
    d = b.values(2)
    full = d["cap"] * d["enr_full"]
    for kids_start, fill in ((0, full), (full / 2, full / 2), (full, 0), (full * 2, 0)):
        s = run(with_values(b, kids_start=kids_start))
        assert s["new_kids"][0] == pytest.approx(fill + full * 0.24 / 12), kids_start
        assert s["rev_reg"][0] == pytest.approx(s["new_kids"][0] * d["reg_fee"])
    # monthly and annual agree on the opening-year fee income
    half = with_values(b, kids_start=full / 2)
    m, a = python_model(half, 2), python_model(half.with_options(granularity="annual"), 2)
    assert m["new_kids"][1] == pytest.approx(a["new_kids"][1])


def test_kids_start_above_month_one_enrolment_counts_the_same_monthly_and_annual():
    """Children already enrolled beyond the month-1 level are capped at it, so the ramp fill-up is counted once."""
    b = CC.default_brief("planning")  # ramp of 12 months from 20% of 60 places
    d = b.values(2)
    month_one = d["cap"] * d["enr_start"]
    for kids_start in (month_one / 2, month_one, 30, 55):
        b2 = with_values(b, kids_start=kids_start)
        m, a = python_model(b2, 2), python_model(b2.with_options(granularity="annual"), 2)
        for k in ("new_kids", "rev_reg", "rev"):
            for y in (1, 2, 3):
                assert m[k][y] == pytest.approx(a[k][y], abs=0.01), (kids_start, k, y)
    capped, at_level = run(with_values(b, kids_start=55)), run(with_values(b, kids_start=month_one))
    assert capped["new_kids"][0] == pytest.approx(at_level["new_kids"][0])


@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_start_with_the_children_already_in_the_centre(stage):
    d = CC.default_brief(stage).values(2)
    expected = 0 if stage == "planning" else d["cap"] * d["enr_start"]
    assert d["kids_start"] == pytest.approx(expected)
    for scen in (1, 3):
        v = CC.default_brief(stage).values(scen)
        assert v["kids_start"] == pytest.approx(0 if stage == "planning" else v["cap"] * v["enr_start"])


@pytest.mark.parametrize(
    "inf_share,ratio_inf,ratio_pre,want",
    [
        (0.5, 5, 10, 6 + 3),  # 30 / 5 = 6 and 30 / 10 = 3 exactly: no extra teacher for float dust
        (0.35, 6, 12, 4 + 4),  # 21 / 6 = 3.5 and 39 / 12 = 3.25 round up
        (0.0, 5, 10, 6),  # no infants: 60 / 10
    ],
)
def test_teachers_follow_the_ratio_rounded_up(inf_share, ratio_inf, ratio_pre, want):
    b = with_values(
        CC.default_brief("established"),
        cap=60,
        cap_add=0,
        places_ready=60,
        enr_full=1.0,
        ramp_months=0,
        inf_share=inf_share,
        ratio_inf=ratio_inf,
        ratio_pre=ratio_pre,
    )
    assert run(b)["tch_pos"][0] == want


def test_teacher_payroll_is_a_cost_of_sales_line_and_not_headcount(cc_brief):
    d = cc_brief.values(2)
    s = run(cc_brief)
    fte = s["tch_pos"][0] * d["coverage"]
    expected = -fte * d["wage_teach"] / 12 * (1 + d["payroll_tax"])
    assert s["teach_pay"][0] == pytest.approx(expected)
    # headcount holds the roles only: director + office + kitchen
    assert s["fte_total"][0] == pytest.approx(d["director_fte"] + d["admin_fte"] + d["cook_fte"])
    roles_pay = -sum(d[f"{r}_fte"] * d[f"{r}_salary"] for r in ("director", "admin", "cook")) / 12
    assert s["payroll_total"][0] == pytest.approx(roles_pay * (1 + d["payroll_tax"]))
    # a different teacher wage moves cost of sales and leaves headcount payroll alone
    pricey = run(with_values(cc_brief, wage_teach=d["wage_teach"] * 2))
    assert pricey["payroll_total"][0] == pytest.approx(s["payroll_total"][0])
    assert pricey["teach_pay"][0] == pytest.approx(2 * s["teach_pay"][0])
    assert pricey["cogs"][0] < s["cogs"][0]


def test_food_depends_on_operating_days_and_supplies_on_children(cc_brief):
    d = cc_brief.values(2)
    s = run(cc_brief)
    kids = s["child_months"][0]
    assert s["cogs"][0] == pytest.approx(
        s["teach_pay"][0] - kids * d["food_day"] * d["days_open"] - kids * d["supplies_child"]
    )
    fewer = run(with_values(cc_brief, days_open=d["days_open"] - 5))
    assert fewer["cogs"][0] > s["cogs"][0]  # fewer operating days, less food


def test_rent_and_utilities_follow_floor_area():
    b = CC.default_brief("growing")
    d = b.values(2)
    s = run(b, years=2)
    area1, area2 = d["cap"] * d["sqft_place"], (d["cap"] + d["cap_add"]) * d["sqft_place"]
    assert s["opex_rent"][0] == pytest.approx(-area1 * d["rent_sqft"] / 12)
    assert s["opex_rent"][12] == pytest.approx(-area2 * d["rent_sqft"] / 12 * (1 + d["rent_g"]))
    assert s["opex_utilities"][0] == pytest.approx(-area1 * d["occ_sqft"] / 12)


def test_added_places_raise_capacity_and_buy_build_out_from_year_two():
    b = CC.default_brief("growing")
    d = b.values(2)
    s = run(b, years=3)
    per_place = d["sqft_place"] * d["fitout_sqft"] + d["equip_place"]
    assert s["cap_total"][0] == d["cap"]
    assert s["cap_total"][12] == d["cap"] + d["cap_add"]
    assert s["cap_total"][24] == d["cap"] + 2 * d["cap_add"]
    assert s["capex"][12] == pytest.approx(d["cap_add"] * per_place + (d["cap"] + d["cap_add"]) * d["refresh_place"])
    assert s["capex"][13] == 0


def test_opening_capex_builds_only_places_not_ready():
    b = CC.default_brief("planning")
    d = b.values(2)
    per_place = d["sqft_place"] * d["fitout_sqft"] + d["equip_place"]
    assert run(b)["capex"][0] == pytest.approx(d["cap"] * per_place + d["capex_play"])
    ready = with_values(b, places_ready=d["cap"] - 10)
    assert run(ready)["capex"][0] == pytest.approx(10 * per_place + d["capex_play"])


@pytest.mark.parametrize("stage", STAGES)
def test_receivable_days_blend_parent_and_subsidy_collection(stage):
    """ar_days applies to total revenue: the subsidy payments' share of revenue x 45 days of subsidy lag plus the rest
    x 2 days of card / ACH settlement, rounded."""
    b = CC.default_brief(stage)
    d = b.values(2)
    m = python_model(b, 2)
    share = m["rev_sub"][2] / m["rev"][2]
    assert 0.2 < share < 0.26
    assert d["ar_days"] == pytest.approx(share * 45 + (1 - share) * 2, abs=0.5)


@pytest.mark.parametrize("g", ["monthly", "annual"])
@pytest.mark.parametrize("scen", [1, 2, 3])
def test_balance_sheet_balances(cc_brief, g, scen):
    s = run(cc_brief, scen, g)
    for ta, tle in zip(s["total_assets"], s["total_le"], strict=True):
        assert ta == pytest.approx(tle, abs=1e-6)


def test_teacher_rounding_matches_excel_roundup():
    assert math.ceil(round(3.0000000000000004, 6)) == 3


# ---------------------------------------------------------------------- workbook (LibreOffice + challenger)
@pytest.mark.libreoffice
@pytest.mark.parametrize("tier,g,years", COMBOS, ids=[f"{t}-{g}-{y}y" for t, g, y in COMBOS])
def test_matrix(build, tier, g, years):
    brief = CC.default_brief("growing").with_options(tier=tier, granularity=g, horizon_years=years)
    path = build(brief)
    res = recalc(path)
    assert res.ok, res.errors
    vals = load_workbook(path, data_only=True)
    verdict = [r[1].value for r in vals["Checks"].iter_rows(min_col=1, max_col=2) if r[0].value == "Overall verdict"]
    assert verdict and verdict[0] in ("PASS", "PASS WITH WARNINGS")
    report = challenge(brief, path)
    assert report["problems"] == []
    assert report["verdict"] == "PASS"


@pytest.mark.libreoffice
@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_pass_the_challenger(build, stage):
    brief = CC.default_brief(stage)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
    assert report["verdict"] == "PASS"
    assert any(k.startswith("kpi:") for k in report["scenarios"][2]["years"][1])


@pytest.mark.libreoffice
def test_monthly_workbook_equals_annual_workbook(build):
    brief = CC.default_brief("established")  # full from month 1: no ramp-up and no rounding difference
    monthly = challenge(brief, build(brief, "m.xlsx"))
    annual_brief = brief.with_options(granularity="annual")
    annual = challenge(annual_brief, build(annual_brief, "a.xlsx"))
    for s in (1, 2, 3):
        for y in (1, 2, 3):
            m, a = monthly["scenarios"][s]["years"][y], annual["scenarios"][s]["years"][y]
            for k in ("rev", "gp", "ebitda", "kpi:child_months", "kpi:new_kids", "kpi:teach_pay"):
                assert m[k] == pytest.approx(a[k], abs=1), (s, y, k)


@pytest.mark.libreoffice
def test_revenue_rises_from_low_to_high_in_the_workbook(build, cc_brief):
    report = challenge(cc_brief, build(cc_brief))
    for y in (1, 2, 3):
        revs = [report["scenarios"][s]["years"][y]["rev"] for s in (1, 2, 3)]
        assert revs[0] < revs[1] < revs[2]
