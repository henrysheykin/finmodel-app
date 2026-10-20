"""E35: business-model pages (app/bizmodel.py, content/business-models/).

The central test rebuilds the sample workbook of every registered template, recalculates it in LibreOffice and
compares every table cell with the workbook cell(s) the build plan names. The same comparison runs against the
committed ``samples/<key>-full-5y.xlsx`` the page offers, so a stale sample fails too.
"""

import importlib.util
import json
import re
from html.parser import HTMLParser
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import bizmodel, seo
from engine.build import build_model
from engine.recalc import recalc
from engine.samples import SampleSpec
from engine.templates import TEMPLATES
from tests.test_browser_wizard import live_server, no_horizontal_scroll, page  # noqa: F401  (fixtures)

ROOT = Path(__file__).resolve().parents[1]
KEYS = sorted(TEMPLATES)
CONTENT = ROOT / "content" / "business-models"
PAGES = list(bizmodel.pages().values())


# --------------------------------------------------------------------------- table vs workbook
def _split(addr: str) -> tuple[str, str]:
    sheet, cell = addr.rsplit("!", 1)
    return sheet.strip("'"), cell.replace("$", "")


def _expected(cell: bizmodel.Cell, get):
    ref = cell.ref
    assert ref is not None, "every table cell must name its workbook cell"
    nums = [(c, get(a)) for c, a in ref.num]
    if any(isinstance(v, str) for _, v in nums):
        assert len(nums) == 1
        return nums[0][1]
    num = sum(c * v for c, v in nums)
    if not ref.den:
        return num
    den = sum(c * get(a) for c, a in ref.den)
    return num / den if den else 0.0


def check_against_workbook(model: bizmodel.BizModel, wb) -> int:
    """Compare every cell of every table; returns how many were checked. Raises AssertionError on a mismatch."""

    def get(addr: str):
        sheet, cell = _split(addr)
        return wb[sheet][cell].value

    checked = 0
    for table in model.tables.values():
        for row in table.rows:
            for (col, _), cell in zip(table.columns, row.cells, strict=True):
                where = f"{model.key} {table.name}.{row.key}.{col}"
                got = _expected(cell, get)
                if isinstance(cell.value, str):
                    assert got == cell.value, f"{where}: workbook {got!r} vs page {cell.value!r}"
                else:
                    assert isinstance(got, int | float), f"{where}: workbook value {got!r}"
                    tol = bizmodel.tolerance(row.fmt)
                    assert abs(got - cell.value) <= tol, f"{where}: workbook {got} vs page {cell.value} (tol {tol})"
                checked += 1
    return checked


@pytest.mark.parametrize("key", KEYS)
def test_tables_equal_a_rebuilt_and_recalculated_workbook(key, tmp_path):
    brief = bizmodel.sample_brief(key)
    path = tmp_path / f"{key}.xlsx"
    man = build_model(brief, path)
    res = recalc(path, 150)
    assert not res.errors, res.errors[:3]
    wb = load_workbook(path, data_only=True)
    sheet, cell = man.selector.split("!")
    assert wb[sheet][cell].value == bizmodel.BASE_SCENARIO  # the tables are the Base case
    model = bizmodel.compute(key)
    assert check_against_workbook(model, wb) > 60
    for name in bizmodel.TABLE_NAMES:
        assert model.tables[name].rows, (key, name)
    assert len(model.tables["unit"].rows) >= 2, key


@pytest.mark.parametrize("key", KEYS)
def test_tables_equal_the_committed_sample_the_page_offers(key):
    path = ROOT / "samples" / bizmodel.sample_spec(key).filename
    assert path.is_file(), f"run: python scripts/build_samples.py {key}"
    wb = load_workbook(path, data_only=True)
    assert check_against_workbook(bizmodel.compute(key), wb) > 60, "stale sample? rebuild it with build_samples.py"


def test_a_wrong_cell_would_be_caught(tmp_path):
    model = bizmodel.compute("cafe")
    wb = load_workbook(ROOT / "samples" / "cafe-full-5y.xlsx", data_only=True)
    sheet, cell = _split(model.tables["pnl"].row("rev").cells[0].ref.num[0][1])
    wb[sheet][cell] = wb[sheet][cell].value + 50
    with pytest.raises(AssertionError, match="pnl.rev.y1"):
        check_against_workbook(model, wb)


def test_default_sample_builds_include_the_page_sample():
    spec = importlib.util.spec_from_file_location("build_samples", ROOT / "scripts" / "build_samples.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    names = {s.name for s in mod.default_samples("cafe")}
    assert names == {"cafe-full", "cafe-full-5y"}
    assert {s.name for s in mod.seo_samples()} >= {"cafe-full-5y"}


def test_every_template_has_its_five_year_sample_committed():
    for key in KEYS:
        assert (ROOT / "samples" / f"{key}-full-5y.xlsx").is_file(), key


# --------------------------------------------------------------------------- numbers and formulas
def test_number_formats_follow_the_workbook():
    f = bizmodel.format_value
    assert f(1234.5, "usd") == "$1,235" and f(-1234.4, "usd") == "($1,234)" and f(0, "usd") == "-"
    assert f(0.3304, "pct") == "33.0%" and f(-0.012, "pct") == "(1.2%)" and f(6.5, "usd2") == "$6.50"
    assert f(6, "period") == "Month 6" and f(6, "period", prose=True) == "month 6"
    assert f("Not within horizon", "period") == "Not within horizon"
    assert f(3.04, "x") == "3.0x" and f(160.04, "dec1") == "160.0"


def _eval(text: str, **refs: float) -> float:
    tree = bizmodel._Parser(bizmodel._tokens(text.lstrip("="))).parse()
    return bizmodel._evaluate(tree, lambda name: refs[name])


def test_kpi_formula_evaluator():
    assert _eval("={a}*12", a=2) == 24
    assert _eval("=IF({n}=0,0,{m}/{n})", n=0, m=5) == 0  # the unchosen branch would divide by zero
    assert _eval("=IF({n}=0,0,{m}/{n})", n=4, m=6) == 1.5
    assert _eval("=IF({x}<=0,0,{y}/({x}*{z}))", x=2, y=8, z=2) == 2
    assert _eval("=MAX(0,{a}-{b})+-1", a=5, b=2) == 2
    assert _eval("=IF(AND({a}>1,{b}<3),1,2)", a=2, b=2) == 1
    with pytest.raises(ValueError):
        _eval("=SQRT({a})", a=4)
    with pytest.raises(ValueError):
        _eval("={a}+", a=1)


# --------------------------------------------------------------------------- placeholders
def _page(tmp_path: Path, body: str, name: str = "cafe", **meta) -> Path:
    fm = {
        "slug": f"/business-models/{name}",
        "title": "T | StartupModel",
        "description": "d",
        "keyword": "k",
        "template": name,
        "sample": f"{name}, full, monthly, 5 years",
    } | meta
    path = tmp_path / f"{name}.md"
    path.write_text("---\n" + "\n".join(f"{k}: {v}" for k, v in fm.items()) + "\n---\n" + body, encoding="utf-8")
    return path


@pytest.mark.parametrize(
    "bad",
    [
        "{{table:nope}}",
        "{{chart:nope}}",
        "{{num:pnl.nope.y1}}",
        "{{num:pnl.rev.y9}}",
        "{{num:nope.rev.y1}}",
        "{{num:pnl}}",
        "{{foo:bar}}",
    ],
)
def test_unknown_placeholder_is_a_build_error(tmp_path, bad):
    path = _page(tmp_path, f"# H\n\nText.\n\n{bad}\n")
    with pytest.raises(bizmodel.PlaceholderError):
        bizmodel.load_page(path)


def test_block_placeholders_must_stand_alone(tmp_path):
    with pytest.raises(bizmodel.PlaceholderError, match="own line"):
        bizmodel.load_page(_page(tmp_path, "# H\n\nSee {{table:pnl}} here.\n"))


def test_page_front_matter_is_checked(tmp_path):
    with pytest.raises(ValueError, match="slug"):
        bizmodel.load_page(_page(tmp_path, "# H\n", slug="/business-models/other"))
    with pytest.raises(ValueError, match="sample"):
        bizmodel.load_page(_page(tmp_path, "# H\n", sample="cafe, full, monthly, 3 years"))
    with pytest.raises(ValueError, match="registered template"):
        bizmodel.load_page(_page(tmp_path, "# H\n", name="nope"))


def test_num_placeholders_read_the_data(tmp_path):
    model = bizmodel.compute("cafe")
    rev1 = model.tables["pnl"].row("rev").cells[0].value
    cogs1 = model.tables["pnl"].row("cogs").cells[0].value
    assert bizmodel.num_text(model, "pnl.rev.y1") == bizmodel.format_value(rev1, "usd")
    assert bizmodel.num_text(model, "pnl.cogs.y1").startswith("($")
    assert bizmodel.num_text(model, "pnl.cogs.y1|abs") == bizmodel.format_value(abs(cogs1), "usd")
    assert bizmodel.num_text(model, "startup.required") == bizmodel.format_value(
        model.tables["startup"].row("required").cells[0].value, "usd"
    )
    bp = bizmodel.load_page(_page(tmp_path, "# H\n\nRevenue was {{num:pnl.rev.y1}} in year 1.\n"))
    assert f"Revenue was {bizmodel.num_text(model, 'pnl.rev.y1')} in year 1." in bp.body


# --------------------------------------------------------------------------- the cafe copy
def _plain(md: str) -> str:
    return bizmodel.PLACEHOLDER_RE.sub("", md)


def test_content_files_are_named_after_registered_templates():
    assert PAGES, "ship at least the cafe page"
    assert {p.key for p in PAGES} == {f.stem for f in CONTENT.glob("*.md")} <= set(TEMPLATES)
    assert "/business-models/cafe" in bizmodel.pages()


@pytest.mark.parametrize("bp", PAGES, ids=lambda p: p.key)
def test_copy_rules(bp):
    md = (CONTENT / f"{bp.key}.md").read_text(encoding="utf-8")
    body = md.split("\n---\n", 1)[1]
    bizmodel.check_placeholders(body, bizmodel.compute(bp.key))  # no unknown placeholder
    for tok in re.findall(r"\[[A-Z][A-Z0-9_ ]{2,}\]", body):
        assert tok.startswith("[PRICE"), f"{bp.key}: placeholder {tok} in the copy"
    assert 900 <= bp.words <= 1500, bp.words
    # a dollar or percent figure written by hand needs a source: its sentence carries an external link
    sentences = re.split(r"(?<=[.!?])\s+|\n+", _plain(body))
    for s in sentences:
        if re.search(r"\$\s?\d|\d\s?%", s):
            assert re.search(r"\]\(https?://", s), f"{bp.key}: unsourced figure in: {s!r}"
    # every table and chart is used once
    used = [m for m in bizmodel.placeholders(body) if m[0] in ("table", "chart")]
    assert len(used) == len(set(used))
    assert {a for k, a in used if k == "table"} == set(bizmodel.TABLE_NAMES)
    assert {a for k, a in used if k == "chart"} == set(bizmodel.CHART_NAMES)
    assert "FinModelsLab" not in md


def test_unsourced_figure_rule_would_fire():
    sent = "Rent is about $4,500 a month."
    assert re.search(r"\$\s?\d|\d\s?%", sent) and not re.search(r"\]\(https?://", sent)


# --------------------------------------------------------------------------- HTML
class _Doc(HTMLParser):
    def __init__(self):
        super().__init__()
        self.hrefs, self.ids, self.headings, self.tags = [], set(), [], []
        self.svgs: list[dict] = []
        self._svg: dict | None = None
        self._capture = ""

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        self.tags.append((tag, a))
        if tag == "a" and a.get("href"):
            self.hrefs.append(a["href"])
        if a.get("id"):
            self.ids.add(a["id"])
        if re.fullmatch(r"h[1-6]", tag):
            self.headings.append(int(tag[1]))
        if tag == "svg":
            self._svg = {"attrs": a, "title": "", "desc": ""}
            self.svgs.append(self._svg)
        if self._svg is not None and tag in ("title", "desc") and a.get("id"):
            self._capture = tag

    def handle_endtag(self, tag):
        if tag == "svg":
            self._svg = None
        if tag in ("title", "desc"):
            self._capture = ""

    def handle_data(self, data):
        if self._svg is not None and self._capture:
            self._svg[self._capture] += data


def parse(html: str) -> _Doc:
    d = _Doc()
    d.feed(html)
    return d


def test_cafe_page_has_six_tables_two_charts_and_the_sample(client):
    r = client.get("/business-models/cafe")
    assert r.status_code == 200
    html, doc = r.text, parse(r.text)
    for name in bizmodel.TABLE_NAMES:
        assert f'id="table-{name}"' in html, name
    for name in bizmodel.CHART_NAMES:
        assert f'id="chart-{name}"' in html, name
    assert html.count("Illustrative Base case of the sample model — your model uses your numbers") >= 8
    assert doc.headings.count(1) == 1
    for prev, cur in zip(doc.headings, doc.headings[1:], strict=False):
        assert cur <= prev + 1
    assert 'rel="canonical" href="http://testserver/business-models/cafe"' in html
    assert "/samples/cafe-full-5y.xlsx" in doc.hrefs
    assert any(h.startswith("/start?type=restaurant_cafe") for h in doc.hrefs)
    assert "/models/restaurant-coffee-shop-financial-projections" in doc.hrefs  # its SEO page
    assert "/business-models" in doc.hrefs
    assert "{{" not in html and "BMBLOCK" not in html and "FinModelsLab" not in html
    for href in doc.hrefs:
        if href.startswith("#"):
            assert href[1:] in doc.ids
        elif href.startswith("/"):
            assert client.get(href.split("#")[0]).status_code == 200, href


def test_tables_are_accessible_html(client):
    doc = parse(client.get("/business-models/cafe").text)
    tables = [a for t, a in doc.tags if t == "table"]
    assert len(tables) == 6
    assert sum(1 for t, a in doc.tags if t == "caption") == 6
    assert all(a.get("scope") in ("row", "col") for t, a in doc.tags if t == "th")
    regions = [a for t, a in doc.tags if t == "div" and a.get("role") == "region"]
    assert all(a.get("tabindex") == "0" and a.get("aria-label") for a in regions)


def test_charts_are_labelled_svg(client):
    doc = parse(client.get("/business-models/cafe").text)
    charts = [s for s in doc.svgs if s["attrs"].get("class") == "bm-svg"]  # not the logo icons
    assert len(charts) == 2
    for svg in charts:
        a = svg["attrs"]
        assert a["role"] == "img" and a["viewbox"]
        ids = a["aria-labelledby"].split()
        assert len(ids) == 2 and set(ids) <= doc.ids
        assert len(svg["title"]) > 10 and len(svg["desc"]) > 50  # the desc carries the data
    assert "Year 1: revenue" in charts[0]["desc"] and "Year 5" in charts[1]["desc"]


def test_chart_marks_match_the_data():
    model = bizmodel.compute("cafe")
    svg = bizmodel.chart_revenue_costs(model.charts["revenue_costs"], "t")
    assert svg.count('class="bm-s1"') == 5 + 1 and svg.count('class="bm-s2"') == 5 + 1  # bars + legend swatch
    data = model.charts["cost_structure"]
    assert len(data["labels"]) == 5 and all(0 <= v < 1 for v in data["first"] + data["last"])
    svg2 = bizmodel.chart_cost_structure(data, "t")
    assert svg2.count('class="bm-y1"') == 5 + 1 and svg2.count('class="bm-y5"') == 5 + 1
    # the tallest revenue bar reaches the top gridline: the axis is scaled to the data
    assert max(model.charts["revenue_costs"]["revenue"]) <= float(re.findall(r">\$([\d.]+)k<", svg)[-1]) * 1000


def test_json_ld_is_article_and_breadcrumbs(client):
    html = client.get("/business-models/cafe").text
    blocks = [json.loads(b) for b in re.findall(r'<script type="application/ld\+json">(.*?)</script>', html, re.S)]
    kinds = {b["@type"] for b in blocks}
    assert kinds == {"Article", "BreadcrumbList"}
    crumbs = next(b for b in blocks if b["@type"] == "BreadcrumbList")
    assert [i["name"] for i in crumbs["itemListElement"]][:2] == ["Home", "Business models"]
    assert crumbs["itemListElement"][2]["item"] == "http://testserver/business-models/cafe"


def test_index_has_a_card_per_page_and_is_listed_everywhere(client):
    r = client.get("/business-models")
    assert r.status_code == 200
    doc = parse(r.text)
    assert doc.hrefs.count("/business-models/cafe") == 1
    assert r.text.count('class="bm-card"') == len(PAGES)
    assert "Funding required" in r.text and "'" not in r.text.split("<title>")[1].split("</title>")[0]
    assert "/business-models" in client.get("/sitemap.xml").text
    assert "http://testserver/business-models/cafe</loc>" in client.get("/sitemap.xml").text
    llms = client.get("/llms.txt").text
    assert "## Business models" in llms and "/business-models/cafe" in llms
    assert client.get("/business-models/nope").status_code == 404
    assert client.get("/business-models/").status_code in (200, 307, 308)


def test_titles_and_canonicals_differ_from_the_model_pages(client):
    bp = bizmodel.pages()["/business-models/cafe"]
    seo_pages = bizmodel.seo_pages_for("cafe")
    assert seo_pages
    for sp in seo_pages:
        assert bp.title != sp.title and bp.slug != sp.slug and bp.keyword != sp.keyword
        html = client.get(sp.slug).text
        assert "/business-models/cafe" in parse(html).hrefs  # cross-link back
        assert (
            f'href="http://testserver{sp.slug}"' in html
            and 'business-models/cafe">' not in html.split("canonical")[1][:80]
        )
    all_titles = [p.title for p in seo.pages().values()] + [p.title for p in PAGES]
    assert len(set(all_titles)) == len(all_titles)


def test_calculator_links_appear_when_the_folder_has_matching_pages(client, tmp_path, monkeypatch):
    assert bizmodel.calculators_for("cafe") == [] or (ROOT / "content" / "calculators").is_dir()
    monkeypatch.setattr(bizmodel, "CALCULATORS", tmp_path / "missing")
    assert bizmodel.calculators_for("cafe") == []  # absent folder is fine
    calc = tmp_path / "calcs"
    calc.mkdir()
    (calc / "menu-price.md").write_text(
        "---\nslug: /calculators/cafe-menu-price\ntitle: Menu price | StartupModel\ntemplate: cafe\n---\n"
        "# Cafe menu price calculator\n",
        encoding="utf-8",
    )
    (calc / "other.md").write_text(
        "---\nslug: /calculators/saas-x\ntitle: X | StartupModel\ntemplate: saas\n---\n# SaaS X\n", encoding="utf-8"
    )
    (calc / "broken.md").write_text("no front matter", encoding="utf-8")
    monkeypatch.setattr(bizmodel, "CALCULATORS", calc)
    assert bizmodel.calculators_for("cafe") == [
        {"slug": "/calculators/cafe-menu-price", "title": "Cafe menu price calculator"}
    ]
    html = client.get("/business-models/cafe").text
    assert 'href="/calculators/cafe-menu-price"' in html and "saas-x" not in html


def test_sample_spec_is_the_5y_full_monthly_file():
    assert bizmodel.sample_spec("cafe") == SampleSpec("cafe", "full", "monthly", 5)
    assert bizmodel.sample_spec("cafe").filename == "cafe-full-5y.xlsx"
    assert bizmodel.pages()["/business-models/cafe"].sample.filename == "cafe-full-5y.xlsx"
