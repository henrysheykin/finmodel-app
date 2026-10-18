"""E17: the whole customer journey in one browser session, nothing faked but the card network.

landing → wizard → checkout → Stripe-shaped hosted page (signed webhook) → duplicate webhook
delivery → worker with the real engine (openpyxl → LibreOffice → challenger) → "model ready"
email (real copy, console sender) → download from the email link → magic-link sign-in from the
sign-in email → My models → download. The file must say PASS on its Checks sheet, and no page
on the way may trip the Content-Security-Policy.
"""

import io
import re
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from adapters.email import ConsoleSender
from app.db import session_scope
from app.emails import EmailNotifier
from app.models import Order
from tests.test_browser_wizard import live_server, page, sync_api  # noqa: F401  (fixtures)
from worker.pipeline import Pipeline

pytestmark = [pytest.mark.browser, pytest.mark.libreoffice]


def _verdict(xlsx: bytes) -> str:
    wb = load_workbook(io.BytesIO(xlsx), data_only=True)
    return next(r[1].value for r in wb["Checks"].iter_rows(min_col=1, max_col=2) if r[0].value == "Overall verdict")


def _mail(outbox, to: str, subject_part: str):
    hits = [m for m in outbox if m.to == to and subject_part.lower() in m.subject.lower()]
    assert len(hits) == 1, [(m.to, m.subject) for m in outbox]
    return hits[0]


def test_full_customer_journey(page, live_server, services):  # noqa: F811
    outbox = ConsoleSender()
    services.cfg = replace(services.cfg, support_email="help@startup-model.test", admin_email="ops@startup-model.test")
    services.notifier = EmailNotifier(outbox, services.cfg, services.sessions)
    problems: list[str] = []
    page.on("console", lambda m: problems.append(m.text) if "Content Security Policy" in m.text else None)
    page.on("pageerror", lambda e: problems.append(str(e)))
    buyer = "journey@example.com"

    # landing → wizard
    page.goto(f"{live_server}/")
    page.locator('a[href="/start?plan=light"]').click()
    page.wait_for_url("**/start?plan=light")
    page.get_by_label("Professional services").check()
    page.get_by_label("Business name").fill("Journey Advisory")
    for _ in range(4):
        page.get_by_role("button", name="Next").click()
    page.get_by_label("Use typical sample values").check()
    page.get_by_role("button", name="Next").click()
    page.get_by_role("button", name="Next").click()
    page.get_by_label("Email").fill(buyer)
    page.get_by_role("button", name="Continue to payment").click()
    page.wait_for_url("**/checkout/*")
    oid = page.url.rsplit("/", 1)[1]

    # checkout → hosted payment page → webhook (signed) → success page
    page.get_by_role("button", name="Pay $149").click()
    page.wait_for_url("**/dev/pay/*")
    sid = page.url.rsplit("/", 1)[1]
    page.get_by_role("button", name="Pay $149 (test)").click()
    page.wait_for_url("**/success")

    # the provider delivers the same payment again: acknowledged, no second receipt or build
    payload, sig = services.payments.signed(services.payments.event(sid))
    r = page.request.post(
        f"{live_server}/webhooks/stripe",
        data=payload,
        headers={"Stripe-Signature": sig, "Content-Type": "application/json"},
    )
    assert r.ok and r.json()["action"] == "already_paid"
    _mail(outbox.outbox, buyer, "got your order")  # exactly one receipt

    # one worker, real engine
    pipe = Pipeline(services.sessions, services.cfg, services.storage, services.notifier, worker_id="journey")
    assert pipe.run_once() is True
    assert pipe.run_once() is False  # exactly one build job was queued
    with session_scope(services.sessions) as s:
        o = s.get(Order, oid)
        assert o.status == "ready" and len(o.jobs) == 1
    sync_api.expect(page.locator("#status-text")).to_have_attribute("data-status", "ready", timeout=10000)

    # the "model ready" email and its link
    ready = _mail(outbox.outbox, buyer, "ready to download")
    link = re.search(r"http://127\.0\.0\.1:\d+/download/[A-Za-z0-9._-]+", ready.text).group(0)
    via_email = page.request.get(link)
    assert via_email.ok and via_email.headers["cache-control"] == "private, no-store"
    assert _verdict(via_email.body()) in ("PASS", "PASS WITH WARNINGS")
    assert len([m for m in outbox.outbox if m.to == buyer]) == 2  # receipt + ready; nothing else

    # sign in with the emailed link → My models → download
    page.goto(f"{live_server}/account")
    page.wait_for_url("**/login")
    page.get_by_label("Email").fill(buyer)
    page.get_by_role("button", name="Email me a sign-in link").click()
    signin = _mail(outbox.outbox, buyer, "sign-in link")
    login = re.search(r"http://127\.0\.0\.1:\d+/login/verify\?token=[A-Za-z0-9_-]+", signin.text).group(0)
    page.goto(login)
    page.get_by_role("button", name="Sign in").click()
    page.wait_for_url("**/account")
    assert page.get_by_role("rowheader", name="Journey Advisory").is_visible()
    with page.expect_download() as dl:
        page.get_by_role("link", name="Download .xlsx").click()
    data = Path(dl.value.path()).read_bytes()
    assert data == via_email.body()  # the same checked file both ways
    assert dl.value.suggested_filename == "journey-advisory-financial-model-light.xlsx"

    # the emailed sign-in link works once
    page.goto(login)
    page.get_by_role("button", name="Sign in").click()
    assert page.get_by_role("heading", name="Sign-in link not valid").is_visible()

    assert problems == []
