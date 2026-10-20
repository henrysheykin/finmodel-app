"""Summary sheet (E28, docs/WORKBOOK_STANDARD.md §1, §3): every year left, a selected year by month right."""

import re
import zipfile

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L

from engine.builder import manifest_for, plan
from engine.challenger import challenge
from engine.dashboard import CORE
from engine.recalc import recalc
from engine.summary import CHART_ROWS, KEEP, TOP, plan_summary, tie_condition
from engine.templates import TEMPLATES, get_template

FMT_K = '$#,##0,;($#,##0,);"-"'


def _brief(key="saas", **kw):
    tpl = TEMPLATES[key]
    return tpl.default_brief(tpl.sample_stage).with_options(**kw)


def _plan(key="saas", **kw):
    return plan(_brief(key, **kw), get_template(key))


def _labels(ws, col=1):
    return {r: ws.cell(row=r, column=col).value for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=col).value}


def _widths(ws) -> dict[str, float]:
    out = {}
    for dim in ws.column_dimensions.values():
        for i in range(dim.min or 1, (dim.max or dim.min or 1) + 1):
            out[L(i)] = dim.width
    return out


@pytest.fixture
def saas(build):
    path = build(_brief("saas", tier="full", horizon_years=5))
    return load_workbook(path), plan(_brief("saas", tier="full", horizon_years=5), get_template("saas"))


# --------------------------------------------------------------------------- layout (standard §1, §3)
def test_blocks_rows_and_positions(saas):
    wb, pl = saas
    ws, lay = wb["Summary"], pl.summary
    lab = _labels(ws)
    tops = {k: b.top for k, b in lay.blocks.items()}
    assert (
        tops
        == {"pl": TOP, "cf": TOP + CHART_ROWS + 1, "bs": TOP + 2 * (CHART_ROWS + 1)}
        == {"pl": 5, "cf": 22, "bs": 39}
    )
    assert lab[5].startswith("Income statement ($'000)") and lab[22].startswith("Cash flow ($'000)")
    assert lab[39].startswith("Balance sheet at year end ($'000)")
    assert [lab[r] for r in range(7, 21)] == [  # 14 lines
        "Revenue",
        lab[8],  # the template's cost-of-revenue label
        "Gross profit",
        "Gross margin",
        "Total operating expenses",
        "EBITDA",
        "EBITDA margin",
        "Depreciation",
        "EBIT",
        "Interest",
        "Pre-tax income",
        "Income tax",
        "Net income",
        "Net margin",
    ]
    assert [lab[r] for r in range(24, 32)] == [  # 8 lines
        "Operating cash flow",
        "Capital spending",
        "Equity raised",
        "Loan drawdown",
        "Loan repayment",
        "Financing cash flow",
        "Net change in cash",
        "Closing cash",
    ]
    bs = [lab[r] for r in range(41, 60) if r in lab]  # SaaS: 13 lines with deferred revenue
    assert bs[:5] == ["Cash", "Accounts receivable", "Inventory", "Net fixed assets", "Total assets"]
    assert bs[-2:] == ["Total equity", "Total liabilities & equity"]  # the check row stays on the Balance Sheet
    assert len(bs) == 13 and lay.end == 53


def test_every_block_is_padded_to_the_chart_height():
    for key in TEMPLATES:
        lay = _plan(key, tier="full").summary
        tops = [b.top for b in lay.blocks.values()]
        for a, b in zip(tops, tops[1:], strict=False):
            assert b - a >= CHART_ROWS + 1
        for b in lay.blocks.values():
            assert b.end >= b.top + 2  # at least one row


def test_columns_switch_and_titles(saas):
    wb, _ = saas
    ws = wb["Summary"]
    w = _widths(ws)
    assert w["A"] == 30 and w["B"] == w["F"] == 11.5 and w["G"] == 2 and w["T"] == 2
    assert w["H"] == w["S"] == 9.2
    assert ws["A1"].value == "Financial statements summary" and ws["A2"].value
    assert ws["U1"].value == '=HYPERLINK("#\'Cover\'!A1","◂ Contents")'
    assert ws["P2"].value == "Select detailed year:"
    assert ws["R2"].value == 1 and ws["R2"].number_format == '"Year "0'
    assert ws["R2"].fill.fgColor.rgb[-6:] == "FFE08A" and ws["R2"].border.left.style == "medium"
    assert "R2:S2" in {str(r) for r in ws.merged_cells.ranges}
    dv = [d for d in ws.data_validations.dataValidation if "R2" in str(d.sqref)]
    assert dv and dv[0].formula1 == '"1,2,3,4,5"' and dv[0].type == "list"
    assert ws["H5"].value == '="Income statement ($\'000) — Year "&$R$2&", by month"'
    assert ws["H39"].value == '="Balance sheet at month end ($\'000) — Year "&$R$2'
    assert ws.freeze_panes == "B5"
    assert ws.sheet_properties.tabColor.rgb[-6:] == "0F6E56"


def test_left_links_statements_and_right_indexes_the_grid(saas):
    wb, pl = saas
    ws = wb["Summary"]
    lay = pl.summary
    r = lay.row("pl", "rev")
    assert ws[f"B{r}"].value == f"='P&L'!B{pl.blocks['pl'].placed['rev']}"
    assert ws[f"F{r}"].value == f"='P&L'!F{pl.blocks['pl'].placed['rev']}"
    assert ws[f"B{r}"].number_format == FMT_K
    assert ws[f"H{r}"].value == f"=INDEX({pl.grid_range('rev')},($R$2-1)*12+1)"
    assert ws[f"S{r}"].value == f"=INDEX({pl.grid_range('rev')},($R$2-1)*12+12)"
    rc = lay.row("cf", "cash_close")
    assert ws[f"C{rc}"].value == f"='Cash Flow'!C{pl.blocks['cf'].placed['cash_close']}"
    assert ws[f"K{rc}"].value == f"=INDEX({pl.grid_range('cash_close')},($R$2-1)*12+4)"
    rb = lay.row("bs", "total_le")
    assert ws[f"B{rb}"].value == f"='Balance Sheet'!B{pl.blocks['bs'].placed['total_le']}"
    gm = lay.row("pl", "gm")
    assert ws[f"B{gm}"].value == f"='P&L'!B{pl.blocks['pl'].placed['gm']}"  # margins link on the left
    assert ws[f"H{gm}"].value == f"=IF(H{r}=0,0,H{lay.row('pl', 'gp')}/H{r})"  # and are computed on the right
    assert ws[f"H{gm}"].font.italic and ws[f"A{gm}"].font.italic
    hdr = lay.blocks["pl"].header
    assert ws[f"B{hdr}"].value == 1 and ws[f"B{hdr}"].number_format == '"Year "0'
    prow = next(g for g in pl.grids if g.sheet == "Monthly").period_row
    assert ws[f"H{hdr}"].value == f"=INDEX(Monthly!$C${prow}:$BJ${prow},($R$2-1)*12+1)"
    assert ws[f"H{hdr}"].number_format == '"M"0'
    # totals: thin line, net income / closing cash / total L&E: double underline
    assert ws[f"B{lay.row('pl', 'gp')}"].border.top.style == "thin"
    assert ws[f"B{lay.row('pl', 'ni')}"].border.bottom.style == "double"
    assert ws[f"M{rc}"].border.bottom.style == "double" and ws[f"M{rb}"].border.bottom.style == "double"


def test_report_sheet_links_are_ink_not_green(saas):
    wb, _ = saas
    for row in wb["Summary"].iter_rows(min_row=5, max_row=60, max_col=19):
        for c in row:
            if c.value is not None:
                assert not c.font.color or c.font.color.rgb[-6:] != "008000", c.coordinate
                assert c.font.name == "Arial"


def test_three_charts_beside_their_blocks(saas):
    wb, pl = saas
    ws = wb["Summary"]
    anchors = [(c.anchor._from.col, c.anchor._from.row) for c in ws._charts]
    assert anchors == [(20, 4), (20, 21), (20, 38)]  # U5 U22 U39
    assert [b.top for b in pl.summary.blocks.values()] == [5, 22, 39]


def _chart_parts(path, sheet="Summary"):
    """The sheet's charts, in anchor order: (chart xml, drawing anchor xml) with namespace prefixes removed."""
    with zipfile.ZipFile(path) as z:
        n = load_workbook(path).sheetnames.index(sheet) + 1
        rels = z.read(f"xl/worksheets/_rels/sheet{n}.xml.rels").decode()
        drawing = re.search(r'Target="[^"]*/(drawing\d+\.xml)"', rels).group(1)
        drel = z.read(f"xl/drawings/_rels/{drawing}.rels").decode()
        charts = re.findall(r'Target="[^"]*/(chart\d+\.xml)"', drel)
        norm = lambda x: re.sub(r"<(/?)\w+:", r"<\1", x)  # noqa: E731
        return norm(z.read(f"xl/drawings/{drawing}").decode()), [
            norm(z.read(f"xl/charts/{c}").decode()) for c in charts
        ]


def test_chart_series_axes_and_size(build):
    path = build(_brief("saas", tier="full", horizon_years=5))
    drawing, charts = _chart_parts(path)
    assert len(charts) == 3
    assert drawing.count('cx="5112000"') == 3 and drawing.count('cy="2988000"') == 3  # 14.2 x 8.3 cm
    refs = [re.findall(r"<f>([^<]+)</f>", c) for c in charts]
    assert all(f.startswith("'Summary'!") for r in refs for f in r)
    assert "'Summary'!$H$7:$S$7" in refs[0]  # revenue by month
    assert "'Summary'!$H$6:$S$6" in refs[0]  # month header
    # IS: lines share the bars' axis; CF: closing cash on a second (right) axis; BS: stacked bars + equity line
    assert charts[0].count("<valAx>") == 1 and charts[1].count("<valAx>") == 2 and charts[2].count("<valAx>") == 1
    assert 'grouping val="stacked"' in charts[2] and 'grouping val="stacked"' not in charts[0]
    assert [c.count("<lineChart>") for c in charts] == [1, 1, 1]
    assert [c.count("<ser>") for c in charts] == [3, 4, 4]
    assert "Income statement by month" in charts[0] and "Cash flow by month" in charts[1]
    assert "Assets and equity at month end" in charts[2]


# --------------------------------------------------------------------------- variants
def test_annual_granularity_has_years_only_and_charts_by_year(build):
    brief = _brief("saas", tier="full", granularity="annual", horizon_years=5)
    path = build(brief)
    wb = load_workbook(path)
    ws = wb["Summary"]
    pl = plan(brief, get_template("saas"))
    assert wb.sheetnames[2] == "Summary" and "Annual" in wb.sheetnames
    assert ws["R2"].value is None and ws["P2"].value is None and not ws.data_validations.dataValidation
    assert all(ws.cell(row=r, column=c).value is None for r in range(5, 54) for c in range(8, 22))  # no monthly side
    assert ws["B7"].value == f"='P&L'!B{pl.blocks['pl'].placed['rev']}"
    assert [(c.anchor._from.col, c.anchor._from.row) for c in ws._charts] == [(7, 4), (7, 21), (7, 38)]  # H5 H22 H39
    _, charts = _chart_parts(path)
    refs = [re.findall(r"<f>([^<]+)</f>", c) for c in charts]
    assert "'Summary'!$B$7:$F$7" in refs[0] and "'Summary'!$B$6:$F$6" in refs[0]  # revenue and the year header
    assert "Income statement by year" in charts[0] and "Assets and equity at year end" in charts[2]
    assert ws["P1"].value == '=HYPERLINK("#\'Cover\'!A1","◂ Contents")'
    assert ws.freeze_panes == "B5"


def test_one_year_horizon_works_monthly_and_annual(build):
    for gran in ("monthly", "annual"):
        brief = _brief("cafe", tier="full", granularity=gran, horizon_years=1)
        path = build(brief)
        ws = load_workbook(path)["Summary"]
        lay = plan(brief, get_template("cafe")).summary
        assert lay.last_col == "B" and ws["B7"].value.startswith("='P&L'!B") and ws["C7"].value is None
        if gran == "monthly":
            dv = ws.data_validations.dataValidation[0]
            assert dv.formula1 == '"1"'
        _, charts = _chart_parts(path)
        refs = [re.findall(r"<f>([^<]+)</f>", c) for c in charts]
        assert len(charts) == 3
        # one year by month: the 12 month cells; annual: the single year cell (openpyxl writes B7:B7 as B7)
        assert ("'Summary'!$H$7:$S$7" if gran == "monthly" else "'Summary'!$B$7") in refs[0]


def test_light_tier_has_no_summary(build):
    brief = _brief("saas", tier="light")
    pl = plan(brief, get_template("saas"))
    assert pl.summary is None and "Summary" not in pl.sheets
    assert manifest_for(pl, "").summary_years == {}
    assert plan_summary(pl) is None


def test_sheet_sits_after_the_dashboard_and_on_the_cover(build):
    for gran, text in (("monthly", "any year by month"), ("annual", "every year, with charts")):
        wb = load_workbook(build(_brief("saas", tier="full", granularity=gran, horizon_years=3)))
        assert wb.sheetnames[:3] == ["Cover", "Dashboard", "Summary"]
        cover = wb["Cover"]
        row = next(
            r for r in range(1, 40) if cover.cell(row=r, column=3).value == '=HYPERLINK("#Summary!A1","Summary")'
        )
        assert text in cover.cell(row=row, column=4).value
        assert wb.sheetnames.index("Summary") == 2


def test_manifest_lists_the_summary_cells():
    pl = _plan("saas", tier="full", horizon_years=3)
    man = manifest_for(pl, "")
    assert set(man.summary_years) == {"sum_rev", "sum_ni", "sum_cash"}
    assert man.summary_years["sum_rev"] == {1: "Summary!B7", 2: "Summary!C7", 3: "Summary!D7"}
    assert man.summary_years["sum_cash"][3] == f"Summary!D{pl.summary.row('cf', 'cash_close')}"
    assert man.to_dict()["summary_years"] == man.summary_years


def test_check_row_and_error_range(build):
    path = build(_brief("saas", tier="full", horizon_years=5))
    ws = load_workbook(path)["Checks"]
    names = [c.value for c in ws["A"]]
    assert "Summary ties to statements" in names
    text = " ".join(str(c.value) for c in ws["B"] if c.value)
    assert "SUMPRODUCT(--ISERROR(Summary!$A$5:$S$53))" in text
    assert "Summary!$B$7:$F$7-'P&L'!$B$" in text and "Summary!$R$2" in text
    annual = _plan("saas", tier="full", granularity="annual", horizon_years=3)
    cond = tie_condition(annual)
    assert "R$2" not in cond and cond.startswith("IFERROR(AND(") and "Summary!$B$7:$D$7" in cond


def test_keep_lists_exist_in_every_template():
    for key in TEMPLATES:
        pl = _plan(key, tier="full")
        have = {r.key for r in pl.blocks["pl"].rows}
        assert set(KEEP["pl"]) <= have, key
        assert set(KEEP["cf"]) <= {r.key for r in pl.blocks["cf"].rows}, key
        assert len(pl.summary.blocks["pl"].rows) == 14 and len(pl.summary.blocks["cf"].rows) == 8


# --------------------------------------------------------------------------- recalculated behaviour
def _verdict(vals, name):
    return next(r[1].value for r in vals["Checks"].iter_rows(min_col=1, max_col=2) if r[0].value == name)


@pytest.mark.libreoffice
@pytest.mark.parametrize("key", sorted(TEMPLATES))
def test_summary_check_passes_and_statements_agree_in_every_template(build, key):
    """Summary ties to statements; the Dashboard Core financials equal the statement cells every year (E27 review)."""
    brief = _brief(key, tier="full", horizon_years=3)
    path = build(brief)
    assert recalc(path).ok
    wb = load_workbook(path, data_only=True)
    assert _verdict(wb, "Summary ties to statements") == "PASS"
    assert _verdict(wb, "No error values") == "PASS"
    pl = plan(brief, get_template(key))
    dash, lay = wb["Dashboard"], pl.dash

    def stmt(block, k, y):
        return wb[pl.blocks[block].sheet].cell(row=pl.blocks[block].placed[k], column=1 + y).value

    core = {  # Dashboard Core financials line -> (statement block, row)
        "rev": ("pl", "rev"),
        "gp": ("pl", "gp"),
        "gm": ("pl", "gm"),
        "opex": ("pl", "opex_total"),
        "ebitda": ("pl", "ebitda"),
        "ebitda_m": ("pl", "ebitda_m"),
        "ni": ("pl", "ni"),
        "ni_m": ("pl", "ni_m"),
        "cfo": ("cf", "cfo"),
        "cfi": ("cf", "cfi"),
        "cff": ("cf", "cff"),
        "net_cf": ("cf", "net_cf"),
        "cash": ("cf", "cash_close"),
    }
    assert set(core) == {c[0] for c in CORE}
    for y in range(1, 4):
        for k, (block, row) in core.items():
            got = dash.cell(row=lay.core[k], column=1 + y).value
            assert got == pytest.approx(stmt(block, row, y), abs=1e-6), (key, k, y)
        assert dash.cell(row=lay.core["cash"], column=1 + y).value == pytest.approx(stmt("bs", "cash", y), abs=1e-6)
    s = wb["Summary"]
    sl = pl.summary
    for y in range(1, 4):  # every left cell is the statement cell
        for block in ("pl", "cf", "bs"):
            for srow, r in sl.blocks[block].rows:
                assert s.cell(row=r, column=1 + y).value == pytest.approx(stmt(block, srow.key, y), abs=1e-6), (
                    block,
                    srow.key,
                )
    # the detailed year (1) by month adds up to its annual figure; a balance ends the year at the last month
    for block, key_, kind in (
        ("pl", "rev", "sum"),
        ("pl", "ni", "sum"),
        ("cf", "cfo", "sum"),
        ("cf", "net_cf", "sum"),
        ("cf", "cash_close", "last"),
        ("bs", "cash", "last"),
    ):
        r = sl.row(block, key_)
        months = [s.cell(row=r, column=c).value for c in range(8, 20)]
        want = sum(months) if kind == "sum" else months[-1]
        assert want == pytest.approx(s.cell(row=r, column=2).value, abs=0.01), (key, block, key_)
    rb = sl.row("bs", "total_le")
    assert s.cell(row=rb, column=19).value == pytest.approx(s.cell(row=sl.row("bs", "total_assets"), column=19).value)


@pytest.mark.libreoffice
@pytest.mark.parametrize(
    "gran,years", [("monthly", 5), ("annual", 3), ("monthly", 1), ("annual", 1)], ids=lambda v: str(v)
)
def test_recalculated_file_keeps_three_live_charts(build, gran, years):
    path = build(_brief("saas", tier="full", granularity=gran, horizon_years=years))
    assert recalc(path).ok
    _, charts = _chart_parts(path)
    assert len(charts) == 3
    for c in charts:
        fs = re.findall(r"<f>([^<]+)</f>", c)
        assert fs and all(f.startswith("Summary!") or f.startswith("'Summary'!") for f in fs), fs  # live references
    assert len(load_workbook(path)["Summary"]._charts) == 3
    vals = load_workbook(path, data_only=True)
    assert _verdict(vals, "Summary ties to statements") == "PASS"
    s = vals["Summary"]
    assert isinstance(s["B7"].value, int | float)  # the cells the charts read hold values
    if gran == "monthly":
        assert isinstance(s["H7"].value, int | float) and s["H6"].value == 1


@pytest.mark.libreoffice
def test_the_detailed_year_switch_moves_the_monthly_side(build):
    brief = _brief("saas", tier="full", horizon_years=3)
    path = build(brief)
    wb = load_workbook(path)
    wb["Summary"]["R2"] = 3
    wb.save(path)
    assert recalc(path).ok
    vals = load_workbook(path, data_only=True)
    s = vals["Summary"]
    assert s["H6"].value == 25 and s["S6"].value == 36  # months of year 3
    assert sum(s.cell(row=7, column=c).value for c in range(8, 20)) == pytest.approx(s["D7"].value, abs=0.01)
    assert _verdict(vals, "Summary ties to statements") == "PASS"
    assert "Year 3" in s["H5"].value


@pytest.mark.libreoffice
@pytest.mark.parametrize("gran", ["monthly", "annual"])
def test_a_broken_summary_link_fails_the_check(build, gran):
    brief = _brief("cafe", tier="full", granularity=gran, horizon_years=3)
    path = build(brief)
    pl = plan(brief, get_template("cafe"))
    lay = pl.summary
    wb = load_workbook(path)
    wb["Summary"][f"C{lay.row('pl', 'ni')}"] = 12345  # a hard-coded number instead of the link
    wb.save(path)
    assert recalc(path).ok
    vals = load_workbook(path, data_only=True)
    assert _verdict(vals, "Summary ties to statements") == "FAIL"
    assert _verdict(vals, "Overall verdict") == "FAIL"
    report = challenge(brief, path)
    assert report["verdict"] == "FAIL"
    assert any("sum_ni" in p for p in report["problems"]), report["problems"][:4]


@pytest.mark.libreoffice
def test_a_broken_monthly_link_fails_the_check(build):
    brief = _brief("cafe", tier="full", horizon_years=3)
    path = build(brief)
    lay = plan(brief, get_template("cafe")).summary
    wb = load_workbook(path)
    wb["Summary"][f"K{lay.row('pl', 'rev')}"] = 1  # one month of the detailed year
    wb["Summary"][f"S{lay.row('cf', 'cash_close')}"] = 5  # and its closing cash
    wb.save(path)
    assert recalc(path).ok
    assert _verdict(load_workbook(path, data_only=True), "Summary ties to statements") == "FAIL"


@pytest.mark.libreoffice
def test_an_error_value_on_the_summary_fails_no_error_values(build):
    brief = _brief("cafe", tier="full", horizon_years=3)
    path = build(brief)
    wb = load_workbook(path)
    wb["Summary"]["M20"] = "=1/0"
    wb.save(path)
    recalc(path)
    vals = load_workbook(path, data_only=True)
    assert _verdict(vals, "No error values") == "FAIL"


@pytest.mark.libreoffice
def test_a_bad_switch_value_is_a_fail_not_an_error_cell(build):
    brief = _brief("cafe", tier="full", horizon_years=3)
    path = build(brief)
    wb = load_workbook(path)
    wb["Summary"]["R2"] = 9
    wb.save(path)
    recalc(path)
    vals = load_workbook(path, data_only=True)
    assert _verdict(vals, "Summary ties to statements") == "FAIL"
    assert _verdict(vals, "Overall verdict") == "FAIL"


@pytest.mark.libreoffice
@pytest.mark.parametrize("key", ["cafe", "real_estate"])
def test_challenger_covers_the_summary(build, key):
    brief = _brief(key, tier="full", horizon_years=3)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
