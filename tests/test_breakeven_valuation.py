"""Break-even chart and Valuation layout (E32, docs/WORKBOOK_STANDARD.md §1, §3, §4)."""

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L

from engine import style as st
from engine.builder import plan
from engine.challenger import challenge
from engine.recalc import recalc
from engine.spec import FORMATS
from engine.templates import TEMPLATES, get_template

pytestmark = pytest.mark.filterwarnings("ignore::DeprecationWarning")
lo = pytest.mark.libreoffice
EMU_CM = 360000


def _brief(key="saas", **kw):
    tpl = TEMPLATES[key]
    return tpl.default_brief(tpl.sample_stage).with_options(**kw)


def _labels(ws):
    return {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value}


def _ref(series_obj):
    return _span(series_obj.val.numRef.f)


def _span(ref):
    """A one-cell range comes back from the file as a single cell: normalise to A:A form."""
    return ref if ":" in ref else f"{ref}:{ref.split('!')[1]}"


@pytest.mark.parametrize("tier", ["light", "full"])
@pytest.mark.parametrize("years", [1, 3, 5])
def test_breakeven_has_one_large_combo_chart(build, tier, years):
    brief = _brief("saas", tier=tier, horizon_years=years)
    wb = load_workbook(build(brief))
    ws = wb["Break-even"]
    assert len(ws._charts) == 1
    chart = ws._charts[0]
    assert (round(chart.anchor.ext.width / EMU_CM, 1), round(chart.anchor.ext.height / EMU_CM, 1)) == (20.0, 10.5)
    anchor = chart.anchor._from
    assert anchor.col == 0 and anchor.row + 1 > max(r for r in _labels(ws).values())  # below the table and notes
    labels = _labels(ws)
    last = L(1 + years)
    bars = chart.series
    assert len(bars) == 1 and len(chart._charts) == 2
    line = chart._charts[1].series
    assert len(line) == 1
    assert _ref(bars[0]) == f"'Break-even'!$B${labels['Revenue']}:${last}${labels['Revenue']}"
    be = labels["Break-even revenue (EBITDA = 0)"]
    assert _ref(line[0]) == f"'Break-even'!$B${be}:${last}${be}"
    assert chart.title.tx.rich.p[0].r[0].t == "Revenue vs break-even revenue (EBITDA = 0)"
    cats = _span(bars[0].cat.numRef.f)
    year_row = cats.rsplit("$", 1)[1]
    assert cats == f"'Break-even'!$B${year_row}:${last}${year_row}"
    assert wb["Break-even"].cell(row=int(year_row), column=2).value == 1  # the "Year #" row


def test_valuation_has_no_chart_and_sheet_tabs_are_pine(build):
    wb = load_workbook(build(_brief("saas", tier="full", horizon_years=3)))
    assert wb["Valuation"]._charts == []
    for name in ("Break-even", "Valuation"):
        assert wb[name].sheet_properties.tabColor.rgb[-6:] == st.PINE
        assert wb[name].sheet_view.showGridLines is False


def test_light_has_no_valuation_sheet(build):
    wb = load_workbook(build(_brief("saas", tier="light", horizon_years=3)))
    assert "Valuation" not in wb.sheetnames and len(wb["Break-even"]._charts) == 1


def test_valuation_assumptions_are_green_links_to_inputs(build):
    brief = _brief("saas", tier="full", horizon_years=3)
    ws = load_workbook(build(brief))["Valuation"]
    labels = _labels(ws)
    expected = {
        "Discount rate (WACC)": "pct",
        "Terminal growth rate": "pct",
        "Exit multiple (EV ÷ final-year EBITDA)": "x",
        "Tax rate (for the interest tax shield)": "pct",
        "Sensitivity step: WACC": "pct",
        "Sensitivity step: terminal growth": "pct",
    }
    for label, fmt in expected.items():
        c = ws.cell(row=labels[label], column=2)
        assert isinstance(c.value, str) and c.value.startswith("=Inputs!"), label
        assert c.font.color.rgb[-6:] == st.GREEN, label
        assert c.number_format == FORMATS[fmt], label
        assert c.font.name == "Arial" and c.font.sz == 9


def test_valuation_and_breakeven_money_in_full_dollars_and_key_total_formats(build):
    brief = _brief("saas", tier="full", horizon_years=3)
    wb = load_workbook(build(brief))
    val, be = wb["Valuation"], wb["Break-even"]
    lv, lb = _labels(val), _labels(be)
    for c in (val.cell(row=lv["Enterprise value"], column=2), val.cell(row=lv["Unlevered free cash flow"], column=2)):
        assert c.number_format == FORMATS["usd"]
    assert val.cell(row=lv["Discount factor"], column=2).number_format == FORMATS["dec3"]
    assert val.cell(row=lv["Terminal value share of enterprise value"], column=2).number_format == FORMATS["pct"]
    assert be.cell(row=lb["Funding required (uses + buffer)"], column=2).number_format == FORMATS["usd"]
    assert be.cell(row=lb["Gross margin (contribution)"], column=2).number_format == FORMATS["pct"]
    ev = val.cell(row=lv["Enterprise value — average of the two methods"], column=2)
    assert ev.border.bottom.style == "double"
    # WACC x g sensitivity: 5 x 5 grid whose middle reproduces the headline value
    sens = next(r for label, r in lv.items() if label.startswith("WACC"))
    assert val.cell(row=sens + 3, column=4).value.startswith("=IF(")
    assert val.cell(row=sens + 3, column=4).number_format == FORMATS["usd"]


def test_challenger_manifest_still_knows_the_cells(build):
    brief = _brief("saas", tier="full", horizon_years=3)
    pl = plan(brief, get_template("saas"))
    assert {"sens_centre", "ev_g", "required"} <= set(pl.modules["dcf"].cells) | set(pl.modules["breakeven"].cells)


@lo
@pytest.mark.parametrize("key", sorted(TEMPLATES))
@pytest.mark.parametrize("tier", ["light", "full"])
def test_every_template_passes_with_the_breakeven_chart(build, key, tier):
    brief = _brief(key, tier=tier, horizon_years=3)
    path = build(brief)
    res = recalc(path)
    assert res.ok, res.errors
    report = challenge(brief, path)
    assert report["problems"] == [] and report["verdict"] == "PASS"
    wb = load_workbook(path)  # the delivered (recalculated) file still holds the chart
    assert len(wb["Break-even"]._charts) == 1
    if tier == "full":
        assert wb["Valuation"]._charts == []
