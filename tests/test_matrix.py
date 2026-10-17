"""Golden matrix: every template × light+full × monthly+annual × 1/3/5 years.

Every combination must generate, recalculate with zero errors, balance (full
tier) and pass the challenger in all three scenarios.
"""

import itertools

import pytest
from openpyxl import load_workbook

from engine.challenger import challenge
from engine.recalc import recalc
from engine.templates import TEMPLATES

pytestmark = pytest.mark.libreoffice

COMBOS = list(itertools.product(["light", "full"], ["monthly", "annual"], [1, 3, 5]))


@pytest.mark.parametrize("tier,g,years", COMBOS, ids=[f"{t}-{g}-{y}y" for t, g, y in COMBOS])
def test_matrix(build, any_brief, tier, g, years):
    brief = any_brief.with_options(tier=tier, granularity=g, horizon_years=years)
    path = build(brief)
    res = recalc(path)
    assert res.ok, res.errors
    vals = load_workbook(path, data_only=True)
    verdict = [r[1].value for r in vals["Checks"].iter_rows(min_col=1, max_col=2) if r[0].value == "Overall verdict"]
    assert verdict and verdict[0] in ("PASS", "PASS WITH WARNINGS")
    report = challenge(brief, path)
    assert report["problems"] == []
    assert report["verdict"] == "PASS"


def test_light_and_full_show_identical_numbers(build, cafe_brief):
    reports = {}
    for tier in ("light", "full"):
        b = cafe_brief.with_options(tier=tier)
        reports[tier] = challenge(b, build(b, f"{tier}.xlsx"))
    for s in (1, 2, 3):
        for y in (1, 2, 3):
            full, light = reports["full"]["scenarios"][s]["years"][y], reports["light"]["scenarios"][s]["years"][y]
            for k in light:
                assert light[k] == full[k], (s, y, k)


@pytest.mark.parametrize("template", [k for k in TEMPLATES if k != "cafe"])
@pytest.mark.parametrize("stage", ["planning", "early", "growing", "established"])
def test_stage_defaults_pass(build, template, stage):
    from engine.templates import get_template

    brief = get_template(template).default_brief(stage)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
    assert any(k.startswith("kpi:") for k in report["scenarios"][2]["years"][1])
