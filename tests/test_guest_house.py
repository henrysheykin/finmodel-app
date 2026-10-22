"""Guest house template: economics in the plain-Python twin, plus the LibreOffice + challenger matrix."""

import itertools

import pytest
from openpyxl import load_workbook

from engine.brief import BriefError
from engine.challenger import challenge, python_model
from engine.core import reference_run
from engine.recalc import recalc
from engine.spec import Timing
from engine.templates import get_template
from engine.templates.guest_house import _ramp, _season

GH = get_template("guest_house")
STAGES = ["planning", "early", "growing", "established"]
#: no ramp-up and occupancy far below the cap: every room-night asked for is sold
ROOMY = {"ramp_months": 0, "occ": 0.5, "peak_uplift": 0.3}


@pytest.fixture
def gh_brief():
    return GH.default_brief("planning")


def run(brief, scen=2, g="monthly", years=3):
    return reference_run(brief.values(scen), brief.roles, Timing(g, years), GH.reference_top, GH.loan_style)


def with_values(brief, **vals):
    d = brief.model_dump()
    for drv in d["drivers"]:
        if drv["key"] in vals:
            v = vals[drv["key"]]
            drv.update(low=v, base=v, high=v)
    return brief.model_validate(d)


def test_registered_with_wizard_hooks():
    assert GH.order == 140 and GH.business_types == ("guest_house",)
    assert GH.loan_style == "annuity"
    assert {"equity", "loan_amount"} <= set(GH.key_numbers) and 5 <= len(GH.key_numbers) <= 8
    assert GH.dashboard_metrics(Timing("monthly", 3)) and GH.dashboard_chart(Timing("monthly", 3))
    assert {"occ_rate", "adr_eff", "revpar"} <= set(GH.kpi_derived)


@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_are_valid_ordered_and_funded(stage):
    b = GH.default_brief(stage)
    low, base, high = (python_model(b, s) for s in (1, 2, 3))
    for y in (1, 2, 3):
        assert low["rev"][y] < base["rev"][y] < high["rev"][y]
        assert 0 < base["gp"][y] / base["rev"][y] < 1
    assert min(run(b)["cash_close"]) > 0  # the Base case is funded in every month; Low is the stress case


def test_needs_the_housekeeper_role(gh_brief):
    no_hk = gh_brief.model_copy(update={"roles": [r for r in gh_brief.roles if r.key != "housekeeper"]})
    with pytest.raises(BriefError, match="housekeeper"):
        GH.validate(no_hk)


def test_ramp_up_is_linear_and_annual_is_the_mean_of_its_months(gh_brief):
    d = {**gh_brief.values(2), "ramp_months": 18, "ramp_start": 0.4}
    m, a = Timing("monthly", 3), Timing("annual", 3)
    factors = [_ramp(d, p, m) for p in range(1, 37)]
    assert factors[0] == pytest.approx(0.4)
    assert factors[9] == pytest.approx(0.4 + 0.6 * 9 / 18)
    assert factors[18:] == pytest.approx([1.0] * 18)
    for y in (1, 2, 3):
        assert _ramp(d, y, a) == pytest.approx(sum(factors[(y - 1) * 12 : y * 12]) / 12)
    assert _ramp({**d, "ramp_months": 0}, 1, m) == 1.0


def test_season_peaks_in_summer_and_nets_to_zero_over_the_year(gh_brief):
    d = {**gh_brief.values(2), "peak_uplift": 0.3}
    f = [_season(d, p, Timing("monthly", 3)) for p in range(1, 13)]
    assert [i + 1 for i, v in enumerate(f) if v > 1] == [6, 7, 8]
    assert f[5] == pytest.approx(1.3) and f[0] == pytest.approx(0.9)
    assert sum(f) / 12 == pytest.approx(1.0)
    assert _season(d, 1, Timing("annual", 3)) == 1.0


def test_revenue_is_rooms_times_occupancy_times_adr_plus_extras(gh_brief):
    d = gh_brief.values(2)
    s = run(with_values(gh_brief, **ROOMY), g="annual")
    nights = d["rooms"] * d["nights_m"] * 12 * 0.5
    assert s["room_nights"][0] == pytest.approx(nights)
    assert s["rev_room"][0] == pytest.approx(nights * d["adr"])
    assert s["rev_extras"][0] == pytest.approx(nights * d["extras_night"])
    assert s["rev"][0] == pytest.approx(s["rev_room"][0] + s["rev_extras"][0])


def test_stays_are_room_nights_over_length_of_stay_and_guarded(gh_brief):
    d = gh_brief.values(2)
    s = run(with_values(gh_brief, **ROOMY))
    assert s["stays"][3] == pytest.approx(s["room_nights"][3] / d["los"])
    longer = run(with_values(gh_brief, **ROOMY, los=d["los"] * 2))
    assert longer["stays"][3] == pytest.approx(s["stays"][3] / 2)
    assert longer["cogs"][3] > s["cogs"][3]  # fewer stays, less cleaning cost, so the cost of sales is less negative
    floor = run(with_values(gh_brief, **ROOMY, los=1))
    assert floor["stays"] == pytest.approx(floor["room_nights"])  # a stay is at least one night


def test_length_of_stay_below_one_night_is_rejected(gh_brief):
    with pytest.raises(BriefError, match="los"):
        GH.validate(with_values(gh_brief, los=0.5))


def test_cost_of_sales_is_housekeeping_ota_extras_cleaning_and_supplies(gh_brief):
    d = gh_brief.values(2)
    s = run(with_values(gh_brief, **ROOMY))
    hk = d["housekeeper_fte"] * d["housekeeper_salary"] / 12 * (1 + d["payroll_tax"])
    expected = (
        -hk
        - s["rev_room"][0] * d["ota_share"] * d["ota_comm"]
        - s["rev_extras"][0] * d["extras_cost_pct"]
        - s["stays"][0] * d["clean_stay"]
        - s["room_nights"][0] * d["supplies_night"]
    )
    assert s["cogs"][0] == pytest.approx(expected)


def test_occupancy_never_exceeds_100_percent(gh_brief):
    b = with_values(gh_brief, occ=0.95, peak_uplift=0.5, occ_g=0.2, ramp_months=0)
    s = run(b)
    avail = b.values(2)["rooms"] * b.values(2)["nights_m"]
    assert max(s["room_nights"]) == pytest.approx(avail)  # peak months are capped at every room full
    assert all(n <= avail + 1e-9 for n in s["room_nights"])
    a = python_model(b, 2)
    assert a["room_nights"][3] <= a["avail_nights"][3] + 1e-9


def test_busier_in_summer_and_year_total_unchanged(gh_brief):
    s = run(with_values(gh_brief, **ROOMY))
    assert s["room_nights"][6] > s["room_nights"][2] > 0  # July vs March
    assert sum(s["room_nights"][:12]) == pytest.approx(
        run(with_values(gh_brief, **ROOMY), g="annual")["room_nights"][0]
    )


def test_rates_grow_each_year(gh_brief):
    d = gh_brief.values(2)
    s = run(with_values(gh_brief, **ROOMY, occ_g=0.1), g="annual")
    assert s["room_nights"][1] == pytest.approx(s["room_nights"][0] * 1.1)
    assert s["rev_room"][1] / s["room_nights"][1] == pytest.approx(d["adr"] * (1 + d["price_g"]))


def test_ota_commission_scales_with_share_and_rate(gh_brief):
    base = run(with_values(gh_brief, **ROOMY, ota_share=0.5, ota_comm=0.2))
    none = run(with_values(gh_brief, **ROOMY, ota_share=0.0))
    assert none["cogs"][5] - base["cogs"][5] == pytest.approx(base["rev_room"][5] * 0.1)


def test_land_share_of_the_purchase_is_not_depreciated(gh_brief):
    b = with_values(gh_brief, capex_buy=600_000, land_pct=0.3, ffe_room=0, capex_fit=0, capex_room=0)
    full = run(with_values(b, dep_years=1), g="annual")
    assert full["ppe_net"][0] == pytest.approx(180_000)  # only the land is left after the one-year write-off
    no_land = run(with_values(b, land_pct=0), g="annual")
    assert run(b, g="annual")["dep"][0] == pytest.approx(0.7 * no_land["dep"][0])


def test_leased_building_has_no_purchase_and_no_land(gh_brief):
    s = run(with_values(gh_brief, capex_buy=0, rent=5000))
    d = gh_brief.values(2)
    assert s["capex"][0] == pytest.approx(
        d["rooms"] * d["capex_room"] + d["capex_fit"] + d["rooms"] * d["ffe_room"] / 12
    )


def test_ffe_replacement_every_period_and_grows_with_costs(gh_brief):
    d = gh_brief.values(2)
    s = run(gh_brief, g="annual")
    repl = d["rooms"] * d["ffe_room"]
    assert s["capex"][1] == pytest.approx(repl * (1 + d["cost_g"]))


def test_property_loan_is_an_annuity_with_level_payments():
    b = with_values(GH.default_brief("planning"), loan_amount=300_000, loan_rate=0.06, loan_term_years=20)
    s = run(b)
    pay = [-x for x in s["debt_service"][:24]]
    assert pay[0] == pytest.approx(pay[23], rel=1e-6)
    r = 0.06 / 12
    assert pay[0] == pytest.approx(300_000 * r / (1 - (1 + r) ** -240), rel=1e-6)
    assert s["loan_close"][11] < 300_000


@pytest.mark.parametrize("stage", STAGES)
def test_monthly_and_annual_agree_where_volume_is_additive(stage):
    b = with_values(GH.default_brief(stage), **ROOMY)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    for k in ("room_nights", "stays", "rev", "cogs", "payroll_total", "capex"):
        for y in (1, 2, 3):
            assert m[k][y] == pytest.approx(a[k][y], abs=0.01), (k, y)


@pytest.mark.parametrize("ramp", [12, 18])
def test_monthly_and_annual_agree_through_the_ramp(ramp):
    # a ramp times a season is a product of two monthly shapes, so only a flat season keeps the year additive
    b = with_values(GH.default_brief("planning"), **{**ROOMY, "ramp_months": ramp, "peak_uplift": 0})
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    assert m["room_nights"][1] < m["room_nights"][2]
    for k in ("room_nights", "rev", "cogs"):
        for y in (1, 2, 3):
            assert m[k][y] == pytest.approx(a[k][y], abs=0.01), (k, y)


def test_annual_is_close_with_a_season_during_the_ramp():
    b = with_values(GH.default_brief("planning"), peak_uplift=0.25, ramp_months=10)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    assert a["rev"][1] == pytest.approx(m["rev"][1], rel=0.03)
    assert a["rev"][3] == pytest.approx(m["rev"][3], abs=0.01)  # ramp over: additive again


@pytest.mark.parametrize("g", ["monthly", "annual"])
@pytest.mark.parametrize("scen", [1, 2, 3])
def test_balance_sheet_balances(gh_brief, g, scen):
    s = run(gh_brief, scen, g)
    for ta, tle in zip(s["total_assets"], s["total_le"], strict=True):
        assert ta == pytest.approx(tle, abs=1e-6)


def test_degenerate_inputs_are_guarded(gh_brief):
    s = run(with_values(gh_brief, rooms=1, occ=0, ramp_months=0, los=1))
    assert all(v == 0 for v in s["room_nights"]) and all(v == 0 for v in s["stays"])
    assert all(v == 0 for v in s["rev"])


# ---------------------------------------------------------------------- LibreOffice + challenger
lo = pytest.mark.libreoffice
COMBOS = list(itertools.product(["light", "full"], ["monthly", "annual"], [1, 3, 5]))


@lo
@pytest.mark.parametrize("tier,g,years", COMBOS, ids=[f"{t}-{g}-{y}y" for t, g, y in COMBOS])
def test_matrix_passes_the_challenger(build, gh_brief, tier, g, years):
    brief = gh_brief.with_options(tier=tier, granularity=g, horizon_years=years)
    path = build(brief)
    res = recalc(path)
    assert res.ok, res.errors
    vals = load_workbook(path, data_only=True)
    verdict = [r[1].value for r in vals["Checks"].iter_rows(min_col=1, max_col=2) if r[0].value == "Overall verdict"]
    assert verdict and verdict[0] in ("PASS", "PASS WITH WARNINGS")
    report = challenge(brief, path)
    assert report["problems"] == []
    assert report["verdict"] == "PASS"
    assert any(k.startswith("kpi:") for k in report["scenarios"][2]["years"][1])


@lo
@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_pass_the_challenger(build, stage):
    brief = GH.default_brief(stage)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
    assert report["verdict"].startswith("PASS")


@lo
def test_full_house_owned_property_and_long_ramp_pass_the_challenger(build, gh_brief):
    brief = with_values(gh_brief, occ=0.95, peak_uplift=0.4, occ_g=0.05, los=1.5, ramp_months=18)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
