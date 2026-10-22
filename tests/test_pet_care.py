"""Pet care template: economics in the plain-Python twin, plus the LibreOffice + challenger matrix."""

import itertools

import pytest
from openpyxl import load_workbook

from engine.brief import BriefError
from engine.challenger import challenge, python_model
from engine.core import reference_run
from engine.recalc import recalc
from engine.spec import Timing
from engine.templates import get_template
from engine.templates.pet_care import _ramp, _season

PC = get_template("pet_care")
STAGES = ["planning", "early", "growing", "established"]
#: no ramp-up, plenty of stations and groomers, occupancy far below the cap: demand is always served
ROOMY = {"tables": 100, "grooms_groomer": 100, "ramp_months": 0, "occ": 0.5, "peak_uplift": 0.3}


@pytest.fixture
def pc_brief():
    return PC.default_brief("early")


def run(brief, scen=2, g="monthly", years=3):
    return reference_run(brief.values(scen), brief.roles, Timing(g, years), PC.reference_top)


def with_values(brief, **vals):
    d = brief.model_dump()
    for drv in d["drivers"]:
        if drv["key"] in vals:
            v = vals[drv["key"]]
            drv.update(low=v, base=v, high=v)
    return brief.model_validate(d)


def test_registered_with_wizard_hooks():
    assert PC.order == 130 and PC.business_types == ("pet_care",)
    assert "equity" in PC.key_numbers and 5 <= len(PC.key_numbers) <= 8
    assert PC.dashboard_metrics(Timing("monthly", 3)) and PC.dashboard_chart(Timing("monthly", 3))


@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_are_valid_ordered_and_funded(stage):
    b = PC.default_brief(stage)
    low, base, high = (python_model(b, s) for s in (1, 2, 3))
    for y in (1, 2, 3):
        assert low["rev"][y] < base["rev"][y] < high["rev"][y]
        assert 0 < base["gp"][y] / base["rev"][y] < 1
    monthly = reference_run(b.values(2), b.roles, Timing("monthly", 3), PC.reference_top)
    assert min(monthly["cash_close"]) > 0  # the Base case is funded in every month; Low is the stress case


def test_needs_the_groomer_role(pc_brief):
    no_groomer = pc_brief.model_copy(update={"roles": [r for r in pc_brief.roles if r.key != "groomer"]})
    with pytest.raises(BriefError, match="groomer"):
        PC.validate(no_groomer)


def test_ramp_up_is_linear_and_annual_is_the_mean_of_its_months(pc_brief):
    d = {**pc_brief.values(2), "ramp_months": 18, "ramp_start": 0.4}
    m, a = Timing("monthly", 3), Timing("annual", 3)
    factors = [_ramp(d, p, m) for p in range(1, 37)]
    assert factors[0] == pytest.approx(0.4)
    assert factors[9] == pytest.approx(0.4 + 0.6 * 9 / 18)
    assert factors[18:] == pytest.approx([1.0] * 18)
    for y in (1, 2, 3):
        assert _ramp(d, y, a) == pytest.approx(sum(factors[(y - 1) * 12 : y * 12]) / 12)
    assert _ramp({**d, "ramp_months": 0}, 1, m) == 1.0


def test_season_peaks_in_summer_and_december_and_leaves_the_year_unchanged(pc_brief):
    d = {**pc_brief.values(2), "peak_uplift": 0.2}
    m = Timing("monthly", 3)
    f = [_season(d, p, m) for p in range(1, 13)]
    assert [i + 1 for i, v in enumerate(f) if v > 1] == [6, 7, 8, 12]
    assert f[5] == pytest.approx(1.2) and f[0] == pytest.approx(0.9)
    assert sum(f) / 12 == pytest.approx(1.0)
    assert _season(d, 1, Timing("annual", 3)) == 1.0


def test_grooms_follow_demand_and_ramp(pc_brief):
    d = pc_brief.values(2)
    s = run(with_values(pc_brief, **{**ROOMY, "ramp_months": 12}))
    assert s["grooms"][0] == pytest.approx(d["grooms_day"] * d["days_open"] * d["ramp_start"])
    assert s["lost_grooms"][0] == pytest.approx(0)


def test_revenue_is_grooming_boarding_and_retail(pc_brief):
    d = pc_brief.values(2)
    s = run(with_values(pc_brief, **ROOMY), g="annual")
    grooms = d["grooms_day"] * d["days_open"] * 12
    nights = d["runs"] * d["nights_m"] * 12 * 0.5
    assert s["rev_groom"][0] == pytest.approx(grooms * d["groom_price"])
    assert s["rev_board"][0] == pytest.approx(nights * d["night_rate"])
    assert s["rev_retail"][0] == pytest.approx(grooms * d["retail_groom"])
    assert s["rev"][0] == pytest.approx(s["rev_groom"][0] + s["rev_board"][0] + s["rev_retail"][0])


def test_cost_of_sales_is_groomers_commission_supplies_boarding_and_retail_cost(pc_brief):
    d = pc_brief.values(2)
    s = run(with_values(pc_brief, **ROOMY))
    base = d["groomer_fte"] * d["groomer_salary"] / 12 * (1 + d["payroll_tax"])
    expected = (
        -base
        - s["rev_groom"][0] * d["commission_pct"]
        - s["grooms"][0] * d["supplies_groom"]
        - s["nights"][0] * d["board_night"]
        - s["rev_retail"][0] * (1 - d["retail_margin"])
    )
    assert s["cogs"][0] == pytest.approx(expected)


def test_station_capacity_caps_grooms(pc_brief):
    d = pc_brief.values(2)
    s = run(with_values(pc_brief, tables=1, grooms_groomer=100, ramp_months=0))
    cap = d["grooms_table"] * d["days_open"]
    assert max(s["grooms"]) == pytest.approx(cap)
    assert s["groom_cap"][5] == pytest.approx(cap)
    assert s["lost_grooms"][5] > 0
    roomy = run(with_values(pc_brief, **ROOMY))
    assert s["rev_groom"][5] < roomy["rev_groom"][5]


def test_groomer_capacity_caps_grooms_when_below_the_stations(pc_brief):
    d = pc_brief.values(2)
    s = run(with_values(pc_brief, tables=100, grooms_day=100, ramp_months=0))
    cap = d["groomer_fte"] * d["grooms_groomer"] * d["days_open"]
    assert s["grooms"][5] == pytest.approx(cap)
    more = run(
        with_values(pc_brief, tables=100, grooms_day=100, ramp_months=0, grooms_groomer=d["grooms_groomer"] * 1.1)
    )
    assert more["grooms"][5] == pytest.approx(cap * 1.1)


def test_occupancy_never_exceeds_100_percent(pc_brief):
    s = run(with_values(pc_brief, occ=0.95, peak_uplift=0.5, occ_g=0.2, ramp_months=0))
    d = pc_brief.values(2)
    avail = d["runs"] * d["nights_m"]
    assert max(s["nights"]) == pytest.approx(avail)  # peak months are capped at every run full
    assert all(n <= avail + 1e-9 for n in s["nights"])
    b = python_model(with_values(pc_brief, occ=0.95, occ_g=0.2, ramp_months=0), 2)
    assert b["nights"][3] <= b["avail_nights"][3] + 1e-9


def test_boarding_is_busier_in_peak_months(pc_brief):
    s = run(with_values(pc_brief, **ROOMY))
    assert s["nights"][6] > s["nights"][2] > 0  # July vs March
    assert s["nights"][11] > s["nights"][10]  # December vs November
    assert sum(s["nights"][:12]) == pytest.approx(run(with_values(pc_brief, **ROOMY), g="annual")["nights"][0])


def test_volume_and_prices_grow_each_year(pc_brief):
    d = pc_brief.values(2)
    s = run(with_values(pc_brief, **ROOMY, groom_g=0.1), g="annual")
    assert s["grooms"][1] == pytest.approx(s["grooms"][0] * 1.1)
    assert s["grooms"][2] == pytest.approx(s["grooms"][0] * 1.21)
    assert s["rev_groom"][1] / s["grooms"][1] == pytest.approx(d["groom_price"] * (1 + d["price_g"]))


def test_land_is_not_depreciated(pc_brief):
    b = with_values(pc_brief, capex_land=500_000, capex_refresh=0)
    no_land = with_values(b, capex_land=0)
    s, n = run(b, g="annual"), run(no_land, g="annual")
    assert s["capex"][0] - n["capex"][0] == pytest.approx(500_000)
    assert s["dep"] == pytest.approx(n["dep"])
    full = run(with_values(b, dep_years=1), g="annual")
    assert full["ppe_net"][-1] == pytest.approx(500_000)


@pytest.mark.parametrize("stage", STAGES)
def test_monthly_and_annual_agree_where_volume_is_additive(stage):
    b = with_values(PC.default_brief(stage), **ROOMY)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    for k in ("grooms", "nights", "rev", "cogs", "payroll_total", "capex"):
        for y in (1, 2, 3):
            assert m[k][y] == pytest.approx(a[k][y], abs=0.01), (k, y)


@pytest.mark.parametrize("ramp", [12, 18])
def test_monthly_and_annual_agree_through_the_ramp(ramp):
    # without a season: ramp x season is a product of two monthly shapes, so only a flat season keeps the year additive
    b = with_values(PC.default_brief("planning"), **{**ROOMY, "ramp_months": ramp, "peak_uplift": 0})
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    assert m["grooms"][1] < m["grooms"][2]
    for k in ("grooms", "nights", "rev", "cogs"):
        for y in (1, 2, 3):
            assert m[k][y] == pytest.approx(a[k][y], abs=0.01), (k, y)


@pytest.mark.parametrize("ramp", [10, 18])
def test_annual_is_close_with_a_season_during_the_ramp(ramp):
    """Exact without a season (above); with one the ramp and the season interact in year 1, so only close."""
    b = with_values(PC.default_brief("planning"), peak_uplift=0.2, ramp_months=ramp)
    m, a = python_model(b, 2), python_model(b.with_options(granularity="annual"), 2)
    for k in ("nights", "rev"):
        assert a[k][1] == pytest.approx(m[k][1], rel=0.03), k
        assert a[k][1] != pytest.approx(m[k][1], rel=1e-6), k
        assert a[k][3] == pytest.approx(m[k][3], abs=0.01), k  # ramp over: additive again


def test_capex_scales_in_month_one_and_refresh_from_year_two(pc_brief):
    d = pc_brief.values(2)
    s = run(pc_brief)
    assert s["capex"][0] == pytest.approx(d["tables"] * d["capex_table"] + d["runs"] * d["capex_run"] + d["capex_fit"])
    assert s["capex"][1] == 0 and s["capex"][11] == 0
    assert s["capex"][12] == pytest.approx(d["capex_refresh"])
    assert run(pc_brief, g="annual")["capex"][1] == pytest.approx(d["capex_refresh"])


@pytest.mark.parametrize("g", ["monthly", "annual"])
@pytest.mark.parametrize("scen", [1, 2, 3])
def test_balance_sheet_balances(pc_brief, g, scen):
    s = run(pc_brief, scen, g)
    for ta, tle in zip(s["total_assets"], s["total_le"], strict=True):
        assert ta == pytest.approx(tle, abs=1e-6)


def test_degenerate_inputs_are_guarded(pc_brief):
    s = run(with_values(pc_brief, grooms_day=0, occ=0, ramp_months=0))
    assert all(v == 0 for v in s["grooms"]) and all(v == 0 for v in s["nights"])
    assert all(v == 0 for v in s["rev"])


# ---------------------------------------------------------------------- LibreOffice + challenger
lo = pytest.mark.libreoffice
COMBOS = list(itertools.product(["light", "full"], ["monthly", "annual"], [1, 3, 5]))


@lo
@pytest.mark.parametrize("tier,g,years", COMBOS, ids=[f"{t}-{g}-{y}y" for t, g, y in COMBOS])
def test_matrix_passes_the_challenger(build, pc_brief, tier, g, years):
    brief = pc_brief.with_options(tier=tier, granularity=g, horizon_years=years)
    path = build(brief)
    res = recalc(path)
    assert res.ok, res.errors
    vals = load_workbook(path, data_only=True)
    verdict = [r[1].value for r in vals["Checks"].iter_rows(min_col=1, max_col=2) if r[0].value == "Overall verdict"]
    assert verdict and verdict[0] in ("PASS", "PASS WITH WARNINGS")
    report = challenge(brief, path)
    assert report["problems"] == []
    assert report["verdict"] == "PASS"
    assert any(k.startswith("kpi:") for k in report["scenarios"][2]["years"][1])


@lo
@pytest.mark.parametrize("stage", STAGES)
def test_stage_defaults_pass_the_challenger(build, stage):
    brief = PC.default_brief(stage)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
    assert report["verdict"].startswith("PASS")


@lo
def test_binding_capacity_full_kennel_and_land_pass_the_challenger(build, pc_brief):
    brief = with_values(pc_brief, tables=2, occ=0.9, peak_uplift=0.4, capex_land=400_000, groom_g=0.05)
    report = challenge(brief, build(brief))
    assert report["problems"] == []
