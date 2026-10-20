"""Build + verify the standard-v1 prototype. Run from the repo root with the venv."""

import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, ".")
from openpyxl import load_workbook  # noqa: E402

from engine.brief import Brief  # noqa: E402
from engine.challenger import challenge  # noqa: E402
from engine.recalc import recalc  # noqa: E402

HERE = Path(__file__).resolve().parent
P = Path(os.environ.get("PROTO_DIR", "var/prototype-standard-v1"))
P.mkdir(parents=True, exist_ok=True)
RAW = P / "raw.xlsx"
from engine.build import build_model  # noqa: E402
from engine.goals import apply_goal  # noqa: E402
from engine.templates import TEMPLATES  # noqa: E402

_tpl = TEMPLATES["saas"]
brief = apply_goal(_tpl.default_brief("established").with_options(tier="full", granularity="monthly", horizon_years=5),
                   "investor_pitch")
build_model(brief, RAW)

# 0. the engine file itself passes the challenger (all 3 scenarios)
rep = challenge(brief, RAW)
print("challenger on engine file:", rep["verdict"], rep["problems"][:3])


def run(src, scen, name):
    dst = P / f"{name}-s{scen}.xlsx"
    shutil.copy(src, dst)
    wb = load_workbook(dst)
    wb["Dashboard"]["B4"] = scen
    wb.save(dst)
    res = recalc(dst, 120)
    assert not res.errors, (name, scen, res.errors)
    return load_workbook(dst, data_only=True)


ENGINE_SHEETS = ["P&L", "Cash Flow", "Balance Sheet", "Break-even", "Valuation", "Monthly", "Headcount", "Inputs"]
raw = {s: run(RAW, s, "raw") for s in (1, 2, 3)}

# 1. snapshot for the Scenarios sheet, from the engine's own recalculated values
snap = {}
for s, wb in raw.items():
    pl, bsh, be, va, mo = wb["P&L"], wb["Balance Sheet"], wb["Break-even"], wb["Valuation"], wb["Monthly"]

    def row(ws, label):
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value == label:
                return r
        raise KeyError(label)

    rv, eb, ni, ca = row(pl, "Revenue"), row(pl, "EBITDA"), row(pl, "Net income"), row(bsh, "Cash")
    em = row(pl, "EBITDA margin")
    cl = row(mo, "Closing cash")
    yrs = "BCDEF"
    snap[str(s)] = {
        "rev": [pl[f"{c}{rv}"].value for c in yrs], "ebitda": [pl[f"{c}{eb}"].value for c in yrs],
        "cash": [bsh[f"{c}{ca}"].value for c in yrs],
        "ni": [pl[f"{c}{ni}"].value for c in yrs],
        "rev_total": sum(pl[f"{c}{rv}"].value for c in yrs), "rev_y5": pl[f"F{rv}"].value,
        "ebitda_total": sum(pl[f"{c}{eb}"].value for c in yrs), "ebitdam_y5": pl[f"F{em}"].value,
        "ni_total": sum(pl[f"{c}{ni}"].value for c in yrs), "cash_y5": bsh[f"F{ca}"].value,
        "min_cash": min(mo.cell(cl, c).value for c in range(3, 63)), "funding": be["B12"].value, "ev": va["B30"].value,
    }
(P / "snap.json").write_text(json.dumps(snap, indent=1))

# 2. build the prototype
pre = P / "proto-pre.xlsx"
subprocess.run([sys.executable, str(HERE / "proto_build.py"), str(RAW), str(pre), str(P / "snap.json")], check=True)

# 3. equivalence: every engine cell equals the engine's own value in all 3 scenarios
bad = 0
checked = 0
for s in (1, 2, 3):
    pw = run(pre, s, "proto")
    rw = raw[s]
    for sh in ENGINE_SHEETS:
        a, b = rw[sh], pw[sh]
        for r in range(3, a.max_row + 1):
            for c in range(2, a.max_column + 1):
                va_, vb_ = a.cell(r, c).value, b.cell(r, c).value
                if isinstance(va_, (int, float)):
                    checked += 1
                    if not isinstance(vb_, (int, float)) or abs(va_ - vb_) > 0.005:
                        bad += 1
                        if bad < 6:
                            print("DIFF", s, sh, a.cell(r, c).coordinate, va_, vb_)
    # renamed KPI block
    a, b = rw["Dashboard"], pw["KPIs"]
    for r in range(25, 45):
        for c in range(2, 7):
            va_, vb_ = a.cell(r, c).value, b.cell(r, c).value
            if isinstance(va_, (int, float)):
                checked += 1
                if not isinstance(vb_, (int, float)) or abs(va_ - vb_) > 0.005:
                    bad += 1
    ch = pw["Checks"]
    print(f"scenario {s}: verdict {ch['B22'].value}; S&U {ch['B21'].value}; errors {ch['B18'].value}; "
          f"summary tie {ch['B17'].value}; scenarios Active=snap: "
          f"{abs(pw['Scenarios']['E' + str(19)].value - snap[str(s)]['rev_total']) < 1 if pw['Scenarios']['E19'].value is not None else '?'}")
print(f"equivalence: {checked} engine cells compared across 3 scenarios, {bad} differ")

# 4. final file = base scenario, recalculated like production
final = P / "SaaS_model_standard_v1_prototype.xlsx"
shutil.copy(pre, final)
res = recalc(final, 120)
print("final recalc errors:", res.errors)
