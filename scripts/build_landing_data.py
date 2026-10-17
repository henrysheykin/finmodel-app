"""Regenerate the landing page's live sheet preview and the cafe sample from the real engine.

    python scripts/build_landing_data.py

Builds the cafe brief as a Full, monthly, 3-year workbook, recalculates it in LibreOffice once per
scenario and writes:
  - app/data/landing_sheet.json  (numbers, check results and formulas shown in the hero preview)
  - samples/cafe-full.xlsx       (the downloadable sample, Base case)
tests/test_landing.py checks the JSON against the engine's independent Python model, so the
marketing preview cannot drift from what the product actually builds.
"""

from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from engine.brief import Brief  # noqa: E402
from engine.build import build_model  # noqa: E402
from engine.recalc import recalc  # noqa: E402
from engine.templates import get_template  # noqa: E402

INPUT_KEYS = ["cust_full", "ticket", "traffic_g", "cogs_pct", "rent", "staff_fte", "staff_salary"]
FORMULA_CELLS = {
    "sum": ("Dashboard", "B8"),
    "inp": ("Inputs", None),  # first formula in the Inputs sheet
    "pl": ("P&L", "B5"),
    "chk": ("Checks", "B16"),
}


def _rows(ws, first: int, last: int, years: int) -> list[dict]:
    out = []
    for r in range(first, last + 1):
        label = ws.cell(r, 1).value
        if not label:
            continue
        vals = [ws.cell(r, 2 + i).value for i in range(years)]
        if all(v is None for v in vals):
            continue
        out.append({"label": str(label), "values": [round(v, 4) if isinstance(v, float) else v for v in vals]})
    return out


def _checks(ws) -> list[dict]:
    out = []
    for r in range(4, ws.max_row + 1):
        label, result, what = (ws.cell(r, c).value for c in (1, 2, 3))
        if label:
            out.append({"label": str(label), "result": str(result), "what": what or ""})
    return out


def main() -> int:
    brief = Brief.load(ROOT / "engine" / "briefs" / "cafe.json").with_options(
        tier="full", granularity="monthly", horizon_years=3
    )
    tpl = get_template(brief.template)
    specs = {s.key: s for s in tpl.all_specs(brief.roles)}
    drivers = {d.key: d for d in brief.drivers}
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "cafe.xlsx"
        build_model(brief, src)
        formulas = load_workbook(src)
        fcells = {}
        for tab, (sheet, cell) in FORMULA_CELLS.items():
            ws = formulas[sheet]
            if cell is None:
                cell = next(
                    c.coordinate
                    for row in ws.iter_rows()
                    for c in row
                    if isinstance(c.value, str) and c.value[:1] == "="
                )
            fcells[tab] = {"cell": f"{sheet}!{cell}", "formula": str(ws[cell].value)}
        scenarios = {}
        for scen in (1, 2, 3):
            path = Path(tmp) / f"s{scen}.xlsx"
            shutil.copy(src, path)
            wb = load_workbook(path)
            wb["Dashboard"]["B4"] = scen
            wb.save(path)
            res = recalc(path, 120)
            if res.errors:
                raise SystemExit(f"scenario {scen}: formula errors {res.errors}")
            wv = load_workbook(path, data_only=True)
            scenarios[str(scen)] = {
                "summary": _rows(wv["Dashboard"], 8, 14, brief.horizon_years),
                "pl": _rows(wv["P&L"], 5, 22, brief.horizon_years),
                "checks": _checks(wv["Checks"]),
            }
            if scen == 2:
                (ROOT / "samples").mkdir(exist_ok=True)
                shutil.copy(path, ROOT / "samples" / "cafe-full.xlsx")
    data = {
        "title": "Neighborhood cafe, three-year plan",
        "template": brief.template,
        "years": brief.horizon_years,
        "inputs": [
            {
                "key": k,
                "label": specs[k].label,
                "unit": specs[k].unit,
                "fmt": specs[k].fmt,
                "values": [drivers[k].low, drivers[k].base, drivers[k].high],
            }
            for k in INPUT_KEYS
        ],
        "formulas": fcells,
        "scenarios": scenarios,
    }
    out = ROOT / "app" / "data" / "landing_sheet.json"
    out.parent.mkdir(exist_ok=True)
    out.write_text(json.dumps(data, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"wrote {out.relative_to(ROOT)} and samples/cafe-full.xlsx")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
