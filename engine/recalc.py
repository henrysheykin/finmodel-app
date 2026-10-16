"""Headless LibreOffice recalculation.

openpyxl writes formulas without cached values. This module opens a workbook in
LibreOffice (headless, throwaway user profile), forces a full recalculation on
load, saves it back in place as .xlsx, and scans the result for error values.

Usage: python -m engine.recalc model.xlsx [timeout_seconds]
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
from dataclasses import asdict, dataclass, field
from pathlib import Path

from openpyxl import load_workbook

ERROR_VALUES = ("#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "Err:")
MAX_LOCATIONS = 100

# Profile settings: always recalculate OOXML/ODF on load, never prompt.
_SETTINGS = [
    ("/org.openoffice.Office.Calc/Formula/Load", "OOXMLRecalcMode", "0"),
    ("/org.openoffice.Office.Calc/Formula/Load", "ODFRecalcMode", "0"),
    ("/org.openoffice.Office.Common/Misc", "FirstRun", "false"),
]
_REGISTRY = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
    'xmlns:xs="http://www.w3.org/2001/XMLSchema" '
    'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
    + "".join(
        f'<item oor:path="{path}"><prop oor:name="{name}" oor:op="fuse"><value>{val}</value></prop></item>\n'
        for path, name, val in _SETTINGS
    )
    + "</oor:items>\n"
)


class RecalcError(RuntimeError):
    """LibreOffice could not recalculate the file (nothing was written)."""


@dataclass
class RecalcResult:
    status: str  # "success" | "errors_found"
    total_formulas: int
    total_errors: int
    errors: dict[str, list[str]] = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return self.status == "success"

    def to_dict(self) -> dict:
        return asdict(self)


def find_soffice() -> str:
    for name in (os.environ.get("SOFFICE_BIN"), "soffice", "libreoffice"):
        if name and shutil.which(name):
            return shutil.which(name)  # type: ignore[return-value]
    raise RecalcError("soffice not found on PATH; LibreOffice is required to recalculate")


def _run_libreoffice(src: Path, outdir: Path, timeout: int) -> Path:
    soffice = find_soffice()
    with tempfile.TemporaryDirectory(prefix="lo-profile-") as prof:
        user = Path(prof) / "user"
        user.mkdir(parents=True)
        (user / "registrymodifications.xcu").write_text(_REGISTRY, encoding="utf-8")
        cmd = [
            soffice,
            f"-env:UserInstallation={Path(prof).as_uri()}",
            "--headless",
            "--norestore",
            "--nologo",
            "--nodefault",
            "--convert-to",
            "xlsx:Calc MS Excel 2007 XML",
            "--outdir",
            str(outdir),
            str(src),
        ]
        env = dict(os.environ, SAL_USE_VCLPLUGIN="svp", HOME=prof)
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, env=env)
        except subprocess.TimeoutExpired as exc:
            raise RecalcError(f"LibreOffice timed out after {timeout}s") from exc
    out = outdir / (src.stem + ".xlsx")
    if proc.returncode != 0 or not out.exists():
        raise RecalcError(f"LibreOffice failed (rc={proc.returncode}): {proc.stderr.strip()[:500]}")
    return out


def scan_errors(path: str | Path) -> RecalcResult:
    """Count formulas and error values in an already-recalculated workbook."""
    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    total_formulas = 0
    errors: dict[str, list[str]] = {}
    total_errors = 0
    for ws in formulas.worksheets:
        vws = values[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    total_formulas += 1
                cached = vws[cell.coordinate].value
                if isinstance(cached, str) and cached.startswith(ERROR_VALUES):
                    total_errors += 1
                    key = cached.split(":")[0] if cached.startswith("Err:") else cached
                    locs = errors.setdefault(key, [])
                    if len(locs) < MAX_LOCATIONS:
                        locs.append(f"{ws.title}!{cell.coordinate}")
    formulas.close()
    values.close()
    status = "success" if total_errors == 0 else "errors_found"
    return RecalcResult(status, total_formulas, total_errors, errors)


def recalc(path: str | Path, timeout: int = 60) -> RecalcResult:
    """Recalculate ``path`` in place and report formula errors.

    Raises RecalcError if LibreOffice is missing or fails (file left untouched).
    """
    src = Path(path).resolve()
    if not src.exists():
        raise RecalcError(f"File {src} does not exist")
    with tempfile.TemporaryDirectory(prefix="recalc-") as tmp:
        tmpdir = Path(tmp)
        work = tmpdir / "in" / src.name
        work.parent.mkdir()
        shutil.copy(src, work)
        outdir = tmpdir / "out"
        outdir.mkdir()
        out = _run_libreoffice(work, outdir, timeout)
        shutil.copy(out, src)
    return scan_errors(src)


def main(argv: list[str]) -> int:
    if not argv:
        print("usage: python -m engine.recalc file.xlsx [timeout]", file=sys.stderr)
        return 2
    timeout = int(argv[1]) if len(argv) > 1 else 60
    try:
        res = recalc(argv[0], timeout)
    except RecalcError as exc:
        print(json.dumps({"error": str(exc)}))
        return 1
    print(json.dumps(res.to_dict(), indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
