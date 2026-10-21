"""Independent reviewer for a generated model.

1. Recalculates the workbook once per scenario (1/2/3) in LibreOffice; any
   formula error fails the model.
2. Recomputes the same economics in plain Python from the brief (the
   template's ``reference_top`` + ``engine.core.reference_run``, which share no
   code with the formula builder) and compares every scenario-year within $1.
   Tier modules (Break-even, Valuation) are recomputed by ``engine.goals.reference`` and the
   property Returns sheet by ``engine.returns.reference``; both are compared cell by cell (money within $1,
   ratios within 1e-6, IRR within 0.01 pp, periods and texts exactly).
3. Requires the in-file Checks verdict to be PASS (or PASS WITH WARNINGS), the
   balance sheet to balance (full tier), and revenue Low < Base < High.

Usage: python -m engine.challenger brief.json model.xlsx    (exit 0 = PASS)
"""

from __future__ import annotations

import json
import math
import re
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from engine.brief import Brief
from engine.builder import manifest_for, plan
from engine.core import annualise, reference_run
from engine.goals import TOLERANCE as GOAL_TOLERANCE
from engine.goals import apply_goal, dashboard_reference
from engine.goals import reference as goal_reference
from engine.recalc import RecalcError, recalc
from engine.returns import TOLERANCE as RETURNS_TOLERANCE
from engine.returns import reference as returns_reference
from engine.spec import Timing
from engine.templates import get_template

TOLERANCE_USD = 1.0

# workbook metric -> python series (sign convention identical: costs negative)
METRIC_SOURCE = {
    "rev": "rev",
    "gp": "gp",
    "ebitda": "ebitda",
    "ni": "ni",
    "cash": "cash_close",
    "loan": "loan_close",
    "total_assets": "total_assets",
    "total_le": "total_assets",
    "cfo": "cfo",
}


#: Summary sheet cells (manifest ``summary_years``) -> python series
SUMMARY_SOURCE = {"sum_rev": "rev", "sum_ni": "ni", "sum_cash": "cash_close"}


def python_model(brief: Brief, scenario: int) -> dict[str, dict[int, float]]:
    tpl = get_template(brief.template)
    t = Timing(brief.granularity, brief.horizon_years)
    series = reference_run(brief.values(scenario), brief.roles, t, tpl.reference_top, tpl.loan_style)
    return annualise(series, t, tpl.kpi_flows)


def python_goal_model(brief: Brief, scenario: int) -> dict:
    """Tier-module values (Break-even, Valuation) and every Dashboard tile / cell recomputed in plain Python."""
    brief = apply_goal(brief)
    tpl = get_template(brief.template)
    t = Timing(brief.granularity, brief.horizon_years)
    d = brief.values(scenario)
    series = reference_run(d, brief.roles, t, tpl.reference_top, tpl.loan_style)
    out = goal_reference(brief, d, series, t)
    dash = dashboard_reference(brief, d, series, t, out["cells"])
    out["cells"].update(dash["cells"])
    out["years"].update(dash["years"])
    from engine.kpis import twin

    out["years"].update(twin(brief, d, series, t))  # KPIs sheet (E31) and the Dashboard's industry lines
    if tpl.returns:
        out["cells"].update(returns_reference(d, series, t)["cells"])
    return out


def _tolerance(name: str, exp=None) -> float:
    """Goal / Returns cell tolerance by name (``dscr_3`` → ``dscr``); money cells default to $1.

    KPI-sheet cells and the Dashboard's industry lines (E31) are ratios and unit figures: 1e-6, relative to the value.
    """
    if name.startswith(("kpi_", "dash_m_")):
        return 1e-6 * max(1.0, abs(exp)) if isinstance(exp, int | float) else 0.0
    base = name.split(" ")[0]
    stem = re.sub(r"_\d+$", "", base)
    for key in (base, stem):
        if key in RETURNS_TOLERANCE:
            return RETURNS_TOLERANCE[key]
    return GOAL_TOLERANCE.get(base, TOLERANCE_USD)


def _same(got, exp, tol: float) -> bool:
    if isinstance(exp, str):
        return got == exp
    return _num(got) and abs(got - exp) <= tol


def _num(v) -> bool:
    return isinstance(v, int | float) and not isinstance(v, bool) and math.isfinite(v)


def _split(ref: str) -> tuple[str, str]:
    sheet, cell = ref.rsplit("!", 1)
    return sheet.strip("'"), cell.replace("$", "")


def _run_scenario(src: Path, scen: int, selector: str, workdir: Path):
    path = workdir / f"s{scen}.xlsx"
    shutil.copy(src, path)
    wb = load_workbook(path)
    sheet, cell = _split(selector)
    wb[sheet][cell] = scen
    wb.save(path)
    try:
        res = recalc(path, 90).to_dict()
    except RecalcError as exc:
        res = {"error": str(exc)}
    return res, load_workbook(path, data_only=True)


def challenge(brief: Brief | dict, model_path: str | Path) -> dict:
    brief = brief if isinstance(brief, Brief) else Brief.model_validate(brief)
    brief = apply_goal(brief)  # same tier module drivers the build added (ADR-013)
    tpl = get_template(brief.template)
    pl = plan(brief, tpl)
    man = manifest_for(pl, verdict="")  # cell addresses only
    verdict_cell = _verdict_cell(model_path)
    report: dict = {
        "template": brief.template,
        "tier": brief.tier,
        "granularity": brief.granularity,
        "years": brief.horizon_years,
        "scenarios": {},
        "problems": [],
    }
    revenue: dict[int, list] = {}
    with tempfile.TemporaryDirectory() as tmp:
        for scen in (1, 2, 3):
            res, wb = _run_scenario(Path(model_path), scen, man.selector, Path(tmp))
            if res.get("status") != "success":
                report["problems"].append(f"scenario {scen}: recalc {res}")
            expected = python_model(brief, scen)

            def val(ref: str, wb=wb):
                sheet, cell = _split(ref)
                return wb[sheet][cell].value

            verdict = val(verdict_cell) if verdict_cell else None
            out = {"verdict": verdict, "years": {}}
            for y in range(1, brief.horizon_years + 1):
                got = {}
                for metric, cells in man.metrics.items():
                    v = val(cells[y])
                    if not _num(v):
                        report["problems"].append(f"scenario {scen} year {y} {metric}: workbook value {v!r}")
                        got[metric] = None
                        continue
                    src = tpl.kpi_checked[metric[4:]] if metric.startswith("kpi:") else METRIC_SOURCE[metric]
                    exp = expected[src][y]
                    if abs(v - exp) > TOLERANCE_USD:
                        report["problems"].append(
                            f"scenario {scen} year {y} {metric}: workbook {v:,.2f} vs python {exp:,.2f}"
                        )
                    got[metric] = round(v)
                if "total_assets" in got and None not in (got["total_assets"], got["total_le"]):
                    ta, tle = val(man.metrics["total_assets"][y]), val(man.metrics["total_le"][y])
                    if abs(ta - tle) > 0.5:
                        report["problems"].append(f"scenario {scen} year {y}: balance sheet off by {ta - tle:,.2f}")
                for name, cells in man.summary_years.items():  # Summary (E28): its links equal the reference
                    v = val(cells[y])
                    exp = expected[SUMMARY_SOURCE[name]][y]
                    if not _num(v) or abs(v - exp) > TOLERANCE_USD:
                        report["problems"].append(
                            f"scenario {scen} year {y} {name}: workbook {v!r} vs python {exp:,.2f}"
                        )
                out["years"][y] = got
            if man.goal_cells or man.goal_years or man.dash_cells or man.dash_years or man.kpi_years:
                gexp = python_goal_model(brief, scen)
                checked = [(n, ref, gexp["cells"][n]) for n, ref in {**man.goal_cells, **man.dash_cells}.items()]
                checked += [
                    (f"{n} year {y}", ref, gexp["years"][n][y])
                    for n, refs in {**man.goal_years, **man.dash_years, **man.kpi_years}.items()
                    for y, ref in refs.items()
                ]
                for name, ref, exp in checked:
                    got = val(ref)
                    if not _same(got, exp, _tolerance(name, exp)):
                        report["problems"].append(f"scenario {scen} goal {name}: workbook {got!r} vs python {exp!r}")
                out["goal"] = {n: val(ref) for n, ref in man.goal_cells.items()}
                out["dash"] = {n: val(ref) for n, ref in man.dash_cells.items()}
            if verdict not in ("PASS", "PASS WITH WARNINGS"):
                report["problems"].append(f"scenario {scen}: in-file checks {verdict!r}")
            revenue[scen] = [out["years"][y].get("rev") for y in range(1, brief.horizon_years + 1)]
            report["scenarios"][scen] = out
    for i in range(brief.horizon_years):
        lo, ba, hi = revenue[1][i], revenue[2][i], revenue[3][i]
        if None in (lo, ba, hi):
            report["problems"].append(f"year {i + 1}: revenue missing in a scenario")
        elif not lo < ba < hi:
            report["problems"].append(f"year {i + 1}: revenue not Low < Base < High ({lo}, {ba}, {hi})")
    report["verdict"] = "PASS" if not report["problems"] else "FAIL"
    return report


def _verdict_cell(model_path) -> str | None:
    """The Overall verdict cell is located by its label, not trusted from the builder."""
    wb = load_workbook(model_path)
    if "Checks" not in wb.sheetnames:
        return None
    ws = wb["Checks"]
    for row in ws.iter_rows(min_col=1, max_col=1):
        c = row[0]
        if c.value == "Overall verdict":
            return f"Checks!B{c.row}"
    return None


def main(argv: list[str]) -> int:
    brief = Brief.load(argv[0])
    report = challenge(brief, argv[1])
    print(json.dumps(report, indent=2, ensure_ascii=False))
    return 0 if report["verdict"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
