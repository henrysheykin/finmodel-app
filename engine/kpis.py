"""KPIs & ratios sheet (E31, docs/WORKBOOK_STANDARD.md §1, §3, §5): Full tier only.

Top to bottom, on the left: Financial ratios (growth, margins, cost shares, liquidity, leverage, returns), the
DuPont decomposition of return on equity, then the template's KPI table (moved here from the Dashboard). The last
column of the first and third table is "Target", a blank input column (a view setting, standard §5: it changes no
figure). Four charts sit to the right, 2 x 2: margins, ROA / ROE and two chosen by the template
(``Template.kpi_charts``).

A ratio whose denominator is zero or negative (no revenue, no assets, equity at or below zero) shows the text
"n/a" instead of a misleading number. The sheet is links and simple ratios over the statements, so the challenger
covers it by recomputing every ratio and the DuPont identity in plain Python (``reference`` below, which shares no
code with the formula writer); the in-file check "DuPont ties to return on equity" proves the identity in the file.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as L

from engine import charts as ch
from engine import style as st
from engine.builder import Plan, _rollup
from engine.dashboard import FMT, YEAR_FMT
from engine.spec import KpiChart

SHEET = "KPIs"
BAND = 4  # "Financial ratios" band; the year header is the next row
HDR = BAND + 1
NA = "n/a"
#: charts: first row, rows per chart row (standard 8 grid cells x 16 rows + one blank)
CHART_ROW, CHART_STEP = BAND, 17
CHART_COLS = 8  # a standard chart is 8 grid columns wide
GAP = 2  # width of the gap column between the tables and the charts
TOL = 1e-9  # DuPont: |net margin x turnover x multiplier - ROE| per year
AXIS = {
    "usd": '[>=1000000]$#,##0.0,,"M";[>=1000]$#,##0,"k";$#,##0',
    "usd2": "$#,##0.00",
    "int": "#,##0",
    "dec1": "#,##0.0",
    "pct": "0%",
    "x": '0.0"x"',
}
CELL_FMT = {**FMT}  # ratios use the report formats; the KPI table keeps the engine's own (full dollars)

#: financial ratios: (key, label, format); the marketing row is dropped for a template without one
RATIOS = (
    ("growth", "Revenue growth", "pct"),
    ("gm", "Gross margin", "pct"),
    ("ebitda_m", "EBITDA margin", "pct"),
    ("ni_m", "Net margin", "pct"),
    ("payroll_pct", "Payroll, % of revenue", "pct"),
    ("mkt_pct", "Marketing, % of revenue", "pct"),
    ("opex_pct", "Operating expenses, % of revenue", "pct"),
    ("current", "Current ratio (current assets ÷ current liabilities)", "x"),
    ("debt_eq", "Debt ÷ equity (loan ÷ total equity)", "x"),
    ("roa", "Return on assets (net income ÷ year-end assets)", "pct"),
    ("roe", "Return on equity (net income ÷ year-end equity)", "pct"),
)
DUPONT = (
    ("dp_nm", "Net margin", "pct"),
    ("dp_turn", "× Asset turnover (revenue ÷ assets)", "x"),
    ("dp_mult", "× Equity multiplier (assets ÷ equity)", "x"),
    ("dp_roe", "Return on equity (product of the three)", "pct"),
    ("dp_diff", "Check: product − return on equity (must be 0)", "pct"),
)
DUPONT_TITLE = "DuPont analysis — return on equity = net margin × asset turnover × equity multiplier"


def ratio_keys(tpl) -> list[str]:
    return [k for k, *_ in RATIOS if k != "mkt_pct" or tpl.marketing_key]


def block_start(tpl) -> int:
    """Year row of the template KPI table: below the ratios and the DuPont table, one blank row and its band."""
    dp_last = HDR + len(ratio_keys(tpl)) + 2 + 1 + len(DUPONT)
    return dp_last + 3


@dataclass
class KpiLayout:
    years: int
    ratios: dict[str, int] = field(default_factory=dict)
    dupont: dict[str, int] = field(default_factory=dict)
    dp_band: int = 0
    dp_hdr: int = 0
    end: int = 0
    cells: dict[str, dict[int, str]] = field(default_factory=dict)  # manifest name -> {year: 'KPIs'!B7}

    @property
    def target_col(self) -> int:
        return 2 + self.years

    @property
    def last_col(self) -> str:
        return L(1 + self.years)

    @property
    def error_range(self) -> str:
        return f"{SHEET}!$A${BAND}:${L(self.target_col)}${self.end}"

    @property
    def chart_cols(self) -> tuple[str, str]:
        first = self.target_col + 2
        return L(first), L(first + CHART_COLS)

    def ref(self, row: int, y: int) -> str:
        return f"{SHEET}!{L(1 + y)}{row}"


def plan_kpis(pl: Plan) -> KpiLayout | None:
    """Rows of the ratio and DuPont tables (Full tier only); the KPI table is the plan's ``kpi`` block."""
    if pl.brief.tier != "full":
        return None
    lay = KpiLayout(pl.t.years)
    r = HDR + 1
    for key in ratio_keys(pl.template):
        lay.ratios[key] = r
        r += 1
    lay.dp_band = r + 1
    lay.dp_hdr = lay.dp_band + 1
    r = lay.dp_hdr + 1
    for key, *_ in DUPONT:
        lay.dupont[key] = r
        r += 1
    lay.end = max([r - 1, *(pl.blocks["kpi"].placed.values() if "kpi" in pl.blocks else [])])
    H = pl.t.years
    for key, row in lay.ratios.items():
        lay.cells[f"kpi_r_{key}"] = {y: lay.ref(row, y) for y in range(1, H + 1)}
    for key, row in lay.dupont.items():
        lay.cells[f"kpi_{key}"] = {y: lay.ref(row, y) for y in range(1, H + 1)}
    if "kpi" in pl.blocks:
        b = pl.blocks["kpi"]
        for key in pl.template.kpi_derived:
            lay.cells[f"kpi_row_{key}"] = {y: pl.stmt_cell("kpi", key, y) for y in range(1, H + 1)}
        assert b.start_row == block_start(pl.template), "KPI table is not where the layout expects it"
    pl.kpis = lay
    return lay


# --------------------------------------------------------------------------- in-file check
def dupont_condition(pl: Plan) -> str:
    """Every year where the DuPont product and the ratio are both numbers: they agree (n/a years are skipped)."""
    lay: KpiLayout = pl.kpis
    r = lay.dupont["dp_diff"]
    return f"SUMPRODUCT(--(ABS({SHEET}!$B${r}:${lay.last_col}${r})>{TOL}))=0"


# --------------------------------------------------------------------------- writing
def _src(pl: Plan, y: int):
    """Excel references of the statement lines the ratios read, for year ``y`` (links from the KPIs sheet)."""

    def c(block: str, key: str) -> str:
        return pl.stmt_cell(block, key, y, from_sheet=SHEET)

    return c


def _formulas(pl: Plan, lay: KpiLayout, y: int) -> dict[str, str]:
    """Formula of every ratio and DuPont cell for year ``y`` (``"n/a"`` where a denominator is zero or negative)."""
    c = _src(pl, y)
    col = L(1 + y)
    rev, gp, ebitda, ni = c("pl", "rev"), c("pl", "gp"), c("pl", "ebitda"), c("pl", "ni")
    assets, equity, loan = c("bs", "total_assets"), c("bs", "total_equity"), c("bs", "loan")
    cur_assets = f"{c('bs', 'cash')}+{c('bs', 'ar')}+{c('bs', 'inv')}"
    cur_liab = f"{c('bs', 'total_liab')}-{loan}"

    def over(num: str, den: str, cond: str | None = None) -> str:
        return f'=IF({cond or den + "<=0"},"{NA}",{num}/{den})'

    out = {
        "gm": over(gp, rev),
        "ebitda_m": over(ebitda, rev),
        "ni_m": over(ni, rev),
        "opex_pct": over(f"-{c('pl', 'opex_total')}", rev),
        "current": over(f"({cur_assets})", f"({cur_liab})"),
        "debt_eq": over(loan, equity),
        "roa": over(ni, assets),
        "roe": over(ni, equity),
        "dp_nm": over(ni, rev),
        "dp_turn": over(rev, assets),
        "dp_mult": over(assets, equity),
    }
    if "payroll_total" in pl.line_at:
        payroll = _rollup(pl, pl.lines["payroll_total"], col, HDR)[1:]
        out["payroll_pct"] = over(f"-{payroll}", rev)
    else:
        out["payroll_pct"] = f'=IF({rev}<=0,"{NA}",0)'
    if pl.template.marketing_key:
        out["mkt_pct"] = over(f"-{c('pl', pl.template.marketing_key)}", rev)
    if y == 1:
        out["growth"] = NA  # no year before the first
    else:
        prev = pl.stmt_cell("pl", "rev", y - 1, from_sheet=SHEET)
        out["growth"] = f'=IF({prev}<=0,"{NA}",{rev}/{prev}-1)'
    nm, turn, mult = (f"{col}{lay.dupont[k]}" for k in ("dp_nm", "dp_turn", "dp_mult"))
    out["dp_roe"] = f'=IF(AND(ISNUMBER({nm}),ISNUMBER({turn}),ISNUMBER({mult})),{nm}*{turn}*{mult},"{NA}")'
    dp, roe = f"{col}{lay.dupont['dp_roe']}", f"{col}{lay.ratios['roe']}"
    out["dp_diff"] = f"=IF(AND(ISNUMBER({dp}),ISNUMBER({roe})),{dp}-{roe},0)"
    return out


def _value(ws, row: int, col: int, value, fmt: str, pct: bool) -> None:
    c = ws.cell(row=row, column=col, value=value)
    c.number_format = fmt
    c.font = st.font(st.GREY, italic=True) if pct else st.font(st.INK)
    c.alignment = Alignment(horizontal="right")


def _target(ws, row: int, col: int, fmt: str) -> None:
    c = ws.cell(row=row, column=col)  # blank on purpose: never an invented benchmark (standard §1)
    st.input_cell(c)
    c.number_format = fmt
    c.alignment = Alignment(horizontal="right")


def _year_header(ws, row: int, lay: KpiLayout, target: bool) -> None:
    last = lay.target_col if target else 1 + lay.years
    for col in range(1, last + 1):
        st.header(ws.cell(row=row, column=col), right=col > 1)
    for y in range(1, lay.years + 1):
        ws.cell(row=row, column=1 + y, value=y).number_format = YEAR_FMT
    if target:
        ws.cell(row=row, column=lay.target_col, value="Target")


def write_kpis(wb, pl: Plan) -> None:
    ws = wb[SHEET]
    lay: KpiLayout = pl.kpis
    H, tc = pl.t.years, lay.target_col
    ws.column_dimensions["A"].width = 46
    for col in range(2, tc + 1):
        ws.column_dimensions[L(col)].width = 11.5
    ws.column_dimensions[L(tc + 1)].width = GAP
    for col in range(tc + 2, tc + 2 + 2 * CHART_COLS):
        ws.column_dimensions[L(col)].width = 9
    st.title(
        ws,
        "KPIs & ratios",
        f"{pl.brief.model_name} — industry metrics and financial ratios for the active scenario",
        8,
        link_col=tc + 2,
    )
    # ---- financial ratios
    st.section(ws, BAND, "Financial ratios", tc)
    _year_header(ws, HDR, lay, target=True)
    labels = {k: (label, fmt) for k, label, fmt in (*RATIOS, *DUPONT)}
    for y in range(1, H + 1):
        formulas = _formulas(pl, lay, y)
        for key, row in lay.ratios.items():
            label, fmt = labels[key]
            _value(ws, row, 1 + y, formulas[key], CELL_FMT[fmt], fmt == "pct")
        for key, row in lay.dupont.items():
            label, fmt = labels[key]
            _value(ws, row, 1 + y, formulas[key], CELL_FMT[fmt], fmt == "pct")
    for key, row in lay.ratios.items():
        label, fmt = labels[key]
        ws.cell(row=row, column=1, value=label).font = st.font(st.GREY, italic=True) if fmt == "pct" else st.font()
        _target(ws, row, tc, CELL_FMT[fmt])
    # ---- DuPont
    st.section(ws, lay.dp_band, DUPONT_TITLE, tc - 1)
    _year_header(ws, lay.dp_hdr, lay, target=False)
    for key, row in lay.dupont.items():
        label, fmt = labels[key]
        small = key == "dp_diff"
        total = key == "dp_roe"
        ws.cell(row=row, column=1, value=label).font = (
            st.font(st.GREY, italic=True, size=8) if small else st.font(st.INK, bold=total)
        )
        if total:
            for col in range(1, tc):
                x = ws.cell(row=row, column=col)
                x.border = st.KEY_TOTAL
                x.font = st.font(st.INK, bold=True)
    # ---- the template's KPI table (written by the builder as the plan's "kpi" block): add Target and a full band
    if "kpi" in pl.blocks:
        b = pl.blocks["kpi"]
        st.section(ws, b.start_row - 1, pl.template.kpi_title, tc)
        st.header(ws.cell(row=b.year_row + 1, column=tc, value="Target"), right=True)
        for row in b.rows:
            _target(ws, b.placed[row.key], tc, CELL_FMT.get(row.fmt) or "General")
    _write_charts(ws, pl, lay)
    ws.freeze_panes = "B4"


# --------------------------------------------------------------------------- charts
def _ser(ws_row: int, last: str, color: str | None = None) -> ch.Ser:
    return ch.Ser(f"'{SHEET}'!$B${ws_row}:${last}${ws_row}", f"'{SHEET}'!$A${ws_row}", color)


def _template_chart(pl: Plan, lay: KpiLayout, spec: KpiChart, cats: str, last: str):
    """Bars (left axis) and optional lines (right axis) from KPI-table or ratio rows; None if no row resolves."""
    rows = pl.blocks["kpi"].placed if "kpi" in pl.blocks else {}
    fmts = {r.key: r.fmt for r in pl.blocks["kpi"].rows} if "kpi" in pl.blocks else {}
    ratio_fmt = {k: f for k, _l, f in RATIOS}

    def resolve(key: str) -> tuple[int, str] | None:
        if key.startswith("r:"):
            k = key[2:]
            return (lay.ratios[k], ratio_fmt[k]) if k in lay.ratios else None
        return (rows[key], fmts[key]) if key in rows else None

    bars = [x for x in map(resolve, spec.bars) if x]
    lines = [x for x in map(resolve, spec.lines) if x]
    if not bars:
        return None
    b_ser = [_ser(r, last) for r, _f in bars]
    fmt = AXIS[bars[0][1]]
    if not lines:
        return ch.bar(spec.title, cats, b_ser, fmt=fmt, stacked=spec.stacked, cat_fmt=YEAR_FMT)
    return ch.combo(
        spec.title,
        cats,
        b_ser,
        [_ser(r, last) for r, _f in lines],
        fmt=fmt,
        line_fmt=AXIS[lines[0][1]],
        stacked=spec.stacked,
        cat_fmt=YEAR_FMT,
    )


def _write_charts(ws, pl: Plan, lay: KpiLayout) -> None:
    last = lay.last_col
    cats = f"'{SHEET}'!$B${HDR}:${last}${HDR}"
    rr = lay.ratios
    charts = [
        ch.line(
            "Margins",
            cats,
            [_ser(rr[k], last) for k in ("gm", "ebitda_m", "ni_m")],
            fmt=AXIS["pct"],
            cat_fmt=YEAR_FMT,
        ),
        ch.bar(
            "Return on assets and return on equity",
            cats,
            [_ser(rr["roa"], last), _ser(rr["roe"], last)],
            fmt=AXIS["pct"],
            cat_fmt=YEAR_FMT,
        ),
    ]
    for spec in pl.template.kpi_charts(pl.t):
        chart = _template_chart(pl, lay, spec, cats, last)
        if chart is not None:
            charts.append(chart)
    for slot, chart in enumerate(charts):
        ch.place(ws, chart, ch.slot_anchor(slot, CHART_ROW, CHART_STEP, lay.chart_cols))


# --------------------------------------------------------------------------- the plain-Python twin
def _div(num: float, den: float):
    """A ratio, or "n/a" when the denominator is zero or negative."""
    return num / den if den > 0 else NA


def reference(tpl, ann: dict, t) -> dict[str, dict[int, float | str]]:
    """Every ratio and DuPont cell recomputed from the annualised reference series (names = manifest names).

    Independent of the formula writer: equity is paid-in capital plus retained earnings, current liabilities are
    total liabilities and equity minus equity minus the loan.
    """
    out: dict[str, dict[int, float | str]] = {}

    def put(name: str, y: int, v) -> None:
        out.setdefault(name, {})[y] = v

    key = tpl.marketing_key
    for y in range(1, t.years + 1):
        rev, ni = ann["rev"][y], ann["ni"][y]
        equity = ann["paid_in"][y] + ann["retained"][y]
        assets, loan = ann["total_assets"][y], ann["loan_close"][y]
        cur_assets = ann["cash_close"][y] + ann["ar"][y] + ann["inv"][y]
        cur_liab = ann["total_le"][y] - equity - loan
        put("kpi_r_growth", y, NA if y == 1 else _growth(ann, y))
        put("kpi_r_gm", y, _div(ann["gp"][y], rev))
        put("kpi_r_ebitda_m", y, _div(ann["ebitda"][y], rev))
        put("kpi_r_ni_m", y, _div(ni, rev))
        put("kpi_r_payroll_pct", y, _div(-ann["payroll_total"][y], rev))
        if key:
            put("kpi_r_mkt_pct", y, _div(-ann[f"opex_{key}"][y], rev))
        put("kpi_r_opex_pct", y, _div(-ann["opex_total"][y], rev))
        put("kpi_r_current", y, _div(cur_assets, cur_liab))
        put("kpi_r_debt_eq", y, _div(loan, equity))
        put("kpi_r_roa", y, _div(ni, assets))
        put("kpi_r_roe", y, _div(ni, equity))
        nm, turn, mult = _div(ni, rev), _div(rev, assets), _div(assets, equity)
        put("kpi_dp_nm", y, nm)
        put("kpi_dp_turn", y, turn)
        put("kpi_dp_mult", y, mult)
        product = nm * turn * mult if all(isinstance(v, float) for v in (nm, turn, mult)) else NA
        put("kpi_dp_roe", y, product)
        put("kpi_dp_diff", y, 0.0)  # the identity: the workbook's difference must be zero
    return out


def _growth(ann: dict, y: int):
    prev = ann["rev"][y - 1]
    return ann["rev"][y] / prev - 1 if prev > 0 else NA


def kpi_row_values(tpl, ann: dict, d: dict, t) -> dict[str, dict[int, float | str]]:
    """Plain-Python value of every checked KPI-table row (series-backed rows and the template's derived rows)."""
    rows: dict[str, dict[int, float | str]] = {k: dict(ann[src]) for k, src in tpl.kpi_checked.items()}
    rows.update(tpl.kpi_reference(ann, d, t))
    return rows


def twin(brief, d: dict, s: dict[str, list[float]], t) -> dict[str, dict[int, float | str]]:
    """Reference values for the manifest's ``kpi_years`` (ratios, DuPont, derived rows) and the Dashboard's
    industry lines (``dash_m_<key>``), recomputed from the per-period series ``s`` (scenario drivers ``d``)."""
    from engine.core import annualise
    from engine.templates import get_template

    tpl = get_template(brief.template)
    ann = annualise(s, t, tpl.kpi_flows)
    out: dict[str, dict[int, float | str]] = {}
    if brief.tier == "full":
        out.update(reference(tpl, ann, t))
    rows = kpi_row_values(tpl, ann, d, t)
    out.update({f"kpi_row_{k}": rows[k] for k in tpl.kpi_derived})
    metrics = list(tpl.dashboard_metrics(t))
    if brief.roles:  # the Dashboard adds the team line (year-end FTE)
        from engine.spec import DashMetric

        metrics.append(DashMetric("team", "Team (FTE, year end)", "fte_total", "dec1"))
    for m in metrics:
        if m.src.startswith("kpi:"):
            out[f"dash_m_{m.key}"] = rows[m.src[4:]]
        else:
            vals = {}
            for y in range(1, t.years + 1):
                chunk = s[m.src][(y - 1) * t.ppy : y * t.ppy]
                vals[y] = sum(chunk) if m.agg == "sum" else chunk[-1]
            out[f"dash_m_{m.key}"] = vals
    return out
