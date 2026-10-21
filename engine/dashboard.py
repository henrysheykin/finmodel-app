"""Dashboard v2 (E27, docs/WORKBOOK_STANDARD.md §2): scenario switch, six tiles, Core financials, industry
metrics, Funding & valuation and four charts on one 68 px grid.

``plan_dashboard`` lays the sheet out while the plan is built (so the manifest knows every cell the challenger
recomputes), ``write_dashboard`` writes it. The sheet holds no financial logic of its own: Core financials roll up
period-grid lines (so Light works without Cash Flow / Balance Sheet sheets), Funding & valuation links to Break-even
and Valuation. The plain-Python twin is ``engine.goals.dashboard_reference`` (the names below are its keys).
"""

from __future__ import annotations

from dataclasses import dataclass, field, replace

from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

from engine import charts as ch
from engine import style as st
from engine.builder import Plan, _rollup, _write_selector, q
from engine.spec import DashChart, DashMetric

SHEET = "Dashboard"
LEFT = 6  # the left column is A:F (label + up to five years); G is the 2-wide gap
LAST_COL = 23  # W
HDR = 11  # Core financials year header (numeric years, shown as "Year n"): the roll-ups read it
CHART_ROW = 10  # charts: H10, P10, H27, P27
YEAR_FMT = '"Year "0'

#: number formats (standard §4): $'000 on report sheets, millions on tiles
MONEY_K = '$#,##0,;($#,##0,);"-"'
TILE_M = '$#,##0.0,,"M";($#,##0.0,,"M");"-"'
PCT = '0.0%;(0.0%);"-"'
FMT = {
    "usdk": MONEY_K,
    "usd2": '$#,##0.00;($#,##0.00);"-"',
    "int": '#,##0;(#,##0);"-"',
    "dec1": '#,##0.0;(#,##0.0);"-"',
    "pct": PCT,
    "x": '0.0"x";(0.0"x");"-"',
}
AXIS = {
    "usdk": ch.MILLIONS,
    "usd2": "$#,##0.00",
    "int": "#,##0",
    "dec1": "#,##0.0",
    "pct": "0%",
    "x": '0.0"x"',
}
#: tiles: first and last column of each of the six 272 px tiles
TILES = (("A", "B"), ("C", "F"), ("H", "K"), ("L", "O"), ("P", "S"), ("T", "W"))

#: Core financials: (key, label, source, kind); kind "" plain, "t" total, "tt" key total, "%" ratio (source = numerator)
CORE = (
    ("rev", "Revenue", "rev", ""),
    ("gp", "Gross profit", "gp", "t"),
    ("gm", "Gross margin", "gp", "%"),
    ("opex", "Operating expenses", "opex_total", ""),
    ("ebitda", "EBITDA", "ebitda", "t"),
    ("ebitda_m", "EBITDA margin", "ebitda", "%"),
    ("ni", "Net income", "ni", "tt"),
    ("ni_m", "Net margin", "ni", "%"),
    ("cfo", "Operating cash flow", "cfo", ""),
    ("cfi", "Investing cash flow", "cfi", ""),
    ("cff", "Financing cash flow", "cff", ""),
    ("net_cf", "Net change in cash", "net_cf", "t"),
    ("cash", "Year-end cash", "cash_close", "tt"),
)


@dataclass
class DashLayout:
    core: dict[str, int] = field(default_factory=dict)
    metrics: list[tuple[DashMetric, int]] = field(default_factory=list)
    ind_band: int = 0
    ind_hdr: int = 0
    fv_band: int = 0
    fv: dict[str, int] = field(default_factory=dict)
    end: int = 0
    chart: DashChart | None = None
    cells: dict[str, str] = field(default_factory=dict)  # manifest name -> 'Dashboard'!B7
    years: dict[str, dict[int, str]] = field(default_factory=dict)


def _ref(cell: str) -> str:
    return f"{SHEET}!{cell}"


def _metrics(pl: Plan) -> list[DashMetric]:
    out = list(pl.template.dashboard_metrics(pl.t))
    if pl.brief.roles and "fte_total" in pl.line_at:
        out.append(DashMetric("team", "Team (FTE, year end)", "fte_total", "dec1"))
    for m in out:
        kpi = pl.blocks["kpi"].placed if "kpi" in pl.blocks else {}
        if m.src.startswith("kpi:") and m.src[4:] not in kpi:
            raise ValueError(f"template {pl.template.key!r}: dashboard metric {m.key!r} needs KPI row {m.src[4:]!r}")
        if not m.src.startswith("kpi:") and m.src not in pl.line_at:
            raise ValueError(f"template {pl.template.key!r}: dashboard metric {m.key!r} needs line {m.src!r}")
    return out


def plan_dashboard(pl: Plan) -> DashLayout:
    """Rows and manifest cells of the Dashboard for this plan (no workbook needed)."""
    lay = DashLayout()
    H = pl.t.years
    full = "dcf" in pl.modules
    r = HDR + 1
    for key, *_ in CORE:
        lay.core[key] = r
        r += 1
    metrics = _metrics(pl)
    lay.chart = pl.template.dashboard_chart(pl.t)
    r += 1  # blank row
    if metrics:
        lay.ind_band, lay.ind_hdr = r, r + 1
        r += 2
        for m in metrics:
            lay.metrics.append((m, r))
            r += 1
        r += 1
    lay.fv_band = r
    r += 1
    names = ["equity", "loan", "required", "gap", "low", "first", "payback"] + (
        ["ev_g", "ev_m", "ev_avg"] if full else []
    )
    for n in names:
        lay.fv[n] = r
        r += 1
    lay.end = r - 1
    # --- manifest cells ---
    for n, key in (
        ("dash_rev", "rev"),
        ("dash_gp", "gp"),
        ("dash_gm", "gm"),
        ("dash_opex", "opex"),
        ("dash_ebitda", "ebitda"),
        ("dash_ebitda_m", "ebitda_m"),
        ("dash_ni", "ni"),
        ("dash_ni_m", "ni_m"),
        ("dash_cfo", "cfo"),
        ("dash_cfi", "cfi"),
        ("dash_cff", "cff"),
        ("dash_net_cf", "net_cf"),
        ("dash_cash", "cash"),
    ):
        lay.years[n] = {y: _ref(f"{L(1 + y)}{lay.core[key]}") for y in range(1, H + 1)}
    for m, row in lay.metrics:  # industry lines: the challenger recomputes each one (KPI-table rows and grid lines)
        lay.years[f"dash_m_{m.key}"] = {y: _ref(f"{L(1 + y)}{row}") for y in range(1, H + 1)}
    for i, (a, _b) in enumerate(TILES):
        lay.cells[f"dash_t{i + 1}"] = _ref(f"{a}7")
        lay.cells[f"dash_t{i + 1}_sub"] = _ref(f"{a}8")
    # tile 6 differs by tier: the names carry the tier so the twin and the layout agree
    if full:
        lay.cells["dash_t6_ev"] = lay.cells.pop("dash_t6")
        lay.cells["dash_t6_ev_sub"] = lay.cells.pop("dash_t6_sub")
    else:
        lay.cells["dash_t6_first"] = lay.cells.pop("dash_t6")
        lay.cells["dash_t6_first_sub"] = lay.cells.pop("dash_t6_sub")
    for n, row in lay.fv.items():
        lay.cells[f"dash_fv_{n}"] = _ref(f"{L(LEFT)}{row}")
    lay.cells["dash_fv_low_label"] = _ref(f"A{lay.fv['low']}")
    pl.dash = lay
    return lay


# --------------------------------------------------------------------------- writing
def _put(ws, ref: str, value, fmt: str | None = None, font=None, align: str | None = None):
    c = ws[ref]
    c.value = value
    c.font = font or st.font(st.INK)
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center")
    return c


def _header(ws, row: int, years: int) -> None:
    """Year header: numeric years shown as "Year n" (the roll-up formulas and the chart categories read them)."""
    for col in range(1, LEFT + 1):
        st.header(ws.cell(row=row, column=col), right=col > 1)
    for y in range(1, years + 1):
        ws.cell(row=row, column=1 + y, value=y).number_format = YEAR_FMT


def _rule(ws, row: int, kind: str) -> None:
    """Total (thin top line) or key total (double underline) across the left column."""
    if not kind or kind == "%":
        return
    border = st.KEY_TOTAL if kind == "tt" else st.TOTAL_TOP
    for col in range(1, LEFT + 1):
        c = ws.cell(row=row, column=col)
        c.border = border
        c.font = st.font(st.INK, bold=True)


def _tile(ws, cols: tuple[str, str], label: str, value, fmt: str, sub, sub_fmt: str | None) -> None:
    a, b = cols
    white = Side(style="thin", color=st.WHITE)
    for row, text, font in (
        (6, label, st.font(st.GREY, bold=True, size=7)),
        (7, value, st.font(st.INK, bold=True, size=18)),
        (8, sub, st.font(st.PINE, size=8)),
    ):
        ws.merge_cells(f"{a}{row}:{b}{row}")
        c = ws[f"{a}{row}"]
        c.value, c.font = text, font
        c.alignment = Alignment(horizontal="center", vertical="center")
        if row == 7:
            c.number_format = fmt
        elif row == 8 and sub_fmt:
            c.number_format = sub_fmt
        for col in range(ws[f"{a}1"].column, ws[f"{b}1"].column + 1):
            x = ws.cell(row=row, column=col)
            x.fill = st.PAPER_FILL
            x.border = Border(top=Side(style="thick", color=st.PINE) if row == 6 else None, left=white, right=white)
    ws.row_dimensions[6].height, ws.row_dimensions[7].height, ws.row_dimensions[8].height = 16, 30, 15


def write_dashboard(wb, pl: Plan, verdict: str) -> None:
    ws = wb[SHEET]
    lay: DashLayout = pl.dash
    H = pl.t.years
    word = pl.t.period_word
    full = "dcf" in pl.modules
    be, dcf = pl.modules["breakeven"], pl.modules.get("dcf")

    def bref(name: str) -> str:
        return f"{q(be.sheet)}!{be.local[name]}"

    # ---- frame: widths, title, selector row, checks verdict
    ws.column_dimensions["A"].width = 28.4
    for col in range(2, LAST_COL + 1):
        ws.column_dimensions[L(col)].width = 9
    ws.column_dimensions["G"].width = 2
    st.title(ws, "Dashboard", pl.brief.model_name, 8, link_col=LAST_COL)
    _write_selector(ws, 4, pl.brief.default_scenario)
    _put(ws, "S4", "Model checks", font=st.font(st.INK, bold=True), align="right")
    ws.merge_cells("T4:W4")
    _put(ws, "T4", f"={verdict}", font=st.font(st.INK, bold=True), align="center")
    st.verdict_colours(ws, "T4:W4", "$T$4")

    # ---- Core financials
    st.section(ws, 10, "Core financials ($'000) — active scenario", LEFT)
    _header(ws, HDR, H)
    for key, label, src, kind in CORE:
        row = lay.core[key]
        pct = kind == "%"
        ws.cell(row=row, column=1, value=label).font = st.font(st.GREY, italic=True) if pct else st.font()
        for y in range(1, H + 1):
            col = L(1 + y)
            if pct:
                f = f"=IF({col}{lay.core['rev']}=0,0,{col}{lay.core[_NUMERATOR[key]]}/{col}{lay.core['rev']})"
            else:
                f = _rollup(pl, pl.lines[src], col, HDR)
            c = ws.cell(row=row, column=1 + y, value=f)
            c.number_format = PCT if pct else MONEY_K
            c.font = st.font(st.GREY, italic=True) if pct else st.font(st.INK)
            c.alignment = Alignment(horizontal="right")
        _rule(ws, row, kind)

    # ---- industry metrics
    if lay.metrics:
        st.section(ws, lay.ind_band, pl.template.kpi_title, LEFT)
        _header(ws, lay.ind_hdr, H)
        for m, row in lay.metrics:
            ws.cell(row=row, column=1, value=m.label).font = st.font()
            for y in range(1, H + 1):
                col = L(1 + y)
                if m.src.startswith("kpi:"):
                    f = "=" + pl.stmt_cell("kpi", m.src[4:], y, from_sheet=SHEET)
                else:
                    f = _rollup(pl, replace(pl.lines[m.src], agg=m.agg), col, lay.ind_hdr)
                c = ws.cell(row=row, column=1 + y, value=f)
                c.number_format = FMT[m.fmt]
                c.font = st.font(st.INK)
                c.alignment = Alignment(horizontal="right")

    # ---- Funding & valuation
    st.section(ws, lay.fv_band, "Funding & valuation ($'000) — active scenario", LEFT)
    month_fmt = f'"{word.capitalize()} "0'
    ev = (lambda n: f"{q(dcf.sheet)}!{dcf.local[n]}") if dcf else None
    low_label = f'="Lowest cash balance ({word} "&{bref("low_p")}&")"'
    rows = [
        ("equity", "Equity raised", f"={pl.input_ref('equity')}", MONEY_K, ""),
        ("loan", "Loan drawn", f"={pl.input_ref('loan_amount')}", MONEY_K, ""),
        ("required", "Funding required (incl. cash buffer)", f"={bref('required')}", MONEY_K, ""),
        ("gap", "Funding surplus / (shortfall)", f"={bref('gap')}", MONEY_K, "t"),
        ("low", low_label, f"={bref('min_cash')}", MONEY_K, ""),
        ("first", f"First {word} with positive EBITDA", f"={bref('first_ebitda')}", month_fmt, ""),
        ("payback", "Cash invested paid back", f"={bref('payback')}", month_fmt, ""),
    ]
    if full:
        rows += [
            ("ev_g", "Enterprise value — perpetual growth", f"={ev('ev_g')}", MONEY_K, ""),
            ("ev_m", "Enterprise value — exit multiple", f"={ev('ev_g_c')}", MONEY_K, ""),
            ("ev_avg", "Enterprise value — average", f"={ev('ev_avg')}", MONEY_K, "tt"),
        ]
    for name, label, f, fmt, kind in rows:
        row = lay.fv[name]
        ws.cell(row=row, column=1, value=label).font = st.font()
        c = ws.cell(row=row, column=LEFT, value=f)
        c.number_format, c.font = fmt, st.font(st.INK)
        c.alignment = Alignment(horizontal="right")
        for col in range(2, LEFT + 1):
            ws.cell(row=row, column=col).border = st.HAIRLINE
        _rule(ws, row, kind)

    # ---- tiles
    def core(key: str, y: int = H) -> str:
        return f"{L(1 + y)}{lay.core[key]}"

    def fv(name: str) -> str:
        return f"{L(LEFT)}{lay.fv[name]}"

    cagr_fmt = '"CAGR "0%;"CAGR "-0%;"CAGR "0%'
    if H > 1:
        cagr = (
            f'=IF(AND({core("rev", 1)}>0,{core("rev")}>=0),({core("rev")}/{core("rev", 1)})^(1/{H - 1})-1,"CAGR n/a")'
        )
    else:
        cagr = "CAGR n/a"
    t5 = f'=IF({fv("gap")}>=0,"Covered by funding raised",-{fv("gap")})'
    if full:
        t6 = ("ENTERPRISE VALUE (DCF)", f"={fv('ev_avg')}", TILE_M, "Average of 2 methods", None)
    else:
        first_ni = bref("first_ni")
        t6 = (
            f"FIRST {word.upper()} WITH POSITIVE EBITDA",
            f"={fv('first')}",
            month_fmt,
            f'=IF(ISNUMBER({first_ni}),"Net income positive from {word} "&{first_ni},'
            '"Net income not positive within horizon")',
            None,
        )
    tiles = [
        (f"REVENUE · YEAR {H}", f"={core('rev')}", TILE_M, cagr, cagr_fmt),
        (
            f"EBITDA MARGIN · YEAR {H}",
            f"={core('ebitda_m')}",
            PCT,
            f"={core('ebitda')}",
            '"EBITDA "$#,##0.0,,"M";"EBITDA "-$#,##0.0,,"M";"EBITDA "$0.0"M"',
        ),
        (
            f"NET INCOME · YEAR {H}",
            f"={core('ni')}",
            TILE_M,
            f"={core('ni_m')}",
            '"Net margin "0%;"Net margin "-0%;"Net margin "0%',
        ),
        ("LOWEST CASH BALANCE", f"={fv('low')}", TILE_M, f"={bref('low_p')}", f'"in {word} "0'),
        ("FUNDING REQUIRED", f"={fv('required')}", TILE_M, t5, '"Shortfall "$#,##0.0,,"M"'),
        t6,
    ]
    for cols, (label, value, fmt, sub, sub_fmt) in zip(TILES, tiles, strict=True):
        _tile(ws, cols, label, value, fmt, sub, sub_fmt)

    _write_charts(ws, pl, lay)
    ws.freeze_panes = None


_NUMERATOR = {"gm": "gp", "ebitda_m": "ebitda", "ni_m": "ni"}


# --------------------------------------------------------------------------- charts
def _row_ser(row: int, last: str, color: str | None = None) -> ch.Ser:
    return ch.Ser(f"'{SHEET}'!$B${row}:${last}${row}", f"'{SHEET}'!$A${row}", color)


def _write_charts(ws, pl: Plan, lay: DashLayout) -> None:
    last = L(1 + pl.t.years)
    cats = f"'{SHEET}'!$B${HDR}:${last}${HDR}"
    core, kw = lay.core, {"cat_fmt": YEAR_FMT}
    charts = [
        ch.combo(
            "Revenue & EBITDA",
            cats,
            [_row_ser(core["rev"], last), _row_ser(core["ebitda"], last)],
            [_row_ser(core["ebitda_m"], last)],
            **kw,
        ),
        ch.bar(
            "Cash flow by activity",
            cats,
            [_row_ser(core[k], last) for k in ("cfo", "cfi", "cff")],
            **kw,
        ),
        ch.bar("Year-end cash", cats, [_row_ser(core["cash"], last, st.INK)], labels=True, legend=False, **kw),
    ]
    charts.append(_industry_chart(pl, lay, cats, last))
    for slot, chart in enumerate(charts):
        ch.place(ws, chart, ch.slot_anchor(slot, CHART_ROW))


def _industry_chart(pl: Plan, lay: DashLayout, cats: str, last: str):
    spec = lay.chart
    by_key = {m.key: (m, row) for m, row in lay.metrics}
    if spec is None or not all(k in by_key for k in (*spec.bars, *spec.lines)):
        spec = DashChart("Gross profit & gross margin", bars=("gp",), lines=("gm",))
        by_key = {
            "gp": (DashMetric("gp", "", "gp", "usdk"), lay.core["gp"]),
            "gm": (DashMetric("gm", "", "gp", "pct"), lay.core["gm"]),
        }
    bars = [_row_ser(by_key[k][1], last) for k in spec.bars]
    lines = [_row_ser(by_key[k][1], last) for k in spec.lines]
    fmt = AXIS[by_key[spec.bars[0]][0].fmt]
    if not lines:
        return ch.bar(spec.title, cats, bars, fmt=fmt, stacked=spec.stacked, cat_fmt=YEAR_FMT)
    return ch.combo(
        spec.title,
        cats,
        bars,
        lines,
        fmt=fmt,
        line_fmt=AXIS[by_key[spec.lines[0]][0].fmt],
        stacked=spec.stacked,
        cat_fmt=YEAR_FMT,
    )
