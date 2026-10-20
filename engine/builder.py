"""Turn a validated brief + template into an .xlsx workbook.

Two phases: ``plan`` assigns every line and statement row to a sheet/row
(pure, no openpyxl), ``build`` writes the workbook from that plan. The plan
also yields the ``Manifest`` the challenger uses to find cells.

Tiers differ in presentation only — the calculation lines (and therefore the
numbers) are identical, so a Light → Full upgrade never changes a figure.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field, replace
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from engine import style as st
from engine.brief import Brief
from engine.core import (
    LAND_LINE,
    balance_sheet_lines,
    headcount_lines,
    liability_delta_lines,
    pnl_head_lines,
    pnl_tail_lines,
)
from engine.goals import apply_goal, module_lines, modules_for, purpose
from engine.spec import FORMATS, DriverSpec, Line, StmtRow, Timing
from engine.templates.base import Template

FIRST_COL = 3  # period 1 is column C on every grid sheet
#: (block, row) shown as key totals with a double underline (docs/WORKBOOK_STANDARD.md §4)
KEY_TOTALS = {("pl", "ni"), ("cf", "cash_close"), ("bs", "total_le"), ("sum", "ni"), ("sum", "cash")}
DASH_KPI_ROW = 47  # template KPI table on the Dashboard (Full) starts here, below Funding & valuation and the charts
#: dollar amounts of the KPI table when it sits on the Dashboard's narrow grid: $135, $4.5k, $8.7M
DASH_USD_K = '[>=1000000]$#,##0.0,,"M";[>=1000]$#,##0.0,"k";$#,##0'
TOKEN = re.compile(r"\{([^{}]+)\}")
PURE_LINK = re.compile(r"^=(?:'[^']+'|[A-Za-z0-9_]+)!\$?[A-Z]{1,3}\$?\d+$")


def q(sheet: str) -> str:
    return sheet if re.fullmatch(r"[A-Za-z0-9_]+", sheet) else f"'{sheet}'"


# --------------------------------------------------------------------------- plan types
@dataclass
class Block:
    name: str
    sheet: str
    start_row: int
    rows: list[StmtRow]
    title: str = ""
    year_row: int = 0  # filled by plan
    placed: dict[str, int] = field(default_factory=dict)


@dataclass
class GridPlan:
    sheet: str
    start_row: int
    lines: list[Line]
    period_row: int = 0
    year_row: int = 0
    rows: dict[str, int] = field(default_factory=dict)
    sections: dict[int, str] = field(default_factory=dict)
    end_row: int = 0


@dataclass
class Manifest:
    tier: str
    granularity: str
    years: int
    sheets: list[str]
    selector: str
    verdict: str
    metrics: dict[str, dict[int, str]]
    grid_ranges: dict[str, str]
    #: goal modules (E07): single cells and per-year cells the challenger recomputes
    goal_cells: dict[str, str] = field(default_factory=dict)
    goal_years: dict[str, dict[int, str]] = field(default_factory=dict)
    #: Dashboard (E27): tiles and Funding & valuation cells, and the Core financials lines per year
    dash_cells: dict[str, str] = field(default_factory=dict)
    dash_years: dict[str, dict[int, str]] = field(default_factory=dict)

    def to_dict(self) -> dict:
        return self.__dict__.copy()


@dataclass
class Plan:
    brief: Brief
    template: Template
    t: Timing
    specs: list[DriverSpec]
    lines: dict[str, Line]
    grids: list[GridPlan]
    blocks: dict[str, Block]
    sheets: list[str]
    selector: tuple[str, str]  # sheet, cell
    input_row: dict[str, int]
    calc_sheet: str
    line_at: dict[str, tuple[str, int]] = field(default_factory=dict)
    #: goal-module layouts (engine/goal_sheets.py), filled by ``plan``
    modules: dict = field(default_factory=dict)
    #: Dashboard layout (engine/dashboard.py), filled by ``plan``
    dash: object = None

    # ---- references ------------------------------------------------------
    def period_cell(self, key: str, p: int, from_sheet: str, absolute: bool = False) -> str:
        sheet, row = self.line_at[key]
        col = L(FIRST_COL + p - 1)
        ref = f"${col}${row}" if absolute else f"{col}{row}"
        return ref if sheet == from_sheet else f"{q(sheet)}!{ref}"

    def grid_range(self, key: str, from_sheet: str | None = None) -> str:
        sheet, row = self.line_at[key]
        rng = f"${L(FIRST_COL)}${row}:${L(FIRST_COL + self.t.periods - 1)}${row}"
        return rng if sheet == from_sheet else f"{q(sheet)}!{rng}"

    def input_ref(self, key: str) -> str:
        return f"Inputs!$E${self.input_row[key]}"

    def input_col(self, key: str, col: str) -> str:
        return f"Inputs!${col}${self.input_row[key]}"

    def stmt_cell(self, block: str, key: str, y: int, from_sheet: str | None = None) -> str:
        b = self.blocks[block]
        ref = f"{L(1 + y)}{b.placed[key]}"
        return ref if b.sheet == from_sheet else f"{q(b.sheet)}!{ref}"

    def stmt_range(self, block: str, key: str) -> str:
        b = self.blocks[block]
        r = b.placed[key]
        return f"{q(b.sheet)}!$B${r}:${L(1 + self.t.years)}${r}"

    def render(self, ln: Line, p: int, sheet: str) -> str:
        text = ln.first if (p == 1 and ln.first is not None) else ln.formula
        grid = next(g for g in self.grids if g.sheet == sheet)

        def sub(m: re.Match) -> str:
            tok = m.group(1)
            if tok.startswith("in:"):
                return self.input_ref(tok[3:])
            if tok == "p":
                return f"{L(FIRST_COL + p - 1)}${grid.period_row}"
            if tok == "y":
                return f"{L(FIRST_COL + p - 1)}${grid.year_row}"
            if tok in ("mpp", "ppy", "dpp"):
                return str(getattr(self.t, tok))
            key, off = tok, 0
            mm = re.fullmatch(r"([a-z0-9_]+)\[(-\d+)\]", tok)
            if mm:
                key, off = mm.group(1), int(mm.group(2))
            if key not in self.line_at:
                raise KeyError(f"line {ln.key!r} references unknown line {key!r}")
            if p + off < 1:
                raise ValueError(f"line {ln.key!r} looks before period 1; give it a `first` formula")
            return self.period_cell(key, p + off, sheet)

        return TOKEN.sub(sub, text)


# --------------------------------------------------------------------------- planning
def _calc_lines(tpl: Template, brief: Brief, t: Timing) -> tuple[list[Line], list[Line], list[Line]]:
    top = tpl.revenue_lines(t) + pnl_head_lines() + tpl.opex_lines(t)
    liabs = tpl.liability_lines(t)
    liab_keys = [ln.key for ln in liabs]
    capex = tpl.capex_lines(t)
    land = any(ln.key == LAND_LINE for ln in capex)
    core = pnl_tail_lines(tpl.opex_keys(t), liab_keys, loan_style=tpl.loan_style, land=land)
    capex[0] = replace(capex[0], section="Fixed assets & capital spending")
    at = next(i for i, ln in enumerate(core) if ln.key == "ppe_gross")
    core = core[:at] + capex + core[at:]
    at = next(i for i, ln in enumerate(core) if ln.key == "d_ap") + 1
    core = core[:at] + liabs + liability_delta_lines(liabs) + core[at:]
    hc = headcount_lines(brief.roles)
    return top + core + tpl.extra_lines(t) + module_lines(brief.tier), hc, balance_sheet_lines(liab_keys)


def _layout_grid(g: GridPlan) -> None:
    g.period_row, g.year_row = g.start_row + 1, g.start_row + 2
    r = g.start_row + 3
    for ln in g.lines:
        if ln.section:
            r += 1
            g.sections[r] = ln.section
            r += 1
        g.rows[ln.key] = r
        r += 1
    g.end_row = r - 1


def _layout_block(b: Block) -> None:
    b.year_row = b.start_row
    r = b.start_row + 2
    for row in b.rows:
        if row.kind == "section":
            r += 1 if r > b.start_row + 2 else 0
        b.placed[row.key] = r
        r += 1


def _pl_rows(tpl: Template, t: Timing, lines: dict[str, Line]) -> list[StmtRow]:
    margin = "=IF({rev}=0,0,{%s}/{rev})"
    rows = [
        StmtRow("rev", "Revenue", "line", "rev", bold=True),
        StmtRow("cogs", lines["cogs"].label, "line", "cogs"),
        StmtRow("gp", "Gross profit", "line", "gp", bold=True, total=True),
        StmtRow("gm", "Gross margin", "formula", margin % "gp", "pct"),
    ]
    rows += [StmtRow(k, lines[k].label, "line", k) for k in tpl.opex_keys(t)]
    rows += [
        StmtRow("opex_total", "Total operating expenses", "line", "opex_total", bold=True, total=True),
        StmtRow("ebitda", "EBITDA", "line", "ebitda", bold=True, total=True),
        StmtRow("ebitda_m", "EBITDA margin", "formula", margin % "ebitda", "pct"),
        StmtRow("dep", "Depreciation", "line", "dep"),
        StmtRow("ebit", "EBIT", "line", "ebit", bold=True),
        StmtRow("interest", "Interest", "line", "interest"),
        StmtRow("ebt", "Pre-tax income", "line", "ebt", total=True),
        StmtRow("tax", "Income tax", "line", "tax"),
        StmtRow("ni", "Net income", "line", "ni", bold=True, total=True),
        StmtRow("ni_m", "Net margin", "formula", margin % "ni", "pct"),
    ]
    return rows


def _cf_rows(liabs: list[Line]) -> list[StmtRow]:
    rows = [
        StmtRow("ni", "Net income", "line", "ni"),
        StmtRow("dep", "Add back depreciation", "line", "dep", neg=True),
        StmtRow("d_ar", "(Increase) / decrease in receivables", "line", "d_ar"),
        StmtRow("d_inv", "(Increase) / decrease in inventory", "line", "d_inv"),
        StmtRow("d_ap", "Increase / (decrease) in payables", "line", "d_ap"),
        *[
            StmtRow(
                f"d_{ln.key}", f"Increase / (decrease) in {ln.label[0].lower() + ln.label[1:]}", "line", f"d_{ln.key}"
            )
            for ln in liabs
        ],
        StmtRow("cfo", "Operating cash flow", "line", "cfo", bold=True, total=True),
        StmtRow("cfi", "Capital spending", "line", "cfi"),
        StmtRow("equity_in", "Equity raised", "line", "equity_in"),
        StmtRow("loan_draw", "Loan drawdown", "line", "loan_draw"),
        StmtRow("loan_repay", "Loan repayment", "line", "loan_repay"),
        StmtRow("cff", "Financing cash flow", "line", "cff", bold=True, total=True),
        StmtRow("net_cf", "Net change in cash", "line", "net_cf", bold=True),
        StmtRow("cash_open", "Opening cash", "line", "cash_open"),
        StmtRow("cash_close", "Closing cash", "line", "cash_close", bold=True, total=True),
    ]
    return rows


def _bs_rows(liabs: list[Line]) -> list[StmtRow]:
    return [
        StmtRow("cash", "Cash", "line", "cash_close"),
        StmtRow("ar", "Accounts receivable", "line", "ar"),
        StmtRow("inv", "Inventory", "line", "inv"),
        StmtRow("ppe_net", "Net fixed assets", "line", "ppe_net"),
        StmtRow("total_assets", "Total assets", "line", "total_assets", bold=True, total=True),
        StmtRow("ap", "Accounts payable", "line", "ap"),
        *[StmtRow(ln.key, ln.label, "line", ln.key) for ln in liabs],
        StmtRow("loan", "Loan", "line", "loan_close"),
        StmtRow("total_liab", "Total liabilities", "line", "total_liab", total=True),
        StmtRow("paid_in", "Paid-in capital", "line", "paid_in"),
        StmtRow("retained", "Retained earnings", "line", "retained"),
        StmtRow("total_equity", "Total equity", "line", "total_equity", total=True),
        StmtRow("total_le", "Total liabilities & equity", "line", "total_le", bold=True, total=True),
        StmtRow("bs_diff", "Check: assets − liabilities − equity (must be 0)", "line", "bs_diff"),
    ]


def _summary_rows() -> list[StmtRow]:
    margin = "=IF({rev}=0,0,{%s}/{rev})"
    return [
        StmtRow("rev", "Revenue", "line", "rev", bold=True),
        StmtRow("gp", "Gross profit", "line", "gp"),
        StmtRow("gm", "Gross margin", "formula", margin % "gp", "pct"),
        StmtRow("opex_total", "Total operating expenses", "line", "opex_total"),
        StmtRow("ebitda", "EBITDA", "line", "ebitda", bold=True, total=True),
        StmtRow("ebitda_m", "EBITDA margin", "formula", margin % "ebitda", "pct"),
        StmtRow("ni", "Net income", "line", "ni", bold=True),
        StmtRow("cash", "Year-end cash", "line", "cash_close", bold=True, total=True),
        StmtRow("loan", "Year-end loan balance", "line", "loan_close"),
    ]


def plan(brief: Brief, tpl: Template) -> Plan:
    brief = apply_goal(brief, tpl=tpl)  # tier module drivers, also for briefs stored before ADR-013
    tpl.validate(brief)
    t = Timing(brief.granularity, brief.horizon_years)
    calc, hc, bs = _calc_lines(tpl, brief, t)
    all_lines = calc + hc + bs
    lines = {ln.key: ln for ln in all_lines}
    if len(lines) != len(all_lines):
        raise ValueError("duplicate line keys in template")
    for k in ("rev", "cogs", "capex"):
        if k not in lines:
            raise ValueError(f"template {tpl.key!r} must define line {k!r}")
    grid_name = "Monthly" if t.granularity == "monthly" else "Annual"
    blocks: dict[str, Block] = {}
    if brief.tier == "full":
        grids = [GridPlan(grid_name, 3, calc + bs), GridPlan("Headcount", 3, hc)]
        blocks["pl"] = Block("pl", "P&L", 3, _pl_rows(tpl, t, lines), "Annual profit & loss")
        liabs = tpl.liability_lines(t)
        blocks["cf"] = Block("cf", "Cash Flow", 3, _cf_rows(liabs), "Annual cash flow statement")
        blocks["bs"] = Block("bs", "Balance Sheet", 3, _bs_rows(liabs), "Balance sheet at year end")
        kpis = tpl.kpi_rows(t)
        if kpis:
            blocks["kpi"] = Block("kpi", "Dashboard", DASH_KPI_ROW, kpis, f"{tpl.kpi_title} — detail")
        calc_sheet = grid_name
    else:
        summary = Block("sum", "P&L", 3, _summary_rows(), "Annual summary")
        _layout_block(summary)
        blocks["sum"] = summary
        start = max(summary.placed.values()) + 3
        kpis = tpl.kpi_rows(t)
        if kpis:
            kb = Block("kpi", "P&L", start, kpis, tpl.kpi_title)
            _layout_block(kb)
            blocks["kpi"] = kb
            start = max(kb.placed.values()) + 3
        grids = [GridPlan("P&L", start, calc + hc)]
        calc_sheet = "P&L"
    selector = ("Dashboard", "B4")  # both tiers (ADR-013)
    sheets = sheet_order(brief.tier, grid_name, tpl.returns)
    for g in grids:
        _layout_grid(g)
    for b in blocks.values():
        if not b.placed:
            _layout_block(b)
    specs = tpl.all_specs(brief.roles, brief.tier)
    input_row = _input_rows(specs)
    pl = Plan(brief, tpl, t, specs, lines, grids, blocks, sheets, selector, input_row, calc_sheet)
    for g in grids:
        for k, r in g.rows.items():
            pl.line_at[k] = (g.sheet, r)
    if modules_for(brief.tier):
        from engine.goal_sheets import plan_modules

        plan_modules(pl)
    if tpl.returns:
        from engine.returns import plan_returns

        plan_returns(pl)
    from engine.dashboard import plan_dashboard

    plan_dashboard(pl)
    return pl


def sheet_order(tier: str, grid_name: str = "Monthly", returns: bool = False) -> list[str]:
    """Sheet set and order per tier — docs/WORKBOOK_STANDARD.md §1 (copied verbatim)."""
    industry = ["Returns"] if returns else []
    if tier == "light":
        return ["Cover", "Dashboard", "Break-even", *industry, "Inputs", "P&L", "Checks"]
    return [
        "Cover",
        "Dashboard",
        "Break-even",
        "Valuation",
        *industry,
        "Inputs",
        "P&L",
        "Cash Flow",
        "Balance Sheet",
        grid_name,
        "Headcount",
        "Checks",
    ]


def _input_rows(specs: list[DriverSpec]) -> dict[str, int]:
    rows, r, last_sec = {}, 6, None
    for s in specs:
        if s.section != last_sec:
            r += 1 if last_sec is not None else 0
            r += 1  # section header row
            last_sec = s.section
        rows[s.key] = r
        r += 1
    return rows


# --------------------------------------------------------------------------- writing
def _write_inputs(wb: Workbook, pl: Plan) -> None:
    ws = wb["Inputs"]
    st.title(
        ws,
        "Inputs — Low / Base / High",
        pl.brief.assumption_note
        or "Every assumption lives here once. Change the yellow cells; the active scenario is chosen on the Dashboard.",
        7,
    )
    sel_sheet, sel_cell = pl.selector
    if sel_sheet == "Inputs":
        _write_selector(ws, 3, pl.brief.default_scenario)
    hdr = ["Driver", "Low", "Base", "High", "Active", "Unit", "Source / comment"]
    for i, h in enumerate(hdr, 1):
        st.header(ws.cell(row=5, column=i, value=h), right=i not in (1, 6, 7))
    sel = f"{q(sel_sheet)}!${sel_cell[0]}${sel_cell[1:]}"
    last_sec = None
    given = {d.key: d for d in pl.brief.drivers}
    for s in pl.specs:
        r = pl.input_row[s.key]
        if s.section != last_sec:
            st.section(ws, r - 1, s.section, 7)
            last_sec = s.section
        d = given[s.key]
        ws.cell(row=r, column=1, value=s.label).font = st.font()
        for col, v in ((2, d.low), (3, d.base), (4, d.high)):
            c = ws.cell(row=r, column=col, value=int(v) if s.integer else v)
            st.input_cell(c)
            c.number_format = FORMATS[s.fmt]
        act = ws.cell(row=r, column=5, value=f"=CHOOSE({sel},B{r},C{r},D{r})")
        act.font = st.font(bold=True)
        act.fill = st.ACTIVE_FILL
        act.number_format = FORMATS[s.fmt]
        ws.cell(row=r, column=6, value=s.unit).font = st.font(st.GREY, italic=True)
        note = d.note or {"asc": "Low ≤ Base ≤ High", "desc": "Low ≥ Base ≥ High (lower is better)"}.get(s.order, "")
        nc = ws.cell(row=r, column=7, value=note)
        nc.font = st.font(st.GREY)
        nc.border = st.HAIRLINE
    last = max(pl.input_row.values())
    ws.cell(
        row=last + 2,
        column=1,
        value=(
            "Legend: yellow cell with blue figure = input you can change · black = formula · "
            "green = link to another sheet · grey = active value (column E follows the scenario on the Dashboard)."
        ),
    ).font = st.font(st.GREY, size=8, italic=True)
    for col, w in zip("ABCDEFG", (46, 12, 12, 12, 12, 12, 34), strict=True):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B6"


def _write_selector(ws, row: int, default: int) -> None:
    ws.cell(row=row, column=1, value="Active scenario").font = st.font(st.INK, bold=True)
    c = ws.cell(row=row, column=2, value=int(default))
    st.switch_cell(c)
    lab = ws.cell(row=row, column=3, value=f'=CHOOSE(B{row},"Low","Base","High")')
    lab.font = st.font(st.PINE, bold=True, size=11)
    note = ws.cell(row=row, column=4, value="1 = Low · 2 = Base · 3 = High — every sheet follows this cell")
    note.font = st.font(st.GREY, italic=True, size=8)
    ws.row_dimensions[row].height = 20
    dv = DataValidation(
        type="list",
        formula1='"1,2,3"',
        allow_blank=False,
        showErrorMessage=True,
        errorTitle="Scenario",
        error="Enter 1, 2 or 3",
        showInputMessage=True,
        promptTitle="Scenario",
        prompt="1 = Low, 2 = Base, 3 = High. Every sheet recalculates.",
    )
    ws.add_data_validation(dv)
    dv.add(f"B{row}")


def _write_grid(wb: Workbook, pl: Plan, g: GridPlan, title: str, sub: str) -> None:
    ws = wb[g.sheet]
    n = pl.t.periods
    width = min(2 + n, 16)
    if g.start_row == 3:
        st.title(ws, title, sub, width, link_col=4)
    else:
        st.section(ws, g.start_row - 1, title, width)
    word = "M" if pl.t.granularity == "monthly" else "Y"
    ws.cell(row=g.start_row, column=1, value=f"{pl.t.period_word.capitalize()}").font = st.font(bold=True)
    st.header(ws.cell(row=g.start_row, column=1))
    ws.cell(row=g.period_row, column=1, value="Period #").font = st.font(st.FAINT, size=7)
    ws.cell(row=g.year_row, column=1, value="Year #").font = st.font(st.FAINT, size=7)
    for p in range(1, n + 1):
        col = FIRST_COL + p - 1
        st.header(ws.cell(row=g.start_row, column=col, value=f"{word}{p}"), right=True)
        c = ws.cell(row=g.period_row, column=col, value=p)
        c.font = st.font(st.FAINT, size=7)
        y = ws.cell(row=g.year_row, column=col, value=f"=ROUNDUP({L(col)}{g.period_row}*{pl.t.mpp}/12,0)")
        y.font = st.font(st.FAINT, size=7)
        ws.column_dimensions[L(col)].width = 11
    for r, text in g.sections.items():
        st.section(ws, r, text, width)
    for ln in g.lines:
        r = g.rows[ln.key]
        helper = ln.label.startswith("Helper")
        pct = ln.fmt == "pct"  # percentages and margins: italic grey (standard §4)
        lab = ws.cell(row=r, column=1, value=ln.label)
        lab.font = st.font(
            st.GREY if helper or pct else "000000", bold=ln.bold, size=8 if helper else st.SIZE, italic=pct
        )
        if ln.total:
            lab.border = st.TOTAL_TOP
        for p in range(1, n + 1):
            f = pl.render(ln, p, g.sheet)
            c = ws.cell(row=r, column=FIRST_COL + p - 1, value=f)
            c.font = st.font(st.GREEN if PURE_LINK.match(f) else st.GREY if pct else "000000", bold=ln.bold, italic=pct)
            c.number_format = FORMATS[ln.fmt]
            if ln.total:
                c.border = st.TOTAL_TOP
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 2
    ws.freeze_panes = ws.cell(row=g.year_row + 1, column=FIRST_COL)


def _write_block(wb: Workbook, pl: Plan, b: Block, sub: str | None) -> None:
    ws = wb[b.sheet]
    years = pl.t.years
    on_dash = b.sheet == "Dashboard"  # the Dashboard grid (engine/dashboard.py) fixes its own widths
    if b.start_row == 3:
        st.title(
            ws,
            b.title if b.sheet != "P&L" or pl.brief.tier == "full" else f"{pl.brief.model_name} — P&L",
            sub,
            1 + years,
        )
    elif b.title:
        st.section(ws, b.start_row - 1, b.title, max(3, 1 + years))
    ws.cell(row=b.year_row, column=1, value="Year #").font = st.font(st.FAINT, size=7)
    st.header(ws.cell(row=b.year_row + 1, column=1))
    for y in range(1, years + 1):
        ws.cell(row=b.year_row, column=1 + y, value=y).font = st.font(st.FAINT, size=7)
        st.header(ws.cell(row=b.year_row + 1, column=1 + y, value=f"Year {y}"), right=True)
        if not on_dash:
            ws.column_dimensions[L(1 + y)].width = 14
    for row in b.rows:
        r = b.placed[row.key]
        pct = row.fmt == "pct"
        key_total = (b.name, row.key) in KEY_TOTALS
        lab = ws.cell(row=r, column=1, value=row.label)
        lab.font = st.font(st.GREY, italic=True) if pct else st.font(bold=row.bold)
        if row.total or key_total:
            lab.border = st.KEY_TOTAL if key_total else st.TOTAL_TOP
        for y in range(1, years + 1):
            col = L(1 + y)
            if row.kind == "line":
                ln = pl.lines[row.src]
                f = _rollup(pl, ln, col, b.year_row)
                if row.neg:
                    f = "=-" + f[1:]
            elif row.kind == "formula":
                # {key} = another row of this table in the same year; {in:driver} = the live input cell
                f = TOKEN.sub(lambda m, col=col: _block_token(pl, b, m.group(1), col), row.src)
            else:
                raise ValueError(row.kind)
            c = ws.cell(row=r, column=1 + y, value=f)
            c.font = st.font(st.GREY, italic=True) if pct else st.font(bold=row.bold)
            c.number_format = DASH_USD_K if on_dash and row.fmt == "usd" else FORMATS[row.fmt]
            if row.total or key_total:
                c.border = st.KEY_TOTAL if key_total else st.TOTAL_TOP
    if not on_dash:
        ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 44)
    if b.start_row == 3:
        ws.freeze_panes = f"B{b.year_row + 2}"


def _block_token(pl: Plan, b: Block, tok: str, col: str) -> str:
    """{key} = another row of this table in the same year; {in:driver} = the live input cell; {year} = year #."""
    if tok.startswith("in:"):
        return pl.input_ref(tok[3:])
    if tok == "year":
        return f"{col}${b.year_row}"
    return f"{col}{b.placed[tok]}"


def _rollup(pl: Plan, ln: Line, col: str, year_row: int) -> str:
    rng = pl.grid_range(ln.key)
    sheet, _ = pl.line_at[ln.key]
    g = next(g for g in pl.grids if g.sheet == sheet)
    yr = f"{q(sheet)}!${L(FIRST_COL)}${g.year_row}:${L(FIRST_COL + pl.t.periods - 1)}${g.year_row}"
    ppy = pl.t.ppy
    if ln.agg == "sum":
        return f"=SUMIF({yr},{col}${year_row},{rng})"
    if ln.agg == "last":
        return f"=INDEX({rng},{col}${year_row}*{ppy})"
    if ln.agg == "first":
        return f"=INDEX({rng},({col}${year_row}-1)*{ppy}+1)"
    raise ValueError(f"line {ln.key!r} has agg={ln.agg!r} and cannot be rolled up")


def _write_checks(wb: Workbook, pl: Plan) -> str:
    ws = wb["Checks"]
    st.title(ws, "Model checks", "Live formulas: they re-run every time an input changes.", 3)
    for i, h in enumerate(["Check", "Result", "What it tests"], 1):
        st.header(ws.cell(row=3, column=i, value=h))
    sel_sheet, sel_cell = pl.selector
    sel = f"{q(sel_sheet)}!{sel_cell}"
    stmt = "pl" if "pl" in pl.blocks else "sum"
    gr = pl.grid_range

    def ok(cond: str, warn: bool = False) -> str:
        return f'=IF({cond},"PASS","{"WARN" if warn else "FAIL"}")'

    checks = [
        (
            "Annual revenue ties to period detail",
            ok(f"ABS(SUM({pl.stmt_range(stmt, 'rev')})-SUM({gr('rev')}))<0.5"),
            "Sum of annual revenue equals the sum of every period",
        ),
        (
            "Annual net income ties to period detail",
            ok(f"ABS(SUM({pl.stmt_range(stmt, 'ni')})-SUM({gr('ni')}))<0.5"),
            "Sum of annual net income equals the sum of every period",
        ),
        (
            "EBITDA reconciliation",
            ok(f"ABS(SUM({gr('ebitda')})-SUM({gr('gp')})-SUM({gr('opex_total')}))<0.5"),
            "EBITDA = gross profit + operating expenses (costs are negative)",
        ),
        (
            "Cash roll-forward",
            ok(f"ABS(INDEX({gr('cash_close')},{pl.t.periods})-SUM({gr('net_cf')}))<0.5"),
            "Closing cash = all net cash flows since the start (opening cash is zero)",
        ),
        (
            "Scenario selector valid",
            ok(f"AND(ISNUMBER({sel}),{sel}>=1,{sel}<=3,INT({sel})={sel})"),
            f"{sel} is 1, 2 or 3",
        ),
    ]
    ordered = [s for s in pl.specs if s.order != "none"]
    if ordered:
        conds = []
        for s in ordered:
            lo, ba, hi = (pl.input_col(s.key, c) for c in "BCD")
            conds.append(f"{lo}<={ba},{ba}<={hi}" if s.order == "asc" else f"{lo}>={ba},{ba}>={hi}")
        checks.append(
            (
                "Scenario ordering of drivers",
                ok(f"AND({','.join(conds)})"),
                "Low / Base / High ordered the right way for every growth and cost driver",
            )
        )
    checks += [
        (
            "Gross margin within 0–100%",
            ok(f"AND(MIN({pl.stmt_range(stmt, 'gm')})>=0,MAX({pl.stmt_range(stmt, 'gm')})<=1)"),
            "Gross margin never negative or above 100%",
        ),
        ("Loan balance never negative", ok(f"MIN({gr('loan_close')})>=-0.5"), "Repayments never exceed the loan"),
    ]
    if pl.modules:
        from engine.goal_sheets import module_checks

        checks += module_checks(pl)
    if "returns" in pl.modules:
        from engine.returns import returns_checks

        checks += returns_checks(pl)
    err_terms = []
    for g in pl.grids:
        rng = f"{q(g.sheet)}!${L(FIRST_COL)}${g.period_row}:${L(FIRST_COL + pl.t.periods - 1)}${g.end_row}"
        err_terms.append(f"SUMPRODUCT(--ISERROR({rng}))")
    for b in pl.blocks.values():
        last = max(b.placed.values())
        err_terms.append(f"SUMPRODUCT(--ISERROR({q(b.sheet)}!$B${b.year_row}:${L(1 + pl.t.years)}${last}))")
    for rng in (m.error_range for m in pl.modules.values()):
        err_terms.append(f"SUMPRODUCT(--ISERROR({rng}))")
    err_terms.append(f"SUMPRODUCT(--ISERROR(Dashboard!$A$6:$W${pl.dash.end}))")  # not row 4: it holds the verdict
    checks.append(
        (
            "No error values",
            ok(f"{'+'.join(err_terms)}=0"),
            "No error values (divide-by-zero, broken references) on any calculation sheet",
        )
    )
    if "bs" in pl.blocks:
        checks.append(
            (
                "Balance sheet balances every period",
                ok(f"SUMPRODUCT(--(ABS({gr('bs_diff')})>0.5))=0"),
                "Assets = liabilities + equity in every period",
            )
        )
    checks.append(
        (
            "Funding sufficient (cash ≥ 0)",
            ok(f"MIN({gr('cash_close')})>=0", warn=True),
            "Business warning, not an error: cash dips below zero in this scenario",
        )
    )
    first = 4
    for i, (name, f, desc) in enumerate(checks):
        r = first + i
        ws.cell(row=r, column=1, value=name).font = st.font()
        ws.cell(row=r, column=2, value=f).font = st.font(bold=True)
        ws.cell(row=r, column=3, value=desc).font = st.font(st.GREY, size=9)
    last = first + len(checks) - 1
    vr = last + 2
    ws.cell(row=vr, column=1, value="Overall verdict").font = st.font(bold=True)
    c = ws.cell(
        row=vr,
        column=2,
        value=(
            f'=IF(COUNTIF(B{first}:B{last},"FAIL")>0,"FAIL",'
            f'IF(COUNTIF(B{first}:B{last},"WARN")>0,"PASS WITH WARNINGS","PASS"))'
        ),
    )
    c.font = st.font(bold=True)
    for col in (1, 2, 3):
        ws.cell(row=vr, column=col).border = st.KEY_TOTAL
    st.verdict_colours(ws, f"B{first}:B{vr}", f"B{first}")
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 64
    return f"Checks!B{vr}"


#: Cover contents: sheet -> what it shows (docs/WORKBOOK_STANDARD.md §1)
SHEET_DESC = {
    "Dashboard": "Scenario switch, key numbers, core financials",
    "Summary": "Income statement, cash flow and balance sheet — every year and any year by month",
    "Scenarios": "Low / Base / High side by side: drivers, outputs, charts",
    "Break-even": "Startup costs, funding need, break-even timing and revenue",
    "Valuation": "Free cash flow, DCF (perpetual growth and exit multiple), WACC × growth sensitivity",
    "Sources & Uses": "Funding structure, sources and uses of cash for the first months",
    "KPIs": "Financial ratios, DuPont analysis, industry metrics",
    "Returns": "Property returns: unlevered and levered IRR, NPV, multiple",
    "Inputs": "Every assumption, Low / Base / High, with units and comments",
    "P&L": "Annual profit & loss",
    "Cash Flow": "Annual cash flow statement",
    "Balance Sheet": "Balance sheet at year end",
    "Monthly": "Month-by-month engine of the model",
    "Annual": "Year-by-year engine of the model",
    "Headcount": "Team and payroll",
    "Checks": "Live integrity checks and the overall verdict",
}
GROUPS = (
    ("Presentation", st.PRESENTATION),
    ("Inputs", {"Inputs"}),
    ("Statements", st.STATEMENTS),
    ("Calculations", {"Monthly", "Annual", "Headcount"}),
    ("Control", {"Checks"}),
)


def _write_cover(wb: Workbook, pl: Plan, verdict: str, built_on: date) -> None:
    ws = wb["Cover"]
    b = pl.brief
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 5
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 72
    ws["B2"] = b.model_name
    ws["B2"].font = st.font(st.INK, bold=True, size=20)
    plan_word = "Full" if b.tier == "full" else "Light"
    ws["B3"] = (
        f"{b.horizon_years}-year financial model · {b.granularity} detail · Low / Base / High scenarios · "
        f"US dollars · {plan_word} plan"
    )
    ws["B3"].font = st.font(st.GREY, size=10)
    rows = [
        ("Purpose", purpose(b.goal), None),
        ("Active scenario", '=Dashboard!C4&"  (change it on the Dashboard, cell B4)"', st.PINE),
        ("Model checks", f"={verdict}", None),
        ("Prepared", f"{built_on.day} {built_on:%B %Y}", None),
    ]
    for i, (label, value, colour) in enumerate(rows):
        r = 5 + i
        ws.cell(row=r, column=2, value=label).font = st.font(st.INK, bold=True)
        c = ws.cell(row=r, column=4, value=value)
        c.font = st.font(colour or (st.INK if label == "Model checks" else "000000"), bold=label != "Prepared")
        if label == "Model checks":
            st.verdict_colours(ws, c.coordinate, c.coordinate)
    st.section(ws, 10, "Contents", 3, first_col=2)
    r, n = 11, 0
    for group, names in GROUPS:
        shown = [s for s in pl.sheets if s in names and s != "Cover"]
        if not shown:
            continue
        ws.cell(row=r, column=3, value=group.upper()).font = st.font(st.PINE, bold=True, size=8)
        r += 1
        for name in shown:
            n += 1
            ws.cell(row=r, column=2, value=n).font = st.font(st.GREY)
            ws.cell(row=r, column=2).alignment = Alignment(horizontal="right")
            link = ws.cell(row=r, column=3, value=f'=HYPERLINK("#{q(name)}!A1","{name}")')
            link.font = st.font(st.PINE, bold=True, underline="single")
            ws.cell(row=r, column=4, value=SHEET_DESC.get(name, "")).font = st.font()
            for col in (2, 3, 4):
                ws.cell(row=r, column=col).border = st.HAIRLINE
            r += 1
    r += 1
    ws.cell(row=r, column=2, value="How to read this file").font = st.font(st.INK, bold=True)
    legend = [
        ("input", "Input you can change — yellow cell with a blue figure (Inputs sheet and the view switches)"),
        ("formula", "Formula — calculated, do not overwrite"),
        ("link", "Link to another sheet"),
        ("active", "Active value the model uses (follows the scenario)"),
    ]
    for kind, text in legend:
        r += 1
        c = ws.cell(row=r, column=3, value=1000)
        c.number_format = "#,##0"
        c.alignment = Alignment(horizontal="center")
        if kind == "input":
            st.input_cell(c)
        else:
            c.font = st.font(st.GREEN if kind == "link" else "000000")
            c.border = st.INPUT_BOX
            if kind == "active":
                c.fill = st.ACTIVE_FILL
        ws.cell(row=r, column=4, value=text).font = st.font()


def manifest_for(pl: Plan, verdict: str) -> Manifest:
    years = range(1, pl.t.years + 1)
    if pl.brief.tier == "full":
        src = {
            "rev": ("pl", "rev"),
            "gp": ("pl", "gp"),
            "ebitda": ("pl", "ebitda"),
            "ni": ("pl", "ni"),
            "cash": ("bs", "cash"),
            "loan": ("bs", "loan"),
            "total_assets": ("bs", "total_assets"),
            "total_le": ("bs", "total_le"),
            "cfo": ("cf", "cfo"),
        }
    else:
        src = {
            "rev": ("sum", "rev"),
            "gp": ("sum", "gp"),
            "ebitda": ("sum", "ebitda"),
            "ni": ("sum", "ni"),
            "cash": ("sum", "cash"),
            "loan": ("sum", "loan"),
        }
    if "kpi" in pl.blocks:
        src.update({f"kpi:{k}": ("kpi", k) for k in pl.template.kpi_checked})
    metrics = {m: {y: pl.stmt_cell(b, k, y) for y in years} for m, (b, k) in src.items()}
    sheet, cell = pl.selector
    goal_cells: dict[str, str] = {}
    goal_years: dict[str, dict[int, str]] = {}
    for m in pl.modules.values():
        goal_cells.update(m.cells)
        goal_years.update({k: {y: pl.stmt_cell(m.block, k, y) for y in years} for k in m.checked_rows})
    return Manifest(
        pl.brief.tier,
        pl.t.granularity,
        pl.t.years,
        pl.sheets,
        f"{sheet}!{cell}",
        verdict,
        metrics,
        {k: pl.grid_range(k) for k in ("rev", "ebitda", "ni", "cash_close")},
        goal_cells,
        goal_years,
        dict(pl.dash.cells),
        {k: dict(v) for k, v in pl.dash.years.items()},
    )


def build(brief: Brief, tpl: Template, out_path: str | Path, built_on: date | None = None) -> Manifest:
    pl = plan(brief, tpl)
    wb = Workbook()
    wb.active.title = pl.sheets[0]
    for name in pl.sheets[1:]:
        wb.create_sheet(name)
    _write_inputs(wb, pl)
    unit = "$ unless stated; costs shown as negatives. Every period uses the same formula."
    for g in pl.grids:
        if g.sheet == "Headcount":
            _write_grid(wb, pl, g, "Headcount & payroll", unit)
        elif pl.brief.tier == "full":
            _write_grid(wb, pl, g, f"{pl.t.period_word.capitalize()}ly projection".replace("Yearly", "Annual"), unit)
        else:
            _write_grid(wb, pl, g, f"{pl.t.period_word.capitalize()}ly detail".replace("Yearly", "Annual"), unit)
    for b in pl.blocks.values():
        _write_block(wb, pl, b, "Rolled up from the period detail (SUMIF on the year row; balances at year end).")
    verdict = _write_checks(wb, pl)
    from engine.dashboard import write_dashboard

    write_dashboard(wb, pl, verdict)
    _write_cover(wb, pl, verdict, built_on or date.today())
    if pl.modules:
        from engine.goal_sheets import write_modules

        write_modules(wb, pl)
    for ws in wb.worksheets:
        st.finish_sheet(ws, landscape=ws.title != "Cover", zoom=90 if ws.title in ("Monthly", "Headcount") else 100)
        ws.sheet_view.tabSelected = ws.title == "Cover"
    wb.active = 0
    wb.properties.creator = ""
    wb.properties.title = pl.brief.model_name
    wb.save(out_path)
    return manifest_for(pl, verdict)
