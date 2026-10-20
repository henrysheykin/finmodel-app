"""Shared workbook style kit — workbook standard v1 (ADR-013, docs/WORKBOOK_STANDARD.md §4).

Arial 9; ink / pine identity; input = yellow fill + blue font; formula = black; link to another sheet = green
(Inputs and calculation sheets) or ink (presentation sheets).
"""

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "Arial"
SIZE = 9
INK = "1B2A2F"
PINE = "0F6E56"
SAGE = "7FB5A3"
AMBER = "D9A21B"
RED = "B8322A"
PAPER = "F7F8F3"
CHART_GREY = "A3ABA8"
BLUE = "0000FF"
GREEN = "008000"
GREY = "6B7280"
FAINT = "A3ABA8"
WHITE = "FFFFFF"

INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
SWITCH_FILL = PatternFill("solid", fgColor="FFE08A")
ACTIVE_FILL = PatternFill("solid", fgColor="F1F3EE")
PAPER_FILL = PatternFill("solid", fgColor=PAPER)
BAND_FILL = PatternFill("solid", fgColor=PINE)
#: kept for callers that still import it; the selector now uses SWITCH_FILL
YELLOW_FILL = SWITCH_FILL

_thin = Side(style="thin", color="C9CEC9")
_thin_ink = Side(style="thin", color=INK)
_med_ink = Side(style="medium", color=INK)
_dbl_ink = Side(style="double", color=INK)
BOTTOM = Border(bottom=_med_ink)
TOTAL_TOP = Border(top=_thin_ink)
KEY_TOTAL = Border(top=_thin_ink, bottom=_dbl_ink)
INPUT_BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
SWITCH_BOX = Border(left=_med_ink, right=_med_ink, top=_med_ink, bottom=_med_ink)
HAIRLINE = Border(bottom=_thin)

#: tab colours by sheet group (§4)
TAB = {"presentation": PINE, "inputs": "E0B43C", "statements": INK, "calc": CHART_GREY, "checks": RED}
PRESENTATION = frozenset(
    {"Cover", "Dashboard", "Summary", "Scenarios", "Break-even", "Valuation", "Sources & Uses", "KPIs", "Returns"}
)
STATEMENTS = frozenset({"P&L", "Cash Flow", "Balance Sheet"})


def font(
    color: str = "000000", bold: bool = False, size: int = SIZE, italic: bool = False, underline: str | None = None
) -> Font:
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic, underline=underline)


def contents_link(ws, column: int) -> None:
    """'◂ Contents' in row 1 of every sheet but the Cover."""
    c = ws.cell(row=1, column=column, value='=HYPERLINK("#\'Cover\'!A1","◂ Contents")')
    c.font = font(PINE, size=8, underline="single")


def title(ws, text: str, sub: str | None = None, width_cols: int = 8, link_col: int | None = None) -> None:
    ws["A1"] = text
    ws["A1"].font = font(INK, bold=True, size=14)
    ws.row_dimensions[1].height = 22
    if sub:
        ws["A2"] = sub
        ws["A2"].font = font(GREY, size=9)
    if ws.title != "Cover":
        contents_link(ws, link_col or width_cols + 2)


def section(ws, row: int, text: str, ncols: int, first_col: int = 1) -> None:
    ws.cell(row=row, column=first_col, value=text)
    for c in range(first_col, first_col + ncols):
        x = ws.cell(row=row, column=c)
        x.fill = BAND_FILL
        x.font = font(WHITE, bold=True)
        x.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 16


def header(cell, right: bool = False) -> None:
    cell.font = font(INK, bold=True)
    cell.border = BOTTOM
    cell.alignment = Alignment(horizontal="right" if right else "left")


def input_cell(cell) -> None:
    cell.font = font(BLUE)
    cell.fill = INPUT_FILL
    cell.border = INPUT_BOX


def switch_cell(cell) -> None:
    """A view switch (scenario selector, detailed year): stronger yellow with a frame."""
    cell.font = font(BLUE, bold=True, size=11)
    cell.fill = SWITCH_FILL
    cell.border = SWITCH_BOX
    cell.alignment = Alignment(horizontal="center", vertical="center")


def verdict_colours(ws, ref: str, anchor: str) -> None:
    """PASS / WARN / FAIL colours on a verdict cell or range; ``anchor`` = its top-left cell, relative refs."""
    for cond, fill, colour in (
        (f'ISNUMBER(SEARCH("FAIL",{anchor}))', "F6D5D2", RED),
        (f'ISNUMBER(SEARCH("WARN",{anchor}))', "FBEFCF", "8A6100"),
        (f'{anchor}="PASS"', "DCEFE8", PINE),
    ):
        ws.conditional_formatting.add(
            ref, FormulaRule(formula=[cond], fill=PatternFill("solid", fgColor=fill), font=font(colour, bold=True))
        )


def tab_group(name: str) -> str:
    if name in PRESENTATION:
        return "presentation"
    if name == "Inputs":
        return "inputs"
    if name in STATEMENTS:
        return "statements"
    if name == "Checks":
        return "checks"
    return "calc"


def finish_sheet(ws, landscape: bool = True, zoom: int = 100) -> None:
    """Gridlines off, tab colour by group, print setup (§4)."""
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = zoom
    ws.sheet_properties.tabColor = TAB[tab_group(ws.title)]
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.left.text, ws.oddFooter.right.text = "&A", "Page &P of &N"
    ws.oddFooter.left.size = ws.oddFooter.right.size = 8
