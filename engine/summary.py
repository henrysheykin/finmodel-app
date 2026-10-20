"""Summary sheet (E28, docs/WORKBOOK_STANDARD.md §1, §3): the three statements on one page.

Left (A:F): every year of the horizon, linked to the P&L / Cash Flow / Balance Sheet cells. Right (H:S, monthly
models only): one selected year by month — ``INDEX`` into the Monthly grid driven by the "Select detailed year"
switch (R2). Each block is padded to the height of a standard chart, which sits to its right (U<top>; annual models,
which have no monthly side, put the chart at H<top> and plot the year columns).

The sheet holds no financial logic: every figure is a link or an ``INDEX`` into a statement / grid line, so the
challenger covers it through the statements; the in-file check "Summary ties to statements" proves the links.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from engine import charts as ch
from engine import style as st
from engine.builder import FIRST_COL, Plan, q
from engine.dashboard import LEFT, MONEY_K, PCT, YEAR_FMT
from engine.spec import StmtRow

SHEET = "Summary"
TOP = 5  # first block (row 2 holds the year switch)
CHART_ROWS = 16  # a 14.2 x 8.3 cm chart is 16 rows high: blocks are padded to it
MONTH_FIRST, MONTH_LAST = 8, 19  # H:S
MONTH_FMT = '"M"0'
K_AXIS = '$#,##0,"K"'

#: rows of each statement block shown here, in order (None = all rows but SKIP)
KEEP = {
    "pl": (
        "rev",
        "cogs",
        "gp",
        "gm",
        "opex_total",
        "ebitda",
        "ebitda_m",
        "dep",
        "ebit",
        "interest",
        "ebt",
        "tax",
        "ni",
        "ni_m",
    ),
    "cf": ("cfo", "cfi", "equity_in", "loan_draw", "loan_repay", "cff", "net_cf", "cash_close"),
    "bs": None,
}
SKIP = {"bs": ("bs_diff",)}
LABELS = {("cf", "cash_close"): "Closing cash"}
#: totals (thin top line) and key totals (double underline)
TOTALS = {
    ("pl", "gp"),
    ("pl", "ebitda"),
    ("pl", "ebit"),
    ("cf", "cfo"),
    ("cf", "cff"),
    ("cf", "net_cf"),
    ("bs", "total_assets"),
    ("bs", "total_liab"),
    ("bs", "total_equity"),
}
KEY_TOTALS = {("pl", "ni"), ("cf", "cash_close"), ("bs", "total_le")}
#: margin rows: numerator row of the same block (denominator = revenue)
MARGIN_OF = {"gm": "gp", "ebitda_m": "ebitda", "ni_m": "ni"}
#: (left band, right band) per block; the right band is "<text> — Year n, by month"
TITLES = {
    "pl": ("Income statement ($'000) — by year", "Income statement ($'000) — Year "),
    "cf": ("Cash flow ($'000) — by year", "Cash flow ($'000) — Year "),
    "bs": ("Balance sheet at year end ($'000)", "Balance sheet at month end ($'000) — Year "),
}
CHART_TITLES = {"pl": "Income statement", "cf": "Cash flow", "bs": "Assets and equity"}


@dataclass
class SumBlock:
    key: str  # pl | cf | bs (the statement block it mirrors)
    top: int
    rows: list[tuple[StmtRow, int]]  # (statement row, sheet row)

    @property
    def header(self) -> int:
        return self.top + 1

    @property
    def end(self) -> int:
        return self.rows[-1][1]


@dataclass
class SumLayout:
    monthly: bool
    years: int
    blocks: dict[str, SumBlock] = field(default_factory=dict)
    end: int = 0
    cells: dict[str, dict[int, str]] = field(default_factory=dict)  # manifest: 'sum_rev' -> {year: 'Summary!B7'}

    @property
    def last_col(self) -> str:
        return L(1 + self.years)

    @property
    def link_col(self) -> int:
        return 21 if self.monthly else 16  # U above the charts / P over the annual chart

    @property
    def error_range(self) -> str:
        return f"{SHEET}!$A${TOP}:${L(MONTH_LAST if self.monthly else LEFT)}${self.end}"

    def row(self, block: str, key: str) -> int:
        return next(r for s, r in self.blocks[block].rows if s.key == key)

    def ref(self, block: str, key: str, y: int) -> str:
        return f"{SHEET}!{L(1 + y)}{self.row(block, key)}"


def _rows_of(pl: Plan, name: str) -> list[StmtRow]:
    keep = KEEP[name]
    rows = pl.blocks[name].rows
    if keep is not None:
        by_key = {r.key: r for r in rows}
        return [by_key[k] for k in keep]
    return [r for r in rows if r.key not in SKIP.get(name, ())]


def plan_summary(pl: Plan) -> SumLayout | None:
    """Rows of the Summary for this plan (Full tier only); stored in ``pl.summary``."""
    if pl.brief.tier != "full":
        return None
    lay = SumLayout(pl.t.granularity == "monthly", pl.t.years)
    top = TOP
    for name in ("pl", "cf", "bs"):
        rows = [(s, top + 2 + i) for i, s in enumerate(_rows_of(pl, name))]
        lay.blocks[name] = SumBlock(name, top, rows)
        top = max(rows[-1][1] + 2, top + CHART_ROWS + 1)  # padded to the chart height, one blank row between
    lay.end = lay.blocks["bs"].end
    for cell, (block, key) in (
        ("sum_rev", ("pl", "rev")),
        ("sum_ni", ("pl", "ni")),
        ("sum_cash", ("cf", "cash_close")),
    ):
        lay.cells[cell] = {y: lay.ref(block, key, y) for y in range(1, pl.t.years + 1)}
    pl.summary = lay
    return lay


# --------------------------------------------------------------------------- the in-file check
def tie_condition(pl: Plan) -> str:
    """Condition of "Summary ties to statements": every year's links, and the selected year's monthly side."""
    lay: SumLayout = pl.summary
    last = lay.last_col
    parts = []
    for block, key in (("pl", "rev"), ("pl", "ni"), ("cf", "cash_close"), ("bs", "cash")):
        r = lay.row(block, key)
        parts.append(f"SUMPRODUCT(--(ABS({SHEET}!$B${r}:${last}${r}-{pl.stmt_range(block, key)})>0.5))=0")
    if lay.monthly:

        def year(r: int) -> str:
            return f"INDEX({SHEET}!$B${r}:${last}${r},{SHEET}!$R$2)"

        for key in ("rev", "ni"):
            r = lay.row("pl", key)
            parts.append(f"ABS(SUM({SHEET}!$H${r}:$S${r})-{year(r)})<0.5")
        r = lay.row("cf", "cash_close")
        parts.append(f"ABS({SHEET}!$S${r}-{year(r)})<0.5")
    return f"IFERROR(AND({','.join(parts)}),FALSE)"  # a bad switch value is a FAIL, not an error cell


# --------------------------------------------------------------------------- writing
def _grid_period_row(pl: Plan) -> str:
    g = next(g for g in pl.grids if g.sheet == pl.calc_sheet)
    return f"{q(g.sheet)}!${L(FIRST_COL)}${g.period_row}:${L(FIRST_COL + pl.t.periods - 1)}${g.period_row}"


def _rule(ws, row: int, cols: range, key_total: bool) -> None:
    border = st.KEY_TOTAL if key_total else st.TOTAL_TOP
    for col in cols:
        c = ws.cell(row=row, column=col)
        c.border = border
        c.font = st.font(st.INK, bold=True)


def _band_header(ws, blk: SumBlock, pl: Plan, lay: SumLayout) -> None:
    left, right = TITLES[blk.key]
    st.section(ws, blk.top, left, LEFT)
    for col in range(1, LEFT + 1):
        st.header(ws.cell(row=blk.header, column=col), right=col > 1)
    for y in range(1, pl.t.years + 1):
        ws.cell(row=blk.header, column=1 + y, value=y).number_format = YEAR_FMT
    if not lay.monthly:
        return
    suffix = '&", by month"' if blk.key != "bs" else ""
    st.section(ws, blk.top, f'="{right}"&$R$2{suffix}', MONTH_LAST - MONTH_FIRST + 1, MONTH_FIRST)
    prow = _grid_period_row(pl)
    for k in range(1, 13):
        c = ws.cell(row=blk.header, column=MONTH_FIRST + k - 1, value=f"=INDEX({prow},($R$2-1)*{pl.t.ppy}+{k})")
        c.number_format = MONTH_FMT
        st.header(c, right=True)


def _put(ws, row: int, col: int, value: str, pct: bool) -> None:
    c = ws.cell(row=row, column=col, value=value)
    c.number_format = PCT if pct else MONEY_K
    c.font = st.font(st.GREY, italic=True) if pct else st.font(st.INK)
    c.alignment = Alignment(horizontal="right")


def _write_block(ws, pl: Plan, lay: SumLayout, blk: SumBlock) -> None:
    _band_header(ws, blk, pl, lay)
    rev_row = lay.row("pl", "rev") if blk.key == "pl" else 0
    for s, r in blk.rows:
        pct = s.kind == "formula"
        total, key_total = (blk.key, s.key) in TOTALS, (blk.key, s.key) in KEY_TOTALS
        label = ws.cell(row=r, column=1, value=LABELS.get((blk.key, s.key), s.label))
        label.font = st.font(st.GREY, italic=True) if pct else st.font(st.INK, bold=total or key_total)
        for y in range(1, pl.t.years + 1):
            _put(ws, r, 1 + y, "=" + pl.stmt_cell(blk.key, s.key, y, from_sheet=SHEET), pct)
        if lay.monthly:
            for k in range(1, 13):
                col = MONTH_FIRST + k - 1
                if pct:
                    num = lay.row(blk.key, MARGIN_OF[s.key])
                    f = f"=IF({L(col)}{rev_row}=0,0,{L(col)}{num}/{L(col)}{rev_row})"
                else:
                    f = f"={'-' if s.neg else ''}INDEX({pl.grid_range(s.src)},($R$2-1)*{pl.t.ppy}+{k})"
                _put(ws, r, col, f, pct)
        if total or key_total:
            _rule(ws, r, range(1, LEFT + 1), key_total)
            if lay.monthly:
                _rule(ws, r, range(MONTH_FIRST, MONTH_LAST + 1), key_total)


def _write_switch(ws, pl: Plan) -> None:
    c = ws["P2"]
    c.value, c.font = "Select detailed year:", st.font(st.INK, bold=True)
    c.alignment = Alignment(horizontal="right", vertical="center")
    ws.merge_cells("R2:S2")
    sw = ws["R2"]
    sw.value, sw.number_format = 1, YEAR_FMT
    st.switch_cell(sw)
    ws["S2"].border = st.SWITCH_BOX
    ws.row_dimensions[2].height = 20
    dv = DataValidation(
        type="list",
        formula1=f'"{",".join(str(y) for y in range(1, pl.t.years + 1))}"',
        allow_blank=False,
        showErrorMessage=True,
        errorTitle="Detailed year",
        error=f"Enter a year from 1 to {pl.t.years}",
        showInputMessage=True,
        promptTitle="Detailed year",
        prompt="Pick the year shown month by month on the right",
    )
    ws.add_data_validation(dv)
    dv.add("R2")


def write_summary(wb, pl: Plan) -> None:
    ws = wb[SHEET]
    lay: SumLayout = pl.summary
    ws.column_dimensions["A"].width = 30
    for col in range(2, LEFT + 1):
        ws.column_dimensions[L(col)].width = 11.5
    ws.column_dimensions["G"].width = 2
    for col in range(8, 37):
        ws.column_dimensions[L(col)].width = 9.2 if col <= MONTH_LAST else 9
    if lay.monthly:
        ws.column_dimensions["T"].width = 2
    st.title(ws, "Financial statements summary", pl.brief.model_name, 8, link_col=lay.link_col)
    if lay.monthly:
        _write_switch(ws, pl)
    for blk in lay.blocks.values():
        _write_block(ws, pl, lay, blk)
    _write_charts(ws, lay)
    ws.freeze_panes = "B5"


# --------------------------------------------------------------------------- charts
def _write_charts(ws, lay: SumLayout) -> None:
    c0, c1 = ("H", "S") if lay.monthly else ("B", lay.last_col)
    fmt = K_AXIS if lay.monthly else ch.MILLIONS
    cat_fmt = MONTH_FMT if lay.monthly else YEAR_FMT
    by = "month" if lay.monthly else "year"

    def ser(block: str, key: str, name: str | None = None, color: str | None = None) -> ch.Ser:
        r = lay.row(block, key)
        return ch.Ser(f"'{SHEET}'!${c0}${r}:${c1}${r}", name or f"'{SHEET}'!$A${r}", color)

    def cats(block: str) -> str:
        r = lay.blocks[block].header
        return f"'{SHEET}'!${c0}${r}:${c1}${r}"

    kw = {"fmt": fmt, "line_fmt": fmt, "cat_fmt": cat_fmt}
    charts = {
        "pl": ch.combo(
            f"{CHART_TITLES['pl']} by {by}",
            cats("pl"),
            [ser("pl", "rev")],
            [ser("pl", "ebitda", color=st.AMBER), ser("pl", "ni", color=st.INK)],
            secondary=False,
            **kw,
        ),
        "cf": ch.combo(
            f"{CHART_TITLES['cf']} by {by}",
            cats("cf"),
            [ser("cf", "cfo"), ser("cf", "cfi"), ser("cf", "cff")],
            [ser("cf", "cash_close", "Closing cash (right axis)", st.INK)],
            **kw,
        ),
        "bs": ch.combo(
            f"{CHART_TITLES['bs']} at {by} end",
            cats("bs"),
            [ser("bs", "cash"), ser("bs", "ar", "Receivables"), ser("bs", "ppe_net", "Fixed assets")],
            [ser("bs", "total_equity", color=st.INK)],
            stacked=True,
            secondary=False,
            **kw,
        ),
    }
    col = "U" if lay.monthly else "H"
    for key, chart in charts.items():
        ch.place(ws, chart, f"{col}{lay.blocks[key].top}")
